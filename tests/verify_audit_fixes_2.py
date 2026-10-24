"""The fixes for the audit's open findings B12-B28 (Task 24).

B12  A second parent's report for one dose is refused in words, not with a server error.
B13  A huge reorder level is refused; B14 "inf", "nan", "1e9" and other non-numbers too.
B24  A PDF export never fails on long text; B25 the Patient Records search runs clean.
B20  One compose form sends one SMS broadcast. B21, B14 maximum lengths.
The sections for the later batches are added below as each batch is done.
Every name, date and number below is made up.
"""
import io
import re
import sqlite3
import sys
import zlib
from datetime import date, timedelta

import _support  # noqa: E402  (first: builds the test database, points the app at it)

from app import create_app  # noqa: E402

sys.stdout.reconfigure(errors="replace")

app = create_app()
_support.attach(app)
TEST_DB = _support.DB_PATH
ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID = 1, 2, 3, 5
REPO = _support.REPO

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok)))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}")
    if detail and not ok:
        for line in str(detail).splitlines()[:12]:
            print(f"       {line}")
    return ok


def section(title):
    print(f"\n--- {title} ---")


def db():
    conn = sqlite3.connect(TEST_DB)
    conn.row_factory = sqlite3.Row
    return conn


def q1(sql, args=()):
    conn = db()
    try:
        return conn.execute(sql, args).fetchone()
    finally:
        conn.close()


def run_sql(sql, args=()):
    conn = db()
    try:
        cur = conn.execute(sql, args)
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def client_as(user_id):
    c = app.test_client()
    with c.session_transaction() as sess:
        sess["user_id"] = user_id
        _support.open_session(sess)
    return c


def body(response):
    return response.get_data(as_text=True)


def pdf_text(data):
    out = []
    for match in re.finditer(rb"stream\r?\n(.*?)endstream", data, re.S):
        try:
            out.append(zlib.decompress(match.group(1)).decode("latin-1"))
        except Exception:
            out.append(match.group(1).decode("latin-1", "replace"))
    return "\n".join(out)


admin, staff, official, doctor = (client_as(i) for i in (ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID))
TODAY = date.today()

# ===========================================================================
section("B12. A SECOND PARENT'S REPORT FOR ONE DOSE IS REFUSED IN WORDS")
# ===========================================================================
born15 = date.fromisoformat(q1("SELECT birth_date FROM patients WHERE id = 15")[0])
first = run_sql("INSERT INTO immunizations (patient_id, vaccine_code, dose_number, evidence, date_administered, "
                "administered_by) VALUES (15, 'PENTA', 1, 'card', ?, 2)", ((born15 + timedelta(days=20)).isoformat(),))
second = run_sql("INSERT INTO immunizations (patient_id, vaccine_code, dose_number, evidence, date_administered, "
                 "administered_by) VALUES (15, 'PENTA', 1, 'card', ?, 2)", ((born15 + timedelta(days=60)).isoformat(),))
CHANGE = "/staff/patient_profile/15/dose-elsewhere/change"
r = staff.post(CHANGE, data={"record_id": str(first), "evidence": "parent_report", "date_given": "",
                             "place_given": "QA place one"})
check("setup: the first card record becomes a parent's report",
      r.status_code == 302 and q1("SELECT evidence FROM immunizations WHERE id = ?", (first,))[0] == "parent_report")
r = staff.post(CHANGE, data={"record_id": str(second), "evidence": "parent_report", "date_given": "",
                             "place_given": "QA typed place"})
page = body(r)
check("the second one is refused without a server error (200, the profile with the dialog open)",
      r.status_code == 200 and "elsewhereModal" in page, r.status_code)
check("the refusal says what is wrong, in plain words",
      "already has a parent&#39;s report" in page or "already has a parent's report" in page)
check("the typed place is kept in the reopened dialog", 'value="QA typed place"' in page)
check("nothing changed in the database",
      q1("SELECT evidence FROM immunizations WHERE id = ?", (second,))[0] == "card")
for rid in (first, second):
    run_sql("DELETE FROM immunizations WHERE id = ?", (rid,))

# ===========================================================================
section("B13, B14. INVENTORY NUMBERS: PLAIN, FINITE AND CAPPED")
# ===========================================================================
ITEM = {"category": "Medicine", "generic_name": "Qanumberol", "dosage_strength": "10 mg", "dosage_form": "tablet",
        "drug_class": "Analgesic / Antipyretic", "unit": "Tablet", "low_stock_threshold": "10", "unit_cost": "",
        "manufacturer": "", "supplier": "", "storage_location": "", "notes": "", "status": "Active"}


def add_item(**change):
    data = dict(ITEM)
    data.update(change)
    before = q1("SELECT COUNT(*) FROM inventory_items")[0]
    r = admin.post("/admin/add_edit_medicine", data=data)
    return r, q1("SELECT COUNT(*) FROM inventory_items")[0] - before


for label, change, message in (
        ("reorder level 999999999999999999999", {"low_stock_threshold": "999999999999999999999"},
         "Reorder level must be a whole number of 0 or more."),
        ("reorder level 100001", {"low_stock_threshold": "100001"}, "Enter a reorder level of 100,000 or less."),
        ("reorder level +5", {"low_stock_threshold": "+5"}, "Reorder level must be a whole number"),
        ("reorder level 1_000", {"low_stock_threshold": "1_000"}, "Reorder level must be a whole number"),
        ("unit cost inf", {"unit_cost": "inf"}, "Unit cost must be a number of 0 or more"),
        ("unit cost nan", {"unit_cost": "nan"}, "Unit cost must be a number of 0 or more"),
        ("unit cost 1e9", {"unit_cost": "1e9"}, "Unit cost must be a number of 0 or more"),
        ("unit cost -3", {"unit_cost": "-3"}, "Unit cost must be a number of 0 or more"),
        ("unit cost 1000001", {"unit_cost": "1000001"}, "Enter a unit cost of 1,000,000 or less.")):
    r, added = add_item(**change)
    check(f"{label}: refused beside the field, no server error, nothing saved",
          r.status_code == 200 and added == 0 and message in body(r), (r.status_code, added))
    check(f"{label}: the typed generic name is kept", 'value="Qanumberol"' in body(r))
r, added = add_item(low_stock_threshold="100000", unit_cost="2.50")
check("reorder level 100,000 and cost 2.50 are accepted", r.status_code == 302 and added == 1)
for label, qty in (("inf", "inf"), ("1e3", "1e3"), ("+5", "+5"), ("0", "0"), ("100001", "100001")):
    item = q1("SELECT id FROM inventory_items WHERE generic_name = 'Qanumberol'")[0]
    n = q1("SELECT COUNT(*) FROM inventory_batches WHERE inventory_item_id = ?", (item,))[0]
    r = admin.post(f"/admin/medicine_inventory/{item}/receive", data={
        "lot_number": "QA-LOT", "expiry_date": (TODAY + timedelta(days=400)).isoformat(), "quantity": qty,
        "date_received": TODAY.isoformat(), "source": "DOH", "brand_name": ""})
    check(f"receive quantity {label}: refused, no server error",
          r.status_code == 200 and q1("SELECT COUNT(*) FROM inventory_batches WHERE inventory_item_id = ?",
                                      (item,))[0] == n)

# ===========================================================================
section("B24. A PDF EXPORT NEVER FAILS ON LONG TEXT")
# ===========================================================================
from utils.document import CELL_CAPS, render  # noqa: E402

LONG = "Purok " + ("Qalongaddress " * 300)
run_sql("UPDATE patients SET address = ? WHERE id = 3", (LONG.strip(),))
for who, client, url in (("Staff", staff, "/staff/patient_records/export.pdf"),
                         ("Administrator", admin, "/admin/patient_records/export.pdf")):
    r = client.get(url)
    check(f"{who}: the Patient Records PDF is produced with a 4,000-character address on file",
          r.status_code == 200 and r.data[:4] == b"%PDF", r.status_code)
r = staff.get("/staff/patient_records/export.xlsx")
from openpyxl import load_workbook  # noqa: E402

cells = [str(c) for ws in load_workbook(io.BytesIO(r.data)).worksheets
         for row in ws.iter_rows(values_only=True) for c in row if c]
check("the Excel file keeps the long address in full", any(LONG.strip() in c for c in cells))
# Straight at the renderer: one cell far taller than a page, and an unbroken word.
sheets = [{"intro": "QA", "blocks": []}, {"heading": "QA", "intro": "", "blocks": [{
    "title": "QA", "columns": ["A", "B", "C", "D", "E", "F"],
    "rows": [["x" * 9000, "y " * 4000, "z", 1, 2.5, None], ["short", "short", "z", 1, 2.5, None]]}]}]
pdf = render(sheets, "QA Org", "QA", "QA", "QA", "QA")
text = pdf_text(pdf)
check("the renderer draws a 9,000-character cell and an 8,000-character cell (shortened with ...)",
      pdf[:4] == b"%PDF" and "short" in text)
normal = render([{"intro": "QA", "blocks": []}, {"heading": "QA", "intro": "", "blocks": [{
    "title": "QA", "columns": ["A", "B"], "rows": [["Ordinary text that fits", 3]]}]}], "QA", "QA", "QA", "QA", "QA")
check("ordinary text is never shortened", "Ordinary text that fits" in pdf_text(normal) and CELL_CAPS[0] is None)
run_sql("UPDATE patients SET address = 'Purok 1, Barangay Poblacion' WHERE id = 3")

# ===========================================================================
section("B25. PATIENT RECORDS SEARCH: NO SCRIPT ERROR")
# ===========================================================================
template = (REPO / "templates" / "patient_records.html").read_text(encoding="utf-8")
check("the search no longer looks for the old 'div[style]' wrapper",
      "closest('div[style]')" not in template and "closest('.table-responsive')" in template)
page = body(staff.get("/staff/patient_records"))
check("the table sits in a .table-responsive box and the empty message is outside it",
      re.search(r'<div class="table-responsive">\s*<table class="patient-table" id="patientTable">', page) is not None
      and page.find('id="emptyState"') > page.find("</table>"))

# ===========================================================================
section("B20. ONE COMPOSE FORM SENDS ONE BROADCAST")
# ===========================================================================
import threading  # noqa: E402

SEND = "/admin/sms_notifications/send"


def page_reference(client, url="/admin/sms_notifications"):
    m = re.search(r'name="send_reference" value="(SMS-\d{4}-\d{4,6})"', body(client.get(url)))
    return m.group(1) if m else None


def broadcasts():
    return q1("SELECT COUNT(*) FROM sms_broadcasts")[0]


ref = page_reference(admin)
check("the compose form carries the next reference as its one-time token", ref is not None, ref)
page = body(admin.get("/admin/sms_notifications"))
check("the Send button disables itself after the first click (script present)",
      'id="composeSend"' in page and "form.dataset.sending" in page and "send.disabled = true" in page)
recipients = None
form = {"message": "QA double press test", "audience": "all_patients", "category": "", "send_reference": ref}
n0 = broadcasts()
second_admin = client_as(ADMIN_ID)
codes = [None, None]


def press(i, client):
    codes[i] = client.post(SEND, data=form).status_code


threads = [threading.Thread(target=press, args=(0, admin)), threading.Thread(target=press, args=(1, second_admin))]
for t in threads:
    t.start()
for t in threads:
    t.join()
row = q1("SELECT id, reference, status, recipient_count, sent_count FROM sms_broadcasts ORDER BY id DESC LIMIT 1")
check("two requests at the same moment: no server error", codes == [302, 302], codes)
check("... exactly one broadcast is recorded, under the form's reference",
      broadcasts() == n0 + 1 and row["reference"] == ref, (broadcasts() - n0, row["reference"]))
check("... its recipients were texted once (one recipient row each) and the row is complete",
      q1("SELECT COUNT(*) FROM sms_broadcast_recipients WHERE broadcast_id = ?", (row["id"],))[0]
      == row["recipient_count"] == row["sent_count"] and row["status"] == "Sent")
r = admin.post(SEND, data=form)
shown = body(admin.get("/account"))
check("the same form sent again (back button, refresh): refused, nothing sent",
      r.status_code == 302 and broadcasts() == n0 + 1)
check("... with a plain message", f"This broadcast was already sent ({ref}). It was not sent again." in shown)
stale = dict(form, send_reference=ref, message="QA stale form")
admin.post(SEND, data=dict(form, send_reference=page_reference(admin), message="QA other sender"))
r = admin.post(SEND, data=dict(stale, send_reference=f"SMS-{TODAY.year}-9999"))
check("a form whose reference was taken by another broadcast: refused, dialog reopened with the typed message",
      r.status_code == 200 and "QA stale form" in body(r) and "was not sent" in body(r)
      and broadcasts() == n0 + 2)
r = admin.post(SEND, data=dict(form, send_reference="not-a-reference", message="QA forged"))
check("a malformed reference: refused, nothing sent", r.status_code == 200 and broadcasts() == n0 + 2)
r = admin.post(SEND, data={"message": "QA no form", "audience": "all_patients", "category": ""})
check("a post with no compose form (the demo seed's path) still sends once, under the next reference",
      r.status_code == 302 and broadcasts() == n0 + 3)

from html.parser import HTMLParser  # noqa: E402


class _FormFields(HTMLParser):
    """The fields the page's first POST form with no action would submit (no JavaScript)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.fields, self.in_form, self.done = [], False, False
        self.select = self.textarea = None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "form" and not self.done and (a.get("method") or "").lower() == "post" and not a.get("action"):
            self.in_form = True
        if not self.in_form:
            return
        name = a.get("name")
        if tag == "input" and name and "disabled" not in a:
            kind = (a.get("type") or "text").lower()
            if kind in ("checkbox", "radio") and "checked" not in a:
                return
            if kind not in ("submit", "button", "file"):
                self.fields.append((name, a.get("value", "")))
        elif tag == "select" and name and "disabled" not in a:
            self.select = [name, None, None]
        elif tag == "option" and self.select is not None:
            if self.select[2] is None:
                self.select[2] = a.get("value", "")
            if "selected" in a:
                self.select[1] = a.get("value", "")
        elif tag == "textarea" and name and "disabled" not in a:
            self.textarea = [name, ""]

    def handle_data(self, data):
        if self.textarea is not None:
            self.textarea[1] += data

    def handle_endtag(self, tag):
        if tag == "select" and self.select is not None:
            value = self.select[1] if self.select[1] is not None else self.select[2]
            if value is not None:
                self.fields.append((self.select[0], value))
            self.select = None
        elif tag == "textarea" and self.textarea is not None:
            self.fields.append((self.textarea[0], self.textarea[1].lstrip("\n")))
            self.textarea = None
        elif tag == "form" and self.in_form:
            self.in_form, self.done = False, True


def form_fields(page):
    parser = _FormFields()
    parser.feed(page)
    data = {}
    for name, value in parser.fields:
        data.setdefault(name, []).append(value)
    return {k: (v[0] if len(v) == 1 else v) for k, v in data.items()}


# ===========================================================================
section("B21, B14. MAXIMUM LENGTHS, ON THE SERVER AND IN THE FORMS")
# ===========================================================================
REG = {"first_name": "Qalength", "middle_name": "", "last_name": "Testperson", "birth_date": "1991-02-03",
       "sex": "Female", "civil_status": "Single", "address": "Purok 9", "municipality": "San Francisco",
       "barangay": "Poblacion", "province": "Surigao del Norte", "zip_code": "8405", "contact_number": "09000004001",
       "emergency_contact_name": "Qa Kin", "emergency_contact_relationship": "Sibling",
       "emergency_contact_number": "09000004002", "medical_conditions_note": ""}


def register(**change):
    data = dict(REG)
    data.update(change)
    before = q1("SELECT COUNT(*) FROM patients")[0]
    r = staff.post("/staff/register-patient", data=data)
    return r, q1("SELECT COUNT(*) FROM patients")[0] - before


serial = [0]
for field, limit in (("first_name", 100), ("middle_name", 100), ("last_name", 100), ("emergency_contact_name", 100),
                     ("address", 200), ("emergency_contact_address", 200), ("medical_conditions_note", 2000)):
    serial[0] += 1
    over = "Q" * (limit + 1)
    change = {field: over, "birth_date": f"1980-01-{serial[0]:02d}"}
    if field not in ("first_name", "last_name"):
        change["first_name"] = f"Qaover{serial[0]:02d}"
    r, added = register(**change)
    check(f"{field}: {limit + 1} characters refused beside the field, nothing saved",
          r.status_code == 200 and added == 0
          and f"Use {limit:,} characters or fewer. This has {limit + 1:,}." in body(r), (r.status_code, added))
    check(f"{field}: the typed text is kept", over in body(r))
    exact = ("Qa" + "q" * (limit - 2)) if field in ("first_name", "last_name") else ("Q" * limit)
    change = {field: exact, "birth_date": f"1981-01-{serial[0]:02d}"}
    if field not in ("first_name", "last_name"):
        change["first_name"] = f"Qaexact{serial[0]:02d}"
    r, added = register(**change)
    check(f"{field}: exactly {limit} characters is accepted", r.status_code == 302 and added == 1,
          (r.status_code, added))
page = body(staff.get("/staff/register-patient"))
check("the form sets maxlength from the same limits (100 / 200 / 2,000)",
      'name="first_name" maxlength="100"' in page and 'name="address" maxlength="200"' in page
      and 'name="medical_conditions_note" maxlength="2000"' in page)
old = TODAY.replace(year=TODAY.year - 121)
r, added = register(first_name="Qaancient", birth_date=old.isoformat())
check("a birth date over 120 years ago is refused", added == 0 and "over 120 years old" in body(r))
edge = TODAY.replace(year=TODAY.year - 120)
r, added = register(first_name="Qacentenarian", birth_date=edge.isoformat())
check("exactly 120 years old is accepted", added == 1, body(r)[:200])

# An old record longer than the limits: shown and exported in full, asked to shorten on edit.
LONGADDR = "Purok 1 " + "Q" * 400
run_sql("UPDATE patients SET address = ? WHERE id = 4", (LONGADDR,))
check("an old over-long address is shown in full on the profile",
      LONGADDR in body(staff.get("/staff/patient_profile/4")))
xl = staff.get("/staff/patient_records/export.xlsx").data
check("... and exported in full", any(LONGADDR in str(c) for ws in load_workbook(io.BytesIO(xl)).worksheets
                                      for row in ws.iter_rows(values_only=True) for c in row if c))
edit_page = body(staff.get("/staff/register-patient/4"))
check("... and opened for editing with its full text", LONGADDR in edit_page)
data = form_fields(edit_page)
r = staff.post("/staff/register-patient/4", data=data)
check("saving it unchanged asks for the address to be shortened (Use 200 characters or fewer)",
      r.status_code == 200 and "Use 200 characters or fewer. This has" in body(r)
      and q1("SELECT address FROM patients WHERE id = 4")[0] == LONGADDR, body(r)[:300])
run_sql("UPDATE patients SET address = 'Purok 1, Barangay Poblacion' WHERE id = 4")

from models.consultations import CLINICAL_TEXT_MAX_LENGTH, consultation_age_context  # noqa: E402
from models.consultations import validate_consultation_form  # noqa: E402

with app.app_context():
    from models.patients import get_patient_by_id
    age_adult = consultation_age_context(get_patient_by_id(1), on_date=TODAY)
    base = {"chief_complaint": "QA", "diagnosis": "QA", "treatment_plan": "QA"}
    e_over = validate_consultation_form(dict(base, remarks="R" * 2001), age_adult)
    e_ok = validate_consultation_form(dict(base, remarks="R" * 2000), age_adult)
    e_seal = validate_consultation_form(dict(base, diagnosis="enc:v1:" + "A" * 4000), age_adult)
check("consultation remarks: 2,001 refused, 2,000 accepted (limit 2,000)",
      CLINICAL_TEXT_MAX_LENGTH == 2000 and "remarks" in e_over and "remarks" not in e_ok, (e_over, e_ok))
check("a sealed field kept as stored (cipher text) is not measured", "diagnosis" not in e_seal, e_seal)
page = body(doctor.get("/staff/consultation/edit/1"))
check("the consultation form sets maxlength 2000 on the written fields",
      page.count('maxlength="2000"') >= 3, page.count('maxlength="2000"'))

for field, limit in (("generic_name", 200), ("dosage_strength", 200), ("manufacturer", 200), ("supplier", 200),
                     ("storage_location", 200), ("notes", 2000)):
    change = {field: "Q" * (limit + 1)}
    if field != "generic_name":
        change["generic_name"] = "Qalong" + field[:4]
    r, added = add_item(**change)
    check(f"item {field}: {limit + 1} characters refused beside the field", r.status_code == 200 and added == 0
          and f"Use {limit:,} characters or fewer." in body(r), (r.status_code, added))
item = q1("SELECT id FROM inventory_items WHERE generic_name = 'Qanumberol'")[0]
for field in ("lot_number", "brand_name"):
    n0 = q1("SELECT COUNT(*) FROM inventory_batches WHERE inventory_item_id = ?", (item,))[0]
    r = admin.post(f"/admin/medicine_inventory/{item}/receive", data={
        "lot_number": "QA-LOT", "expiry_date": (TODAY + timedelta(days=400)).isoformat(), "quantity": "50",
        "date_received": TODAY.isoformat(), "source": "DOH", "brand_name": "", field: "L" * 201})
    check(f"receive {field}: 201 characters refused beside the field, typed value kept",
          r.status_code == 200 and "Use 200 characters or fewer." in body(r) and "L" * 201 in body(r)
          and q1("SELECT COUNT(*) FROM inventory_batches WHERE inventory_item_id = ?", (item,))[0] == n0)
admin.post(f"/admin/medicine_inventory/{item}/receive", data={
    "lot_number": "QA-LOT", "expiry_date": (TODAY + timedelta(days=400)).isoformat(), "quantity": "50",
    "date_received": TODAY.isoformat(), "source": "DOH", "brand_name": ""})
batch = q1("SELECT id, quantity_remaining FROM inventory_batches WHERE inventory_item_id = ?", (item,))
admin.post(f"/admin/medicine_inventory/{item}/remove", data={"batch_id": str(batch["id"]), "quantity": "1",
                                                             "reason": "Damaged", "note": "N" * 201})
check("remove stock: a 201-character note is refused (the form's limit is 200)",
      q1("SELECT quantity_remaining FROM inventory_batches WHERE id = ?", (batch["id"],))[0]
      == batch["quantity_remaining"])
page = body(admin.get("/admin/add_edit_medicine"))
check("the item form sets maxlength 200 / 2000", 'name="generic_name" maxlength="200"' in page
      and 'name="notes" maxlength="2000"' in page)
run_sql("UPDATE appointments SET status = 'Completed' WHERE patient_id = 12 AND status IN "
        "('Waiting', 'In Consultation', 'Waiting for the Doctor')")
for length, expected in ((201, 0), (200, 1)):
    n0 = q1("SELECT COUNT(*) FROM appointments WHERE patient_id = 12")[0]
    staff.post("/staff/appointment_scheduling/check-in", data={"patient_id": "12", "service": "Consultation",
                                                               "reason": "R" * length})
    check(f"check-in reason of {length} characters: {'accepted' if expected else 'refused'} (limit 200 on both paths)",
          q1("SELECT COUNT(*) FROM appointments WHERE patient_id = 12")[0] - n0 == expected)
run_sql("UPDATE appointments SET status = 'Completed' WHERE patient_id = 12 AND status = 'Waiting'")


# ===========================================================================
def add_patient(birth_date, first="Zzqabatch"):
    pid = q1("SELECT MAX(id) FROM patients")[0] + 1
    run_sql("INSERT INTO patients (id, patient_id, first_name, last_name, birth_date, sex, civil_status, "
            "address, municipality, province, contact_number, status) VALUES (?, ?, ?, 'Qafour', ?, "
            "'Female', 'Single', 'Purok 1', 'San Francisco', 'Surigao del Norte', ?, 'Active')",
            (pid, f"PAT-{pid:03d}", first, birth_date, f"0900000{pid:04d}"))
    return pid


def flat(text):
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", text))


section("B19. ONE NAME FOR THE IMMUNIZATION FIGURE EVERYWHERE")
from openpyxl import load_workbook  # noqa: E402

screens = {
    "Staff dashboard": flat(body(staff.get("/staff/dashboard"))),
    "Doctor dashboard": flat(body(doctor.get("/staff/dashboard"))),
    "Administrator dashboard": flat(body(admin.get("/admin/dashboard"))),
    "Vaccination Campaigns": flat(body(official.get("/official/vaccination-campaigns"))),
}
xl = official.get("/official/community-reports/export.xlsx?type=all").data
report_cells = " ".join(str(c) for ws in load_workbook(io.BytesIO(xl)).worksheets
                        for row in ws.iter_rows(values_only=True) for c in row if c is not None)
screens["Community Report (Excel)"] = report_cells
screens["Community Report (PDF)"] = re.sub(r"\s+", " ", pdf_text(
    official.get("/official/community-reports/export.pdf?type=all").data))
for where, text in screens.items():
    check(f"{where}: the figure is called 'due for immunization'",
          "due for immunization" in text.lower(), text[:300])
    check(f"{where}: the old names 'Behind Schedule' and 'needing attention' are gone",
          "behind schedule" not in text.lower() and "needing attention" not in text.lower())
for where in ("Staff dashboard", "Vaccination Campaigns"):
    check(f"{where}: the no-record sub-line is the shared one: '<N> with a dose with no record (28 days past due)'",
          re.search(r"\d+ with a dose with no record \(28 days past due\)", screens[where]) is not None,
          screens[where][:300])

section("B18. THE QUEUE EXPORT HOLDS THE OPEN VISITS FROM EARLIER DAYS")
earlier_pid = add_patient("1986-02-02")
earlier_day = (TODAY - timedelta(days=3)).isoformat()
earlier_visit = run_sql(
    "INSERT INTO appointments (patient_id, scheduled_by, appointment_type, appointment_date, appointment_time, "
    "status, is_walkin, checked_in_at) VALUES (?, ?, 'Consultation', ?, '09:10', 'Waiting', 1, ?)",
    (earlier_pid, STAFF_ID, earlier_day, f"{earlier_day} 09:10:00"))
earlier_ref = f"APT-{earlier_visit:05d}"
check("set-up: the screen lists the earlier-day visit", earlier_ref in body(staff.get("/staff/appointment_scheduling")))
xl = staff.get("/staff/appointment_scheduling/export.xlsx").data
book = load_workbook(io.BytesIO(xl))
check("the Excel file has an 'Earlier Days' sheet", "Earlier Days" in book.sheetnames, book.sheetnames)
earlier_cells = [str(c) for row in book["Earlier Days"].iter_rows(values_only=True) for c in row
                 if c is not None] if "Earlier Days" in book.sheetnames else []
check("... listing the visit still open from an earlier day", earlier_ref in earlier_cells, earlier_cells[:20])
cover_cells = " ".join(str(c) for row in book.worksheets[0].iter_rows(values_only=True) for c in row if c)
check("the cover counts it ('Still open from earlier days')", "Still open from earlier days" in cover_cells)
check("the PDF lists it too", earlier_ref in pdf_text(staff.get("/staff/appointment_scheduling/export.pdf").data))
run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (earlier_visit,))
book = load_workbook(io.BytesIO(staff.get("/staff/appointment_scheduling/export.xlsx").data))
check("with no earlier open visit the sheet says so in words",
      any("No visit from an earlier day is still open." == str(c)
          for row in book["Earlier Days"].iter_rows(values_only=True) for c in row))

section("B26. SORT BY AGE ORDERS BABIES BY BIRTH DATE")
babies = {label: add_patient((TODAY - timedelta(days=days)).isoformat(), first=f"Zzqababy{label}")
          for label, days in (("w5", 38), ("m5", 160), ("m7", 220), ("y30", 30 * 365 + 10))}
page = body(staff.get("/staff/patient_records"))
order = {int(pid): (int(years) if years else None, int(key) if key else None)  # (whole years, sort key)
         for years, key, pid in re.findall(r'data-age-years="([^"]*)"\s+data-age-order="([^"]*)"'
                                           r'[\s\S]*?name="patient_ids" value="(\d+)"', page)}
keys = [order.get(babies[label], (None, None)) for label in ("w5", "m5", "m7", "y30")]
check("set-up: the three babies are all 0 whole years (the old key tied them)",
      [k[0] for k in keys[:3]] == [0, 0, 0], keys)
check("the sort key grows with age: 5 weeks < 5 months < 7 months < 30 years",
      all(a[1] is not None for a in keys) and keys[0][1] < keys[1][1] < keys[2][1] < keys[3][1], keys)
check("the Age column sorts on that key", "attr: 'data-age-order'" in page)
check("a patient with no birth date gets an empty key (sorted last, as before)",
      'data-age-order=""' in page or q1("SELECT COUNT(*) FROM patients WHERE birth_date IS NULL OR "
                                        "birth_date = ''")[0] == 0)

section("B17. 'NO STOCK' INSTEAD OF 'NO EXPIRY' WHEN EVERY BATCH IS EMPTY")
item_id = q1("SELECT i.id FROM inventory_items i WHERE i.status = 'Active' AND EXISTS "
             "(SELECT 1 FROM inventory_batches b WHERE b.inventory_item_id = i.id) ORDER BY i.id LIMIT 1")[0]
run_sql("UPDATE inventory_batches SET quantity_remaining = 0 WHERE inventory_item_id = ?", (item_id,))
item_page = flat(body(staff.get(f"/staff/medicine_inventory/{item_id}")))
check("item page: the Nearest expiry card reads 'No stock'", "No stock Nearest expiry" in item_page,
      item_page[item_page.find("Nearest expiry") - 80:][:120])
check("... and not 'No expiry'", "No expiry Nearest expiry" not in item_page)
book = load_workbook(io.BytesIO(staff.get("/staff/medicine_inventory/export.xlsx").data))
name = q1("SELECT name FROM inventory_items WHERE id = ?", (item_id,))[0]
rows = [row for ws in book.worksheets for row in ws.iter_rows(values_only=True) if name in row]
check("the inventory export says 'No stock' for that item", any("No stock" in row for row in rows), rows[:2])
check("the medicine list says 'No stock' on that item's row",
      "No stock" in body(staff.get("/staff/medicine_inventory")))
run_sql("UPDATE inventory_batches SET quantity_remaining = 5, expiry_date = NULL WHERE inventory_item_id = ?",
        (item_id,))
item_page = flat(body(staff.get(f"/staff/medicine_inventory/{item_id}")))
check("stock on hand with no expiry date still reads 'No expiry'", "No expiry Nearest expiry" in item_page,
      item_page[item_page.find("Nearest expiry") - 80:][:120])

section("B27. A LONG UNBROKEN BROADCAST PREVIEW WRAPS ON THE OFFICIAL DASHBOARD")
dash = body(official.get("/official/dashboard"))
check("the preview line breaks long words (overflow-wrap:anywhere)",
      re.search(r"\.list-sub\{[^}]*overflow-wrap:anywhere", dash) is not None)
# The browser half (no sideways scroll at 390 px) is in verify_audit_fixes_2_browser.py.

section("B28. THE README GIVES THE MEASURED TEST TIMES")
readme = (REPO / "README.md").read_text(encoding="utf-8")
check("full run 'about 30 minutes', quick run 'about 9 minutes'",
      "every suite (about 30 minutes)" in readme and "(about 9 minutes)" in readme)


# ===========================================================================
section("B22. A STORED UTC TIMESTAMP IS SHOWN ON ITS MANILA DATE")
from utils.dates import clinic_now_iso, now as clinic_now, stamp_date, stamp_long_date  # noqa: E402

check("16:30 UTC on Oct 31 is 00:30 on Nov 1 in Manila: 'Nov 01, 2026' / 'November 1, 2026'",
      stamp_date("2026-10-31 16:30:00") == "Nov 01, 2026" and stamp_long_date("2026-10-31 16:30:00")
      == "November 1, 2026")
check("the boundary: 15:59:59 UTC is still Oct 31; 16:00:00 UTC is Nov 1",
      stamp_date("2026-10-31 15:59:59") == "Oct 31, 2026" and stamp_date("2026-10-31 16:00:00") == "Nov 01, 2026")
check("SQLite's fractional seconds are read too", stamp_date("2026-10-31 16:30:00.123") == "Nov 01, 2026")
check("a plain date is shown as it is (never moved across a timezone)", stamp_date("2026-10-31") == "Oct 31, 2026")
check("nothing -> the fallback", stamp_date(None, "—") == "—" and stamp_long_date("", "—") == "—")

late_pid = add_patient("1990-04-04", first="Zzqalate")
run_sql("UPDATE patients SET created_at = '2026-10-31 16:30:00', updated_at = '2026-10-31 16:30:00' WHERE id = ?",
        (late_pid,))
late_ref = f"PAT-{late_pid:03d}"
profile = flat(body(staff.get(f"/staff/patient_profile/{late_pid}")))
check("Staff profile: 'Registered November 1, 2026' and 'Date Registered November 1, 2026'",
      "Registered November 1, 2026" in profile and "Date Registered November 1, 2026" in profile
      and "October 31, 2026" not in profile, profile[profile.find("Registered"):][:120])
admin_profile = flat(body(admin.get(f"/admin/patient_profile/{late_pid}")))
check("Administrator's profile: 'Registered November 1, 2026' and 'Nov 01, 2026'",
      "Registered November 1, 2026" in admin_profile and "Nov 01, 2026" in admin_profile
      and "Oct 31, 2026" not in admin_profile, admin_profile[admin_profile.find("Registered"):][:120])
records = body(admin.get("/admin/patient_records"))
row = records[records.find(late_ref):][:1500]
check("Administrator's Patient Records: the Registered column says Nov 01, 2026",
      "Nov 01, 2026" in row and "Oct 31, 2026" not in row, flat(row)[:200])
book = load_workbook(io.BytesIO(admin.get("/admin/patient_records/export.xlsx").data))
late_rows = [r for ws in book.worksheets for r in ws.iter_rows(values_only=True) if late_ref in r]
check("... and its Excel file", late_rows and "Nov 01, 2026" in late_rows[0] and "Oct 31, 2026" not in late_rows[0],
      late_rows[:1])

cons = q1("SELECT c.id, c.patient_id FROM consultations c WHERE c.restricted_at IS NULL ORDER BY c.id LIMIT 1")
run_sql("UPDATE consultations SET created_at = '2026-10-31 16:10:00', updated_at = '2026-10-31 16:45:00.250' "
        "WHERE id = ?", (cons["id"],))
profile = flat(body(doctor.get(f"/staff/patient_profile/{cons['patient_id']}?tab=history")))
check("profile, consultation: 'Recorded by … on Nov 01, 2026'",
      re.search(r"Recorded by [^.]+ on Nov 01, 2026\.", profile) is not None,
      re.findall(r"Recorded by [^.]+ on [^.]+\.", profile)[:3])
check("... 'Last edited Nov 01, 2026' (stored with fractional seconds, as an edit saves it)",
      "Last edited Nov 01, 2026" in profile)
edit = flat(body(doctor.get(f"/staff/consultation/edit/{cons['id']}")))
check("Edit Consultation: 'Last changed Nov 01, 2026'", "Last changed Nov 01, 2026" in edit,
      edit[edit.find("Last changed"):][:80] if "Last changed" in edit else edit[:200])

batch = q1("SELECT id, inventory_item_id FROM inventory_batches ORDER BY id LIMIT 1")
run_sql("INSERT INTO inventory_removals (batch_id, quantity, reason, note, removed_by, created_at) "
        "VALUES (?, 1, 'Damaged', 'QA', ?, '2026-10-31 16:30:00')", (batch["id"], STAFF_ID))
item_page = flat(body(staff.get(f"/staff/medicine_inventory/{batch['inventory_item_id']}")))
check("item page, Recent removals: Nov 01, 2026", "Nov 01, 2026" in item_page and "Oct 31, 2026" not in item_page)

unconverted = []
for path in sorted((REPO / "templates").glob("*.html")):
    for number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        if re.search(r"\w+_at\s*\|\s*(display_date|long_date|full_date)\b", line):
            unconverted.append(f"{path.name}:{number}")
check("no template prints a *_at column through display_date / long_date / full_date",
      not unconverted, "\n".join(unconverted))

section("B23. THE PAGE CARRIES THE SERVER'S PHILIPPINE DATE AND TIME")
page = body(staff.get("/staff/dashboard"))
stamp = re.search(r'data-clinic-now="(\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d)"', page)
check("<body data-clinic-now> holds the server's Manila time (within 2 minutes)",
      stamp is not None and abs((clinic_now().replace(tzinfo=None)
                                 - __import__("datetime").datetime.fromisoformat(stamp.group(1)))
                                .total_seconds()) < 120, stamp and stamp.group(1))
check("clinic_now_iso() is the clinic's now(), to the second, with no offset",
      re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d", clinic_now_iso()) is not None
      and clinic_now_iso()[:10] == clinic_now().date().isoformat())
js = (REPO / "static" / "js" / "time-format.js").read_text(encoding="utf-8")
check("the print stamps and the age preview no longer read the device clock",
      not [p.name for p in (REPO / "templates").glob("*.html")
           if "formatDateTime(new Date())" in p.read_text(encoding="utf-8")]
      and "new Date();" not in (REPO / "templates" / "register_patient.html").read_text(encoding="utf-8")
      and "clinicStamp" in js and "clinicToday" in js)
# The browser half (a device in Los Angeles with its clock two days behind) is in
# verify_audit_fixes_2_browser.py.

# ---------------------------------------------------------------------------
failed = [name for name, ok in results if not ok]
print("\n" + "=" * 78)
print(f"RESULT: {len(results) - len(failed)}/{len(results)} checks passed")
for name in failed:
    print(f"  FAILED: {name}")
print("=" * 78)
sys.exit(1 if failed else 0)
