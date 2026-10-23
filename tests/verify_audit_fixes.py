"""The fixes for the 2026-10-23 audit (Task 23: findings B01-B06, B08-B11).

B08  Community Reports: every type, card and Excel/PDF export counts over the chosen
     range, checked against counts taken straight from the database.
B09  Coverage is never "as of" a future day; a later range end reads "(today)".
B10  The "Doses Given Here" sub-line names only children dosed here.
B11  The attendance rate says which visits count as attended.
B01  The Doctor saves a visit that needs no diagnosis without one.
B03  A patient with an open visit today cannot be checked in again.
B04  No clinical action for a patient marked Deceased; no death over an open visit.
B05  SMS segments in UTF-16 units. B06  The schedule uses the queue's service list.
B02  schema.sql says who may change a patient's status.
Every name, date and number below is made up.
"""
import io
import re
import sqlite3
import sys
import zlib
from datetime import date, timedelta

import _support  # noqa: E402  (first: builds the test database, points the app at it)

from openpyxl import load_workbook  # noqa: E402

from app import create_app  # noqa: E402
from models.doctor_worklist import VISIT_SERVICES  # noqa: E402
from models.service_schedule import SCHEDULE_SERVICES  # noqa: E402
from utils.dates import display_date, long_date, today  # noqa: E402
from utils.sms import message_metrics  # noqa: E402

sys.stdout.reconfigure(errors="replace")

app = create_app()
_support.attach(app)
TEST_DB = _support.DB_PATH
ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID = 1, 2, 3, 5

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok)))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}")
    if detail and not ok:
        for line in str(detail).splitlines()[:12]:
            print(f"       {line}")
    return ok


def section(title):
    print(f"\n--- {title} ---")


def db():
    conn = sqlite3.connect(TEST_DB)
    conn.row_factory = sqlite3.Row
    return conn


def q1(sql, args=()):
    conn = db()
    try:
        return conn.execute(sql, args).fetchone()
    finally:
        conn.close()


def run_sql(sql, args=()):
    conn = db()
    try:
        cur = conn.execute(sql, args)
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def client_as(user_id):
    c = app.test_client()
    with c.session_transaction() as sess:
        sess["user_id"] = user_id
        _support.open_session(sess)
    return c


def body(response):
    return response.get_data(as_text=True)


def flashed(client):
    """The flash messages the next page shows, as text."""
    page = body(client.get("/account"))
    return " ".join(re.sub(r"<[^>]+>", " ", m)
                    for m in re.findall(r'(?s)<div class="flash-ids">.*?<div>(.*?)</div></div>', page))


def pdf_text(data):
    """The text of every page of a PDF made by utils/document.py (fpdf2), in order."""
    out = []
    for match in re.finditer(rb"stream\r?\n(.*?)endstream", data, re.S):
        chunk = match.group(1)
        try:
            out.append(zlib.decompress(chunk).decode("latin-1"))
        except Exception:
            out.append(chunk.decode("latin-1", "replace"))
    raw = "\n".join(out)
    # fpdf2 writes each run of text as (...) Tj; join the runs with spaces.
    runs = re.findall(r"\(((?:\\.|[^\\)])*)\)\s*Tj", raw)
    return " ".join(r.replace("\\(", "(").replace("\\)", ")").replace("\\\\", "\\") for r in runs)


def xlsx_rows(data):
    book = load_workbook(io.BytesIO(data), data_only=True)
    return {sheet.title: [tuple("" if v is None else str(v) for v in row)
                          for row in sheet.iter_rows(values_only=True)]
            for sheet in book.worksheets}


def figure(rows, label):
    """The value in the row whose first cell is `label` (the Figure | Value | Basis rows)."""
    for row in rows:
        if row and row[0] == label:
            return row[1]
    return None


def card(page, label):
    m = re.search(r'stat-value">\s*([^<]+?)\s*</div>\s*<div class="stat-label">' + re.escape(label), page)
    return m.group(1) if m else None


def fig_row(page, label):
    m = re.search(re.escape(label) + r'.*?<span class="fig-value[^"]*">\s*([^<]+?)\s*</span>', page, re.S)
    return m.group(1) if m else None


admin, staff, official, doctor = (client_as(i) for i in (ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID))
TODAY = today()

# ===========================================================================
section("B08. COMMUNITY REPORTS: EVERY TYPE COUNTS OVER THE CHOSEN RANGE")
# ===========================================================================
# Fictional rows dated inside January-March 2026, and others inside the default
# window (the last few weeks), so the two windows hold different figures.
START, END = date(2026, 1, 1), date(2026, 3, 31)
RECENT = TODAY - timedelta(days=10)
for n, created in enumerate(["2026-02-10 04:00:00", "2026-02-11 04:00:00", "2026-03-02 04:00:00",
                             f"{RECENT.isoformat()} 04:00:00"]):
    run_sql("INSERT INTO patients (id, patient_id, first_name, last_name, birth_date, sex, "
            "civil_status, address, municipality, province, barangay, status, created_at) "
            "VALUES (?, ?, ?, 'Rangecase', '1991-01-01', 'Female', 'Single', 'Purok 9', "
            "'San Francisco', 'Surigao del Norte', 'Poblacion', 'Active', ?)",
            (300 + n, f"PAT-{300 + n:03d}", f"Qa{n}", created))
# Two children, each with a valid BCG given here: one in February, one recently.
for pid, born, given in ((310, "2025-12-01", "2026-02-02"), (311, (RECENT - timedelta(days=40)).isoformat(),
                                                             RECENT.isoformat())):
    run_sql("INSERT INTO patients (id, patient_id, first_name, last_name, birth_date, sex, "
            "civil_status, address, municipality, province, barangay, status) VALUES "
            "(?, ?, 'QaChild', 'Rangecase', ?, 'Male', 'Single', 'Purok 9', 'San Francisco', "
            "'Surigao del Norte', 'Poblacion', 'Active')", (pid, f"PAT-{pid:03d}", born))
    run_sql("INSERT INTO immunizations (patient_id, vaccine_code, dose_number, evidence, "
            "date_administered, administered_by) VALUES (?, 'BCG', 1, 'given_here', ?, 2)", (pid, given))
# Consultations in the range (the fixture's own are all June to July 2026).
for n, day in enumerate(("2026-01-20", "2026-02-21", "2026-03-22")):
    run_sql("INSERT INTO appointments (id, patient_id, scheduled_by, appointment_type, appointment_date, "
            "status, is_walkin, checked_in_at) VALUES (?, 1, 2, 'Consultation', ?, 'Completed', 1, ?)",
            (700 + n, day, f"{day} 09:00:00"))
    run_sql("INSERT INTO consultations (patient_id, recorded_by, appointment_id, consultation_date, "
            "chief_complaint, diagnosis, treatment_plan) VALUES (1, 5, ?, ?, 'Range complaint', "
            "'Range diagnosis', 'Range plan')", (700 + n, f"{day} 09:30:00"))
# Broadcasts: two in February, one recently.
for n, (sent_at, sent) in enumerate((("2026-02-15 02:00:00", 5), ("2026-02-16 02:00:00", 7),
                                     (f"{RECENT.isoformat()} 02:00:00", 11))):
    run_sql("INSERT INTO sms_broadcasts (reference, message, audience, audience_label, recipient_count, "
            "sent_count, segments, delivery_mode, sent_by, sent_at) VALUES (?, 'Range test', "
            "'all_patients', 'All Registered Patients', ?, ?, 1, 'simulated', 1, ?)",
            (f"BC-RANGE-{n}", sent, sent, sent_at))

# Independent counts, straight from the tables. Manila is UTC+8, so the range's
# stored UTC bounds are 16:00 of the day before each end.
lo_utc, hi_utc = "2025-12-31 16:00:00", "2026-03-31 16:00:00"
EXPECT = {
    "registered": q1("SELECT COUNT(*) FROM patients WHERE created_at >= ? AND created_at < ?",
                     (lo_utc, hi_utc))[0],
    "consultations": q1("SELECT COUNT(*) FROM consultations WHERE substr(consultation_date, 1, 10) "
                        "BETWEEN ? AND ?", (START.isoformat(), END.isoformat()))[0],
    "doses_here": q1("SELECT COUNT(*) FROM immunizations WHERE evidence = 'given_here' AND "
                     "substr(date_administered, 1, 10) BETWEEN ? AND ?",
                     (START.isoformat(), END.isoformat()))[0],
    "broadcasts": q1("SELECT COUNT(*) FROM sms_broadcasts WHERE sent_at >= ? AND sent_at < ?",
                     (lo_utc, hi_utc))[0],
    "messages": q1("SELECT COALESCE(SUM(sent_count), 0) FROM sms_broadcasts WHERE sent_at >= ? "
                   "AND sent_at < ?", (lo_utc, hi_utc))[0],
}
check("setup: the range holds 3 registrations, 3 consultations, 1 dose given here, 2 broadcasts "
      "(12 messages)", EXPECT == {"registered": 3, "consultations": 3, "doses_here": 1,
                                  "broadcasts": 2, "messages": 12}, EXPECT)

RANGE = f"from={START.isoformat()}&to={END.isoformat()}"
DIAGNOSIS_NAMES = {d for d, _ in _support.DIAGNOSES}  # the fixture's own, all dated June to July
covers = f"{display_date(START)} &ndash; {display_date(END)}"
for rtype, wanted in (
        ("all", ("registered", "consultations", "doses_here", "broadcasts")),
        ("registrations", ("registered",)), ("consultations", ("consultations",)),
        ("vaccination", ("doses_here",)), ("notifications", ("broadcasts",)),
        ("concerns", ())):
    page = body(official.get(f"/official/community-reports?type={rtype}&{RANGE}"))
    xlsx = xlsx_rows(official.get(f"/official/community-reports/export.xlsx?type={rtype}&{RANGE}").data)
    pdf = pdf_text(official.get(f"/official/community-reports/export.pdf?type={rtype}&{RANGE}").data)
    summary = xlsx.get("Summary", [])
    cover = xlsx.get("Report", [])
    check(f"{rtype}: the workbook says it covers January 1 to March 31, 2026",
          figure(cover, "Covers") == f"{long_date(START)} to {long_date(END)}", cover[:16])
    if "registered" in wanted:
        check(f"{rtype}: registrations = {EXPECT['registered']} on the screen, in Excel and in the PDF",
              card(page, "New Patient Registrations") == str(EXPECT["registered"])
              and figure(summary, "Registered in this range") == str(EXPECT["registered"])
              and f"Registered in this range {EXPECT['registered']} " in pdf,
              (card(page, "New Patient Registrations"), figure(summary, "Registered in this range")))
    if "consultations" in wanted:
        check(f"{rtype}: consultations = {EXPECT['consultations']} on the screen, in Excel and in the PDF",
              card(page, "Consultations Recorded") == str(EXPECT["consultations"])
              and figure(summary, "Consultations recorded") == str(EXPECT["consultations"])
              and f"Consultations recorded {EXPECT['consultations']} " in pdf,
              (card(page, "Consultations Recorded"), figure(summary, "Consultations recorded")))
    if "doses_here" in wanted:
        check(f"{rtype}: doses given here = {EXPECT['doses_here']} on the screen, in Excel and in the PDF",
              card(page, "Immunization Doses Given Here") == str(EXPECT["doses_here"])
              and figure(summary, "Immunisation doses given here") == str(EXPECT["doses_here"])
              and f"Immunisation doses given here {EXPECT['doses_here']} " in pdf,
              (card(page, "Immunization Doses Given Here"), figure(summary, "Immunisation doses given here")))
        check(f"{rtype}: coverage is as of March 31, 2026 on the screen and in Excel",
              f"as of {display_date(END)}" in page
              and (figure(summary, "Children up to date on the DOH schedule (%)") is not None)
              and any(f"as of {display_date(END)}" in " ".join(r) for r in summary))
    if "broadcasts" in wanted:
        check(f"{rtype}: broadcasts = {EXPECT['broadcasts']} and messages = {EXPECT['messages']} "
              "on the screen, in Excel and in the PDF",
              fig_row(page, "Broadcasts published in this range") == str(EXPECT["broadcasts"])
              and fig_row(page, "Messages handed to the gateway") == str(EXPECT["messages"])
              and figure(summary, "SMS broadcasts published") == str(EXPECT["broadcasts"])
              and figure(summary, "Messages handed to the gateway") == str(EXPECT["messages"])
              and f"SMS broadcasts published {EXPECT['broadcasts']} " in pdf,
              (fig_row(page, "Broadcasts published in this range"), figure(summary, "SMS broadcasts published")))
    if rtype == "concerns":
        concerns = xlsx.get("Health Concerns", [])
        all_concerns = xlsx_rows(official.get(
            f"/official/community-reports/export.xlsx?type=all&{RANGE}").data).get("Health Concerns", [])
        # The diagnosis rows (the fixture's own diagnoses are all outside the range).
        # The Official's table folds rows of 1 to 4 into "Other", so the range's 3 show there.
        check("concerns: the diagnoses table counts only the range's 3 consultations, as type=all does",
              concerns == all_concerns and any(r[:2] == ("Other (each fewer than 5)", "3") for r in concerns)
              and not any(r and r[0] in DIAGNOSIS_NAMES for r in concerns), concerns[:30])
    check(f"{rtype}: the screen shows no range error", "not one this screen offers" not in page)

# ===========================================================================
section("B09. A RANGE ENDING AFTER TODAY: COVERAGE AS OF TODAY, AND SAID SO")
# ===========================================================================
future_end = TODAY + timedelta(days=45)
future = f"from={(TODAY - timedelta(days=60)).isoformat()}&to={future_end.isoformat()}"
page = body(official.get(f"/official/community-reports?type=all&{future}"))
summary = xlsx_rows(official.get(f"/official/community-reports/export.xlsx?type=all&{future}").data)["Summary"]
vacc = xlsx_rows(official.get(f"/official/community-reports/export.xlsx?type=vaccination&{future}").data)
pdf = pdf_text(official.get(f"/official/community-reports/export.pdf?type=vaccination&{future}").data)
today_label = f"as of {display_date(TODAY)} (today)"
check("screen: coverage card reads 'as of <today> (today)'", today_label in page, page[page.find("registered children"):][:300])
check("screen: no 'as of' date after today", f"as of {display_date(future_end)}" not in page)
check("Excel summary: coverage basis reads 'as of <today> (today)'",
      any(today_label in " ".join(r) for r in summary), summary)
check("Excel Immunisation sheet: 'As of <today> (today)', never the future end",
      any(f"As of {display_date(TODAY)} (today)" in " ".join(r) for r in vacc["Immunisation"])
      and not any(display_date(future_end) in " ".join(r[1:]) for r in vacc["Immunisation"]
                  if r and r[0].startswith("Children")), vacc["Immunisation"])
check("PDF: coverage as of today (today)", f"As of {display_date(TODAY)} (today)" in pdf)
default_page = body(official.get("/official/community-reports?type=all"))
check("the coverage figure for the future range equals today's figure",
      card(page, "Children Up to Date on the DOH Schedule") == card(default_page, "Children Up to Date on the DOH Schedule"))
check("the default window (ending today) shows no '(today)' note", "(today)" not in default_page)
past = body(official.get(f"/official/community-reports?type=vaccination&{RANGE}"))
check("a range ending in the past keeps its own end date, with no '(today)'",
      f"as of {display_date(END)}" in past and "(today)" not in past)

# ===========================================================================
section("B10. DOSES GIVEN HERE: THE SUB-LINE NAMES ONLY CHILDREN DOSED HERE")
# ===========================================================================
# April 2026: one card dose given elsewhere, no dose given here.
run_sql("INSERT INTO immunizations (patient_id, vaccine_code, dose_number, evidence, date_administered, "
        "administered_by) VALUES (310, 'HEPB', 1, 'card', '2026-04-10', 2)")
april = "from=2026-04-01&to=2026-04-30"
page = body(official.get(f"/official/community-reports?type=vaccination&{april}"))
sub = page[page.find("Immunization Doses Given Here"):][:900]
check("April: 0 doses given here", card(page, "Immunization Doses Given Here") == "0")
check("April: the sub-line says no doses were given here, not 'to 1 child'",
      "no doses given here in this range" in sub and "to 1 child" not in sub, sub)
check("April: the dose given elsewhere is its own sentence",
      "Also 1 dose given at another facility in this range (card seen), not counted above" in sub, sub)
page = body(official.get(f"/official/community-reports?type=vaccination&{RANGE}"))
sub = page[page.find("Immunization Doses Given Here"):][:900]
check("January-March: '1 dose ... given to 1 child'", "given to 1 child" in sub, sub)
imm = xlsx_rows(official.get(f"/official/community-reports/export.xlsx?type=vaccination&{april}").data)["Immunisation"]
check("Excel: children given a dose here (0) and children with any counted dose (1) are separate rows",
      figure(imm, "Children given a dose here in this range") == "0"
      and figure(imm, "Children with any counted dose in this range") == "1", imm)

# ===========================================================================
section("B11. THE ATTENDANCE RATE SAYS WHAT COUNTS AS ATTENDED")
# ===========================================================================
for who, c, url in (("Official", official, "/official/community-reports/export.xlsx?type=all"),
                    ("Administrator", admin, "/admin/reports_analytics/export.xlsx?type=all"),
                    ("Doctor", doctor, "/staff/reports_analytics/export.xlsx?type=all")):
    text = " ".join(" ".join(r) for rows in xlsx_rows(c.get(url).data).values() for r in rows)
    check(f"{who} workbook: attendance basis names attended, waiting for the Doctor and referred",
          "Patients who came (attended, waiting for the Doctor or referred)" in text
          and "Of the resolved appointments only" not in text and "that have resolved" not in text)
page = body(official.get("/official/community-reports?type=all"))
check("Official screen: the attendance line uses the same words",
      "Patients who came (attended, waiting for the Doctor or referred)" in page
      and "that have resolved" not in page)

# ===========================================================================
section("B01. THE DOCTOR SAVES A VISIT THAT NEEDS NO DIAGNOSIS WITHOUT ONE")
# ===========================================================================
CHECK_IN = "/staff/appointment_scheduling/check-in"
START_URL = "/staff/appointment_scheduling/start-consultation"
INTAKE = {"chief_complaint": "QA intake", "temperature": "36.9"}


def open_visit(client, pid, service):
    """Check a patient in and start the visit through the real routes; return its id."""
    run_sql("UPDATE appointments SET status = 'Completed' WHERE status IN ('Waiting', 'In Consultation', "
            "'Waiting for the Doctor')")
    client.post(CHECK_IN, data={"patient_id": str(pid), "reason": "QA visit", "service": service})
    visit = q1("SELECT id FROM appointments WHERE patient_id = ? ORDER BY id DESC LIMIT 1", (pid,))[0]
    client.post(START_URL, data={"appointment_id": str(visit)})
    return visit


def status_of(visit):
    return q1("SELECT status FROM appointments WHERE id = ?", (visit,))[0]


for service in ("Immunization", "Prenatal Check-up", "Family Planning"):
    v = open_visit(doctor, 9, service)
    page = body(doctor.get(f"/staff/consultation/record/{v}"))
    check(f"{service}: the Doctor's form marks diagnosis and plan optional, with no asterisk",
          page.count("(optional for this service)") == 2
          and not re.search(r'Diagnosis<span class="req">', page), page[page.find('name="diagnosis"') - 400:][:400])
    doctor.post(f"/staff/consultation/record/{v}", data=INTAKE)
    row = q1("SELECT id, diagnosis FROM consultations WHERE appointment_id = ?", (v,))
    check(f"{service}: the Doctor saves without a diagnosis and the visit is Completed",
          row is not None and not row["diagnosis"] and status_of(v) == "Completed", status_of(v))
    r = doctor.post(f"/staff/consultation/edit/{row['id']}", data=dict(INTAKE, temperature="37.2"))
    check(f"{service}: the Doctor's edit of only a vital sign saves (no diagnosis asked)",
          r.status_code == 302 and q1("SELECT temperature FROM consultations WHERE id = ?",
                                      (row["id"],))[0] == 37.2)
# An old Dental visit (no longer a service, still in the data) needs no diagnosis either.
run_sql("INSERT INTO appointments (id, patient_id, scheduled_by, appointment_type, appointment_date, status, "
        "is_walkin) VALUES (760, 10, 2, 'Dental', '2026-05-05', 'Completed', 1)")
dental = run_sql("INSERT INTO consultations (patient_id, recorded_by, appointment_id, consultation_date, "
                 "chief_complaint) VALUES (10, 2, 760, '2026-05-05 09:00:00', 'QA dental')")
r = doctor.post(f"/staff/consultation/edit/{dental}", data={"chief_complaint": "QA dental", "weight": "60"})
check("an old Dental visit: the Doctor's edit saves without a diagnosis",
      r.status_code == 302 and q1("SELECT weight FROM consultations WHERE id = ?", (dental,))[0] == 60)
# A Consultation still needs the Doctor's diagnosis.
for service in ("Consultation", "Follow-up Consultation", "Hypertension / Diabetes Monitoring"):
    v = open_visit(staff, 9, service)
    staff.post(f"/staff/consultation/record/{v}", data=INTAKE)
    cid = q1("SELECT id FROM consultations WHERE appointment_id = ?", (v,))[0]
    page = body(doctor.get(f"/staff/consultation/edit/{cid}"))
    check(f"{service}: the Doctor's form marks diagnosis and plan required (asterisk, no 'optional')",
          "(optional for this service)" not in page and page.count('<span class="req">*</span></label>') >= 3)
    r = doctor.post(f"/staff/consultation/edit/{cid}", data=INTAKE)
    check(f"{service}: saving without a diagnosis is refused and the visit stays Waiting for the Doctor",
          r.status_code == 200 and "Diagnosis is required." in body(r)
          and status_of(v) == "Waiting for the Doctor")
    doctor.post(f"/staff/consultation/edit/{cid}", data=dict(INTAKE, diagnosis="QA dx", treatment_plan="QA plan"))
    check(f"{service}: with a diagnosis it is Completed", status_of(v) == "Completed")
    run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (v,))
v = open_visit(staff, 9, "Consultation")
staff.post("/staff/appointment_scheduling/complete-consultation", data={"appointment_id": str(v)})
check("Mark Completed on a Consultation with no diagnosis is still refused", status_of(v) == "In Consultation")

# ===========================================================================
section("B03. ONE OPEN VISIT PER PATIENT PER DAY")
# ===========================================================================
run_sql("UPDATE appointments SET status = 'Completed' WHERE status IN ('Waiting', 'In Consultation', "
        "'Waiting for the Doctor')")


def open_count(pid):
    return q1("SELECT COUNT(*) FROM appointments WHERE patient_id = ? AND appointment_date = ? AND status IN "
              "('Waiting', 'In Consultation', 'Waiting for the Doctor')", (pid, TODAY.isoformat()))[0]


staff.post(CHECK_IN, data={"patient_id": "12", "reason": "QA", "service": "Consultation"})
first = q1("SELECT id FROM appointments WHERE patient_id = 12 ORDER BY id DESC LIMIT 1")[0]
for state, prepare in (("Waiting", None),
                       ("In Consultation", lambda: staff.post(START_URL, data={"appointment_id": str(first)})),
                       ("Waiting for the Doctor", lambda: staff.post(
                           "/staff/appointment_scheduling/send-to-doctor", data={"appointment_id": str(first)}))):
    if prepare:
        prepare()
    check(f"setup: the first visit is {state}", status_of(first) == state, status_of(first))
    r = staff.post(CHECK_IN, data={"patient_id": "12", "reason": "QA again", "service": "Immunization"})
    shown = body(staff.get("/account"))
    check(f"{state}: a second check-in is refused; still one open visit", open_count(12) == 1)
    check(f"{state}: the message says so and links to the open visit",
          "This patient already has an open visit today. Open it from the queue instead." in shown
          and f'href="/staff/appointment_scheduling#visit-{first}"' in shown,
          shown[shown.find("flash-ids"):][:400])
r = staff.post(CHECK_IN, data={"patient_id": "12", "reason": "QA again", "service": "Immunization"})
page = body(staff.get("/staff/appointment_scheduling"))
check("the refusal carries a working link to the open visit's row",
      f'href="/staff/appointment_scheduling#visit-{first}"' in page and f'id="visit-{first}"' in page)
check("the Check In list marks the patient 'Already Waiting for the Doctor'",
      "Already Waiting for the Doctor" in page)
for closed in ("Completed", "Left before being seen", "Referred"):
    run_sql("UPDATE appointments SET status = ? WHERE id = ?", (closed, first))
    staff.post(CHECK_IN, data={"patient_id": "12", "reason": "QA new", "service": "Consultation"})
    newest = q1("SELECT id, status FROM appointments WHERE patient_id = 12 ORDER BY id DESC LIMIT 1")
    check(f"after the visit is {closed}, a new check-in is allowed",
          newest["id"] != first and newest["status"] == "Waiting" and open_count(12) == 1)
    run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (newest["id"],))
with app.app_context():
    from models.appointments import check_in_patient
    run_sql("INSERT INTO appointments (id, patient_id, scheduled_by, appointment_type, appointment_date, "
            "status, is_walkin) VALUES (780, 13, 2, 'Consultation', ?, 'Waiting for the Doctor', 1)",
            (TODAY.isoformat(),))
    vid, outcome = check_in_patient(13, "model", STAFF_ID, service="Consultation")
check("the model refuses too: check_in_patient() answers 'already' with the open visit's id",
      outcome == "already" and vid == 780)

# ===========================================================================
section("B04. A PATIENT MARKED DECEASED")
# ===========================================================================
run_sql("UPDATE appointments SET status = 'Completed' WHERE status IN ('Waiting', 'In Consultation', "
        "'Waiting for the Doctor')")
# (2) Declaring death over an open visit is refused, with a link to the visit.
staff.post(CHECK_IN, data={"patient_id": "14", "reason": "QA", "service": "Consultation"})
visit14 = q1("SELECT id FROM appointments WHERE patient_id = 14 ORDER BY id DESC LIMIT 1")[0]
doctor.post("/staff/patient_records/archive", data={"status": "Deceased", "single_patient_id": "14"})
message = flashed(doctor)
check("Declare death with an open visit: refused, nothing changed",
      q1("SELECT status FROM patients WHERE id = 14")[0] == "Active")
doctor.post("/staff/patient_records/archive", data={"status": "Deceased", "single_patient_id": "14"})
page = body(doctor.get("/staff/patient_records"))
check("the message says to close the open visits first and links to the visit",
      "Close this patient&#39;s open visits first." in page or "Close this patient's open visits first." in page,
      page[page.find("flash-ids"):][:400])
check("... with a link to the visit's queue row",
      f'href="/staff/appointment_scheduling#visit-{visit14}"' in page)
doctor.post("/staff/patient_records/archive", data={"status": "Deceased", "patient_ids": ["14", "16"]})
check("a bulk declaration including that patient is refused for all",
      q1("SELECT COUNT(*) FROM patients WHERE id IN (14, 16) AND status = 'Deceased'")[0] == 0)
staff.post("/staff/appointment_scheduling/mark-left", data={"appointment_id": str(visit14)})
doctor.post("/staff/patient_records/archive", data={"status": "Deceased", "single_patient_id": "14"})
check("after the visit is closed (Left before being seen) the death is declared",
      q1("SELECT status FROM patients WHERE id = 14")[0] == "Deceased")

# (1) No clinical action on an open visit of a patient already marked Deceased (data from before
# this rule): put an open visit on patient 14 directly.
run_sql("UPDATE appointments SET status = 'Completed' WHERE status IN ('Waiting', 'In Consultation', "
        "'Waiting for the Doctor')")
run_sql("INSERT INTO appointments (id, patient_id, scheduled_by, appointment_type, appointment_date, status, "
        "is_walkin, checked_in_at) VALUES (790, 14, 2, 'Consultation', ?, 'Waiting', 1, ?)",
        (TODAY.isoformat(), f"{TODAY.isoformat()} 08:00:00"))
page = body(staff.get("/staff/appointment_scheduling"))
row = page[page.find('id="visit-790"'):][:6000]
row = row[:row.find("</tr>")]
check("queue: the row says Deceased", "Deceased</span>" in row)
check("queue: no Start and no Refer button on that row",
      "start-consultation" not in row and 'data-bs-target="#referModal"' not in row
      and "Start" not in re.sub(r"<[^>]+>", " ", row).split("Deceased")[-1], row[-1500:])
check("queue: Left before being seen is still offered", "went home before being seen" in row)
for name, path, data in (("Start", START_URL, {"appointment_id": "790"}),
                         ("Refer", "/staff/appointment_scheduling/refer",
                          {"appointment_id": "790", "destination": "Rural Health Unit (RHU)", "reason": ""})):
    staff.post(path, data=data)
    check(f"server: {name} refused for a Deceased patient", status_of(790) == "Waiting")
check("server: the refusal message is plain", "This patient is marked Deceased" in flashed(staff))
run_sql("UPDATE appointments SET status = 'In Consultation' WHERE id = 790")
r = staff.get("/staff/consultation/record/790")
check("server: Record Consultation (intake) refused", r.status_code == 302)
r = staff.post("/staff/consultation/record/790", data=INTAKE)
check("server: a posted intake is refused and nothing is written",
      q1("SELECT COUNT(*) FROM consultations WHERE appointment_id = 790")[0] == 0)
staff.post("/staff/appointment_scheduling/send-to-doctor", data={"appointment_id": "790"})
check("server: Send to the Doctor refused", status_of(790) == "In Consultation")
run_sql("UPDATE appointments SET status = 'Waiting for the Doctor' WHERE id = 790")
r = doctor.get("/staff/appointment_scheduling/see-patient/790")
check("server: the Doctor's See patient refused", status_of(790) == "Waiting for the Doctor"
      and "/consultation/record/" not in (r.headers.get("Location") or ""))
n14 = q1("SELECT COUNT(*) FROM appointments WHERE patient_id = 14")[0]
staff.post(CHECK_IN, data={"patient_id": "14", "reason": "QA", "service": "Consultation"})
check("server: check-in refused", q1("SELECT COUNT(*) FROM appointments WHERE patient_id = 14")[0] == n14)
staff.post("/staff/appointment_scheduling/mark-left", data={"appointment_id": "790"})
check("Left before being seen closes the visit", status_of(790) == "Left before being seen")
# Doses and dispensing.
born15 = date.fromisoformat(q1("SELECT birth_date FROM patients WHERE id = 15")[0])
n_doses = q1("SELECT COUNT(*) FROM immunizations WHERE patient_id = 15")[0]
staff.post("/staff/patient_profile/15/dose-elsewhere",
           data={"dose_key": "PENTA:1", "evidence": "card",
                 "date_given": (born15 + timedelta(days=60)).isoformat(), "place_given": "QA"})
check("control: while the child is Active, the same form records a dose given elsewhere",
      q1("SELECT COUNT(*) FROM immunizations WHERE patient_id = 15")[0] == n_doses + 1)
run_sql("UPDATE patients SET status = 'Deceased' WHERE id = 15")
staff.post("/staff/patient_profile/15/dose-elsewhere",
           data={"dose_key": "PENTA:2", "evidence": "card",
                 "date_given": (born15 + timedelta(days=100)).isoformat(), "place_given": "QA"})
check("server: once Deceased, a dose given elsewhere is refused, with a plain message",
      q1("SELECT COUNT(*) FROM immunizations WHERE patient_id = 15")[0] == n_doses + 1
      and "no vaccine dose can be recorded" in flashed(staff))
profile = body(staff.get("/staff/patient_profile/15?tab=vaccine"))
check("profile: no 'Record dose given elsewhere' button", "Record dose given elsewhere</button>" not in profile)
old = run_sql("INSERT INTO consultations (patient_id, recorded_by, consultation_date, chief_complaint) "
              "VALUES (14, 2, '2026-06-01 09:00:00', 'QA old')")
stock = q1("SELECT SUM(quantity_remaining) FROM inventory_batches WHERE inventory_item_id = 1")[0]
r = doctor.post(f"/staff/consultation/edit/{old}", data={"chief_complaint": "QA old", "diagnosis": "x",
                                                        "treatment_plan": "y", "dispense_item_0": "1",
                                                        "dispense_qty_0": "2"})
check("server: an edit that adds dispensing is refused; no stock moves",
      q1("SELECT COUNT(*) FROM medicine_dispensing WHERE consultation_id = ?", (old,))[0] == 0
      and q1("SELECT SUM(quantity_remaining) FROM inventory_batches WHERE inventory_item_id = 1")[0] == stock
      and "no dose, vaccination or medicine can be added" in body(r))
r = doctor.post(f"/staff/consultation/edit/{old}", data={"chief_complaint": "QA old", "diagnosis": "x",
                                                        "treatment_plan": "y", "weight": "55"})
check("an edit that only corrects the record still saves",
      r.status_code == 302 and q1("SELECT weight FROM consultations WHERE id = ?", (old,))[0] == 55)

# ===========================================================================
section("B05. SMS SEGMENTS IN UTF-16 UNITS, THE SAME ON SCREEN AND SERVER")
# ===========================================================================
EMOJI = "\U0001F600"
for label, text, length, segments, encoding in (
        ("160 GSM characters", "a" * 160, 160, 1, "GSM-7"),
        ("161 GSM characters", "a" * 161, 161, 2, "GSM-7"),
        ("70 UCS-2 characters", "á" * 69 + "’", 70, 1, "UCS-2"),
        ("71 UCS-2 characters", "’" * 71, 71, 2, "UCS-2"),
        ("35 emoji (70 units)", EMOJI * 35, 70, 1, "UCS-2"),
        ("36 emoji (72 units)", EMOJI * 36, 72, 2, "UCS-2"),
        ("67 emoji (134 units, 2 parts of 67)", EMOJI * 67, 134, 2, "UCS-2"),
        ("1 emoji + 133 letters (135 units)", EMOJI + "a" * 133, 135, 3, "UCS-2"),
        ("40 emoji (80 units)", EMOJI * 40, 80, 2, "UCS-2")):
    m = message_metrics(text)
    check(f"{label}: length {length}, {segments} segment(s), {encoding}",
          (m["length"], m["segments"], m["encoding"]) == (length, segments, encoding),
          (m["length"], m["segments"], m["encoding"]))
page = body(admin.get("/admin/sms_notifications"))
check("the screen's counter counts UCS-2 in UTF-16 units (text.length), like the server",
      "if(unicodeForced){ length = text.length; }" in page)
admin.post("/admin/sms_notifications/send", data={"message": "Line one\r\nLine two ’", "audience": "all_patients",
                                                  "category": ""})
sent = q1("SELECT message, segments FROM sms_broadcasts ORDER BY id DESC LIMIT 1")
check("a posted line break (CRLF) is stored as LF, so it counts once, as on the screen",
      sent["message"] == "Line one\nLine two ’" and "\r" not in sent["message"], repr(sent["message"]))

# ===========================================================================
section("B06. THE SERVICE SCHEDULE USES THE PATIENT QUEUE'S SERVICE LIST")
# ===========================================================================
check("one source: SCHEDULE_SERVICES is the Patient Queue's VISIT_SERVICES", SCHEDULE_SERVICES is VISIT_SERVICES)
for typed, stored in (("  family   planning ", "Family Planning"), ("IMMUNIZATION", "Immunization")):
    admin.post("/admin/service_schedule/add", data={"service": typed, "weekdays": ["2"]})
    row = q1("SELECT service FROM service_schedule ORDER BY id DESC LIMIT 1")
    check(f"'{typed.strip()}' is saved as the list's '{stored}'", row is not None and row[0] == stored)
n = q1("SELECT COUNT(*) FROM service_schedule")[0]
for typed in ("Laboratory", "Medical Certificate", "Prenatal"):
    r = admin.post("/admin/service_schedule/add", data={"service": typed, "weekdays": ["1"]})
    check(f"'{typed}' (not a queue service) is refused", q1("SELECT COUNT(*) FROM service_schedule")[0] == n
          and "The schedule uses the same services as the Patient Queue." in body(r))
legacy = run_sql("INSERT INTO service_schedule (service, weekdays) VALUES ('Laboratory', '1,2,3,4,5')")
page = body(admin.get("/admin/service_schedule"))
check("an entry saved before the rule is still shown, flagged 'Not a queue service'",
      "Laboratory" in page and "Not a queue service" in page)
r = admin.post(f"/admin/service_schedule/{legacy}/edit", data={"service": "Laboratory", "weekdays": ["1"]})
check("it cannot be saved again under its old name",
      q1("SELECT weekdays FROM service_schedule WHERE id = ?", (legacy,))[0] == "1,2,3,4,5")
admin.post(f"/admin/service_schedule/{legacy}/edit", data={"service": "Consultation", "weekdays": ["1"]})
check("edited to a list service, it saves",
      q1("SELECT service FROM service_schedule WHERE id = ?", (legacy,))[0] == "Consultation")
page = body(staff.get("/staff/appointment_scheduling"))
check("the queue's Check In list offers exactly the same services",
      all(f'value="{s}"' in page for s in VISIT_SERVICES))

# ===========================================================================
section("B02. schema.sql SAYS WHO MAY CHANGE A PATIENT'S STATUS")
# ===========================================================================
schema = (_support.REPO / "schema.sql").read_text(encoding="utf-8")
check("the comment names the Doctor for Deceased and excludes the Administrator",
      "only the\n-- Doctor sets or reverses Deceased" in schema
      and "Only\n-- Health Center Staff and Administrator may change it" not in schema)

# ---------------------------------------------------------------------------
failed = [name for name, ok in results if not ok]
print("\n" + "=" * 78)
print(f"RESULT: {len(results) - len(failed)}/{len(results)} checks passed")
for name in failed:
    print(f"  FAILED: {name}")
print("=" * 78)
sys.exit(1 if failed else 0)
