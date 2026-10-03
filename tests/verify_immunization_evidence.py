"""Honest immunization status (Task 9): a dose says what its record rests on.

The health center cannot know what happened at another hospital. So a dose that is not
given here counts only when a vaccination card was SEEN (written, dated evidence); a
parent's word alone is recorded but does not count; and a dose with nothing on file reads
"No record", never "Overdue". Staff and the Doctor record a dose given elsewhere from the
Immunization Record tab, outside a consultation. Evidence types map to HL7 FHIR:
given_here = primarySource true; card = false + reportOrigin "record"; parent_report =
false + reportOrigin "recall". Basis: CDC General Best Practices.
"""
import io
import re
import shutil
import sqlite3
import sys
import zlib
from datetime import timedelta
from pathlib import Path

import _support  # noqa: E402  (first: builds the test database, points the app at it)

from app import create_app  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

sys.stdout.reconfigure(errors="replace")

app = create_app()
_support.attach(app)
TEST_DB = _support.DB_PATH
ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID = 1, 2, 3, 5

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok)))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}")
    if detail and not ok:
        for line in str(detail).splitlines()[:12]:
            print(f"       {line}")
    return ok


def section(title):
    print(f"\n--- {title} ---")


def db():
    conn = sqlite3.connect(TEST_DB)
    conn.row_factory = sqlite3.Row
    return conn


def q1(sql, args=()):
    conn = db()
    try:
        return conn.execute(sql, args).fetchone()
    finally:
        conn.close()


def qall(sql, args=()):
    conn = db()
    try:
        return conn.execute(sql, args).fetchall()
    finally:
        conn.close()


def run_sql(sql, args=()):
    conn = db()
    try:
        cur = conn.execute(sql, args)
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def client_as(user_id):
    c = app.test_client()
    with c.session_transaction() as sess:
        sess["user_id"] = user_id
        _support.open_session(sess)
    return c


def body(response):
    return response.get_data(as_text=True)


def flat(text):
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", text))


def pdf_text(data):
    out = []
    for match in re.finditer(rb"stream\r?\n(.*?)endstream", data, re.S):
        try:
            out.append(zlib.decompress(match.group(1)).decode("latin-1"))
        except Exception:
            out.append(match.group(1).decode("latin-1", "replace"))
    return "\n".join(out).replace("\\(", "(").replace("\\)", ")")


def xlsx_text(data):
    book = load_workbook(io.BytesIO(data))
    return "\n".join(str(cell.value) for sheet in book for row in sheet.iter_rows()
                     for cell in row if cell.value is not None)


from models.immunizations import (  # noqa: E402
    get_dose_coverage, get_immunization_alert, get_immunization_coverage,
    get_immunization_status, record_elsewhere_dose,
)
from models.patients import get_patient_by_id  # noqa: E402
from utils.dates import today  # noqa: E402

admin, staff, official, doctor = (client_as(i) for i in
                                  (ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID))
TODAY = today()

print("=" * 78)
print("IDS-BPHCSN — honest immunization status (Task 9)")
print("=" * 78)

# ---------------------------------------------------------------------------
section("1. SET UP THREE CHILDREN, EACH WITH NOTHING ON FILE")
CHILDREN = {}
base_id = q1("SELECT MAX(id) FROM patients")[0] + 1
for offset, tag in enumerate(("A", "B", "C")):
    pid = base_id + offset
    birth = (TODAY - timedelta(days=150)).isoformat()
    run_sql("INSERT INTO patients (id, patient_id, first_name, last_name, birth_date, sex, civil_status, "
            "address, municipality, province, contact_number, status) VALUES (?, ?, ?, 'Evidence', ?, "
            "'Female', 'Single', 'Purok 1', 'San Francisco', 'Surigao del Norte', ?, 'Active')",
            (pid, f"PAT-{pid:03d}", f"Zzev{tag}", birth, f"0900000{pid:04d}"))
    CHILDREN[tag] = pid
A, B, C = CHILDREN["A"], CHILDREN["B"], CHILDREN["C"]
BIRTH = TODAY - timedelta(days=150)


def status_of(pid):
    with app.app_context():
        return get_immunization_status(get_patient_by_id(pid))


def dose(pid, key):
    return next(item for item in status_of(pid)["schedule"] if item["key"] == key)


def coverage_numbers():
    with app.app_context():
        c, d = get_immunization_coverage(), get_dose_coverage()
        return {"given": d["total_given"], "elsewhere": d["total_elsewhere"],
                "parent": d["total_parent_report"], "no_record": d["total_no_record"],
                "doses_elsewhere": c["doses_elsewhere"], "doses_parent": c["doses_parent_report"],
                "doses_given": c["doses_given"], "enrolled": c["enrolled"]}


def alert_ids():
    with app.app_context():
        return get_immunization_alert()["patient_ids"]


start = status_of(A)
check("a child with nothing on file: every due dose reads 'No record' (not 'Overdue'), with its due date",
      start["no_record_count"] >= 8 and start["completed_count"] == 0
      and all(i["due_date"] for i in start["schedule"])
      and {i["status"] for i in start["schedule"]} <= {"No record", "Upcoming"}, start["summary"])
check("the child is on the follow-up list", A in alert_ids())

URL = "/staff/patient_profile/{}/dose-elsewhere"
CHANGE = "/staff/patient_profile/{}/dose-elsewhere/change"
REMOVE = "/staff/patient_profile/{}/dose-elsewhere/remove"
IMM_AUDIT = ("Recorded a dose given elsewhere", "Recorded a parent report of a dose",
             "Upgraded a parent report to a card record", "Changed a dose record given elsewhere",
             "Removed a dose record given elsewhere")


def audit_rows(*actions):
    marks = ", ".join("?" for _ in actions)
    return [dict(r) for r in qall(f"SELECT * FROM audit_logs WHERE action IN ({marks}) ORDER BY id", actions)]


def record(client, pid, key, evidence, day="", place=""):
    return client.post(URL.format(pid), data={"dose_key": key, "evidence": evidence,
                                              "date_given": day, "place_given": place})


def n_doses():
    return q1("SELECT COUNT(*) FROM immunizations")[0]


# ---------------------------------------------------------------------------
section("2. A CARD DOSE, RECORDED OUTSIDE A CONSULTATION, COUNTS AS GIVEN")
date_ok = (BIRTH + timedelta(days=1)).isoformat()
for who, client, pid, key in (("Staff", staff, A, "BCG:1"), ("the Doctor", doctor, A, "HEPB:1")):
    before, cov_before, n_before = status_of(pid)["completed_count"], coverage_numbers(), n_doses()
    appts = q1("SELECT COUNT(*) FROM appointments")[0]
    r = record(client, pid, key, "card", date_ok, "District hospital")
    row = q1("SELECT * FROM immunizations WHERE patient_id = ? AND vaccine_code = ?", (pid, key.split(":")[0]))
    item = dose(pid, key)
    cov_after = coverage_numbers()
    check(f"{who}: a card dose is recorded with no consultation and no visit (302 to the profile)",
          r.status_code == 302 and row and row["consultation_id"] is None and row["evidence"] == "card"
          and row["place_given"] == "District hospital" and n_doses() == n_before + 1
          and q1("SELECT COUNT(*) FROM appointments")[0] == appts, (r.status_code, dict(row) if row else None))
    check(f"{who}: it shows 'Given elsewhere' and counts as given for completion",
          item["status"] == "Given elsewhere" and item["given"] and item["badge_class"] == "badge-completed"
          and status_of(pid)["completed_count"] == before + 1)
    check(f"{who}: it counts for coverage (the dose, the programme total, and the patient count)",
          cov_after["given"] == cov_before["given"] + 1 and cov_after["elsewhere"] == cov_before["elsewhere"] + 1
          and cov_after["doses_given"] == cov_before["doses_given"] + 1
          and cov_after["doses_elsewhere"] == cov_before["doses_elsewhere"] + 1, (cov_before, cov_after))
page = body(staff.get(f"/staff/patient_profile/{A}"))
check("the profile shows 'Given elsewhere', the place, and a Change button; no 'Overdue'",
      "Given elsewhere" in page and "District hospital" in page and "overdue" not in page.lower()
      and 'data-mode="change"' in page)
entry = audit_rows("Recorded a dose given elsewhere")
label = q1("SELECT first_name || ' ' || last_name AS n FROM patients WHERE id = ?", (A,))["n"]
check("each card dose wrote one audit entry, worded as decided, with no vaccine",
      len(entry) == 2 and all(e["details"] == f"Recorded 1 dose given elsewhere for {label} (PAT-{A:03d})."
                              for e in entry), [e["details"] for e in entry])
check("the BCG dose is no longer on the child's no-record list",
      dose(A, "BCG:1")["status"] != "No record" and status_of(A)["no_record_count"] == start["no_record_count"] - 2)

# The button is on every scheduled dose that has no given_here or card record.
buttons = len(re.findall(r'data-mode="record"', page))
expected_buttons = sum(1 for i in status_of(A)["schedule"] if i["can_record_elsewhere"])
check("the 'Record dose given elsewhere' button is on every dose with no given-here or card record",
      buttons == expected_buttons and buttons == len(status_of(A)["schedule"]) - 2
      and "Record dose given elsewhere" in page, (buttons, expected_buttons))

# ---------------------------------------------------------------------------
section("3. A PARENT'S WORD IS RECORDED BUT DOES NOT COUNT")
before_status, cov_before = status_of(B), coverage_numbers()
r = record(staff, B, "PENTA:1", "parent_report", "", "")
item = dose(B, "PENTA:1")
after = status_of(B)
cov_after = coverage_numbers()
check("a parent report with no date is accepted (the date is optional)",
      r.status_code == 302 and item["evidence"] == "parent_report" and item["date_administered"] is None)
check("it shows 'Parent says given — no card', amber, and is not counted as given",
      item["status"] == "Parent says given — no card" and item["badge_class"] == "badge-waiting"
      and not item["given"] and after["completed_count"] == before_status["completed_count"]
      and after["parent_report_count"] == 1)
check("it does not count for coverage; it has its own line",
      cov_after["given"] == cov_before["given"] and cov_after["doses_given"] == cov_before["doses_given"]
      and cov_after["parent"] == cov_before["parent"] + 1 and cov_after["doses_parent"] == cov_before["doses_parent"] + 1
      and cov_after["enrolled"] == cov_before["enrolled"], (cov_before, cov_after))
check("the next doses are unaffected: PENTA 2 is still graded by its due date, and the child is still followed up",
      dose(B, "PENTA:2")["status"] == before_status["schedule"][5]["status"] != "Completed"
      and B in alert_ids() and after["no_record_count"] == before_status["no_record_count"] - 1)
mine = audit_rows("Recorded a parent report of a dose")
check("it wrote one entry, worded as decided",
      len(mine) == 1 and mine[0]["details"].startswith("Recorded a parent report of 1 dose for ")
      and mine[0]["details"].endswith(f"(PAT-{B:03d})."), [m["details"] for m in mine])
page = body(doctor.get(f"/staff/patient_profile/{B}"))
check("the profile says 'Parent says given — no card' and asks for the card",
      "Parent says given — no card" in page and "vaccination card" in page)
# Upgrade to card.
r = record(staff, B, "PENTA:1", "card", (BIRTH + timedelta(days=45)).isoformat(), "Rural health unit")
item = dose(B, "PENTA:1")
up = audit_rows("Upgraded a parent report to a card record")
check("when the card is shown, recording it upgrades the parent report to a card record that counts",
      item["status"] == "Given elsewhere" and item["given"] and item["place_given"] == "Rural health unit"
      and q1("SELECT COUNT(*) FROM immunizations WHERE patient_id = ? AND vaccine_code = 'PENTA'", (B,))[0] == 1
      and status_of(B)["completed_count"] == before_status["completed_count"] + 1)
check("the upgrade wrote 'Upgraded a parent report to a card record for <Name> (PAT-nnn).'",
      len(up) == 1 and up[0]["details"].startswith("Upgraded a parent report to a card record for ")
      and up[0]["details"].endswith(f"(PAT-{B:03d})."), [u["details"] for u in up])
# Upgrade through a change (edit) instead.
record(staff, B, "OPV:1", "parent_report")
opv = q1("SELECT id FROM immunizations WHERE patient_id = ? AND vaccine_code = 'OPV'", (B,))["id"]
r = staff.post(CHANGE.format(B), data={"record_id": str(opv), "evidence": "card",
                                       "date_given": (BIRTH + timedelta(days=50)).isoformat(), "place_given": ""})
check("a parent report is also upgraded from the Change dialog",
      dose(B, "OPV:1")["status"] == "Given elsewhere" and len(audit_rows("Upgraded a parent report to a card record")) == 2)
# Recording a card where a card already is: refused, nothing changes.
n = n_doses()
record(staff, B, "PENTA:1", "card", (BIRTH + timedelta(days=60)).isoformat())
check("a dose that already counts as given is not overwritten by a second record",
      n_doses() == n and dose(B, "PENTA:1")["date_administered"] == (BIRTH + timedelta(days=45)).isoformat())
# A dose given here stays given here.
here_before = q1("SELECT evidence, date_administered FROM immunizations WHERE patient_id = 17 AND vaccine_code = 'BCG'")
record(staff, 17, "BCG:1", "card", date_ok)
check("a dose given here is never replaced by a card record",
      tuple(q1("SELECT evidence, date_administered FROM immunizations WHERE patient_id = 17 "
               "AND vaccine_code = 'BCG'")) == tuple(here_before) and here_before["evidence"] == "given_here")

# ---------------------------------------------------------------------------
section("4. VALIDATION: EVERY RULE IS CHECKED ON THE SERVER")
n, audits_before = n_doses(), len(audit_rows(*IMM_AUDIT))
cases = [
    ("a card with no date", dict(dose_key="MR:1", evidence="card", date_given=""), "date"),
    ("a date in the future", dict(dose_key="MR:1", evidence="card",
                                  date_given=(TODAY + timedelta(days=1)).isoformat()), "future"),
    ("a date before the child was born", dict(dose_key="MR:1", evidence="card",
                                              date_given=(BIRTH - timedelta(days=1)).isoformat()), "born"),
    ("a parent report dated in the future", dict(dose_key="MR:1", evidence="parent_report",
                                                 date_given=(TODAY + timedelta(days=3)).isoformat()), "future"),
    ("a parent report dated before birth", dict(dose_key="MR:1", evidence="parent_report",
                                                date_given=(BIRTH - timedelta(days=9)).isoformat()), "born"),
    ("a missing proof", dict(dose_key="MR:1", evidence="", date_given=date_ok), "know"),
    ("a proof that is not one of the two", dict(dose_key="MR:1", evidence="given_here", date_given=date_ok), "know"),
    ("a made-up proof", dict(dose_key="MR:1", evidence="guess", date_given=date_ok), "know"),
    ("a date that is not a date", dict(dose_key="MR:1", evidence="card", date_given="31-31-2026"), "valid date"),
]
for label_text, data, word in cases:
    for who, client in (("Staff", staff), ("the Doctor", doctor)):
        r = client.post(URL.format(C), data=dict(data, place_given=""))
        check(f"{who}: {label_text} is refused, the dialog comes back with the reason, and nothing is saved",
              r.status_code == 200 and "This was not saved" in body(r) and word in flat(body(r))
              and n_doses() == n and len(audit_rows(*IMM_AUDIT)) == audits_before, (r.status_code, word))
r = staff.post(URL.format(C), data={"dose_key": "PENTA:9", "evidence": "card", "date_given": date_ok})
check("a dose the schedule does not list is refused", r.status_code == 302 and n_doses() == n)
r = staff.post(URL.format(C), data={"dose_key": "MR:1", "evidence": "card", "date_given": date_ok,
                                    "place_given": "x" * 400})
placed = q1("SELECT place_given FROM immunizations WHERE patient_id = ? AND vaccine_code = 'MR'", (C,))
check("an over-long place is cut, not stored whole", placed and len(placed["place_given"]) == 120)
run_sql("DELETE FROM immunizations WHERE patient_id = ? AND vaccine_code = 'MR'", (C,))
with app.app_context():
    check("the model refuses a card with no date and an unknown proof, called directly",
          record_elsewhere_dose(C, "MR", 1, "card", None, None, STAFF_ID) == "invalid"
          and record_elsewhere_dose(C, "MR", 1, "given_here", date_ok, None, STAFF_ID) == "invalid")

# ---------------------------------------------------------------------------
section("5. THE ADMINISTRATOR AND THE OFFICIAL ARE REFUSED ON THE SERVER")
n = n_doses()
opv_id = q1("SELECT id FROM immunizations WHERE patient_id = ? AND vaccine_code = 'OPV'", (B,))["id"]
refused = []
for who, client in (("Administrator", admin), ("Barangay Official", official)):
    for label_text, response in (
            ("record", client.post(URL.format(C), data={"dose_key": "MR:1", "evidence": "card", "date_given": date_ok})),
            ("change", client.post(CHANGE.format(B), data={"record_id": str(opv_id), "evidence": "parent_report"})),
            ("remove", client.post(REMOVE.format(B), data={"record_id": str(opv_id)}))):
        if not (response.status_code == 302 and response.headers.get("Location", "").endswith("dashboard")):
            refused.append((who, label_text, response.status_code))
check("POST record, change and remove (valid token) are redirected away for both roles", not refused, refused)
check("nothing was written or changed by them",
      n_doses() == n and dose(B, "OPV:1")["evidence"] == "card")
page = body(admin.get(f"/admin/patient_profile/{A}"))
check("the Administrator's own profile route is not the clinical profile (no dialog, no button)",
      "elsewhereModal" not in page and "Record dose given elsewhere" not in page)
check("there is no /admin/ or /official/ route for any of the three",
      all(client.post(path, data={"record_id": "1"}).status_code == 404 for client, path in (
          (admin, f"/admin/patient_profile/{C}/dose-elsewhere"),
          (admin, f"/admin/patient_profile/{C}/dose-elsewhere/change"),
          (official, f"/official/patient_profile/{C}/dose-elsewhere/remove"))))
staff_page = body(staff.get(f"/staff/patient_profile/{C}"))
check("Staff and the Doctor do get the dialog and the button",
      "elsewhereModal" in staff_page and "Record dose given elsewhere" in staff_page
      and "elsewhereModal" in body(doctor.get(f"/staff/patient_profile/{C}")))

# ---------------------------------------------------------------------------
section("6. CORRECTIONS: CHANGE AND REMOVE, AUDITED, NAMING NO VACCINE")
record(staff, C, "PENTA:1", "parent_report", "", "")
record(doctor, C, "OPV:1", "card", date_ok, "Hospital one")
penta = q1("SELECT id FROM immunizations WHERE patient_id = ? AND vaccine_code = 'PENTA'", (C,))["id"]
opv = q1("SELECT id FROM immunizations WHERE patient_id = ? AND vaccine_code = 'OPV'", (C,))["id"]
n_changed = len(audit_rows("Changed a dose record given elsewhere"))
r = doctor.post(CHANGE.format(C), data={"record_id": str(opv), "evidence": "card",
                                        "date_given": (BIRTH + timedelta(days=3)).isoformat(),
                                        "place_given": "Hospital two"})
row = q1("SELECT * FROM immunizations WHERE id = ?", (opv,))
changed = audit_rows("Changed a dose record given elsewhere")
check("the Doctor changes a card record's date and place; one audit entry, worded as decided",
      r.status_code == 302 and row["place_given"] == "Hospital two"
      and row["date_administered"] == (BIRTH + timedelta(days=3)).isoformat()
      and len(changed) == n_changed + 1
      and changed[-1]["details"] == f"Changed 1 dose record given elsewhere for Zzev{'C'} Evidence (PAT-{C:03d}).",
      [c["details"] for c in changed])
r = staff.post(CHANGE.format(C), data={"record_id": str(opv), "evidence": "card", "date_given": ""})
check("a change that would leave a card with no date is refused and shown again, nothing changed",
      r.status_code == 200 and "This was not saved" in body(r)
      and q1("SELECT date_administered FROM immunizations WHERE id = ?", (opv,))[0] == row["date_administered"])
r = staff.post(CHANGE.format(C), data={"record_id": str(opv), "evidence": "card",
                                       "date_given": row["date_administered"], "place_given": "Hospital two"})
check("a change that changes nothing writes nothing",
      len(audit_rows("Changed a dose record given elsewhere")) == n_changed + 1)
# change card back to a parent's word
staff.post(CHANGE.format(C), data={"record_id": str(opv), "evidence": "parent_report", "date_given": "", "place_given": ""})
check("a card record can be changed to a parent's word (it then stops counting)",
      dose(C, "OPV:1")["status"] == "Parent says given — no card" and not dose(C, "OPV:1")["given"])
# remove
n = n_doses()
r = staff.post(REMOVE.format(C), data={"record_id": str(penta)})
removed = audit_rows("Removed a dose record given elsewhere")
check("Staff remove a parent report; the dose reads 'No record' again; one audit entry",
      r.status_code == 302 and n_doses() == n - 1 and dose(C, "PENTA:1")["status"] == "No record"
      and len(removed) == 1 and removed[0]["details"] == f"Removed 1 dose record given elsewhere for Zzev{'C'} Evidence (PAT-{C:03d}).",
      [x["details"] for x in removed])
given_here_id = q1("SELECT id FROM immunizations WHERE patient_id = 17 AND vaccine_code = 'BCG'")["id"]
n = n_doses()
for label_text, response in (
        ("change", staff.post(CHANGE.format(17), data={"record_id": str(given_here_id), "evidence": "card",
                                                       "date_given": date_ok})),
        ("remove", doctor.post(REMOVE.format(17), data={"record_id": str(given_here_id)}))):
    check(f"a dose given here cannot be {label_text}d from the given-elsewhere route",
          response.status_code == 302 and n_doses() == n
          and q1("SELECT evidence FROM immunizations WHERE id = ?", (given_here_id,))[0] == "given_here")
from models.immunizations import remove_elsewhere_dose, update_elsewhere_dose  # noqa: E402

with app.app_context():
    check("the model refuses to remove or change a dose given here, called directly",
          remove_elsewhere_dose(given_here_id) == "not-allowed"
          and update_elsewhere_dose(given_here_id, "card", date_ok, None) == "not-allowed"
          and remove_elsewhere_dose(999999) == "missing" and q1("SELECT COUNT(*) FROM immunizations WHERE id = ?",
                                                                (given_here_id,))[0] == 1)
check("a record that does not exist is handled",
      staff.post(REMOVE.format(C), data={"record_id": "999999"}).status_code == 302
      and staff.post(CHANGE.format(C), data={"record_id": "abc"}).status_code == 302)
texts = " | ".join(r["details"] or "" for r in audit_rows(*IMM_AUDIT))
from models.immunizations import EPI_SCHEDULE, VACCINE_NAMES, dose_label  # noqa: E402

names = set(VACCINE_NAMES.values()) | {dose_label(r[0], r[3]) for r in EPI_SCHEDULE}
check("no audit entry about a dose given elsewhere names a vaccine",
      texts and not [n_ for n_ in names if n_ in texts], texts[:200])
all_entries = " | ".join(r["details"] or "" for r in qall("SELECT details FROM audit_logs WHERE module = 'Immunization Records'"))
check("nor does any other entry of the Immunization Records module",
      not [n_ for n_ in names if n_ in all_entries], all_entries[:200])
check("the Administrator's Audit Logs screen shows the entries",
      "Recorded 1 dose given elsewhere" in body(admin.get("/admin/audit_logs")))

# ---------------------------------------------------------------------------
section("7. THE CONSULTATION FORM ASKS FOR THE PROOF")
run_sql("INSERT INTO appointments (patient_id, scheduled_by, appointment_type, appointment_date, status, "
        "is_walkin, checked_in_at) VALUES (?, ?, 'Consultation', ?, 'In Consultation', 1, ?)",
        (B, STAFF_ID, TODAY.isoformat(), f"{TODAY.isoformat()} 08:00:00"))
visit = q1("SELECT MAX(id) FROM appointments")[0]
FORM = {"chief_complaint": "check", "temperature": "36.8", "blood_pressure": "100/60", "weight": "6"}
r = staff.post(f"/staff/consultation/record/{visit}", data=dict(
    FORM, immunizations_elsewhere=["IPV:1"], elsewhere_date_IPV_1=date_ok))
check("ticking Given elsewhere with no proof is refused and nothing is saved",
      r.status_code == 200 and "proof" in flat(body(r)).lower()
      and q1("SELECT COUNT(*) FROM consultations WHERE appointment_id = ?", (visit,))[0] == 0)
r = staff.post(f"/staff/consultation/record/{visit}", data=dict(
    FORM, immunizations_elsewhere=["IPV:1", "PCV:1"], elsewhere_proof_IPV_1="card",
    # IPV's minimum age is 14 weeks (Task 13): a card date the day after birth would count as too early.
    elsewhere_date_IPV_1=(BIRTH + timedelta(days=100)).isoformat(),
    elsewhere_proof_PCV_1="parent_report", elsewhere_date_PCV_1=""))
check("a card dose and a parent report are both saved at the visit",
      dose(B, "IPV:1")["status"] == "Given elsewhere" and dose(B, "PCV:1")["status"] == "Parent says given — no card"
      and dose(B, "IPV:1")["consultation_id"] is not None)
run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (visit,))
run_sql("INSERT INTO appointments (patient_id, scheduled_by, appointment_type, appointment_date, status, "
        "is_walkin, checked_in_at) VALUES (?, ?, 'Consultation', ?, 'In Consultation', 1, ?)",
        (B, STAFF_ID, TODAY.isoformat(), f"{TODAY.isoformat()} 09:00:00"))
visit2 = q1("SELECT MAX(id) FROM appointments")[0]
page = body(staff.get(f"/staff/consultation/record/{visit2}"))
check("the form lists the proof choice on each dose that can still be recorded, and no 'overdue'",
      'name="elsewhere_proof_PENTA_2"' in page and "overdue" not in page.lower())
r = staff.post(f"/staff/consultation/record/{visit2}", data=dict(FORM, immunizations=["PCV:1"]))
check("a dose given here replaces a parent's report of the same dose",
      q1("SELECT evidence FROM immunizations WHERE patient_id = ? AND vaccine_code = 'PCV'", (B,))[0] == "given_here"
      and dose(B, "PCV:1")["status"] == "Completed")

# ---------------------------------------------------------------------------
section("8. NO SCREEN, BADGE, REPORT, EXPORT OR PRINT SAYS 'OVERDUE' FOR IMMUNIZATIONS")
# A child is in today's queue so the queue badge is drawn.
run_sql("INSERT INTO appointments (patient_id, scheduled_by, appointment_type, appointment_date, status, "
        "is_walkin, checked_in_at) VALUES (?, ?, 'Immunization', ?, 'Waiting', 1, ?)",
        (A, STAFF_ID, TODAY.isoformat(), f"{TODAY.isoformat()} 10:00:00"))
PAGES = [
    (staff, "/staff/dashboard"), (doctor, "/staff/dashboard"), (admin, "/admin/dashboard"),
    (official, "/official/dashboard"), (staff, "/staff/appointment_scheduling"),
    (doctor, "/staff/appointment_scheduling"), (admin, "/admin/appointment_scheduling"),
    (staff, f"/staff/patient_profile/{A}"), (staff, f"/staff/patient_profile/{A}?print=1"),
    (doctor, f"/staff/patient_profile/{B}"), (doctor, f"/staff/patient_profile/{C}"),
    (doctor, "/staff/reports_analytics"), (admin, "/admin/reports_analytics"),
    (official, "/official/reports_analytics"), (official, "/official/community-reports"),
    (official, "/official/health-programs"), (official, "/official/vaccination-campaigns"),
]
found = []
for client, url in PAGES:
    r = client.get(url)
    text = body(r)
    if r.status_code != 200:
        found.append((url, r.status_code))
    elif "overdue" in text.lower():
        found.append((url, "says overdue"))
check(f"none of {len(PAGES)} screens says 'overdue' (dashboards, queue, profile and its print, reports, campaigns)",
      not found, found)
exports = [(doctor, "/staff/reports_analytics/export"), (admin, "/admin/reports_analytics/export"),
           (official, "/official/reports_analytics/export"), (official, "/official/community-reports/export")]
bad = []
for client, stem in exports:
    for ext, convert in ((".xlsx", xlsx_text), (".pdf", pdf_text)):
        r = client.get(stem + ext)
        if r.status_code != 200 or "overdue" in convert(r.get_data()).lower():
            bad.append((stem + ext, r.status_code))
check(f"none of the {len(exports) * 2} Excel and PDF exports says 'overdue'", not bad, bad)
queue = body(staff.get("/staff/appointment_scheduling"))
check("the queue badge for a child with doses with no record says '<n> vaccines: no record'",
      re.search(r"\d+ vaccines?: no record", queue) is not None and "immun-norecord" in queue)
no_record_page = body(staff.get(f"/staff/patient_profile/{A}"))
check("the profile alert reads as decided: '<N> doses have no record here. Ask for the child's vaccination card. "
      "If the card shows the dose, record it as given elsewhere.'",
      re.search(r"\d+ doses? ha(?:s|ve) no record here", flat(no_record_page))
      and "Ask for the child's vaccination card. If the card shows the dose, record it as given elsewhere."
      in flat(no_record_page))
check("the staff dashboard card and Administrator's say 'no record'",
      "with no record" in body(staff.get("/staff/dashboard"))
      and "with no record" in body(admin.get("/admin/dashboard")))
camp = body(official.get("/official/vaccination-campaigns"))
check("Vaccination Campaigns has No Record, Given Elsewhere (Card) and Parent Says Given columns",
      all(word in camp for word in ("No Record", "Given Elsewhere (Card)", "Parent Says Given")))
community = xlsx_text(official.get("/official/community-reports/export.xlsx").get_data())
check("the Community Reports file counts parent reports on their own line, and cards as given elsewhere",
      "Parent says given, no card" in community and "Doses given elsewhere (card seen)" in community
      and "Children with a dose with no record" in community, community[:200])

# ---------------------------------------------------------------------------
section("9. THE MIGRATION CLASSIFIES EXISTING DOSES, IS SAFE TO RUN TWICE, AND RECORDS ITSELF")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "migrations"))
import migrate_immunization_evidence as mig  # noqa: E402

work = Path(_support.work_path("evidence_migration"))
shutil.rmtree(work, ignore_errors=True)
work.mkdir(parents=True)
old = work / "old.db"
shutil.copy(TEST_DB, old)
con = sqlite3.connect(old)
con.executescript("""
CREATE TABLE immunizations_old (
    id INTEGER PRIMARY KEY AUTOINCREMENT, patient_id INTEGER NOT NULL, consultation_id INTEGER,
    vaccine_code TEXT NOT NULL, dose_number INTEGER NOT NULL,
    source TEXT NOT NULL DEFAULT 'administered' CHECK (source IN ('administered', 'reported')),
    date_administered TEXT, administered_by INTEGER NOT NULL, notes TEXT,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    CHECK (source = 'reported' OR date_administered IS NOT NULL),
    UNIQUE (patient_id, vaccine_code, dose_number));
DELETE FROM immunizations;
""")
rows_in = [(1, 17, "BCG", 1, "administered", "2026-01-02"), (2, 17, "HEPB", 1, "administered", "2026-01-02"),
           (3, 17, "PENTA", 1, "reported", "2026-03-01"), (4, 17, "OPV", 1, "reported", None),
           (5, 17, "PCV", 1, "administered", "2026-03-02")]
for rid, pid, code, dose_no, src, day in rows_in:
    con.execute("INSERT INTO immunizations_old (id, patient_id, vaccine_code, dose_number, source, "
                "date_administered, administered_by, notes) VALUES (?, ?, ?, ?, ?, ?, 2, ?)",
                (rid, pid, code, dose_no, src, day, f"n{rid}"))
con.executescript("DROP TABLE immunizations; ALTER TABLE immunizations_old RENAME TO immunizations;")
con.commit()
audit_before = con.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0]
con.close()
result = mig.run(old, work / "backups")
con = sqlite3.connect(old)
con.row_factory = sqlite3.Row
after = {r["id"]: dict(r) for r in con.execute("SELECT * FROM immunizations")}
check("the migration reports no problem, makes one backup, and carries every row across",
      not result["problems"] and result["backup"] and Path(result["backup"]).exists()
      and result["rows"] == 5 and len(after) == 5, result)
check("doses given here became 'given_here' and doses reported as given elsewhere became 'card'",
      [after[i]["evidence"] for i in (1, 2, 3, 4, 5)] == ["given_here", "given_here", "card", "card", "given_here"]
      and (result["given_here"], result["card"]) == (3, 2))
check("every other value is unchanged, an undated reported dose stays undated, and place_given is empty",
      after[4]["date_administered"] is None and after[3]["date_administered"] == "2026-03-01"
      and all(after[i]["notes"] == f"n{i}" and after[i]["place_given"] is None for i in after))
check("it wrote exactly one audit entry about itself, naming no patient and no vaccine",
      con.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0] == audit_before + 1
      and "migrate_immunization_evidence.py" in con.execute(
          "SELECT details FROM audit_logs ORDER BY id DESC LIMIT 1").fetchone()[0])
con.execute("PRAGMA foreign_keys = ON")


def refused_by_check(evidence, day):
    try:
        con.execute("INSERT INTO immunizations (patient_id, vaccine_code, dose_number, evidence, "
                    "date_administered, administered_by) VALUES (17, 'MR', 1, ?, ?, 2)", (evidence, day))
    except sqlite3.IntegrityError:
        return True
    return False


check("the new CHECKs hold: a dose given here must have a date, an unknown evidence is refused",
      refused_by_check("given_here", None) and refused_by_check("guessed", "2026-01-01")
      and not refused_by_check("parent_report", None))
con.rollback()
con.close()
second = mig.run(old, work / "backups")
con = sqlite3.connect(old)
check("a second run does nothing: no backup, no change, no second record",
      second["backup"] is None and not second["problems"]
      and con.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0] == audit_before + 1
      and len(list((work / "backups").iterdir())) == 1)
con.close()
shutil.rmtree(work, ignore_errors=True)
check("schema.sql matches: the table has `evidence` and `place_given` and no `source`",
      {"evidence", "place_given"} <= {r[1] for r in qall("PRAGMA table_info(immunizations)")}
      and "source" not in {r[1] for r in qall("PRAGMA table_info(immunizations)")})

# ---------------------------------------------------------------------------
passed = sum(1 for _, ok in results if ok)
print("\n" + "=" * 78)
print(f"RESULT: {passed}/{len(results)} checks passed")
failed = [name for name, ok in results if not ok]
for name in failed:
    print(f"  FAILED: {name}")
sys.exit(1 if failed else 0)
