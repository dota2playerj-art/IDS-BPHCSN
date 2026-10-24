"""Every time a person reads is on the 12-hour clock: '2:21 PM' (Task 21, display only).

D1. utils/dates.py format_time() and format_datetime() build the 12-hour text without strftime's
    %p (locale-dependent) or %-I / %#I (one platform each). Registered as Jinja filters.
D2. The screens, the exports and the printed referral slip use them: the Patient Queue's Arrived
    column and the open visits from earlier days, the Doctor's list, the Appointment History, the
    referral slip, the Audit Logs, the clinic hours, and the "Exported ..." stamp in every file.
D3. Stored values, form values and <input type="time"> stay 24-hour ISO; the queue is still
    ordered by the stored time, never by the "2:21 PM" text.
D4. A structural check: no display path formats a time with %H/%I/%p, slices an ISO time, or
    asks the browser's locale for a time. Each storage or file-name use is listed below.
The browser's half (static/js/time-format.js) is tested in verify_time_format_browser.py.
"""
import html
import io
import re
import sqlite3
import sys
import zlib
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path

import _support  # noqa: E402  (first: builds the test database, points the app at it)

from app import create_app  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

sys.stdout.reconfigure(errors="replace")

app = create_app()
_support.attach(app)
TEST_DB = _support.DB_PATH
REPO = Path(_support.REPO)
ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID = 1, 2, 3, 5

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok)))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}")
    if detail and not ok:
        for line in str(detail).splitlines()[:12]:
            print(f"       {line}")
    return ok


def section(title):
    print(f"\n--- {title} ---")


def db():
    conn = sqlite3.connect(TEST_DB)
    conn.row_factory = sqlite3.Row
    return conn


def q1(sql, args=()):
    conn = db()
    try:
        return conn.execute(sql, args).fetchone()
    finally:
        conn.close()


def run_sql(sql, args=()):
    conn = db()
    try:
        cur = conn.execute(sql, args)
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def client_as(user_id):
    c = app.test_client()
    with c.session_transaction() as sess:
        sess["user_id"] = user_id
        _support.open_session(sess)
    return c


def body(response):
    return response.get_data(as_text=True)


def flat(text):
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", text)))


def pdf_text(data):
    out = []
    for match in re.finditer(rb"stream\r?\n(.*?)endstream", data, re.S):
        try:
            out.append(zlib.decompress(match.group(1)).decode("latin-1"))
        except Exception:
            out.append(match.group(1).decode("latin-1", "replace"))
    return "\n".join(out).replace("\\(", "(").replace("\\)", ")")


from utils.dates import (  # noqa: E402
    CLINIC_TIMEZONE, clock_time, display_date, display_timestamp, format_datetime, format_time,
)

# The 12-hour clock, as a person reads it: no leading zero, one space, AM or PM.
TWELVE = re.compile(r"^(?:[1-9]|1[0-2]):[0-5]\d(?::[0-5]\d)? (?:AM|PM)$")

print("=" * 78)
print("IDS-BPHCSN — the 12-hour clock on every screen, export and printout (Task 21)")
print("=" * 78)

# ---------------------------------------------------------------------------
section("1. format_time(): the cases in the task")
CASES = [
    ("00:00", "12:00 AM"), ("00:05", "12:05 AM"), ("11:59", "11:59 AM"), ("12:00", "12:00 PM"),
    ("12:30", "12:30 PM"), ("13:05", "1:05 PM"), ("23:59", "11:59 PM"),
]
for value, expected in CASES:
    got = format_time(value)
    check(f"{value} -> {expected}", got == expected, got)
check("seconds are kept when asked: 14:21:05 -> 2:21:05 PM",
      format_time("14:21:05", seconds=True) == "2:21:05 PM", format_time("14:21:05", seconds=True))
check("seconds are kept at midnight and noon: 12:00:00 AM, 12:00:09 PM",
      format_time("00:00:00", seconds=True) == "12:00:00 AM"
      and format_time("12:00:09", seconds=True) == "12:00:09 PM")
check("without seconds asked, a value with seconds shows none: 14:21:05 -> 2:21 PM",
      format_time("14:21:05") == "2:21 PM")
check("an ISO date and time without seconds: '2026-10-12 11:34' -> 11:34 AM",
      format_time("2026-10-12 11:34") == "11:34 AM")
check("an ISO date and time with seconds: '2026-10-12 14:21:05' -> 2:21 PM / 2:21:05 PM",
      format_time("2026-10-12 14:21:05") == "2:21 PM"
      and format_time("2026-10-12 14:21:05", seconds=True) == "2:21:05 PM")
check("the T form and SQLite's fractional seconds: '2026-10-12T00:05:00', '... 23:59:59.123'",
      format_time("2026-10-12T00:05:00") == "12:05 AM"
      and format_time("2026-10-12 23:59:59.123") == "11:59 PM")
check("None and '' -> '' (never a crash)", format_time(None) == "" and format_time("") == "")
check("blank, unreadable and impossible values -> ''",
      all(format_time(v) == "" for v in ("   ", "noon", "24:00", "12:60", "2026-13-01 10:00")))
check("a plain date has no time of day -> '' (no invented midnight)",
      format_time("2026-10-12") == "" and format_time(date(2026, 10, 12)) == "")
check("datetime and time objects", format_time(datetime(2026, 1, 1, 0, 0)) == "12:00 AM"
      and format_time(time(12, 0)) == "12:00 PM" and format_time(time(9, 7)) == "9:07 AM")
check("an aware UTC datetime is shown at the health center: 04:05 UTC -> 12:05 PM",
      format_time(datetime(2026, 10, 12, 4, 5, tzinfo=timezone.utc)) == "12:05 PM")
check("a hand-typed '9:30' reads as 9:30 AM", format_time("9:30") == "9:30 AM")
every_minute = [format_time(f"{h:02d}:{m:02d}") for h in range(24) for m in range(60)]
check("all 1,440 minutes of a day match the 12-hour pattern, and none repeats",
      all(TWELVE.match(t) for t in every_minute) and len(set(every_minute)) == 1440)
check("the hour never has a leading zero", not any(t.startswith("0") for t in every_minute))

# ---------------------------------------------------------------------------
section("2. format_datetime(): the date keeps its style, only the time changes")
check("default: 'Oct 12, 2026 · 11:34 AM'",
      format_datetime("2026-10-12 11:34:00") == "Oct 12, 2026 · 11:34 AM",
      format_datetime("2026-10-12 11:34:00"))
check("the exports' style: 'October 22, 2026 at 2:21 PM'",
      format_datetime(datetime(2026, 10, 22, 14, 21), "%B %d, %Y", " at ")
      == "October 22, 2026 at 2:21 PM")
check("the Audit Logs' style with seconds: 'Oct 12, 2026, 12:05:09 AM'",
      format_datetime("2026-10-12 00:05:09", separator=", ", seconds=True)
      == "Oct 12, 2026, 12:05:09 AM")
check("a plain date returns the date alone", format_datetime("2026-10-12") == "Oct 12, 2026")
check("None, '' and unreadable -> ''",
      format_datetime(None) == "" and format_datetime("") == "" and format_datetime("x") == "")
check("a time with no date -> '' (there is no date to show)", format_datetime("14:21") == "")

# ---------------------------------------------------------------------------
section("3. The older helpers now go through format_time(), and the filters are registered")
check("display_timestamp(): stored UTC 06:21 -> 'Oct 12, 2026 · 2:21 PM' (Manila)",
      display_timestamp("2026-10-12 06:21:00") == "Oct 12, 2026 · 2:21 PM",
      display_timestamp("2026-10-12 06:21:00"))
check("display_timestamp(): UTC 16:00 the day before -> 'Oct 12, 2026 · 12:00 AM'",
      display_timestamp("2026-10-11 16:00:00") == "Oct 12, 2026 · 12:00 AM")
check("display_timestamp() keeps its em dash for nothing, and a plain date stays a date",
      display_timestamp(None) == "—" and display_timestamp("2026-10-12") == "Oct 12, 2026")
check("clock_time() (the 'Loaded' pill): 12:00 Manila -> 12:00 PM, 00:05 -> 12:05 AM",
      clock_time(datetime(2026, 10, 12, 12, 0, tzinfo=CLINIC_TIMEZONE)) == "12:00 PM"
      and clock_time(datetime(2026, 10, 12, 0, 5, tzinfo=CLINIC_TIMEZONE)) == "12:05 AM")
check("clock_time() with no value is now, on the 12-hour clock", TWELVE.match(clock_time()))
from models.settings import display_time  # noqa: E402

check("clinic hours: 08:00 -> 8:00 AM, 12:00 -> 12:00 PM, 17:30 -> 5:30 PM, 00:00 -> 12:00 AM",
      [display_time(v) for v in ("08:00", "12:00", "17:30", "00:00")]
      == ["8:00 AM", "12:00 PM", "5:30 PM", "12:00 AM"])
check("clinic hours: a hand-edited value still shows what it holds",
      display_time("8am") == "8am" and display_time(None) == "")
env = app.jinja_env
check("format_time and format_datetime are Jinja filters",
      env.filters.get("format_time") is format_time
      and env.filters.get("format_datetime") is format_datetime)
check("the filters render: '14:21' | format_time, a value | format_datetime, None | format_time",
      env.from_string("{{ '14:21' | format_time }}|{{ v | format_datetime }}|"
                      "{{ none | format_time or '—' }}").render(v="2026-10-12 00:05:00", none=None)
      == "2:21 PM|Oct 12, 2026 · 12:05 AM|—")
from models.appointments import _clock  # noqa: E402

check("the queue's clock: checked_in_at '... 14:21:00' -> 2:21 PM, booked '09:05' -> 9:05 AM, "
      "nothing -> None", _clock("2026-10-12 14:21:00") == "2:21 PM"
      and _clock("09:05") == "9:05 AM" and _clock(None) is None and _clock("") is None)
from routes.audit_views import _decorate  # noqa: E402

audit_row = _decorate({"created_at": "2026-10-12 04:05:09", "details": "Signed in.",
                       "module": "Authentication", "user_name": "Test"})
check("Audit Logs: stored UTC 04:05:09 -> time cell 12:05 PM, date Oct 12, 2026, full text "
      "with seconds 'Oct 12, 2026, 12:05:09 PM'",
      audit_row["time_text"] == "12:05 PM" and audit_row["date_text"] == "Oct 12, 2026"
      and audit_row["full_text"] == "Oct 12, 2026, 12:05:09 PM",
      (audit_row["time_text"], audit_row["full_text"]))
audit_empty = _decorate({"created_at": None, "details": "", "module": "", "user_name": ""})
check("Audit Logs: no stored time shows the em dash", audit_empty["time_text"] == "—"
      and audit_empty["full_text"] == "—")

# ---------------------------------------------------------------------------
section("4. Visits at known times")
today = date.today()
run_sql("UPDATE appointments SET status = 'Completed' "
        "WHERE status IN ('In Consultation', 'Waiting', 'Waiting for the Doctor')")
admin, staff, doctor = (client_as(i) for i in (ADMIN_ID, STAFF_ID, DOCTOR_ID))
from models.referrals import RHU  # noqa: E402

_next = [100]


def add_visit(tag, status, clock, days_ago=0, booked=None, consultation=False):
    """A visit for a fresh fictional patient, checked in at `clock` (clinic-local, as stored)."""
    pid = _next[0]
    _next[0] += 1
    run_sql("INSERT INTO patients (id, patient_id, first_name, last_name, birth_date, sex, "
            "civil_status, address, municipality, province, contact_number, status) VALUES "
            "(?, ?, ?, 'Timemark', '1985-01-01', 'Female', 'Single', 'Purok 1', 'San Francisco', "
            "'Surigao del Norte', ?, 'Active')", (pid, f"PAT-{pid:03d}", f"Zztm{tag}",
                                                  f"0900000{pid:04d}"))
    day = (today - timedelta(days=days_ago)).isoformat()
    visit = run_sql(
        "INSERT INTO appointments (patient_id, scheduled_by, appointment_type, appointment_date, "
        "appointment_time, status, is_walkin, checked_in_at) VALUES (?, ?, 'Consultation', ?, ?, ?, "
        "1, ?)", (pid, STAFF_ID, day, booked, status, f"{day} {clock}"))
    if consultation:
        run_sql("INSERT INTO consultations (patient_id, recorded_by, appointment_id, "
                "consultation_date, chief_complaint) VALUES (?, ?, ?, ?, 'cough')",
                (pid, STAFF_ID, visit, day))
    return visit, pid, day


early, _, _ = add_visit("Early", "Waiting", "09:45:00", booked="09:45")
late, _, _ = add_visit("Late", "Waiting", "13:10:00")
old_visit, _, old_day = add_visit("Old", "Waiting", "00:05:00", days_ago=3)
wfd, wfd_pid, _ = add_visit("Wfd", "Waiting for the Doctor", "14:21:00", booked="14:21",
                            consultation=True)
referred, _, _ = add_visit("Ref", "Waiting", "12:00:00")
run_sql("UPDATE appointments SET status = 'Referred', status_before_referred = 'Waiting', "
        "referral_destination = ?, referred_by = ?, referred_at = ? WHERE id = ?",
        (RHU, STAFF_ID, f"{today.isoformat()} 12:35:00", referred))
stored_before = {r["id"]: (r["checked_in_at"], r["appointment_time"], r["referred_at"]) for r in
                 db().execute("SELECT id, checked_in_at, appointment_time, referred_at "
                              "FROM appointments").fetchall()}
check("five visits at 09:45, 13:10, 00:05 (three days ago), 14:21 and 12:00 (referred 12:35)",
      len({early, late, old_visit, wfd, referred}) == 5)

# ---------------------------------------------------------------------------
section("5. The Patient Queue")
for label, client, path in (("Staff", staff, "/staff/appointment_scheduling"),
                            ("the Doctor", doctor, "/staff/appointment_scheduling"),
                            ("the Administrator", admin, "/admin/appointment_scheduling")):
    response = client.get(path)
    page = body(response)
    text = flat(page)
    today_part, _, earlier_part = text.partition("Open visits from earlier days")
    check(f"{label}: the queue opens", response.status_code == 200)
    check(f"{label}: Arrived shows 9:45 AM, 1:10 PM, 2:21 PM and 12:00 PM",
          all(t in today_part for t in ("9:45 AM", "1:10 PM", "2:21 PM", "12:00 PM")),
          today_part[:600])
    check(f"{label}: an earlier day's open visit shows '{display_date(old_day)} · 12:05 AM'",
          f"{display_date(old_day)} · 12:05 AM" in earlier_part, earlier_part[:600])
    check(f"{label}: no 24-hour time is left (13:10, 14:21, 00:05, 09:45)",
          not any(t in text for t in ("13:10", "14:21", "00:05", "09:45")))
    check(f"{label}: the queue is in arrival order by the stored time (9:45 AM before 1:10 PM, "
          "though '1:10 PM' sorts first as text)",
          0 <= today_part.find("ZztmEarly") < today_part.find("ZztmLate"),
          (today_part.find("ZztmEarly"), today_part.find("ZztmLate")))

# ---------------------------------------------------------------------------
section("6. The dashboards, the record page and the Appointment History")
page = flat(body(doctor.get("/staff/dashboard")))
check("the Doctor's list: Checked in shows today's date · 2:21 PM",
      f"{display_date(today)} · 2:21 PM" in page, page[:400])
check("the Doctor's dashboard has no 24-hour arrival", "14:21" not in page)
page = flat(body(staff.get("/staff/dashboard")))
check("the Staff dashboard's queue list shows 9:45 AM and 1:10 PM, no 09:45 or 13:10",
      "9:45 AM" in page and "1:10 PM" in page and "09:45" not in page and "13:10" not in page)
page = flat(body(staff.get(f"/staff/patient_profile/{wfd_pid}")))
check("the Appointment History shows the booked time 14:21 as 2:21 PM",
      "2:21 PM" in page and "14:21" not in page)
page = flat(body(admin.get("/admin/dashboard")))
check("every time on the Administrator's dashboard is on the 12-hour clock",
      not re.search(r"\b(?:[01]\d|2[0-3]):[0-5]\d\b(?! ?[AP]M)", page),
      re.findall(r".{30}\b(?:[01]\d|2[0-3]):[0-5]\d\b.{10}", page)[:5])

# ---------------------------------------------------------------------------
section("7. The printed referral slip")
for label, client in (("Staff", staff), ("the Doctor", doctor)):
    page = flat(body(client.get(f"/staff/appointment_scheduling/referral-slip/{referred}")))
    check(f"{label}: the slip prints '{display_date(today)} · 12:35 PM'",
          f"{display_date(today)} · 12:35 PM" in page, page[:400])

# ---------------------------------------------------------------------------
section("8. Exports show the same time text as the screen")
response = staff.get("/staff/appointment_scheduling/export.xlsx")
book = load_workbook(io.BytesIO(response.data), data_only=True)
queue_sheet = book["Today's Queue"]
rows = [list(r) for r in queue_sheet.iter_rows(values_only=True)]
header = next((r for r in rows if "Arrived" in r), None)
arrived = []
if header:
    column = header.index("Arrived")
    arrived = [r[column] for r in rows[rows.index(header) + 1:] if len(r) > column and r[column]]
check("Excel: the queue's Arrived column is 9:45 AM, 1:10 PM, 2:21 PM, 12:00 PM",
      {"9:45 AM", "1:10 PM", "2:21 PM", "12:00 PM"} <= set(arrived), arrived)
check("Excel: every Arrived cell is 12-hour text, as on the screen (not a 24-hour value)",
      arrived and all(isinstance(v, str) and TWELVE.match(v) for v in arrived), arrived)
cells = [c for sheet in book.worksheets for row in sheet.iter_rows(values_only=True) for c in row]
check("Excel: no cell holds a time or datetime value (each time is the screen's text)",
      not any(isinstance(c, (time, datetime)) for c in cells))
stamps = [row[1] for sheet in book.worksheets for row in sheet.iter_rows(values_only=True)
          if row and row[0] == "Exported" and len(row) > 1]
STAMP = re.compile(r"^[A-Z][a-z]+ \d{2}, \d{4} at (?:[1-9]|1[0-2]):[0-5]\d (?:AM|PM)$")
check("Excel: the 'Exported' stamp reads 'October 22, 2026 at 2:21 PM' (no '02:21')",
      stamps and all(STAMP.match(str(s)) for s in stamps), stamps)
text = pdf_text(staff.get("/staff/appointment_scheduling/export.pdf").data)
check("PDF: the queue shows 9:45 AM, 1:10 PM, 2:21 PM and 12:00 PM",
      all(t in text for t in ("9:45 AM", "1:10 PM", "2:21 PM", "12:00 PM")))
check("PDF: the 'Exported' footer is on the 12-hour clock with no leading zero",
      re.search(r"Exported [A-Z][a-z]+ \d{2}, \d{4} at (?:[1-9]|1[0-2]):[0-5]\d (?:AM|PM)", text)
      and not re.search(r"at 0\d:\d\d", text))
for label, path in (("Reports & Analytics", "/admin/reports_analytics/export.xlsx"),
                    ("Audit Logs", "/admin/audit_logs/export.xlsx")):
    response = admin.get(path)
    if response.status_code != 200:
        check(f"Excel ({label}) downloads", False, response.status_code)
        continue
    book = load_workbook(io.BytesIO(response.data), data_only=True)
    stamps = [row[1] for sheet in book.worksheets for row in sheet.iter_rows(values_only=True)
              if row and row[0] == "Exported" and len(row) > 1]
    check(f"Excel ({label}): the 'Exported' stamp is on the 12-hour clock",
          stamps and all(STAMP.match(str(s)) for s in stamps), stamps)

# ---------------------------------------------------------------------------
section("9. Stored values, form values and time inputs stay 24-hour")
stored_after = {r["id"]: (r["checked_in_at"], r["appointment_time"], r["referred_at"]) for r in
                db().execute("SELECT id, checked_in_at, appointment_time, referred_at "
                             "FROM appointments").fetchall()}
check("rendering and exporting changed no stored value", stored_after == stored_before)
from models.appointments import check_in_patient  # noqa: E402

with app.test_request_context():
    run_sql("INSERT INTO patients (id, patient_id, first_name, last_name, birth_date, sex, "
            "civil_status, address, municipality, province, contact_number, status) VALUES "
            "(150, 'PAT-150', 'ZztmNew', 'Timemark', '1985-01-01', 'Female', 'Single', 'Purok 1', "
            "'San Francisco', 'Surigao del Norte', '09000000150', 'Active')")
    check_in_patient(150, "", STAFF_ID, service="Consultation")
new_row = q1("SELECT checked_in_at FROM appointments WHERE patient_id = 150")
check("a new check-in is still stored as 'YYYY-MM-DD HH:MM:SS' (24-hour)",
      new_row and re.match(r"^\d{4}-\d{2}-\d{2} [0-2]\d:[0-5]\d:[0-5]\d$", new_row["checked_in_at"]),
      new_row and new_row["checked_in_at"])
page = body(admin.get("/admin/settings"))
inputs = re.findall(r'<input type="time"[^>]*?value="([^"]*)"', page, re.S)
check("Settings: both <input type=\"time\"> fields keep a 24-hour HH:MM value",
      len(inputs) == 2 and all(re.match(r"^[0-2]\d:[0-5]\d$", v) for v in inputs), inputs)
opening = q1("SELECT setting_value FROM settings WHERE setting_key = 'opening_time'")
check("Settings: the stored opening time is still 24-hour HH:MM",
      opening is None or re.match(r"^[0-2]\d:[0-5]\d$", opening["setting_value"] or ""))
labels = re.search(r"Currently (.+?) to (.+?)\.", flat(page))
check("Settings: 'Currently ... to ...' reads on the 12-hour clock",
      labels and TWELVE.match(labels.group(1)) and TWELVE.match(labels.group(2)),
      labels and labels.groups())
page = flat(body(staff.get("/staff/dashboard")))
pill = re.search(r"Loaded (\S+ [AP]M)", page)
check("the 'Loaded' pill is on the 12-hour clock", pill and TWELVE.match(pill.group(1)),
      pill and pill.group(0))

# ---------------------------------------------------------------------------
section("10. Structural check: no display path formats a time any other way")
# Where to look: everything that renders a page, a file or a printout, plus the scripts.
SCAN = ["app.py", "config.py", "routes", "models", "utils", "templates", "static/js", "database",
        "migrations", "seed_demo_data.py", "restore_backup.py", "archive_plain_databases.py",
        "setup_env.py", "setup_https.py", "export_sealed_key.py"]
# strftime's hour and AM/PM directives, and slicing the time out of stored ISO text.
SERVER = re.compile(r"%H|%I|%p\b|%-I|%#I|\[11:1[69]\]|\[11:\]")
# The browser's locale clock. Only static/js/time-format.js (the rule itself) may read hours.
BROWSER = re.compile(r"toLocaleTimeString|toLocaleString\(|getHours\(|hour12|\bhour\s*:")
HELPER_JS = "static/js/time-format.js"

# Every remaining use, listed explicitly. None of them is shown to a person as a time.
ALLOWED = {
    ("models/appointments.py", 'stamp = now().strftime("%Y-%m-%d %H:%M:%S")'):
        (2, "checked_in_at / left_at: stored"),
    ("models/appointments.py",
     'since = datetime.strptime((entry["checked_in_at"] or "")[:19], "%Y-%m-%d %H:%M:%S")'):
        (1, "parsing the stored value"),
    ("models/referrals.py", 'stamp = now().strftime("%Y-%m-%d %H:%M:%S")'):
        (1, "referred_at: stored"),
    ("models/consultations.py",
     '"restricted_at = strftime(\'%Y-%m-%d %H:%M:%f\', \'now\') WHERE id = ?",'):
        (1, "SQL, stored"),
    ("models/consultations.py", '"updated_at = strftime(\'%Y-%m-%d %H:%M:%f\', \'now\') "'):
        (1, "SQL, stored"),
    ("models/consultations.py",
     '"SET restricted_at = strftime(\'%Y-%m-%d %H:%M:%f\', \'now\') WHERE id = ?",'):
        (1, "SQL, stored"),
    ("models/doctor_worklist.py",
     'for pattern, size in (("%Y-%m-%d %H:%M:%S", 19), ("%Y-%m-%d", 10)):'):
        (1, "parsing the stored value"),
    ("utils/dates.py", 'SQLITE_TIMESTAMP = "%Y-%m-%d %H:%M:%S"'): (1, "the stored shape"),
    ("utils/dates.py", 'for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d"):'):
        (1, "parsing the stored value"),
    ("utils/dates.py",
     "strftime's %p, which follows the computer's locale and can come back empty or"):
        (1, "format_time()'s docstring, naming what it avoids"),
    ("utils/dates.py",
     "in another language, or %-I / %#I, which work on only one platform each. The"):
        (1, "format_time()'s docstring, naming what it avoids"),
    ("utils/security.py",
     'return datetime.fromtimestamp(seconds, CLINIC_TIMEZONE).strftime("%Y-%m-%d %H:%M")'):
        (1, "text written into a stored audit entry; stored data is not rewritten"),
    ("routes/settings_views.py",
     "f\"{config.SYSTEM_SHORT_NAME} backup {stamped.strftime('%Y-%m-%d %H%M')}{BACKUP_EXTENSION}\""):
        (1, "a backup file name"),
    ("seed_demo_data.py", 'return moment.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")'):
        (1, "demo data: stored"),
    ("restore_backup.py", 'stamp = datetime.now().strftime("%Y%m%d_%H%M%S")'): (1, "a file name"),
    ("archive_plain_databases.py", 'stamp = datetime.now().strftime("%Y%m%d_%H%M%S")'):
        (1, "a file name"),
    ("migrations/migrate_add_must_change_password.py",
     '["git", "log", "--all", "--format=%H", "--", "models/users.py"],'): (1, "a git commit hash"),
}
for name in ("migrate_add_left_before_seen", "migrate_add_must_change_password",
             "migrate_add_patient_barangay", "migrate_civil_status_minor_to_single",
             "migrate_add_referrals", "migrate_admin_minimum_necessary",
             "migrate_immunization_evidence", "migrate_immunization_repeats",
             "migrate_inventory_batches", "migrate_redact_clinical_audit_text",
             "migrate_remove_service_names_from_audit", "migrate_seal_free_text_fields",
             "migrate_visit_service_and_doctor_status"):
    ALLOWED[(f"migrations/{name}.py", 'stamp = datetime.now().strftime("%Y%m%d_%H%M%S")')] = (
        1, "a backup file name")


def scan_text(relative, text):
    """Every (relative path, stripped line) in `text` that formats a time for display."""
    hits = []
    browser = relative.endswith((".html", ".js")) and relative != HELPER_JS
    for line in text.splitlines():
        if SERVER.search(line) or (browser and BROWSER.search(line)):
            hits.append((relative, line.strip()))
    return hits


def scan_files():
    hits = []
    for entry in SCAN:
        path = REPO / entry
        files = [path] if path.is_file() else sorted(
            p for p in path.rglob("*") if p.suffix in (".py", ".html", ".js")
            and "__pycache__" not in p.parts and not p.name.endswith(".min.js"))
        for file in files:
            relative = file.relative_to(REPO).as_posix()
            hits += scan_text(relative, file.read_text(encoding="utf-8", errors="replace"))
    return hits


found = scan_files()
counts = {}
for hit in found:
    counts[hit] = counts.get(hit, 0) + 1
unexpected = [f"{path}: {line}" for (path, line), n in counts.items()
              if n > ALLOWED.get((path, line), (0, ""))[0]]
stale = [f"{path}: {line}" for (path, line), (n, _) in ALLOWED.items()
         if counts.get((path, line), 0) != n]
check("no display path uses %H, %I, %p, an ISO time slice or the browser's locale clock",
      not unexpected, "\n".join(unexpected))
check("every listed storage / file-name use is still there exactly as listed (the list is current)",
      not stale, "\n".join(stale))
check("the scan really looks: it sees the files it should",
      any(p == "models/appointments.py" for p, _ in found)
      and any(p.startswith("migrations/") for p, _ in found))
probe_hits = scan_text("templates/probe.html",
                       "{{ t.strftime('%H:%M') }}\n<script>d.toLocaleTimeString()</script>\n"
                       "x = s[11:16]\nnew Date().toLocaleString(undefined, {hour:'2-digit'})\n"
                       "{{ t.strftime('%I:%M %p') }}")
check("the scan catches each kind it forbids (%H:%M, toLocaleTimeString, [11:16], a locale hour, "
      "%I:%M %p)", len(probe_hits) == 5, probe_hits)
check("the scan does not flag the helper itself or a plain date format",
      not scan_text(HELPER_JS, "hour: value.getHours(),")
      and not scan_text("routes/x.py", "stamped.strftime('%B %d, %Y')"))
base = (REPO / "templates/base.html").read_text(encoding="utf-8")
check("base.html loads time-format.js before the page scripts block",
      0 <= base.find("js/time-format.js") < base.find("{% block scripts %}"))
check("time-format.js is served", app.test_client().get("/static/js/time-format.js").status_code == 200)
for name in ("community_reports", "patient_profile", "patient_profile_demographics",
             "patient_records", "reports_analytics"):
    source = (REPO / f"templates/{name}.html").read_text(encoding="utf-8")
    # The clinic's clock, not the device's (audit B23, Task 24): IDSTime.clinicStamp().
    check(f"{name}.html stamps its printout with IDSTime.clinicStamp()",
          "IDSTime.clinicStamp()" in source and "new Date()" not in source and "toLocaleString" not in source)

# ---------------------------------------------------------------------------
failed = [name for name, ok in results if not ok]
print("\n" + "=" * 78)
print(f"RESULT: {len(results) - len(failed)}/{len(results)} checks passed")
for name in failed:
    print(f"  FAILED: {name}")
print("=" * 78)
sys.exit(1 if failed else 0)
