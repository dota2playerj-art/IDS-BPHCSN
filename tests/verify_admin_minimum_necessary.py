"""The Administrator sees no clinical data (Task 5c; RA 11166, RA 10173, minimum necessary).

Clinical data is for the care team: the Doctor and Health Center Staff. The
Administrator runs the system and keeps accounts, settings, inventory, reports,
SMS broadcasts and the audit trail, plus a demographics-only view of patients and
a read-only Patient Queue.
"""
import contextlib
import io
import re
import shutil
import sqlite3
import sys
import zlib
from pathlib import Path

import _support  # noqa: E402  (first: builds the test database, points the app at it)

from app import create_app  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

sys.stdout.reconfigure(errors="replace")

app = create_app()
_support.attach(app)
TEST_DB = _support.DB_PATH
ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID = 1, 2, 3, 5

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok)))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}")
    if detail and not ok:
        for line in str(detail).splitlines()[:12]:
            print(f"       {line}")
    return ok


def section(title):
    print(f"\n--- {title} ---")


def db():
    conn = sqlite3.connect(TEST_DB)
    conn.row_factory = sqlite3.Row
    return conn


def q1(sql, args=()):
    conn = db()
    try:
        return conn.execute(sql, args).fetchone()
    finally:
        conn.close()


def run_sql(sql, args=()):
    conn = db()
    try:
        cur = conn.execute(sql, args)
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def client_as(user_id):
    c = app.test_client()
    with c.session_transaction() as sess:
        sess["user_id"] = user_id
        _support.open_session(sess)
    return c


def body(response):
    return response.get_data(as_text=True)


def pdf_text(data):
    out = []
    for match in re.finditer(rb"stream\r?\n(.*?)endstream", data, re.S):
        try:
            out.append(zlib.decompress(match.group(1)).decode("latin-1"))
        except Exception:
            out.append(match.group(1).decode("latin-1", "replace"))
    return "\n".join(out)


def xlsx_text(data):
    book = load_workbook(io.BytesIO(data), data_only=True)
    return "\n".join(str(cell) for sheet in book.worksheets
                     for row in sheet.iter_rows(values_only=True) for cell in row
                     if cell is not None)


def markup(text):
    """The page as a person sees and uses it: no style, script or comment text."""
    return re.sub(r"<style.*?</style>|<script.*?</script>|<!--.*?-->|\{#.*?#\}", "",
                  text, flags=re.S)


def readable(response):
    kind = response.headers.get("Content-Type", "")
    if "spreadsheetml" in kind:
        return xlsx_text(response.data)
    if "pdf" in kind:
        return pdf_text(response.data) + response.data.decode("latin-1")
    return response.data.decode("utf-8", "replace")


admin, staff, official, doctor = (client_as(i) for i in
                                  (ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID))

print("=" * 78)
print("IDS-BPHCSN — the Administrator sees no clinical data (Task 5c)")
print("=" * 78)

# ===========================================================================
section("1. PUT A MARKER INTO EVERY CLINICAL FIELD (unsealed records)")
P = 3                      # Ana Cruz, PAT-003: an adult
P_CHILD = 17               # the infant
P_QUEUE = 4
MK = {
    "chief_complaint": "ZZMKCHIEFCOMPLAINT",
    "diagnosis": "ZZMKDIAGNOSIS",
    "treatment_plan": "ZZMKTREATMENT",
    "remarks": "ZZMKNOTES",
    "visit_reason": "ZZMKVISITREASON",
    "queue_reason": "ZZMKQUEUEREASON",
    "conditions": "ZZMKCONDITION",
    "allergies": "ZZMKALLERGY",
    "dose_notes": "ZZMKDOSENOTE",
    "vaccine_name": "ZZMKVACCINE",
    "medicine_name": "ZZMKMEDNAME",
    "sms_message": "ZZMKSMSMESSAGE",
    "emergency_name": "ZZMKEMERGENCYNAME",
}
CONTACT, EMERGENCY_CONTACT = "09007775311", "09007775322"
VITALS = {"blood_pressure": "211/137", "weight": "77.77", "temperature": "41.3"}
# Markers that are values rather than words. Checked as they would print.
VITAL_MARKS = ("211/137", "77.77 kg", "41.3 °C")
FIXTURE_NUMBER = re.compile(r"0900000\d{4}")

run_sql("UPDATE patients SET contact_number = ?, emergency_contact_name = ?, "
        "emergency_contact_number = ?, medical_conditions = ?, allergies = ?, "
        "blood_type = 'AB-' WHERE id = ?",
        (CONTACT, MK["emergency_name"], EMERGENCY_CONTACT,
         f'["{MK["conditions"]}"]', f'["{MK["allergies"]}"]', P))

appt = run_sql(
    "INSERT INTO appointments (patient_id, scheduled_by, appointment_date, status, "
    "is_walkin, checked_in_at, notes) VALUES (?, ?, date('now','localtime'), "
    "'In Consultation', 1, datetime('now','localtime'), ?)",
    (P, STAFF_ID, MK["visit_reason"]))
from models.general_vaccinations import OTHER_KEY, OTHER_NAME_FIELD, VACCINE_FIELD  # noqa: E402

mark = q1("SELECT COALESCE(MAX(id),0) FROM audit_logs")[0]
response = doctor.post(f"/staff/consultation/record/{appt}", data={
    "chief_complaint": MK["chief_complaint"], "diagnosis": MK["diagnosis"],
    "treatment_plan": MK["treatment_plan"], "remarks": MK["remarks"],
    "temperature": VITALS["temperature"], "blood_pressure": VITALS["blood_pressure"],
    "weight": VITALS["weight"], "pulse_rate": "80", "respiratory_rate": "18",
    "dispense_item_0": "1", "dispense_qty_0": "3",
    VACCINE_FIELD: OTHER_KEY, OTHER_NAME_FIELD: MK["vaccine_name"],
})
CID = q1("SELECT id FROM consultations WHERE appointment_id = ?", (appt,))
check("the Doctor records a marked consultation through the real form",
      CID is not None and q1("SELECT chief_complaint FROM consultations WHERE id=?",
                             (CID[0],))[0] == MK["chief_complaint"], response.status_code)
CID = CID[0] if CID else None
NEW_ENTRIES = [dict(r) for r in db().execute(
    "SELECT action, details FROM audit_logs WHERE id > ?", (mark,))]
run_sql("INSERT INTO medicine_dispensing (patient_id, consultation_id, inventory_item_id, "
        "medicine_name, unit, quantity_dispensed, date_dispensed, dispensed_by) "
        "VALUES (?, ?, 2, ?, 'capsule', 4, date('now','localtime'), ?)",
        (P, CID, MK["medicine_name"], STAFF_ID))
run_sql("INSERT INTO immunizations (patient_id, vaccine_code, dose_number, evidence, "
        "date_administered, administered_by, notes) VALUES (?, 'OPV', 1, 'given_here', "
        "date('now','localtime'), ?, ?)", (P_CHILD, STAFF_ID, MK["dose_notes"]))
QUEUE_APPT = run_sql(
    "INSERT INTO appointments (patient_id, scheduled_by, appointment_date, status, "
    "is_walkin, checked_in_at, notes) VALUES (?, ?, date('now','localtime'), 'Waiting', "
    "1, datetime('now','localtime'), ?)", (P_QUEUE, STAFF_ID, MK["queue_reason"]))
run_sql("INSERT INTO appointment_sms_messages (reference, appointment_id, patient_id, event, "
        "recipient_name, mobile_number, message, segments, encoding, status, delivery_mode, "
        "sent_by) VALUES ('SMS-ZZ01', ?, ?, 'booked', 'Ana Cruz', ?, ?, 1, 'GSM-7', 'Sent', "
        "'simulated', ?)", (appt, P, CONTACT, MK["sms_message"], ADMIN_ID))
response = admin.post("/admin/sms_notifications/send", data={
    "message": "Health center open on Saturday.", "audience": "all_patients",
    "category": "General Reminder"})
check("the Administrator still sends a community broadcast (KEEP)",
      q1("SELECT COUNT(*) FROM sms_broadcast_recipients WHERE mobile_number = ?",
         (CONTACT,))[0] == 1, response.status_code)
row = q1("SELECT * FROM consultations WHERE id = ?", (CID,)) if CID else None
check("every marker is in the database, unsealed",
      row is not None and row["is_restricted"] == 0 and row["remarks"] == MK["remarks"]
      and row["blood_pressure"] == VITALS["blood_pressure"]
      and q1("SELECT COUNT(*) FROM medicine_dispensing WHERE medicine_name=?",
             (MK["medicine_name"],))[0] == 1
      and q1("SELECT COUNT(*) FROM general_vaccinations WHERE consultation_id=?",
             (CID,))[0] == 1)

# A sealed visit too, for the backup check. Its text must not be in the file either.
SEALED_CC = "ZZMKSEALEDCC"
appt2 = run_sql(
    "INSERT INTO appointments (patient_id, scheduled_by, appointment_date, status, "
    "is_walkin, checked_in_at) VALUES (5, ?, date('now','localtime'), 'In Consultation', "
    "1, datetime('now','localtime'))", (STAFF_ID,))
doctor.post(f"/staff/consultation/record/{appt2}", data={
    "chief_complaint": SEALED_CC, "diagnosis": "ZZMKSEALEDDX", "treatment_plan": "ZZMKSEALEDTX",
    "temperature": "36.6", "is_restricted": "1", "dispense_item_0": "3", "dispense_qty_0": "2"})
CID_SEALED = q1("SELECT id FROM consultations WHERE appointment_id = ?", (appt2,))[0]

WORD_MARKS = tuple(MK.values()) + (CONTACT, EMERGENCY_CONTACT) + VITAL_MARKS
# Reports & Analytics keep their rankings by decision (Task 5c): a ranked complaint,
# diagnosis or medicine name is shown with a count and no patient. Those three are
# therefore not looked for there; every other marker is.
RANKED = (MK["chief_complaint"], MK["diagnosis"], MK["medicine_name"])

# ===========================================================================
section("2. CRAWL EVERYTHING THE ADMINISTRATOR CAN REACH")
plain_gets = sorted({rule.rule for rule in app.url_map.iter_rules()
                     if "GET" in rule.methods and not rule.arguments
                     and not rule.rule.startswith("/static")
                     and rule.rule != "/logout"})
BACKUP_URL = "/admin/settings/backup.idsbackup"
pages = exports = jsons = 0
found = []
for url in plain_gets:
    if url == BACKUP_URL:
        continue  # an encrypted file; checked byte by byte in section 6
    response = admin.get(url)
    text = readable(response)
    kind = response.headers.get("Content-Type", "")
    if "spreadsheetml" in kind or "pdf" in kind:
        exports += 1
    elif "json" in kind:
        jsons += 1
    else:
        pages += 1
    marks = [m for m in WORD_MARKS
             if m.lower() in text.lower() and not ("reports_analytics" in url and m in RANKED)]
    if FIXTURE_NUMBER.search(text):
        marks.append("a patient contact number")
    if marks:
        found.append(f"{url}: {marks}")
patient_ids = [r[0] for r in db().execute("SELECT id FROM patients")]
for pid in patient_ids:
    for url in (f"/admin/patient_profile/{pid}", f"/admin/patient_profile/{pid}?print=1"):
        text = body(admin.get(url))
        pages += 1
        marks = [m for m in WORD_MARKS if m.lower() in text.lower()]
        if FIXTURE_NUMBER.search(text):
            marks.append("a patient contact number")
        if marks:
            found.append(f"{url}: {marks}")
for kind in ("xlsx", "pdf"):
    response = admin.post(f"/admin/patient_records/export-selection.{kind}",
                          data={"patient_ids": [str(i) for i in patient_ids]})
    exports += 1
    text = readable(response)
    marks = [m for m in WORD_MARKS if m.lower() in text.lower()]
    if FIXTURE_NUMBER.search(text):
        marks.append("a patient contact number")
    if marks:
        found.append(f"selection export {kind}: {marks}")
response = admin.post("/session/keep-alive", headers={"X-Requested-With": "fetch"})
jsons += 1
if any(m.lower() in body(response).lower() for m in WORD_MARKS):
    found.append("keep-alive JSON")
CRAWLED = (pages, exports, jsons)
print(f"       crawled {pages} pages and print views, {exports} exports, {jsons} JSON responses")
check("*** no clinical marker, contact number or medicine name reaches the Administrator "
      "on any page, export, print view or JSON response ***", not found, "\n".join(found))
check("the crawl is wide: over 30 screens, every patient's profile and print view, "
      "and the exports", pages >= 30 + 2 * len(patient_ids) and exports >= 10 and jsons >= 1,
      CRAWLED)

profile = body(admin.get(f"/admin/patient_profile/{P}"))
check("his profile shows the demographics: ID, name, sex, age, birth date, address, "
      "status and registration date",
      all(t in profile for t in ("PAT-003", "Ana Cruz", "Female", "Nov 20, 1992",
                                 "Purok 1, Barangay Poblacion", "Active", "Registered")))
check("and one plain line where the clinical sections were",
      "Clinical details are shown only to the Doctor and Health Center Staff." in profile)
check("it carries no contact number, blood type, emergency contact or tab",
      CONTACT not in profile and "AB-" not in profile and "tab-consult" not in profile
      and MK["emergency_name"] not in profile)
roster = body(admin.get("/admin/patient_records"))
check("his patient list shows the demographic columns and no contact or visit column",
      "PAT-003" in roster and "Ana Cruz" in roster and "<th>Registered</th>" in roster
      and "<th>Contact Number</th>" not in roster and "<th>Last Visit</th>" not in roster)
book = load_workbook(io.BytesIO(admin.get("/admin/patient_records/export.xlsx").data))
header = next((list(r) for s in book.worksheets for r in s.iter_rows(values_only=True)
               if r and r[0] == "Patient ID"), [])
check("his patient export is demographics only",
      [h for h in header if h] == ["Patient ID", "Full Name", "Age", "Sex", "Address",
                                   "Status", "Registered"], header)
dash = markup(body(admin.get("/admin/dashboard")))
check("his dashboard keeps counts and lists no consultation or child",
      "Recent Patient Consultations" not in dash and "Total Registered Patients" in dash
      and "Immunization Follow-up" in dash and "Paolo Morales" not in dash)
queue = markup(body(admin.get("/admin/appointment_scheduling")))
check("his Patient Queue shows who is waiting, read-only",
      "Morales" in queue or "Bautista" in queue)
check("with no Check In, Start, Record or Complete control, and no dialog behind them",
      not any(t in queue for t in ("checkInModal", "startConsultationForm",
                                   "completeConsultationForm", "Record Consultation",
                                   "Start Consultation", "immun-tag")))
sms = body(admin.get("/admin/sms_notifications"))
check("SMS Notifications: he sees recipient counts only, no name and no number (final fixes, M37)",
      "Ana Cruz" not in sms and "Total" in sms and CONTACT not in sms and not FIXTURE_NUMBER.search(sms)
      and MK["sms_message"] not in sms)
base = markup(body(admin.get("/admin/patient_records")))
check("no Register Patient, Add Patient, Edit or Change Status control anywhere for him",
      "Register Patient" not in base and "Add Patient" not in base
      and "archivePatientModal" not in base and "Change Status" not in base)

# ===========================================================================
section("3. EVERY REMOVED ROUTE REFUSES HIM ON THE SERVER, POST INCLUDED")


def snapshot():
    conn = db()
    try:
        return tuple(conn.execute(sql).fetchone()[0] for sql in (
            "SELECT COUNT(*) FROM patients",
            "SELECT group_concat(status) FROM (SELECT status FROM patients ORDER BY id)",
            "SELECT COUNT(*) FROM consultations",
            "SELECT group_concat(status) FROM (SELECT status FROM appointments ORDER BY id)",
            "SELECT COUNT(*) FROM appointments",
            "SELECT COUNT(*) FROM immunizations",
            "SELECT COUNT(*) FROM general_vaccinations",
            "SELECT COUNT(*) FROM medicine_dispensing",
            "SELECT group_concat(temperature) FROM consultations",
        ))
    finally:
        conn.close()


before = snapshot()
REMOVED_ADMIN = [
    ("GET", "/admin/register_patient", None),
    ("POST", "/admin/register_patient", {"first_name": "Zed", "last_name": "Zz"}),
    ("GET", f"/admin/register_patient/{P}", None),
    ("POST", f"/admin/register_patient/{P}", {"first_name": "Zed"}),
    ("POST", "/admin/patient_records/archive", {"patient_ids": [str(P)], "status": "Inactive"}),
    ("POST", f"/admin/patient_profile/{P}/correct-vaccination", {"record_id": "1"}),
    ("POST", "/admin/appointment_scheduling/check-in", {"patient_id": "6", "reason": "x", "service": "Consultation"}),
    ("POST", "/admin/appointment_scheduling/start-consultation", {"appointment_id": str(QUEUE_APPT)}),
    ("POST", "/admin/appointment_scheduling/complete-consultation", {"appointment_id": str(QUEUE_APPT)}),
    ("GET", f"/admin/consultation/record/{QUEUE_APPT}", None),
    ("POST", f"/admin/consultation/record/{QUEUE_APPT}", {"chief_complaint": "x"}),
    ("GET", f"/admin/consultation/edit/{CID}", None),
    ("POST", f"/admin/consultation/edit/{CID}", {"temperature": "39"}),
]
for method, url, data in REMOVED_ADMIN:
    response = admin.get(url) if method == "GET" else admin.post(url, data=data)
    check(f"{method} {url} no longer exists for him (404)", response.status_code == 404,
          response.status_code)
STAFF_EQUIVALENTS = [
    ("GET", "/staff/register-patient", None),
    ("POST", "/staff/register-patient", {"first_name": "Zed", "last_name": "Zz"}),
    ("POST", f"/staff/register-patient/{P}", {"first_name": "Zed"}),
    ("POST", "/staff/patient_records/archive", {"patient_ids": [str(P)], "status": "Inactive"}),
    ("POST", f"/staff/patient_profile/{P}/correct-vaccination", {"record_id": "1"}),
    ("GET", f"/staff/patient_profile/{P}", None),
    ("POST", "/staff/appointment_scheduling/check-in", {"patient_id": "6", "reason": "x", "service": "Consultation"}),
    ("POST", "/staff/appointment_scheduling/start-consultation", {"appointment_id": str(QUEUE_APPT)}),
    ("POST", "/staff/appointment_scheduling/complete-consultation", {"appointment_id": str(QUEUE_APPT)}),
    ("GET", f"/staff/consultation/record/{QUEUE_APPT}", None),
    ("POST", f"/staff/consultation/record/{QUEUE_APPT}", {"chief_complaint": "x"}),
    ("GET", f"/staff/consultation/edit/{CID}", None),
    ("POST", f"/staff/consultation/edit/{CID}", {"temperature": "39"}),
    ("GET", "/staff/patient_records/export.xlsx", None),
]
for method, url, data in STAFF_EQUIVALENTS:
    response = admin.get(url) if method == "GET" else admin.post(url, data=data)
    check(f"{method} {url} sends him back to his dashboard (POST carries a valid form token)",
          response.status_code == 302 and response.headers.get("Location", "").endswith("/dashboard"),
          (response.status_code, response.headers.get("Location")))
check("*** none of those requests changed a patient, a visit, a dose, a dispensing or "
      "the queue ***", snapshot() == before)
with app.test_request_context():
    from flask import session as _session
    _session["user_id"] = ADMIN_ID
    _session["role"] = "Admin"
    from routes.patient_views import patient_profile_context
    try:
        patient_profile_context(P)
        refused = False
    except PermissionError:
        refused = True
check("the full-profile builder itself refuses the Administrator, whatever route asks",
      refused)

# ===========================================================================
section("4. THE DOCTOR, HEALTH CENTER STAFF AND THE OFFICIAL KEEP THEIR ACCESS")
page = body(staff.get(f"/staff/patient_profile/{P}"))
check("Staff: the full profile, with the chief complaint, contact and medicine",
      MK["chief_complaint"] in page and CONTACT in page and MK["medicine_name"] in page
      and MK["conditions"] in page)
check("Staff: Register Patient opens", staff.get("/staff/register-patient").status_code == 200)
check("Staff: the queue shows the check-in reason and the Check In control",
      MK["queue_reason"] in body(staff.get("/staff/appointment_scheduling"))
      and "checkInModal" in body(staff.get("/staff/appointment_scheduling")))
check("Staff: the patient list shows contact numbers",
      CONTACT in body(staff.get("/staff/patient_records")))
check("Staff: the edit form opens", staff.get(f"/staff/consultation/edit/{CID}").status_code == 200)
check("Staff: the dashboard lists the children due for a dose",
      "Paolo Morales" in body(staff.get("/staff/dashboard"))
      or "Morales" in body(staff.get("/staff/dashboard")))
check("Staff: the appointment-SMS panel still shows",
      MK["sms_message"] in body(staff.get("/staff/sms_notifications")))
check("Doctor: the edit form shows the diagnosis",
      MK["diagnosis"] in body(doctor.get(f"/staff/consultation/edit/{CID}")))
check("Doctor: Reports & Analytics opens", doctor.get("/staff/reports_analytics").status_code == 200)
check("Doctor: the full profile with the vaccine given at the visit",
      MK["vaccine_name"].lower() in body(doctor.get(f"/staff/patient_profile/{P}")).lower())
response = staff.post("/staff/appointment_scheduling/check-in", data={"patient_id": "6",
                                                                     "reason": "Cough",
                                                                     "service": "Consultation"})
check("Staff: check-in still writes",
      q1("SELECT COUNT(*) FROM appointments WHERE patient_id = 6 AND status = 'Waiting' "
         "AND date(appointment_date) = date('now','localtime')")[0] == 1)
check("Official: dashboard and reports unchanged, and no /admin page",
      official.get("/official/dashboard").status_code == 200
      and official.get("/official/reports_analytics").status_code == 200
      and official.get("/admin/patient_records").status_code == 302)

# ===========================================================================
section("5. AUDIT ENTRIES NAME NO MEDICINE OR VACCINE")
dispensed = [e for e in NEW_ENTRIES if e["action"] == "Dispensed medicine"]
doses = [e for e in NEW_ENTRIES if e["action"] == "Recorded other vaccinations"]
check("a new dispensing entry says who, which patient, which consultation and how many",
      len(dispensed) == 1 and re.fullmatch(
          rf"Dispensed 1 item for Ana Cruz \(PAT-003\) at consultation #{CID}\.",
          dispensed[0]["details"]), dispensed)
check("and names no medicine", dispensed and "Paracetamol" not in dispensed[0]["details"])
check("a new dose entry says how many, for whom, at which consultation, and no vaccine",
      len(doses) == 1 and re.fullmatch(
          rf"Recorded 1 dose for Ana Cruz \(PAT-003\) at consultation #{CID}\.",
          doses[0]["details"]), doses)
mark = q1("SELECT COALESCE(MAX(id),0) FROM audit_logs")[0]
form = {"chief_complaint": MK["chief_complaint"], "remarks": MK["remarks"],
        "temperature": "39.4", "blood_pressure": VITALS["blood_pressure"],
        "weight": VITALS["weight"], "pulse_rate": "80", "respiratory_rate": "18"}
staff.post(f"/staff/consultation/edit/{CID}", data=form)
edit_entry = q1("SELECT details FROM audit_logs WHERE id > ? AND action = "
                "'Updated consultation record'", (mark,))
check("an edit entry names the vital that changed and quotes no value",
      edit_entry is not None and edit_entry[0].endswith("— Changed: Temperature")
      and "39.4" not in edit_entry[0] and "41.3" not in edit_entry[0],
      edit_entry[0] if edit_entry else None)
details = [r[0] or "" for r in db().execute("SELECT details FROM audit_logs")]
check("no audit entry in the database holds a clinical marker, the medicine or the number",
      not [d for d in details if any(m.lower() in d.lower() for m in WORD_MARKS)
           or "Paracetamol" in d and "PAT-" in d],
      [m for d in details for m in WORD_MARKS if m.lower() in d.lower()])

from migrations import migrate_admin_minimum_necessary as migration  # noqa: E402

copy = Path(_support.work_path("migrate5c.db"))
shutil.copyfile(TEST_DB, copy)
OLD = {
    "Dispensed medicine": "For Ana Cruz (PAT-003), consultation #3: 10 tablets of ZZOLDMED1, 1 bottle of ZZOLDMED2.",
    "Recorded EPI immunisation doses": "2 doses for Ana Cruz (PAT-003) at consultation #3: ZZOLDVAX1, ZZOLDVAX2 dose 1.",
    "Recorded doses given elsewhere": "1 dose for Ana Cruz (PAT-003) at consultation #3: ZZOLDVAX3 (date not known).",
    "Recorded other vaccinations": "1 dose for Ana Cruz (PAT-003) at consultation #3: ZZOLDVAX4.",
    "Updated consultation record": "Consultation #3 for Ana Cruz (PAT-003), visit of Jun 04, 2026, originally recorded by X. 3 fields changed — Temperature: 36 °C → ZZOLDVITAL °C; Medicine dispensed: ZZOLDMED3 — 2 tablets (stock now 5); Immunisations added: ZZOLDVAX5",
    "Corrected other vaccination record": "Other vaccination #1 for Ana Cruz (PAT-003), originally recorded by X outside any consultation. 1 field corrected — Vaccine: \"ZZOLDVAX6\" → \"ZZOLDVAX7\"",
    "Sent appointment SMS to patient": "SMS-0001 (booked) to Ana Cruz (PAT-003) on 09007775399, about appointment APT-00003; Sent; 1 GSM-7 segment; delivery mode simulated (recorded only, nothing transmitted); message: \"ZZOLDSMS\"",
}
conn = sqlite3.connect(copy)
for action, text in OLD.items():
    conn.execute("INSERT INTO audit_logs (user_id, action, module, details, created_at) "
                 "VALUES (1, ?, 'X', ?, '2026-08-01 01:00:00')", (action, text))
conn.commit()
before_entries = conn.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0]
conn.close()
backups = Path(_support.work_path("migration5c-backups"))
with contextlib.redirect_stdout(io.StringIO()):
    first = migration.run(copy, backup_dir=backups)
conn = sqlite3.connect(copy)
rewritten = {a: conn.execute("SELECT details FROM audit_logs WHERE action = ? AND "
                             "created_at = '2026-08-01 01:00:00'", (a,)).fetchone()[0]
             for a in OLD}
records = conn.execute("SELECT details FROM audit_logs WHERE action = ?",
                       (migration.ACTION,)).fetchall()
after_entries = conn.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0]
conn.close()
OLD_MARKS = ("ZZOLDMED", "ZZOLDVAX", "ZZOLDVITAL", "ZZOLDSMS", "09007775399")
check("the migration made a backup first", first["backup"] and Path(first["backup"]).exists()
      and Path(first["backup"]).parent == backups, first)
check("it rewrote every old dispensing, dose, edit, correction and SMS entry",
      first["entries_rewritten"] >= len(OLD)
      and not [a for a, t in rewritten.items() if any(m in t for m in OLD_MARKS)], rewritten)
check("dispensing reads 'Dispensed 2 items for <Name> (PAT-003) at consultation #3.'",
      rewritten["Dispensed medicine"] == "Dispensed 2 items for Ana Cruz (PAT-003) at consultation #3.",
      rewritten["Dispensed medicine"])
check("doses read 'Recorded 2 doses for <Name> (PAT-003) at consultation #3.'",
      rewritten["Recorded EPI immunisation doses"]
      == "Recorded 2 doses for Ana Cruz (PAT-003) at consultation #3.")
check("an edit keeps who, which record and which fields, as names",
      rewritten["Updated consultation record"].endswith(
          "3 fields changed — Changed: Temperature, immunisations added, items dispensed"),
      rewritten["Updated consultation record"])
check("an SMS entry keeps the reference and the visit, without the number or the message",
      rewritten["Sent appointment SMS to patient"].startswith(
          "SMS-0001 (booked) to Ana Cruz (PAT-003), about appointment APT-00003")
      and rewritten["Sent appointment SMS to patient"].endswith("message removed"))
check("it wrote exactly one record of itself, citing the law",
      len(records) == 1 and after_entries == before_entries + 1
      and "RA 11166" in records[0][0] and "RA 10173" in records[0][0], records)
check("after VACUUM the file holds none of the old text",
      not [m for m in OLD_MARKS if m.encode() in copy.read_bytes()])
with contextlib.redirect_stdout(io.StringIO()):
    second = migration.run(copy, backup_dir=backups)
conn = sqlite3.connect(copy)
again = conn.execute("SELECT COUNT(*) FROM audit_logs WHERE action = ?",
                     (migration.ACTION,)).fetchone()[0]
conn.close()
check("running it twice changes nothing and makes no second backup or record",
      second["backup"] is None and again == 1, second)

# ===========================================================================
section("6. THE BACKUP DOWNLOAD IS AN ENCRYPTED FILE")
response = admin.get(BACKUP_URL)
data = response.data
check("the Administrator still downloads a backup, as an .idsbackup file",
      response.status_code == 200
      and ".idsbackup" in response.headers.get("Content-Disposition", ""))
check("it does not start with 'SQLite format 3'; it starts with the IDSBAK1 header",
      not data.startswith(b"SQLite format 3") and data.startswith(b"IDSBAK1"))
plain_db = Path(TEST_DB).read_bytes()
check("control: the live test database does hold the unsealed markers",
      MK["chief_complaint"].encode() in plain_db and CONTACT.encode() in plain_db)
check("*** the backup holds no marker, sealed or unsealed ***",
      not [m for m in WORD_MARKS + (SEALED_CC, "Ana Cruz", "PAT-003") if m.encode() in data])
old_settings = body(admin.get("/admin/settings"))
check("Settings → Backup says the file is encrypted",
      "The backup file is encrypted. It can only be opened with this system" in old_settings)
check("the old .db download is gone", admin.get("/admin/settings/backup.db").status_code == 404)

import restore_backup  # noqa: E402

bk = Path(_support.work_path("restore"))
bk.mkdir(exist_ok=True)
backup_file = bk / "test.idsbackup"
backup_file.write_bytes(data)
lines = []
code = restore_backup.restore(backup_file, backup_dir=bk, say=lines.append)
restored = sorted(bk.glob("restored_*.db"))
check("restore_backup.py decrypts into a new file and reports integrity_check: ok",
      code == 0 and restored and "integrity_check: ok" in lines, lines)


def dump(path):
    conn = sqlite3.connect(path)
    try:
        tables = [r[0] for r in conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' "
            "ORDER BY name")]
        return {t: sorted(map(repr, conn.execute(f'SELECT * FROM "{t}"').fetchall()))
                for t in tables}
    finally:
        conn.close()


live = dump(TEST_DB)
live["audit_logs"] = None   # the download itself adds one audit entry after the copy
back = dump(restored[0]) if restored else {}
back["audit_logs"] = None
check("*** the restored database has identical table contents ***", live == back,
      [t for t in live if live.get(t) != back.get(t)])
lines = []
target = Path(_support.work_path("target.db"))
shutil.copyfile(TEST_DB, target)
run_sql_target = sqlite3.connect(target)
run_sql_target.execute("DELETE FROM inventory_items WHERE id = 13")
run_sql_target.commit()
run_sql_target.close()
code = restore_backup.restore(backup_file, replace=True, backup_dir=bk,
                              database=str(target), say=lines.append)
prerestore = sorted(bk.glob("target.db.PRERESTORE_BACKUP_*"))
check("--replace copies the current database aside first",
      code == 0 and prerestore
      and sqlite3.connect(prerestore[0]).execute(
          "SELECT COUNT(*) FROM inventory_items WHERE id = 13").fetchone()[0] == 0,
      lines)
check("then replaces it with the backup and says to restart the app",
      sqlite3.connect(target).execute(
          "SELECT COUNT(*) FROM inventory_items WHERE id = 13").fetchone()[0] == 1
      and any("Restart the app" in line for line in lines), lines)
import os  # noqa: E402

right_key = os.environ["SEALED_DATA_KEY"]
from utils import sealed_crypto  # noqa: E402

os.environ["SEALED_DATA_KEY"] = sealed_crypto.new_key()
lines = []
code = restore_backup.restore(backup_file, backup_dir=bk, say=lines.append)
os.environ["SEALED_DATA_KEY"] = right_key
check("a wrong key fails with one plain line and no traceback",
      code == 1 and lines == [restore_backup.CANNOT_OPEN], lines)
damaged = bk / "damaged.idsbackup"
damaged.write_bytes(data[:-10] + b"0123456789")
lines = []
check("a damaged file fails with the same plain line",
      restore_backup.restore(damaged, backup_dir=bk, say=lines.append) == 1
      and lines == [restore_backup.CANNOT_OPEN], lines)
process = __import__("subprocess").run(
    [sys.executable, "restore_backup.py", str(damaged)], cwd=_support.REPO,
    capture_output=True, text=True, encoding="utf-8",
    env=dict(os.environ, PYTHONIOENCODING="utf-8"))
check("run from the command line, it prints one line and no traceback",
      process.returncode == 1 and "Traceback" not in process.stderr
      and process.stdout.strip() == restore_backup.CANNOT_OPEN,
      process.stdout + process.stderr)

# ===========================================================================
section("7. THE MOST-DISPENSED RANKING COUNTS SEALED VISITS AS CONFIDENTIAL")
with app.app_context():
    from models.dispensing import get_dispensing_stats
    stats = get_dispensing_stats(viewer_role="Doctor")
ranking = stats["window"]["all_medicines"]
sealed_units = q1("SELECT COALESCE(SUM(d.quantity_dispensed),0) FROM medicine_dispensing d "
                  "JOIN consultations c ON c.id = d.consultation_id WHERE c.is_restricted = 1")[0]
confidential = [e for e in ranking if e["medicine_name"] == "Confidential"]
check("sealed-visit dispensing is one Confidential row with its units",
      sealed_units > 0 and len(confidential) == 1 and confidential[0]["units"] == sealed_units,
      (confidential, sealed_units))
check("the sealed visit's medicine is not ranked by its name",
      not [e for e in ranking if e["medicine_name"] == "Cetirizine 10mg"])
check("and the totals still match the table",
      stats["total_units"] == q1("SELECT SUM(quantity_dispensed) FROM medicine_dispensing")[0]
      and sum(e["units"] for e in ranking) == stats["window"]["units"])
text = readable(admin.get("/admin/reports_analytics/export.xlsx"))
# Task 5d: for the Administrator a row of fewer than 5 issues folds into Other, the
# Confidential row included. The Doctor's own ranking (above) holds the exact row,
# so the expected label is read from its issue count.
from utils.small_counts import OTHER_LABEL
sealed_issues = confidential[0]["events"]
if sealed_issues >= 5:
    check("the Administrator's report export lists Confidential in the dispensing ranking",
          "Confidential" in text and "not named" in text)
else:
    check("the Administrator's export folds the small Confidential dispensing row into Other",
          OTHER_LABEL in text and "Cetirizine" not in text,
          f"sealed issues={sealed_issues}")

# ===========================================================================
passed = sum(1 for _, ok in results if ok)
print("\n" + "=" * 78)
print(f"RESULT: {passed}/{len(results)} checks passed")
failed = [name for name, ok in results if not ok]
for name in failed:
    print("  FAILED:", name)
print("=" * 78)
sys.exit(0 if not failed else 1)
