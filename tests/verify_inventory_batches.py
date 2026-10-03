"""Stock in batches, and medicines named by generic name (Task 12).

The adviser's two concerns. First, a restock edited one quantity and one expiry date per
medicine, so the system lost which boxes expire when: each delivery is now a batch with its
own lot, expiry date, source and brand; stock is the sum of the usable batches; dispensing
takes from the batch that expires first (FEFO); and Edit details, Receive stock and Remove
from stock are three separate actions. Second, RA 6675 (the Generics Act) asks government
health agencies to use generic names, so a medicine is defined by generic name, strength and
dosage form, and the display name is built from them. The brand belongs to the batch.
The demo-data checks for this task (several batches per item, an expiring-soon and an
expired batch, FEFO replay, determinism) are in verify_demo_data.py, which already builds
the demo data twice and compares every table.
"""
import io
import re
import shutil
import sqlite3
import sys
import zlib
from datetime import timedelta
from pathlib import Path

import _support  # noqa: E402  (first: builds the test database, points the app at it)

import config  # noqa: E402
from app import create_app  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

sys.stdout.reconfigure(errors="replace")

app = create_app()
_support.attach(app)
TEST_DB = _support.DB_PATH
ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID = 1, 2, 3, 5

import models.inventory as inv  # noqa: E402
from html import unescape  # noqa: E402
from utils.dates import display_date  # noqa: E402
from models.dispensing import record_dispensing  # noqa: E402
from utils.dates import today  # noqa: E402

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok)))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}")
    if detail and not ok:
        for line in str(detail).splitlines()[:12]:
            print(f"       {line}")
    return ok


def section(title):
    print(f"\n--- {title} ---")


def db():
    conn = sqlite3.connect(TEST_DB)
    conn.row_factory = sqlite3.Row
    return conn


def qall(sql, args=()):
    conn = db()
    try:
        return conn.execute(sql, args).fetchall()
    finally:
        conn.close()


def q1(sql, args=()):
    rows = qall(sql, args)
    return rows[0] if rows else None


def run_sql(sql, args=()):
    conn = db()
    try:
        cur = conn.execute(sql, args)
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def client_as(user_id):
    c = app.test_client()
    with c.session_transaction() as sess:
        sess["user_id"] = user_id
        _support.open_session(sess)
    return c


def body(response):
    return response.get_data(as_text=True)


def flat(text):
    return unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", text)))


def pdf_text(data):
    out = []
    for match in re.finditer(rb"stream\r?\n(.*?)endstream", data, re.S):
        try:
            out.append(zlib.decompress(match.group(1)).decode("latin-1"))
        except Exception:
            out.append(match.group(1).decode("latin-1", "replace"))
    return "\n".join(out).replace("\\(", "(").replace("\\)", ")")


def xlsx_text(data):
    book = load_workbook(io.BytesIO(data))
    return "\n".join(str(cell.value) for sheet in book for row in sheet.iter_rows()
                     for cell in row if cell.value is not None)


def mark():
    return q1("SELECT COALESCE(MAX(id), 0) FROM audit_logs")[0]


def entries_since(start):
    return [dict(r) for r in qall("SELECT * FROM audit_logs WHERE id > ? ORDER BY id", (start,))]


def day(offset):
    return (today() + timedelta(days=offset)).isoformat()


def new_item(generic, strength="100 mg", form="tablet", category="Medicine", threshold=10):
    name = inv.build_display_name(category, generic, strength, form)
    return run_sql(
        "INSERT INTO inventory_items (item_code, name, generic_name, category, drug_class, "
        "dosage_strength, dosage_form, unit, low_stock_threshold, status) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, 'Tablet', ?, 'Active')",
        (f"TST-{generic[:6].upper()}-{strength.replace(' ', '')}", name, generic, category,
         "Antibiotic" if category == "Medicine" else None, strength or None,
         form if category == "Medicine" else None, threshold))


def add_batch(item_id, lot, expiry_offset, quantity, received_offset=0, brand=None,
              source="DOH"):
    return run_sql(
        "INSERT INTO inventory_batches (inventory_item_id, lot_number, expiry_date, "
        "quantity_received, quantity_remaining, date_received, source, brand_name) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        (item_id, lot, None if expiry_offset is None else day(expiry_offset), quantity, quantity,
         day(received_offset), source, brand))


def remaining(batch_id):
    return q1("SELECT quantity_remaining FROM inventory_batches WHERE id = ?", (batch_id,))[0]


def usable_stock(item_id):
    with app.app_context():
        return inv.get_inventory_item_by_id(item_id)["stock_quantity"]


def dispense(item_id, quantity, consultation_id=None):
    with app.app_context():
        return record_dispensing(
            patient_id=8, entries=[{"index": 0, "item_id": item_id, "quantity": quantity}],
            dispensed_by=STAFF_ID, date_dispensed=today(), consultation_id=consultation_id)


def links_for(dispensing_id):
    return [tuple(r) for r in qall(
        "SELECT batch_id, quantity FROM medicine_dispensing_batches WHERE dispensing_id = ? "
        "ORDER BY id", (dispensing_id,))]


def last_dispensing_id():
    return q1("SELECT MAX(id) FROM medicine_dispensing")[0]


staff, doctor, admin, official = (client_as(STAFF_ID), client_as(DOCTOR_ID),
                                  client_as(ADMIN_ID), client_as(OFFICIAL_ID))

# ===========================================================================
section("1. THE ITEM FORM: NO SEPARATE MEDICINE NAME, THE NAME IS BUILT, NO DUPLICATES")
page = body(staff.get("/staff/add_edit_medicine"))
check("the add form has no separate 'Medicine name' field",
      'name="name"' not in page and "Medicine Name" not in page and "Medicine name" not in page)
check("it asks for generic name, strength and dosage form, in that order",
      0 < page.index('name="generic_name"') < page.index('name="dosage_strength"')
      < page.index('name="dosage_form"'))
check("the brand name is not on the item form (it belongs to the batch)",
      'name="brand_name"' not in page)
check("the dosage form list is the one constant: every form is offered and nothing else",
      inv.DOSAGE_FORMS == ("tablet", "capsule", "syrup", "suspension", "drops", "sachet",
                           "injection", "cream", "ointment", "inhaler", "other")
      and all(f'<option value="{form}"' in page for form in inv.DOSAGE_FORMS)
      and len(re.findall(r'<option value="[a-z]+"[^>]*>[a-z]+</option>',
                         page.split('name="dosage_form"')[1].split("</select>")[0]))
      == len(inv.DOSAGE_FORMS))
check("the form shows the name it will build", 'id="namePreview"' in page)
item_form = {
    "category": "Medicine", "generic_name": "Ibuprofen", "dosage_strength": "200mg",
    "dosage_form": "tablet", "item_name": "", "size": "", "drug_class": inv.DRUG_CLASSES[1],
    "unit": "Tablet", "low_stock_threshold": "30", "manufacturer": "", "supplier": "",
    "unit_cost": "", "storage_location": "", "notes": "", "status": "Active",
}
before = qall("SELECT COUNT(*) FROM inventory_items")[0][0]
staff.post("/staff/add_edit_medicine", data=dict(item_form, name="Advil Forte"))
added = q1("SELECT * FROM inventory_items WHERE generic_name = 'Ibuprofen'")
check("the display name is built from generic name, strength and form, strength spaced",
      added is not None and added["name"] == "Ibuprofen 200 mg tablet"
      and added["dosage_strength"] == "200 mg" and added["dosage_form"] == "tablet", added and dict(added))
check("a stray 'name' field in the POST is ignored",
      added is not None and "Advil" not in added["name"])
check("adding an item adds no stock: it has no batch until stock is received",
      added is not None and q1("SELECT COUNT(*) FROM inventory_batches WHERE inventory_item_id = ?",
                               (added["id"],))[0] == 0)
check("the item table has no quantity, expiry, batch or brand column",
      not {"stock_quantity", "expiry_date", "batch_number", "date_received", "brand_name"}
      & {r[1] for r in qall("PRAGMA table_info(inventory_items)")})
response = staff.post("/staff/add_edit_medicine", data=dict(
    item_form, generic_name="ibuprofen", dosage_strength="200 MG", dosage_form="tablet"))
check("the same generic name, strength and form is refused as a duplicate (any case, any spacing)",
      "already in the inventory" in body(response) and "Ibuprofen 200 mg tablet" in body(response)
      and qall("SELECT COUNT(*) FROM inventory_items")[0][0] == before + 1)
staff.post("/staff/add_edit_medicine", data=dict(item_form, dosage_strength="400 mg"))
staff.post("/staff/add_edit_medicine", data=dict(item_form, dosage_form="capsule"))
check("a different strength, or a different form, is a different item",
      q1("SELECT COUNT(*) FROM inventory_items WHERE generic_name = 'Ibuprofen'")[0] == 3)
for field, bad, label in (("generic_name", "", "generic name"), ("dosage_strength", "", "strength"),
                          ("dosage_form", "", "dosage form"), ("dosage_form", "lozenge", "listed form")):
    response = staff.post("/staff/add_edit_medicine", data=dict(item_form, **{field: bad}))
    check(f"a medicine with no {label} is refused" if bad == "" else "a form that is not on the list is refused",
          "was not saved" in body(response)
          and qall("SELECT COUNT(*) FROM inventory_items WHERE generic_name = 'Ibuprofen'")[0][0] == 3)
supply_form = dict(item_form, category="Supply", generic_name="", dosage_strength="", dosage_form="",
                   drug_class="", item_name="Gauze pad", size="4x4 in")
staff.post("/staff/add_edit_medicine", data=supply_form)
supply = q1("SELECT * FROM inventory_items WHERE category = 'Supply' AND generic_name = 'Gauze pad'")
check("a supply keeps a simple item name and size; no generic name, strength or form is asked",
      supply is not None and supply["name"] == "Gauze pad 4x4 in" and supply["dosage_form"] is None,
      supply and dict(supply))
response = staff.post("/staff/add_edit_medicine", data=supply_form)
check("a duplicate supply is refused too",
      "already in the inventory" in body(response)
      and q1("SELECT COUNT(*) FROM inventory_items WHERE generic_name = 'Gauze pad'")[0] == 1)
response = staff.post("/staff/add_edit_medicine", data=dict(supply_form, item_name=""))
check("a supply with no item name is refused", "was not saved" in body(response))

# ===========================================================================
section("2. RECEIVE STOCK ADDS A BATCH WITH ITS OWN EXPIRY; EDIT DETAILS NEVER TOUCHES A QUANTITY")
item_id = added["id"]
receive = {"lot_number": "IBU-001", "expiry_date": day(400), "quantity": "100",
           "date_received": day(-5), "source": "DOH", "brand_name": "Febrex"}
staff.post(f"/staff/medicine_inventory/{item_id}/receive", data=receive)
staff.post(f"/staff/medicine_inventory/{item_id}/receive",
           data=dict(receive, lot_number="IBU-002", expiry_date=day(700), quantity="50",
                     source="Donation", brand_name="Paracare"))
batches = qall("SELECT * FROM inventory_batches WHERE inventory_item_id = ? ORDER BY id", (item_id,))
check("two deliveries are two batches, each with its own lot, expiry, brand, source and date",
      len(batches) == 2 and [b["lot_number"] for b in batches] == ["IBU-001", "IBU-002"]
      and [b["expiry_date"] for b in batches] == [day(400), day(700)]
      and [b["brand_name"] for b in batches] == ["Febrex", "Paracare"]
      and [b["source"] for b in batches] == ["DOH", "Donation"]
      and batches[0]["date_received"] == day(-5) and batches[0]["quantity_remaining"] == 100)
check("the second delivery did not change the first batch's expiry or quantity",
      batches[0]["expiry_date"] == day(400) and batches[0]["quantity_remaining"] == 100)
check("the item's stock is the sum of its usable batches", usable_stock(item_id) == 150)
check("the brand is stored on the batch, and the item has none",
      "brand_name" not in dict(q1("SELECT * FROM inventory_items WHERE id = ?", (item_id,))).keys())
check("the item page lists both batches with lot, brand and source",
      all(s in flat(body(staff.get(f"/staff/medicine_inventory/{item_id}")))
          for s in ("IBU-001", "IBU-002", "Febrex", "Paracare", "Donation")))
r = staff.post(f"/staff/medicine_inventory/{item_id}/receive", data=dict(receive, lot_number=""))
check("a lot number is required", "Enter the lot number" in body(r)
      and q1("SELECT COUNT(*) FROM inventory_batches WHERE inventory_item_id = ?", (item_id,))[0] == 2)
staff.post(f"/staff/medicine_inventory/{item_id}/receive",
           data=dict(receive, lot_number="", no_lot="1", quantity="5"))
nolot = q1("SELECT * FROM inventory_batches WHERE inventory_item_id = ? AND quantity_received = 5",
           (item_id,))
check("'No lot number' is allowed, only with the tick box",
      nolot is not None and nolot["lot_number"] == "No lot number")
for label, change, expected in (
        ("an expiry date is required for a medicine", {"expiry_date": ""}, "Expiry date is required"),
        ("an expiry date in the past is refused", {"expiry_date": day(-1)}, "already passed"),
        ("a bad expiry date is refused", {"expiry_date": "31-12-2030"}, "valid date"),
        ("a quantity of 0 is refused", {"quantity": "0"}, "whole number"),
        ("a quantity that is not a number is refused", {"quantity": "ten"}, "whole number"),
        ("a date received in the future is refused", {"date_received": day(3)}, "future"),
        ("a source that is not on the list is refused", {"source": "A stranger"}, "from the list"),
        ("no source is refused", {"source": ""}, "Source is required")):
    count = q1("SELECT COUNT(*) FROM inventory_batches")[0]
    r = staff.post(f"/staff/medicine_inventory/{item_id}/receive", data=dict(receive, **change))
    check(label, expected in body(r) and q1("SELECT COUNT(*) FROM inventory_batches")[0] == count,
          flat(body(r))[:200])
check("every listed source is accepted",
      inv.STOCK_SOURCES == ("DOH", "Municipal Health Office / RHU", "Barangay purchase",
                            "Donation", "Other"))
staff.post(f"/staff/medicine_inventory/{item_id}/receive",
           data=dict(receive, lot_number="IBU-003", date_received="", quantity="7"))
check("a blank date received is today",
      q1("SELECT date_received FROM inventory_batches WHERE lot_number = 'IBU-003'")[0] == day(0))
supply_id = supply["id"]
staff.post(f"/staff/medicine_inventory/{supply_id}/receive",
           data=dict(receive, lot_number="G-1", expiry_date="", quantity="40"))
check("an expiry date is optional for a supply",
      q1("SELECT expiry_date FROM inventory_batches WHERE lot_number = 'G-1'")[0] is None)

stock_before = usable_stock(item_id)
batches_before = [tuple(r) for r in qall(
    "SELECT id, lot_number, expiry_date, quantity_received, quantity_remaining "
    "FROM inventory_batches WHERE inventory_item_id = ? ORDER BY id", (item_id,))]
start = mark()
r = staff.post(f"/staff/add_edit_medicine/{item_id}", data=dict(
    item_form, manufacturer="Demo Pharma", stock_quantity="9999", quantity="9999",
    expiry_date=day(5), batch_number="HACK"))
check("Edit details saves the item fields",
      q1("SELECT manufacturer FROM inventory_items WHERE id = ?", (item_id,))[0] == "Demo Pharma")
check("Edit details never changes a quantity or an expiry, even when a quantity is posted",
      usable_stock(item_id) == stock_before
      and batches_before == [tuple(r) for r in qall(
          "SELECT id, lot_number, expiry_date, quantity_received, quantity_remaining "
          "FROM inventory_batches WHERE inventory_item_id = ? ORDER BY id", (item_id,))])
edit_rows = entries_since(start)
check("Edit details writes one entry that names the field, not a quantity",
      len(edit_rows) == 1 and edit_rows[0]["action"] == "Edited inventory item"
      and edit_rows[0]["details"].endswith("— Changed: Manufacturer")
      and "Stock" not in edit_rows[0]["details"], edit_rows)
r = staff.post(f"/staff/add_edit_medicine/{item_id}",
               data=dict(item_form, generic_name="Ibuprofen", dosage_strength="400 mg"))
check("Edit details cannot turn an item into a duplicate of another",
      "already in the inventory" in body(r)
      and q1("SELECT dosage_strength FROM inventory_items WHERE id = ?", (item_id,))[0] == "200 mg")
staff.post(f"/staff/add_edit_medicine/{item_id}",
           data=dict(item_form, manufacturer="Demo Pharma", dosage_strength="250 mg"))
check("renaming the strength rebuilds the display name",
      q1("SELECT name FROM inventory_items WHERE id = ?", (item_id,))[0] == "Ibuprofen 250 mg tablet")
staff.post(f"/staff/add_edit_medicine/{item_id}",
           data=dict(item_form, manufacturer="Demo Pharma", dosage_strength="200 mg"))

# ===========================================================================
section("3. FEFO: FIRST EXPIRY, FIRST OUT")
fefo = new_item("Fefotest", "10 mg")
late = add_batch(fefo, "LATE", 300, 50, received_offset=-30)
early = add_batch(fefo, "EARLY", 100, 10, received_offset=-5)
none_b = add_batch(fefo, "NOEXP", None, 20, received_offset=-60)
stored, errors = dispense(fefo, 4)
did = last_dispensing_id()
check("dispensing takes from the earliest expiry, even if it arrived later",
      not errors and remaining(early) == 6 and remaining(late) == 50 and remaining(none_b) == 20
      and links_for(did) == [(early, 4)], (errors, links_for(did)))
stored, errors = dispense(fefo, 10)
did = last_dispensing_id()
check("when the earliest batch runs out it splits across batches, earliest first",
      remaining(early) == 0 and remaining(late) == 46 and links_for(did) == [(early, 6), (late, 4)],
      links_for(did))
check("the batches used are recorded and add up to the amount dispensed",
      sum(qty for _, qty in links_for(did)) == q1(
          "SELECT quantity_dispensed FROM medicine_dispensing WHERE id = ?", (did,))[0] == 10)
dispense(fefo, 46)
check("a batch with no expiry date is used last",
      remaining(late) == 0 and remaining(none_b) == 20 and links_for(last_dispensing_id()) == [(late, 46)])
check("the dispensing row names the item, not a batch (staff do not pick batches)",
      q1("SELECT inventory_item_id FROM medicine_dispensing WHERE id = ?", (last_dispensing_id(),))[0] == fefo)
html = body(doctor.get("/staff/consultation/record/%d" % (run_sql(
    "INSERT INTO appointments (patient_id, scheduled_by, appointment_type, appointment_date, "
    "appointment_time, status, is_walkin, checked_in_at) VALUES (16, 1, 'Consultation', ?, '09:00', "
    "'In Consultation', 1, datetime('now'))", (day(0),)))))
check("the consultation form offers no batch choice",
      "batch" not in html.split('id="dispenseRows"')[1].split("</tbody>")[0].lower())

exp = new_item("Expiredtest", "20 mg")
dead = add_batch(exp, "DEAD", -3, 500, received_offset=-200)
live = add_batch(exp, "LIVE", 200, 30, received_offset=-10)
dispense(exp, 12)
check("an expired batch is never used, even when it expires first and holds more",
      remaining(dead) == 500 and remaining(live) == 18 and links_for(last_dispensing_id()) == [(live, 12)])
check("expired stock is not counted: stock is the usable batches only", usable_stock(exp) == 18)
before_id = last_dispensing_id()
stored, errors = dispense(exp, 19)
check("when usable stock is short, dispensing is refused with one plain message showing the amount available",
      not stored and len(errors) == 1 and "Only 18 tablets" in list(errors.values())[0]
      and "available" in list(errors.values())[0], errors)
check("a refusal changes nothing: no dispensing row, no batch moved",
      last_dispensing_id() == before_id and remaining(dead) == 500 and remaining(live) == 18)
only_dead = new_item("Onlydead", "30 mg")
dead2 = add_batch(only_dead, "OLD", -1, 80)
stored, errors = dispense(only_dead, 1)
check("an item whose only stock is expired is refused as having no usable stock",
      not stored and "no usable stock" in " ".join(errors.values()) and remaining(dead2) == 80
      and usable_stock(only_dead) == 0)
with app.app_context():
    from models.dispensing import dispensable_items
    offered = {entry["id"] for entry in dispensable_items()}
check("the form does not offer an item with no usable stock, and does offer one that has some",
      only_dead not in offered and exp in offered)
today_batch = new_item("Expirestoday", "40 mg")
tb = add_batch(today_batch, "TODAY", 0, 10)
stored, errors = dispense(today_batch, 3)
check("a batch is usable through the whole of its expiry date", not errors and remaining(tb) == 7)

# end to end, through the form, as the Doctor
e2e = new_item("Endtoend", "5 mg")
e_early = add_batch(e2e, "E-EARLY", 120, 5)
e_late = add_batch(e2e, "E-LATE", 365, 40)
appointment = run_sql(
    "INSERT INTO appointments (patient_id, scheduled_by, appointment_type, appointment_date, "
    "appointment_time, status, is_walkin, checked_in_at) VALUES (12, 1, 'Consultation', ?, '09:00', "
    "'In Consultation', 1, datetime('now'))", (day(0),))
form = {"chief_complaint": "Cough", "diagnosis": "Cough", "treatment_plan": "Rest",
        "temperature": "37", "blood_pressure": "110/70", "pulse_rate": "80",
        "respiratory_rate": "18", "weight": "60", "remarks": "",
        "dispense_item_0": str(e2e), "dispense_qty_0": "9"}
r = doctor.post("/staff/consultation/record/%d" % appointment, data=form, follow_redirects=True)
check("the Doctor dispenses through the form: 9 units come from two batches, earliest first",
      remaining(e_early) == 0 and remaining(e_late) == 36
      and links_for(last_dispensing_id()) == [(e_early, 5), (e_late, 4)], flat(body(r))[:300])

# ===========================================================================
section("4. REMOVE FROM STOCK: A REASON, ONE BATCH, AN AUDIT ENTRY")
rem = new_item("Removetest", "50 mg")
ra = add_batch(rem, "R-A", 100, 40)
rb = add_batch(rem, "R-B", 200, 25)
rc = add_batch(rem, "R-C", -2, 15, received_offset=-100)
start = mark()
r = staff.post(f"/staff/medicine_inventory/{rem}/remove",
               data={"batch_id": str(ra), "quantity": "5", "reason": ""})
check("a reason is required", remaining(ra) == 40 and not entries_since(start) and "reason" in body(
    staff.get(f"/staff/medicine_inventory/{rem}")).lower())
r = staff.post(f"/staff/medicine_inventory/{rem}/remove",
               data={"batch_id": str(ra), "quantity": "5", "reason": "Because"}, follow_redirects=True)
check("a reason that is not on the list is refused", remaining(ra) == 40 and not entries_since(start))
check("the reasons are Expired, Damaged, Lost, Returned and Count correction",
      inv.REMOVAL_REASONS == ("Expired", "Damaged", "Lost", "Returned", "Count correction"))
r = staff.post(f"/staff/medicine_inventory/{rem}/remove",
               data={"batch_id": str(ra), "quantity": "41", "reason": "Lost"}, follow_redirects=True)
check("more than the batch holds is refused", remaining(ra) == 40 and "holds only 40" in flat(body(r))
      and not entries_since(start))
for value in ("0", "-3", "x"):
    staff.post(f"/staff/medicine_inventory/{rem}/remove",
               data={"batch_id": str(ra), "quantity": value, "reason": "Lost"})
check("an amount of 0 or less, or not a number, is refused",
      remaining(ra) == 40 and not entries_since(start))
r = staff.post(f"/staff/medicine_inventory/{rem}/remove",
               data={"batch_id": str(ra), "quantity": "7", "reason": "Damaged",
                     "note": "Zznote dropped"}, follow_redirects=True)
check("it takes the amount out of the ONE batch chosen, and no other",
      remaining(ra) == 33 and remaining(rb) == 25 and remaining(rc) == 15, (remaining(ra), remaining(rb)))
removal = q1("SELECT * FROM inventory_removals WHERE batch_id = ?", (ra,))
check("the removal keeps its reason and its optional note, and who did it",
      removal["quantity"] == 7 and removal["reason"] == "Damaged" and removal["note"] == "Zznote dropped"
      and removal["removed_by"] == STAFF_ID)
rows = entries_since(start)
item_name = q1("SELECT name || ' (' || item_code || ')' FROM inventory_items WHERE id = ?", (rem,))[0]
check("one short audit entry: who, the item name, the lot number and the amount (and the reason)",
      len(rows) == 1 and rows[0]["action"] == "Removed stock from a batch"
      and rows[0]["module"] == "Medicine Inventory" and rows[0]["user_id"] == STAFF_ID
      and rows[0]["details"] == f"{item_name}, lot R-A: removed 7 (Damaged).", rows)
check("the note is not copied into the audit entry", "Zznote" not in rows[0]["details"])
# The routes check the reason and the batch too; the model must hold them on its own, because
# the database CHECK and the route are not the rule itself.
with app.app_context():
    other_batch = rb
    removed_none, errs = inv.remove_from_batch(rem, ra, {"quantity": "1", "reason": ""})
    check("the model itself refuses a removal with no reason, whatever the route does",
          removed_none is None and "reason" in errs and remaining(ra) == 33, errs)
    removed_none, errs = inv.remove_from_batch(rem, ra, {"quantity": "1", "reason": "Because"})
    check("and one with a reason that is not on the list",
          removed_none is None and "reason" in errs and remaining(ra) == 33, errs)
    wrong_item = new_item("Wrongitem", "5 mg")
    removed_none, errs = inv.remove_from_batch(wrong_item, ra, {"quantity": "1", "reason": "Lost"})
    check("and a batch that belongs to another item, however the route is called",
          removed_none is None and "batch" in errs and remaining(ra) == 33, errs)
start = mark()
staff.post(f"/staff/medicine_inventory/{rem}/remove",
           data={"batch_id": str(rc), "quantity": "15", "reason": "Expired"})
check("an expired batch can be removed completely, and then reads Empty",
      remaining(rc) == 0 and "Empty" in flat(body(staff.get(f"/staff/medicine_inventory/{rem}"))))
check("removing from an expired batch does not change usable stock", usable_stock(rem) == 33 + 25)
staff.post(f"/staff/medicine_inventory/{rem}/remove",
           data={"batch_id": str(q1("SELECT MAX(id) FROM inventory_batches")[0]), "quantity": "1",
                 "reason": "Lost"})
check("a batch of another item is refused for this item",
      remaining(ra) == 33 and q1("SELECT COUNT(*) FROM inventory_removals")[0] == 2)
start = mark()
staff.post(f"/staff/medicine_inventory/{rem}/receive",
           data={"lot_number": "R-D", "expiry_date": day(500), "quantity": "11",
                 "date_received": day(0), "source": "Barangay purchase", "brand_name": ""})
rows = entries_since(start)
check("Receive stock writes one entry: the item, the lot and the amount",
      len(rows) == 1 and rows[0]["action"] == "Received stock"
      and rows[0]["details"] == f"{item_name}, lot R-D: received 11.", rows)

# ===========================================================================
section("5. STOCK, THE LOW-STOCK ALERT AND THE EXPIRY COUNTS USE USABLE BATCHES ONLY")
check("the expiring-soon window is one constant in config.py: 90 days",
      config.EXPIRING_SOON_DAYS == 90 and inv.EXPIRING_SOON_DAYS == 90
      and not hasattr(inv, "NEAR_EXPIRY_DAYS"))
with app.app_context():
    status = {days: inv.batch_status(day(days), 5) for days in (-1, 0, 89, 90, 91)}
check("a batch is Expired the day after its date, Expiring soon from today to 90 days, "
      "Usable after that, and Empty when nothing is left",
      status == {-1: "Expired", 0: "Expiring soon", 89: "Expiring soon", 90: "Expiring soon",
                 91: "Usable"} and inv.batch_status(day(-9), 0) == "Empty"
      and inv.batch_status(None, 5) == "Usable", status)
soon_item = new_item("Soontest", "15 mg")
add_batch(soon_item, "S-SOON", 45, 12)
add_batch(soon_item, "S-LATER", 400, 30)
low = new_item("Lowtest", "60 mg", threshold=50)
add_batch(low, "L-DEAD", -10, 500, received_offset=-300)
add_batch(low, "L-OK", 200, 20)
with app.app_context():
    alert = inv.get_low_stock_alert()
    low_item = inv.get_inventory_item_by_id(low)
check("500 expired units do not keep an item off the low-stock alert",
      low in {i["id"] for i in alert["low_stock_items"]} and low_item["stock_quantity"] == 20
      and low_item["is_low_stock"] and low_item["expired_batches"] == 1
      and low_item["expired_quantity"] == 500)
out = new_item("Outtest", "70 mg", threshold=5)
add_batch(out, "O-DEAD", -1, 90, received_offset=-300)
with app.app_context():
    out_item = inv.get_inventory_item_by_id(out)
check("an item holding only expired stock is Out of Stock",
      out_item["stock_quantity"] == 0 and out_item["is_out_of_stock"]
      and out_item["stock_status"] == "Out of Stock")
check("the item shows its nearest expiry among the usable batches",
      inv.get_inventory_item_by_id(low)["nearest_expiry"] == day(200)
      and inv.get_inventory_item_by_id(fefo)["nearest_expiry"] is None)

with app.app_context():
    items = inv.get_all_inventory_items()
    stats = inv.get_inventory_stats(items)
    soon_by_hand = sum(1 for i in items if not i["is_retired"] for b in i["batches"]
                       if b["quantity_remaining"] > 0 and b["expiry_date"]
                       and 0 <= (inv._parse_date(b["expiry_date"]) - today()).days <= 90)
    expired_by_hand = sum(1 for i in items if not i["is_retired"] for b in i["batches"]
                          if b["quantity_remaining"] > 0 and b["expiry_date"]
                          and inv._parse_date(b["expiry_date"]) < today())
check("the stat figures count batches with stock: expiring soon and expired",
      stats["expiring_soon_count"] == soon_by_hand and stats["expired_count"] == expired_by_hand
      and soon_by_hand >= 2 and expired_by_hand >= 3, (stats, soon_by_hand, expired_by_hand))
SOON, EXPIRED = stats["expiring_soon_count"], stats["expired_count"]

for who, client, prefix in (("Administrator", admin, "/admin"), ("Doctor", doctor, "/staff"),
                            ("Health Center Staff", staff, "/staff")):
    page = body(client.get(f"{prefix}/medicine_inventory"))
    check(f"{who}: the inventory screen shows the expiring-soon and expired counts",
          f'id="statExpiringSoon">{SOON}<' in page and f'id="statExpired">{EXPIRED}<' in page
          and "Expiring Soon (90 days)" in page and "Expired Batches" in page)
    row = re.search(r'data-sort-name="lowtest 60 mg tablet".*?</tr>', page, re.S)
    check(f"{who}: the list shows usable stock, nearest expiry and the expired batch count per item",
          row is not None and ">20 <" in row.group(0) and 'badge-expired">1<' in row.group(0)
          and display_date(day(200)) in row.group(0), row and row.group(0)[:300])
    dash = flat(body(client.get(f"{prefix}/dashboard")))
    check(f"{who}: the dashboard adds Expiring soon and Expired counts",
          f"Expiring soon: {SOON} · Expired: {EXPIRED}" in dash, dash[dash.find("Expiring soon") - 30:][:120])

report = flat(body(admin.get("/admin/reports_analytics")))
check("the report page lists batches expiring soon and expired, with item, lot and expiry",
      "Batches Expiring Soon and Expired" in report and "L-DEAD" in report and "O-DEAD" in report
      and "Lowtest 60 mg tablet" in report and f"Expiring soon: {SOON}" in report
      and f"Expired: {EXPIRED}" in report)
check("the Doctor sees the same batches in the report",
      "L-DEAD" in flat(body(doctor.get("/staff/reports_analytics"))))
xl = xlsx_text(admin.get("/admin/reports_analytics/export.xlsx").data)
check("the report's Excel file lists the batches (item, lot, expiry, amount) and the counts",
      all(s in xl for s in ("Batches expiring soon and expired", "L-DEAD", "O-DEAD",
                            "Lowtest 60 mg tablet", "Expired batches holding stock")), xl[:300])
pdf = pdf_text(admin.get("/admin/reports_analytics/export.pdf").data)
check("the report's PDF lists the batches too", "L-DEAD" in pdf and "O-DEAD" in pdf, pdf[:200])
inv_xl = xlsx_text(admin.get("/admin/medicine_inventory/export.xlsx").data)
check("the inventory list's Excel file has the new columns and a Batches sheet",
      all(s in inv_xl for s in ("Usable Stock", "Nearest Expiry", "Batches Expiring Soon",
                                "Expired Batches", "L-DEAD")), inv_xl[:200])
check("the inventory list's PDF works", admin.get("/admin/medicine_inventory/export.pdf").data[:4] == b"%PDF")
with app.app_context():
    stats_after_retire = None
run_sql("UPDATE inventory_items SET status = 'Inactive' WHERE id = ?", (low,))
with app.app_context():
    stats_after_retire = inv.get_inventory_stats()
    still_listed = inv.get_inventory_item_by_id(low)
check("a retired item is left out of the expired and low-stock figures but stays listed",
      stats_after_retire["expired_count"] == EXPIRED - 1 and low not in {
          i["id"] for i in inv.get_low_stock_alert()["low_stock_items"]} and still_listed is not None)
run_sql("UPDATE inventory_items SET status = 'Active' WHERE id = ?", (low,))

# ===========================================================================
section("6. ACCESS IS UNCHANGED, AND THE ADMINISTRATOR NEVER SEES WHO RECEIVED WHAT")
target = new_item("Accesstest", "80 mg")
good = {"lot_number": "ACC-1", "expiry_date": day(300), "quantity": "9", "date_received": day(0),
        "source": "Other", "brand_name": ""}
for who, client, prefix, allowed in (("Administrator", admin, "/admin", True),
                                     ("Health Center Staff", staff, "/staff", True),
                                     ("Doctor", doctor, "/staff", True),
                                     ("Barangay Official", official, "/official", False),
                                     ("Barangay Official (via /staff)", official, "/staff", False),
                                     ("Administrator (via /staff)", admin, "/staff", False)):
    count = q1("SELECT COUNT(*) FROM inventory_batches WHERE inventory_item_id = ?", (target,))[0]
    page_status = client.get(f"{prefix}/medicine_inventory/{target}").status_code
    client.post(f"{prefix}/medicine_inventory/{target}/receive", data=dict(good, lot_number=f"A-{who[:5]}"))
    received = q1("SELECT COUNT(*) FROM inventory_batches WHERE inventory_item_id = ?", (target,))[0] - count
    if allowed:
        batch_id = q1("SELECT MAX(id) FROM inventory_batches WHERE inventory_item_id = ?", (target,))[0]
        client.post(f"{prefix}/medicine_inventory/{target}/remove",
                    data={"batch_id": str(batch_id), "quantity": "1", "reason": "Lost"})
        removed = remaining(batch_id) == 8
    else:
        removed = None
    check(f"{who}: " + ("can open the item page, receive stock and remove from stock"
                        if allowed else "is refused: no page, no stock received"),
          (page_status == 200 and received == 1 and removed) if allowed
          else (page_status in (302, 303, 404) and received == 0), (page_status, received, removed))
import flask.testing  # noqa: E402

removals_before_token = q1("SELECT COUNT(*) FROM inventory_removals")[0]
plain = app.test_client()
with plain.session_transaction() as sess:
    sess["user_id"] = STAFF_ID
    _support.open_session(sess)
count = q1("SELECT COUNT(*) FROM inventory_batches")[0]
flask.testing.FlaskClient.open(plain, f"/staff/medicine_inventory/{target}/receive",
                               method="POST", data=good)
flask.testing.FlaskClient.open(plain, f"/staff/medicine_inventory/{target}/remove",
                               method="POST", data={"batch_id": "1", "quantity": "1", "reason": "Lost"})
check("a POST without the form token is refused, for receiving and for removing",
      q1("SELECT COUNT(*) FROM inventory_batches")[0] == count
      and q1("SELECT COUNT(*) FROM inventory_removals")[0] == removals_before_token)

patient = q1("SELECT patient_id, first_name, last_name FROM patients WHERE id = 8")
dispense(target, 2, consultation_id=None)
secrets_ = [patient["patient_id"], patient["first_name"], patient["last_name"]]
pages = {
    "the item page": body(admin.get(f"/admin/medicine_inventory/{target}")),
    "the inventory list": body(admin.get("/admin/medicine_inventory")),
    "the report": body(admin.get("/admin/reports_analytics")),
    "the audit logs": body(admin.get("/admin/audit_logs")),
    "the inventory Excel file": xlsx_text(admin.get("/admin/medicine_inventory/export.xlsx").data),
    "the report Excel file": xlsx_text(admin.get("/admin/reports_analytics/export.xlsx").data),
}
for label, text in pages.items():
    check(f"the Administrator's {label} names no patient and no dispensing",
          not any(s in text for s in secrets_) and "Dispensed to" not in text)
check("the Administrator has no route to the batches a dispensing used",
      all(admin.get(url).status_code in (302, 303, 404) for url in (
          "/admin/patient_profile/8/dispensing", "/staff/patient_profile/8", "/admin/dispensing")))
check("the Administrator's audit view of the receipts and removals names no patient",
      not any(s in " ".join(e["details"] for e in entries_since(0)
                            if e["action"] in ("Received stock", "Removed stock from a batch"))
              for s in secrets_))

# ===========================================================================
section("7. THE MIGRATION")
sys.path.insert(0, str(_support.REPO / "migrations"))
import migrate_inventory_batches as mig  # noqa: E402

work = Path(_support.work_path("migration-scratch"))
shutil.rmtree(work, ignore_errors=True)
work.mkdir(parents=True)
old = str(work / "old.db")
shutil.copy(TEST_DB, old)
c = sqlite3.connect(old)
c.execute("PRAGMA foreign_keys = OFF")
for table in ("medicine_dispensing_batches", "inventory_removals", "inventory_batches",
              "medicine_dispensing", "inventory_items"):
    c.execute(f"DROP TABLE {table}")
c.executescript("""
CREATE TABLE inventory_items (
    id INTEGER PRIMARY KEY AUTOINCREMENT, item_code TEXT NOT NULL UNIQUE, name TEXT NOT NULL,
    generic_name TEXT, brand_name TEXT,
    category TEXT NOT NULL CHECK (category IN ('Medicine', 'Supply', 'Equipment')),
    drug_class TEXT, dosage_strength TEXT, dosage_form TEXT, unit TEXT,
    stock_quantity INTEGER NOT NULL DEFAULT 0, low_stock_threshold INTEGER NOT NULL DEFAULT 0,
    expiry_date TEXT, date_received TEXT, batch_number TEXT, manufacturer TEXT, supplier TEXT,
    unit_cost REAL, storage_location TEXT, notes TEXT,
    status TEXT NOT NULL DEFAULT 'Active' CHECK (status IN ('Active', 'Inactive')),
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
CREATE INDEX idx_inventory_low_stock ON inventory_items(stock_quantity, low_stock_threshold);
CREATE TABLE medicine_dispensing (
    id INTEGER PRIMARY KEY AUTOINCREMENT, patient_id INTEGER NOT NULL, consultation_id INTEGER,
    inventory_item_id INTEGER NOT NULL, medicine_name TEXT NOT NULL, unit TEXT,
    quantity_dispensed INTEGER NOT NULL CHECK (quantity_dispensed > 0),
    date_dispensed TEXT NOT NULL, dispensed_by INTEGER NOT NULL,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    UNIQUE (consultation_id, inventory_item_id),
    FOREIGN KEY (patient_id) REFERENCES patients(id) ON DELETE CASCADE,
    FOREIGN KEY (inventory_item_id) REFERENCES inventory_items(id) ON DELETE RESTRICT,
    FOREIGN KEY (dispensed_by) REFERENCES users(id) ON DELETE RESTRICT);
""")
OLD_ITEMS = [
    # code, name, generic, brand, category, strength, form, stock, expiry, batch no
    ("OLD-1", "Paracetamol 500mg", "Paracetamol", "Biobrand", "Medicine", "500mg", "Tablet", 120, "2030-05-31", "PX-1"),
    ("OLD-2", "Salbutamol Nebule 2.5mg", "Salbutamol sulfate", None, "Medicine", "2.5mg/2.5mL", "Nebule", 33, "2030-01-31", "SX-9"),
    ("OLD-3", "Cotrimoxazole 400/80mg Tablet", None, None, "Medicine", None, None, 70, "2029-09-30", None),
    ("OLD-4", "Mystery tonic", None, None, "Medicine", None, None, 8, "2029-01-31", None),
    ("OLD-5", "Cotton balls", None, None, "Supply", None, None, 300, None, None),
    ("OLD-6", "Digital thermometer", None, None, "Equipment", None, None, 4, None, None),
]
for code, name, generic, brand, category, strength, form, stock, expiry, lot in OLD_ITEMS:
    c.execute("INSERT INTO inventory_items (item_code, name, generic_name, brand_name, category, "
              "dosage_strength, dosage_form, unit, stock_quantity, low_stock_threshold, expiry_date, "
              "date_received, batch_number) VALUES (?, ?, ?, ?, ?, ?, ?, 'Tablet', ?, 10, ?, "
              "'2026-01-15', ?)", (code, name, generic, brand, category, strength, form, stock, expiry, lot))
for item_no, qty in ((1, 10), (1, 20), (2, 3), (4, 1)):
    c.execute("INSERT INTO medicine_dispensing (patient_id, inventory_item_id, medicine_name, unit, "
              "quantity_dispensed, date_dispensed, dispensed_by) VALUES (8, ?, 'old snapshot', 'tablet', ?, "
              "'2026-02-01', 2)", (item_no, qty))
c.commit()
stock_total = c.execute("SELECT SUM(stock_quantity) FROM inventory_items").fetchone()[0]
audit_before = c.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0]
c.close()
result = mig.run(old, work / "backups")
check("the migration backs the database up first", result["backup"] and Path(result["backup"]).exists()
      and "PRETASK_BACKUP_inventory-batches" in result["backup"])
check("counts: items parsed, items needing a fix, supplies, opening batches, dispensing rows linked",
      (result["parsed"], result["needs_fix"], result["supplies"], result["opening_batches"],
       result["dispensing_linked"]) == (3, 1, 2, 6, 4) and not result["problems"], result)
con = sqlite3.connect(old)
con.row_factory = sqlite3.Row
names = {r["item_code"]: r for r in con.execute("SELECT * FROM inventory_items")}
check("names are split into generic name, strength and form, and the display name is built",
      names["OLD-1"]["name"] == "Paracetamol 500 mg tablet"
      and (names["OLD-1"]["generic_name"], names["OLD-1"]["dosage_strength"],
           names["OLD-1"]["dosage_form"]) == ("Paracetamol", "500 mg", "tablet"))
check("a nebule becomes the form 'other', and the generic name already on file is used",
      names["OLD-2"]["name"] == "Salbutamol sulfate 2.5 mg/2.5 mL other"
      and names["OLD-2"]["dosage_form"] == "other")
check("a name with nothing on file but the name itself is parsed from the name",
      names["OLD-3"]["name"] == "Cotrimoxazole 400/80 mg tablet", names["OLD-3"]["name"])
check("a name that cannot be parsed is left as it was, to be fixed with Edit details",
      names["OLD-4"]["name"] == "Mystery tonic" and names["OLD-4"]["dosage_form"] is None)
check("a supply keeps its item name; equipment too",
      names["OLD-5"]["name"] == "Cotton balls" and names["OLD-5"]["generic_name"] == "Cotton balls"
      and names["OLD-6"]["name"] == "Digital thermometer")
opening = {r["item_code"]: r for r in con.execute(
    "SELECT i.item_code, b.* FROM inventory_batches b JOIN inventory_items i "
    "ON i.id = b.inventory_item_id")}
check("every item has one opening batch, lot 'Opening stock', with its old quantity and expiry",
      len(opening) == 6 and all(b["lot_number"] == "Opening stock" for b in opening.values())
      and opening["OLD-1"]["quantity_remaining"] == 120 and opening["OLD-1"]["expiry_date"] == "2030-05-31"
      and opening["OLD-5"]["expiry_date"] is None and opening["OLD-1"]["date_received"] == "2026-01-15"
      and opening["OLD-1"]["brand_name"] == "Biobrand" and opening["OLD-1"]["source"] == "Other")
check("no stock is lost or made up",
      con.execute("SELECT SUM(quantity_remaining) FROM inventory_batches").fetchone()[0] == stock_total)
links = con.execute("SELECT d.id, d.quantity_dispensed, l.batch_id, l.quantity "
                    "FROM medicine_dispensing d JOIN medicine_dispensing_batches l "
                    "ON l.dispensing_id = d.id ORDER BY d.id").fetchall()
check("every old dispensing row is linked to its item's opening batch, for the same amount",
      len(links) == 4 and all(l["quantity"] == l["quantity_dispensed"] for l in links)
      and links[0]["batch_id"] == links[1]["batch_id"] == opening["OLD-1"]["id"])
check("the five columns that moved to the batches are gone from the item table",
      not {"stock_quantity", "expiry_date", "batch_number", "date_received", "brand_name"}
      & {r[1] for r in con.execute("PRAGMA table_info(inventory_items)")})
check("exactly one audit entry records the run",
      con.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0] == audit_before + 1
      and con.execute("SELECT COUNT(*) FROM audit_logs WHERE action = 'Changed how stock is recorded'"
                      ).fetchone()[0] == 1)
con.close()
second = mig.run(old, work / "backups")
con = sqlite3.connect(old)
check("a second run does nothing: no backup, no change, no second record",
      second["backup"] is None and not second["problems"]
      and con.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0] == audit_before + 1
      and con.execute("SELECT COUNT(*) FROM inventory_batches").fetchone()[0] == 6
      and len(list((work / "backups").iterdir())) == 1)
con.close()
fresh = mig.run(TEST_DB, work / "backups-fresh")
check("a database already built from schema.sql needs nothing", fresh["backup"] is None)
shutil.rmtree(work, ignore_errors=True)
check("schema.sql matches the migrated tables: batches, removals and dispensing links exist",
      {"inventory_batches", "inventory_removals", "medicine_dispensing_batches"}
      <= {r[0] for r in qall("SELECT name FROM sqlite_master WHERE type='table'")})

# ---------------------------------------------------------------------------
passed = sum(1 for _, ok in results if ok)
print("\n" + "=" * 78)
print(f"RESULT: {passed}/{len(results)} checks passed")
failed = [name for name, ok in results if not ok]
for name in failed:
    print(f"  FAILED: {name}")
sys.exit(1 if failed else 0)
