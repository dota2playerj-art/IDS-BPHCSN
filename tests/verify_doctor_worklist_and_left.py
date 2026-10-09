"""The Doctor's list, "Left before being seen", the service at check-in and the status
"Waiting for the Doctor" (Task 6 and Task 6b; M4, M5, M33, L31).

M4. A visit whose intake is saved and that needs a diagnosis waits in its own status,
"Waiting for the Doctor". The Doctor's dashboard lists exactly those visits, oldest first,
across all days. Staff and the Administrator see the count only. One function decides.
"""
import io
import re
import sqlite3
import sys
import zlib
from datetime import date, datetime, timedelta

import _support  # noqa: E402  (first: builds the test database, points the app at it)

from app import create_app  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

sys.stdout.reconfigure(errors="replace")

app = create_app()
_support.attach(app)
TEST_DB = _support.DB_PATH
ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID = 1, 2, 3, 5

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok)))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}")
    if detail and not ok:
        for line in str(detail).splitlines()[:12]:
            print(f"       {line}")
    return ok


def section(title):
    print(f"\n--- {title} ---")


def db():
    conn = sqlite3.connect(TEST_DB)
    conn.row_factory = sqlite3.Row
    return conn


def q1(sql, args=()):
    conn = db()
    try:
        return conn.execute(sql, args).fetchone()
    finally:
        conn.close()


def run_sql(sql, args=()):
    conn = db()
    try:
        cur = conn.execute(sql, args)
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def client_as(user_id):
    c = app.test_client()
    with c.session_transaction() as sess:
        sess["user_id"] = user_id
        _support.open_session(sess)
    return c


def body(response):
    return response.get_data(as_text=True)


def flat(text):
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", text))


def pdf_text(data):
    out = []
    for match in re.finditer(rb"stream\r?\n(.*?)endstream", data, re.S):
        try:
            out.append(zlib.decompress(match.group(1)).decode("latin-1"))
        except Exception:
            out.append(match.group(1).decode("latin-1", "replace"))
    return "\n".join(out).replace("\\(", "(").replace("\\)", ")")


admin, staff, official, doctor = (client_as(i) for i in
                                  (ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID))
today = date.today()
LEFT = "Left before being seen"
MARK, UNDO = "/staff/appointment_scheduling/mark-left", "/staff/appointment_scheduling/undo-left"

print("=" * 78)
print("IDS-BPHCSN — the Doctor's list and Left before being seen (Task 6)")
print("=" * 78)


# ---------------------------------------------------------------------------
section("1. SET UP VISITS IN KNOWN STATES")
NAMES = {}
CONSULT, FOLLOW, HTN = "Consultation", "Follow-up Consultation", "Hypertension / Diabetes Monitoring"
WFD = "Waiting for the Doctor"


def name_patient(pid, tag):
    run_sql("UPDATE patients SET first_name = ?, last_name = 'Wlmark' WHERE id = ?",
            (f"Zzwl{tag}", pid))
    NAMES[tag] = f"Zzwl{tag} Wlmark"
    return pid


def add_visit(pid, tag, days_ago, status, kind=CONSULT, time="09:30",
              consultation=None, restricted=False, diagnosis=None):
    """One visit. `consultation` True writes an intake; `diagnosis` fills the diagnosis."""
    name_patient(pid, tag)
    day = (today - timedelta(days=days_ago)).isoformat()
    appointment = run_sql(
        "INSERT INTO appointments (patient_id, scheduled_by, appointment_type, appointment_date, "
        "status, is_walkin, checked_in_at) VALUES (?, ?, ?, ?, ?, 1, ?)",
        (pid, STAFF_ID, kind, day, status, f"{day} {time}:00"))
    cid = None
    if consultation:
        cid = run_sql(
            "INSERT INTO consultations (patient_id, recorded_by, appointment_id, "
            "consultation_date, chief_complaint, diagnosis, is_restricted) VALUES (?,?,?,?,?,?,?)",
            (pid, STAFF_ID, appointment, day, "intake complaint", diagnosis, 1 if restricted else 0))
    return appointment, cid


# The fixture has 20 patients; five more are needed.
for extra in range(21, 26):
    run_sql("INSERT INTO patients (id, patient_id, first_name, last_name, birth_date, sex, civil_status, "
            "address, municipality, province, contact_number, status) VALUES (?, ?, 'Extra', 'Patient', "
            "'1985-01-01', 'Female', 'Single', 'Purok 1', 'San Francisco', 'Surigao del Norte', ?, 'Active')",
            (extra, f"PAT-{extra:03d}", f"0900000{extra:04d}"))

# Nobody is in consultation at the start, so the single-room rule is free to test.
run_sql("UPDATE appointments SET status = 'Completed' WHERE status = 'In Consultation'")
V = {}
V["A"] = add_visit(6, "A", 0, "Waiting", time="08:10")                          # waiting today
V["B"] = add_visit(7, "B", 3, "Waiting", time="08:00")                          # earlier day, still waiting
V["D"] = add_visit(8, "D", 2, WFD, consultation=True, time="10:15")             # intake, no diagnosis
V["E"] = add_visit(9, "E", 1, "Completed", consultation=True, diagnosis="Seen by the Doctor")
V["F"] = add_visit(10, "F", 4, "Completed")                                     # closed, nothing written
V["G"] = add_visit(11, "G", 1, "Waiting", kind="Immunization")                  # not a Doctor service
V["H"] = add_visit(12, "H", 5, "Waiting", time="07:50")
run_sql("UPDATE appointments SET status = ?, status_before_left = 'Waiting', left_by = ?, "
        "left_at = ? WHERE id = ?", (LEFT, STAFF_ID, f"{today.isoformat()} 07:00:00", V["H"][0]))
V["J"] = add_visit(13, "J", 2, "Completed", kind="Immunization", consultation=True, time="11:00")
V["P"] = add_visit(14, "P", 1, "Completed", kind="Prenatal Check-up", consultation=True)
V["K"] = add_visit(15, "K", 0, "In Consultation", time="08:40")                  # with the clinician now
# A sealed visit whose diagnosis is not written yet: its complaint is stored encrypted.
V["S"] = add_visit(16, "S", 2, WFD, consultation=True, restricted=True, time="12:00")
V["L"] = add_visit(17, "L", 1, WFD, kind=FOLLOW, consultation=True, time="09:00")
# Half an hour ago, never before midnight, so the wait is in the past whatever time the suite runs.
M_TIME = max(datetime.now() - timedelta(minutes=30),
             datetime.combine(date.today(), datetime.min.time())).strftime("%H:%M")
V["M"] = add_visit(18, "M", 0, WFD, kind=HTN, time=M_TIME)                      # today; no record found
from utils import sealed_crypto  # noqa: E402

SEALED_MARK = "ZZSEALEDCOMPLAINTMARK"
run_sql("UPDATE consultations SET chief_complaint = ? WHERE id = ?",
        (sealed_crypto.encrypt(SEALED_MARK, V["S"][1], "chief_complaint"), V["S"][1]))
ids = {tag: v[0] for tag, v in V.items()}
check("thirteen visits set up in known states", len(ids) == 13, ids)

# ---------------------------------------------------------------------------
section("2. WHO IS ON THE LIST (D4)")
from models.doctor_worklist import (  # noqa: E402
    SERVICES_NEEDING_DOCTOR, VISIT_SERVICES, WAITING_FOR_DOCTOR, clean_service,
    count_waiting_for_doctor, get_waiting_for_doctor, needs_doctor,
)

with app.app_context():
    listed = get_waiting_for_doctor()
    count = count_waiting_for_doctor()
on_list = {v["id"] for v in listed}
exactly = {r[0] for r in db().execute("SELECT id FROM appointments WHERE status = ?", (WFD,))}
check("the Doctor's list is exactly the visits with the status 'Waiting for the Doctor'",
      on_list == exactly == {ids[t] for t in ("D", "S", "L", "M")}, (sorted(on_list), sorted(exactly)))
check("a visit that is only Waiting or In Consultation is not on it (it is on the queue)",
      not ({ids["A"], ids["B"], ids["K"], ids["G"]} & on_list))
check("a visit with a diagnosis, a closed one and one marked left are not on it",
      not ({ids["E"], ids["F"], ids["H"]} & on_list))
check("immunization and prenatal visits are never on it, waiting or completed",
      not ({ids["G"], ids["J"], ids["P"]} & on_list))
check("every visit on the list has no diagnosis",
      all(not (r[0] or "").strip() for r in db().execute(
          "SELECT c.diagnosis FROM appointments a JOIN consultations c ON c.appointment_id = a.id "
          "WHERE a.status = ?", (WFD,))))
check("VISIT_SERVICES is one fixed list; the services needing the Doctor are a subset; a blank "
      "type still needs one; Immunization, Prenatal Check-up, Family Planning and an old Dental "
      "visit do not",
      VISIT_SERVICES == (CONSULT, FOLLOW, HTN, "Prenatal Check-up", "Family Planning", "Immunization")
      and set(SERVICES_NEEDING_DOCTOR) - {"Walk-in Consultation"} <= set(VISIT_SERVICES)
      and "Walk-in Consultation" not in VISIT_SERVICES and needs_doctor(None) and needs_doctor("")
      and needs_doctor("Walk-in Consultation")
      and not any(needs_doctor(s) for s in ("Immunization", "Prenatal Check-up", "Family Planning", "Dental"))
      and WAITING_FOR_DOCTOR == WFD)
check("a service is accepted only if it is in the list, exactly; there is no default",
      clean_service("Family Planning") == "Family Planning" and clean_service(" Family Planning ") == "Family Planning"
      and clean_service("family planning") is None and clean_service("") is None
      and clean_service("Dental") is None and clean_service("dental") is None
      and clean_service(None) is None and clean_service("Walk-in Consultation") is None)
order = [v["id"] for v in listed]
check("the order is oldest first, across all days",
      order == [ids["D"], ids["S"], ids["L"], ids["M"]]
      and [v["checked_in"] for v in listed] == sorted(v["checked_in"] for v in listed), order)
by_id = {v["id"]: v for v in listed}
check("visits from earlier days are marked, and today's are not",
      all(by_id[ids[t]]["from_earlier_day"] for t in ("D", "S", "L"))
      and not by_id[ids["M"]]["from_earlier_day"])


def plain_days(visit):
    """The wait as the page should say it: whole elapsed days, or hours and minutes."""
    days = (datetime.now() - visit["checked_in"]).days
    if days >= 1:
        return f"{days} day{'' if days == 1 else 's'}"
    # Less than a day since a check-in "yesterday" (the suite ran before that visit's time of day).
    hours = int((datetime.now() - visit["checked_in"]).total_seconds() // 3600)
    return f"{hours} hour{'' if hours == 1 else 's'}"


check("each row says how long it has waited, in plain words",
      all(by_id[ids[t]]["waiting"] == plain_days(by_id[ids[t]]) for t in ("D", "S", "L"))
      and re.fullmatch(r"(Just now|\d+ minutes?|\d+ hours?)", by_id[ids["M"]]["waiting"]),
      {k: v["waiting"] for k, v in by_id.items()})
check("each row carries the name, patient ID and the chosen service",
      by_id[ids["D"]]["name"] == NAMES["D"] and by_id[ids["D"]]["patient_reference"] == "PAT-008"
      and by_id[ids["L"]]["service"] == FOLLOW and by_id[ids["M"]]["service"] == HTN)
check("every row has the same action, 'see patient', with or without an intake record (Task 6d)",
      all(v["action"] == "see_patient" for v in by_id.values())
      and by_id[ids["D"]]["consultation_id"] and by_id[ids["M"]]["consultation_id"] is None)

# ---------------------------------------------------------------------------
section("3. THE DOCTOR'S LIST, THE STAFF COUNT AND THE ADMINISTRATOR'S COUNT AGREE")
doctor_page = body(doctor.get("/staff/dashboard"))
staff_page = body(staff.get("/staff/dashboard"))
admin_page = body(admin.get("/admin/dashboard"))
card = re.compile(r"Waiting for the Doctor: (\d+)")
staff_n = card.search(flat(staff_page))
admin_n = card.search(flat(admin_page))
check("the Doctor's dashboard has the list, with every visit's patient and a 'See patient' button",
      'id="waitingForDoctor"' in doctor_page
      and all(NAMES[t] in doctor_page for t in ("D", "S", "L", "M"))
      and doctor_page.count("See patient") >= 4 and "Open Patient Queue" not in doctor_page
      and "Open consultation" not in doctor_page)
doctor_list = doctor_page.split('id="waitingForDoctor"')[1].split("</table>")[0]
check("a visit that is only Waiting, or In Consultation, is not on the Doctor's list "
      "(it is on the queue)", not any(NAMES[t] in doctor_list for t in ("A", "B", "K")))
check("the Doctor's list has no 'From an earlier day' tag; the Waiting column shows the days (Task 10)",
      "From an earlier day" not in doctor_page and "Patients who need a diagnosis, oldest first." in doctor_page)
check("the Staff dashboard shows 'Waiting for the Doctor: N', the Administrator's too, and they "
      "equal the list",
      staff_n and admin_n and int(staff_n.group(1)) == int(admin_n.group(1)) == count == len(listed) == 4,
      (staff_n and staff_n.group(1), admin_n and admin_n.group(1), count))
check("the Doctor's list length is the same number",
      f'<span class="badge-status badge-waiting">{len(listed)}</span>' in doctor_page)
check("the Staff dashboard has no list; the Administrator's has no list and no names of the "
      "visits that are waiting for the Doctor",
      'id="waitingForDoctor"' not in staff_page and 'id="waitingForDoctor"' not in admin_page
      and not any(NAMES[t] in admin_page for t in ("D", "S", "L", "M")),
      [t for t in NAMES if NAMES[t] in admin_page])
check("a sealed visit is on the Doctor's list by name and shows no sealed text",
      NAMES["S"] in doctor_page and SEALED_MARK not in doctor_page
      and "intake complaint" not in doctor_page)
check("the Official's dashboard has neither the list nor the card",
      "waitingForDoctor" not in body(official.get("/official/dashboard"))
      and "Waiting for the Doctor" not in body(official.get("/official/dashboard")))
with app.app_context():
    from models.doctor_worklist import _FROM_SQL, _WHERE_SQL
check("the count is read without any name: its query holds no name column",
      "first_name" not in "SELECT COUNT(*) FROM appointments a" + _WHERE_SQL)
run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (ids["K"],))


def status_of(appointment):
    return q1("SELECT status FROM appointments WHERE id = ?", (appointment,))[0]


def audits(action):
    conn = db()
    try:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM audit_logs WHERE action = ? ORDER BY id", (action,))]
    finally:
        conn.close()


# ---------------------------------------------------------------------------
section("4. SERVICE AT CHECK-IN IS REQUIRED, CHOSEN FROM THE LIST, AND VALIDATED ON THE SERVER (D1)")
CHECK_IN = "/staff/appointment_scheduling/check-in"
PATIENT = 4          # an adult who is not otherwise used above
name_patient(PATIENT, "C")
run_sql("UPDATE appointments SET status = 'Completed' WHERE patient_id = ? "
        "AND status IN ('Waiting', 'In Consultation')", (PATIENT,))


def visits_of(pid):
    return q1("SELECT COUNT(*) FROM appointments WHERE patient_id = ?", (pid,))[0]


start = visits_of(PATIENT)
for label, data in (("no service at all", {"patient_id": str(PATIENT), "reason": "Ubo"}),
                    ("an empty service", {"patient_id": str(PATIENT), "reason": "Ubo", "service": ""}),
                    ("an unknown service", {"patient_id": str(PATIENT), "reason": "Ubo",
                                            "service": "Massage"}),
                    ("the retired name", {"patient_id": str(PATIENT), "reason": "Ubo",
                                          "service": "Walk-in Consultation"}),
                    ("a wrong case", {"patient_id": str(PATIENT), "reason": "Ubo",
                                      "service": "family planning"})):
    r = staff.post(CHECK_IN, data=data)
    check(f"check-in with {label} is refused and nobody is queued",
          r.status_code == 302 and visits_of(PATIENT) == start)
check("a refused check-in writes no audit entry",
      not [a for a in audits("Checked patient in") if f"PAT-{PATIENT:03d}" in a["details"]])
for service in VISIT_SERVICES:
    run_sql("UPDATE appointments SET status = 'Completed' WHERE patient_id = ? "
            "AND status IN ('Waiting', 'In Consultation', ?)", (PATIENT, WFD))
    r = staff.post(CHECK_IN, data={"patient_id": str(PATIENT), "reason": "Ubo", "service": service})
    row = q1("SELECT * FROM appointments WHERE patient_id = ? ORDER BY id DESC LIMIT 1", (PATIENT,))
    check(f"check-in with '{service}' is accepted and the chosen service is stored",
          r.status_code == 302 and row["appointment_type"] == service and row["status"] == "Waiting",
          dict(row))
page = body(staff.get("/staff/appointment_scheduling"))
check("the Check In dialog offers every service, with no default (the first option is a blank prompt)",
      all(f'<option value="{s}">{s}</option>' in page for s in VISIT_SERVICES)
      and '<option value="" selected disabled>Choose a service</option>' in page
      and page.count('<option value="" selected disabled>Choose a service</option>') == 2)
# The new-patient path: the dialog hands the service over with the reason.
run_sql("UPDATE appointments SET status = 'Completed' WHERE patient_id = ? AND status = 'Waiting'", (PATIENT,))
r = staff.post("/staff/register-patient", data={"check_in_start": "1", "check_in_reason": "Ubo"})
check("the new-patient path refuses a missing service before a registration form opens",
      r.status_code == 302 and "check_in=1" not in r.headers.get("Location", ""))
r = staff.post("/staff/register-patient", data={"check_in_start": "1", "check_in_reason": "Ubo",
                                                "check_in_service": "Massage"})
check("... and an unknown one", r.status_code == 302 and "check_in=1" not in r.headers.get("Location", ""))
r = staff.post("/staff/register-patient", data={"check_in_start": "1", "check_in_reason": "Ubo",
                                                "check_in_service": "Family Planning"})
check("a valid service is carried into the registration form",
      r.status_code == 302 and r.headers["Location"].endswith("check_in=1")
      and 'name="check_in_service" value="Family Planning"' in body(staff.get(r.headers["Location"])))
registered = q1("SELECT COUNT(*) FROM patients")[0]
form = {"first_name": "Zzwlnew", "last_name": "Wlmark", "birth_date": "1990-05-05", "sex": "Female",
        "civil_status": "Single", "address": "Purok 1", "municipality": "San Francisco",
        "province": "Surigao del Norte", "zip_code": "8405", "emergency_contact_name": "Test Relative",
        "emergency_contact_relationship": "Sibling", "emergency_contact_number": "09000000001",
        "confirm_duplicate": "yes", "check_in": "1", "check_in_reason": "Ubo"}
r = staff.post("/staff/register-patient", data=dict(form))
check("a forged submit that asks to be checked in with no valid service registers nobody",
      q1("SELECT COUNT(*) FROM patients")[0] == registered)
r = staff.post("/staff/register-patient", data=dict(form, check_in_service="Family Planning"))
new_patient = q1("SELECT id FROM patients WHERE first_name = 'Zzwlnew'")
row = q1("SELECT * FROM appointments WHERE patient_id = ? ", (new_patient[0],)) if new_patient else None
check("a new patient registered from the dialog is queued with the chosen service",
      row is not None and row["appointment_type"] == "Family Planning" and row["status"] == "Waiting")

# ---------------------------------------------------------------------------
section("5. CORRECTING THE SERVICE OF A VISIT THAT IS NOT FINISHED (D2)")
CHANGE = "/staff/appointment_scheduling/change-service"
fix = add_visit(19, "Q", 0, "Waiting", kind=CONSULT, time="10:00")[0]
before_audit = q1("SELECT COALESCE(MAX(id),0) FROM audit_logs")[0]
r = staff.post(CHANGE, data={"appointment_id": str(fix), "service": "Immunization"})
entries = [a for a in audits("Changed the service of a visit") if a["id"] > before_audit]
check("Staff correct the service of a Waiting visit",
      r.status_code == 302 and q1("SELECT appointment_type FROM appointments WHERE id = ?", (fix,))[0]
      == "Immunization")
check("one short audit entry, in the words the task gave, and it names no patient",
      len(entries) == 1 and entries[0]["user_id"] == STAFF_ID
      and entries[0]["details"] == f"Corrected the service of visit APT-{fix:05d}."
      and NAMES["Q"] not in entries[0]["details"], entries)
r = doctor.post(CHANGE, data={"appointment_id": str(fix), "service": "Family Planning"})
check("the Doctor can correct it too, and the entry names the Doctor",
      q1("SELECT appointment_type FROM appointments WHERE id = ?", (fix,))[0] == "Family Planning"
      and (audits("Changed the service of a visit") or [{}])[-1].get("user_id") == DOCTOR_ID)
n = len(audits("Changed the service of a visit"))
staff.post(CHANGE, data={"appointment_id": str(fix), "service": "Family Planning"})
check("choosing the service the visit already has changes nothing and writes nothing",
      len(audits("Changed the service of a visit")) == n)
for label, value in (("an unknown service", "Massage"), ("no service", ""), ("a wrong case", "family planning")):
    staff.post(CHANGE, data={"appointment_id": str(fix), "service": value})
check("an unknown, empty or wrongly written service is refused on the server",
      q1("SELECT appointment_type FROM appointments WHERE id = ?", (fix,))[0] == "Family Planning"
      and len(audits("Changed the service of a visit")) == n)
refused = {"a Completed visit": ids["E"], "a visit marked left before being seen": ids["H"]}
for label, appointment in refused.items():
    old = q1("SELECT appointment_type FROM appointments WHERE id = ?", (appointment,))[0]
    staff.post(CHANGE, data={"appointment_id": str(appointment), "service": "Family Planning"})
    check(f"the service of {label} cannot be changed",
          q1("SELECT appointment_type FROM appointments WHERE id = ?", (appointment,))[0] == old)
check("the refusals wrote no audit entry", len(audits("Changed the service of a visit")) == n)
for who, client in (("Administrator", admin), ("Barangay Official", official)):
    r = client.post(CHANGE, data={"appointment_id": str(fix), "service": "Immunization"})
    check(f"{who}: POST change-service (valid token) is refused on the server",
          r.status_code == 302 and r.headers.get("Location", "").endswith("dashboard")
          and q1("SELECT appointment_type FROM appointments WHERE id = ?", (fix,))[0] == "Family Planning")
check("the Administrator has no /admin/ route for it",
      admin.post("/admin/appointment_scheduling/change-service",
                 data={"appointment_id": str(fix), "service": "Immunization"}).status_code == 404)
# A correction never completes a visit (Task 6e): waiting for the Doctor, corrected to a
# service that does not need him, goes back In Consultation so the service is recorded.
run_sql("UPDATE appointments SET status = ?, appointment_type = ? WHERE id = ?", (WFD, CONSULT, fix))
staff.post(CHANGE, data={"appointment_id": str(fix), "service": "Immunization"})
check("a visit waiting for the Doctor, corrected to Immunization, is In Consultation and never "
      "Completed, and the entry says so",
      status_of(fix) == "In Consultation" and (audits("Changed the service of a visit") or [{}])[-1].get(
          "details", "").endswith("The visit is now In Consultation."))
run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (fix,))
run_sql("UPDATE appointments SET status = ?, appointment_type = 'Dental' WHERE id = ?", (WFD, fix))
staff.post(CHANGE, data={"appointment_id": str(fix), "service": CONSULT})
check("corrected to a service that does need him, it stays waiting for the Doctor",
      status_of(fix) == WFD)
run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (fix,))
page = body(staff.get("/staff/appointment_scheduling"))
check("Staff and the Doctor see a Change service button and its dialog; the Administrator does not",
      "serviceModal" in page and "Change service" in page
      and "serviceModal" not in body(admin.get("/admin/appointment_scheduling"))
      and "serviceModal" in body(doctor.get("/staff/appointment_scheduling")))

# ---------------------------------------------------------------------------
section("6. THE FLOW: INTAKE, THE DOCTOR'S DIAGNOSIS, A FULL DOCTOR CONSULTATION (D3)")
START = "/staff/appointment_scheduling/start-consultation"
INTAKE = {"chief_complaint": "flow intake", "temperature": "37.0", "blood_pressure": "120/80",
          "weight": "61"}
FULL = dict(INTAKE, diagnosis="Flow diagnosis", treatment_plan="Flow plan")


def new_visit(service, tag, pid=None):
    """Check the flow patient in through the real route, then start the consultation."""
    pid = pid or PATIENT
    run_sql("UPDATE appointments SET status = 'Completed' WHERE status IN ('Waiting', 'In Consultation')"
            " AND patient_id = ?", (pid,))
    staff.post(CHECK_IN, data={"patient_id": str(pid), "reason": tag, "service": service})
    appointment = q1("SELECT id FROM appointments WHERE patient_id = ? ORDER BY id DESC LIMIT 1",
                     (pid,))[0]
    staff.post(START, data={"appointment_id": str(appointment)})
    return appointment


def listed_now():
    with app.app_context():
        return {v["id"] for v in get_waiting_for_doctor()}


def consultation_of(appointment):
    return q1("SELECT id, diagnosis FROM consultations WHERE appointment_id = ?", (appointment,))


for service in (CONSULT, FOLLOW, HTN):
    a = new_visit(service, "intake")
    check(f"{service}: the visit is In Consultation after Start", status_of(a) == "In Consultation")
    r = staff.post(f"/staff/consultation/record/{a}", data=INTAKE)
    check(f"{service}: Staff save the intake with no diagnosis: the visit is 'Waiting for the Doctor', "
          "not Completed", status_of(a) == WFD, status_of(a))
    with app.app_context():
        check(f"{service}: it is on the Doctor's list now",
              a in {v["id"] for v in get_waiting_for_doctor()})
check("the room is free again after the intake is saved",
      q1("SELECT COUNT(*) FROM appointments WHERE status = 'In Consultation'")[0] == 0)
last_entries = [e for e in audits("Intake saved, waiting for the Doctor")]
check("saving the intake writes one entry saying so (patient and visit only), not 'Completed consultation'",
      last_entries and last_entries[-1]["details"].startswith(f"{NAMES['C']} (PAT-{PATIENT:03d}), visit APT-"))
cid = consultation_of(a)[0]
before_completed = len(audits("Completed consultation"))
r = doctor.post(f"/staff/consultation/edit/{cid}", data=FULL)
check("the Doctor saves the diagnosis: the visit is Completed and the diagnosis is stored",
      status_of(a) == "Completed" and consultation_of(a)[1] == "Flow diagnosis")
with app.app_context():
    check("... and it is gone from the Doctor's list",
          a not in {v["id"] for v in get_waiting_for_doctor()})
check("... with one 'Completed consultation' entry by the Doctor",
      len(audits("Completed consultation")) == before_completed + 1
      and audits("Completed consultation")[-1]["user_id"] == DOCTOR_ID)
r = doctor.post(f"/staff/consultation/edit/{cid}", data=dict(FULL, diagnosis="A later correction"))
check("the Doctor may change a diagnosis on a Completed visit; it stays Completed",
      status_of(a) == "Completed" and len(audits("Completed consultation")) == before_completed + 1)
# A visit that needs no Doctor.
for service in ("Immunization", "Prenatal Check-up", "Family Planning", "Dental"):
    b = new_visit(service, "staff only")
    staff.post(f"/staff/consultation/record/{b}", data=INTAKE)
    check(f"{service}: Staff save the intake and the visit is Completed, with no diagnosis and "
          "never on the Doctor's list",
          status_of(b) == "Completed" and not (consultation_of(b)[1] or "").strip()
          and b not in listed_now())
# The Doctor records the whole consultation himself.
c = new_visit(CONSULT, "doctor himself")
r = doctor.post(f"/staff/consultation/record/{c}", data=FULL)
check("the Doctor records the whole consultation, with a diagnosis: straight to Completed",
      status_of(c) == "Completed" and consultation_of(c)[1] == "Flow diagnosis"
      and c not in listed_now())
# The Doctor writes a diagnosis on an immunization visit: it stays Completed.
d = new_visit("Immunization", "doctor adds")
staff.post(f"/staff/consultation/record/{d}", data=INTAKE)
doctor.post(f"/staff/consultation/edit/{consultation_of(d)[0]}", data=FULL)
check("the Doctor may write a diagnosis on any visit; an immunization visit stays Completed",
      status_of(d) == "Completed")
# Task 6c: 'Mark Completed' from the queue is refused for a service that needs the Doctor.
COMPLETE = "/staff/appointment_scheduling/complete-consultation"
SEND = "/staff/appointment_scheduling/send-to-doctor"
e = new_visit(CONSULT, "closed")
n_complete = len(audits("Completed consultation"))
for who, client in (("Staff", staff), ("the Doctor", doctor)):
    r = client.post(COMPLETE, data={"appointment_id": str(e)})
    check(f"{who}: 'Mark Completed' on a Consultation visit with no diagnosis is refused on the server",
          status_of(e) == "In Consultation"
          and "This visit needs the Doctor first." in body(client.get("/staff/appointment_scheduling")))
check("the refusal wrote no entry", len(audits("Completed consultation")) == n_complete)
page = body(staff.get("/staff/appointment_scheduling"))
check("the queue offers 'Send to the Doctor' there, not 'Mark Completed'",
      "Send to the Doctor" in page and 'form="sendToDoctorForm"' in page
      and f'form="completeConsultationForm" name="appointment_id" value="{e}"' not in page)
before_sent = len(audits("Sent visit to the Doctor"))
r = staff.post(SEND, data={"appointment_id": str(e)})
sent = audits("Sent visit to the Doctor")
check("'Send to the Doctor' sets Waiting for the Doctor, and the room is free",
      status_of(e) == WFD and q1("SELECT COUNT(*) FROM appointments WHERE status = 'In Consultation'")[0] == 0
      and e in listed_now())
check("it writes one short entry: the visit only, no patient and no service",
      len(sent) == before_sent + 1 and sent[-1]["user_id"] == STAFF_ID
      and sent[-1]["details"] == f"Sent visit APT-{e:05d} to the Doctor.",
      sent[-1:])
n = len(audits("Sent visit to the Doctor"))
staff.post(SEND, data={"appointment_id": str(e)})
check("sending it a second time changes nothing and writes nothing",
      status_of(e) == WFD and len(audits("Sent visit to the Doctor")) == n)
# Every other way to Completed is refused for a visit waiting for the Doctor, without a diagnosis.
for who, client in (("Staff", staff), ("the Doctor", doctor), ("the Administrator", admin),
                    ("the Official", official)):
    client.post(COMPLETE, data={"appointment_id": str(e)})
check("Mark Completed, by Staff, the Doctor, the Administrator or the Official, cannot complete a visit "
      "waiting for the Doctor", status_of(e) == WFD)
staff.post(f"/staff/consultation/record/{e}", data=FULL)
check("the record route refuses Staff for a visit waiting for the Doctor (only the Doctor sees him; "
      "Task 6d)", status_of(e) == WFD and consultation_of(e) is None)
for who, client in (("the Administrator", admin), ("the Official", official)):
    client.post(f"/staff/consultation/record/{e}", data=FULL)
check("... and the Administrator and the Official", status_of(e) == WFD and consultation_of(e) is None)
e_start = new_visit(CONSULT, "blank")
staff.post(f"/staff/consultation/record/{e_start}", data=INTAKE)
cid_e = consultation_of(e_start)[0]
doctor.post(f"/staff/consultation/edit/{cid_e}", data=dict(FULL, diagnosis=""))
check("the Doctor saving with a blank diagnosis does not complete it", status_of(e_start) == WFD)
staff.post(f"/staff/consultation/edit/{cid_e}", data=FULL)
check("Staff saving a diagnosis (refused on the server) does not complete it",
      status_of(e_start) == WFD and not (consultation_of(e_start)[1] or "").strip())
staff.post("/staff/appointment_scheduling/undo-left", data={"appointment_id": str(e_start)})
check("'Undo left' on a visit that is not marked changes nothing", status_of(e_start) == WFD)
staff.post("/staff/appointment_scheduling/mark-left", data={"appointment_id": str(e_start)})
staff.post("/staff/appointment_scheduling/undo-left", data={"appointment_id": str(e_start)})
check("a visit marked left from Waiting for the Doctor comes back to it, never to Completed",
      status_of(e_start) == WFD)
from models.appointments import complete_consultation  # noqa: E402

with app.app_context():
    check("complete_consultation() itself does not complete a visit waiting for the Doctor "
          "(the rule is in the model, not the buttons)",
          complete_consultation(e_start) == "not-in-consultation" and status_of(e_start) == WFD)
run_sql("UPDATE appointments SET status = 'Completed' WHERE id IN (?, ?)", (e, e_start))
# A service that does not need the Doctor still completes from the queue.
g = new_visit("Family Planning", "closed family planning")
staff.post(COMPLETE, data={"appointment_id": str(g)})
check("Mark Completed still works for a service that does not need the Doctor",
      status_of(g) == "Completed")
h = new_visit("Immunization", "send refused")
staff.post(SEND, data={"appointment_id": str(h)})
check("'Send to the Doctor' is refused for a service that does not need him",
      status_of(h) == "In Consultation")
page = body(staff.get("/staff/appointment_scheduling"))
check("... and the queue offers 'Mark Completed' for it", f'form="completeConsultationForm" name="appointment_id" value="{h}"' in page)
staff.post(COMPLETE, data={"appointment_id": str(h)})
for who, client in (("Administrator", admin), ("Barangay Official", official)):
    spare2 = add_visit(24, "N" + who[:1], 0, "In Consultation", time="17:00")[0]
    client.post(SEND, data={"appointment_id": str(spare2)})
    check(f"{who}: POST send-to-doctor is refused on the server", status_of(spare2) == "In Consultation")
    run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (spare2,))
check("the Administrator has no /admin/ route for it",
      admin.post("/admin/appointment_scheduling/send-to-doctor", data={"appointment_id": "1"}).status_code == 404)
# A stale or forged second save does nothing.
f = new_visit(CONSULT, "stale")
staff.post(f"/staff/consultation/record/{f}", data=INTAKE)
n = len(audits("Intake saved, waiting for the Doctor"))
staff.post(f"/staff/consultation/record/{f}", data=INTAKE)
check("saving the intake a second time changes nothing and writes nothing",
      status_of(f) == WFD and len(audits("Intake saved, waiting for the Doctor")) == n)
run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (f,))

# ---------------------------------------------------------------------------
section("7. LEFT BEFORE BEING SEEN: FROM WAITING AND FROM WAITING FOR THE DOCTOR (D5)")
MARK, UNDO = "/staff/appointment_scheduling/mark-left", "/staff/appointment_scheduling/undo-left"
before_audit = q1("SELECT COALESCE(MAX(id),0) FROM audit_logs")[0]
with app.app_context():
    count = count_waiting_for_doctor()
r = staff.post(MARK, data={"appointment_id": str(ids["A"])})
row = q1("SELECT * FROM appointments WHERE id = ?", (ids["A"],))
check("Staff mark a Waiting visit as left before being seen",
      r.status_code == 302 and row["status"] == LEFT, (r.status_code, row["status"]))
check("who, when and the previous status are recorded",
      row["left_by"] == STAFF_ID and row["status_before_left"] == "Waiting"
      and (row["left_at"] or "")[:10] == today.isoformat(), dict(row))
entries = [a for a in audits("Marked as left before being seen") if a["id"] > before_audit]
check("exactly one short audit entry is written, in the existing style",
      len(entries) == 1 and entries[0]["user_id"] == STAFF_ID
      and entries[0]["details"] == f"{NAMES['A']} (PAT-006), visit APT-{ids['A']:05d}.", entries)
r = staff.post(MARK, data={"appointment_id": str(ids["D"])})
check("Staff mark a visit that is Waiting for the Doctor (the patient left before he came)",
      status_of(ids["D"]) == LEFT and q1("SELECT status_before_left FROM appointments "
                                         "WHERE id = ?", (ids["D"],))[0] == WFD)
with app.app_context():
    check("it leaves the Doctor's list at once, and the count drops with it",
          ids["D"] not in {v["id"] for v in get_waiting_for_doctor()}
          and count_waiting_for_doctor() == count - 1)
r = doctor.post(MARK, data={"appointment_id": str(ids["L"])})
check("the Doctor can mark one too, and the entry names the Doctor",
      status_of(ids["L"]) == LEFT and audits("Marked as left before being seen")[-1]["user_id"] == DOCTOR_ID)
k = add_visit(15, "K2", 0, "In Consultation", time="08:41")[0]
staff.post(MARK, data={"appointment_id": str(k)})
check("a visit In Consultation with nothing recorded can be marked, and it frees the room",
      status_of(k) == LEFT and q1("SELECT COUNT(*) FROM appointments WHERE status = 'In Consultation'")[0] == 0)
refused = {
    "a Completed visit with a diagnosis": ids["E"],
    "a Completed visit closed with nothing written": ids["F"],
    "a Completed immunization with intake": ids["J"],
    "a Completed prenatal visit with intake": ids["P"],
}
marks_before = len(audits("Marked as left before being seen"))
for label, appointment in refused.items():
    staff.post(MARK, data={"appointment_id": str(appointment)})
    check(f"refused on {label}: the status does not change", status_of(appointment) == "Completed")
check("every refusal wrote no audit entry", len(audits("Marked as left before being seen")) == marks_before)
race_visit = add_visit(20, "R", 0, WFD, consultation=True, time="13:00")
from models.appointments import mark_left  # noqa: E402

run_sql("UPDATE consultations SET diagnosis = 'Written just now' WHERE id = ?", (race_visit[1],))
with app.app_context():
    outcome = mark_left(race_visit[0], STAFF_ID)
check("a diagnosis written a moment earlier stops the mark (the rule is in the model)",
      outcome == "not-allowed" and status_of(race_visit[0]) == WFD, outcome)
run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (race_visit[0],))
spare = add_visit(21, "X", 0, "Waiting", time="14:00")
for who, client in (("Administrator", admin), ("Barangay Official", official)):
    r = client.post(MARK, data={"appointment_id": str(spare[0])})
    check(f"{who}: POST mark-left (valid token) is refused on the server",
          r.status_code == 302 and r.headers.get("Location", "").endswith("dashboard")
          and status_of(spare[0]) == "Waiting")
    r = client.post(UNDO, data={"appointment_id": str(ids["A"])})
    check(f"{who}: POST undo-left is refused too", r.status_code == 302 and status_of(ids["A"]) == LEFT)
check("the Administrator has no /admin/ route for either action",
      admin.post("/admin/appointment_scheduling/mark-left",
                 data={"appointment_id": str(spare[0])}).status_code == 404)
# Undo.
undo_before = len(audits("Undid left before being seen"))
staff.post(UNDO, data={"appointment_id": str(ids["A"])})
row = q1("SELECT * FROM appointments WHERE id = ?", (ids["A"],))
check("Staff undo the same day: the visit is Waiting again, and who/when are cleared",
      row["status"] == "Waiting" and row["left_by"] is None and row["left_at"] is None
      and row["status_before_left"] is None)
undone = audits("Undid left before being seen")
check("the undo writes one audit entry, in the same style",
      len(undone) == undo_before + 1
      and undone[-1]["details"] == f"{NAMES['A']} (PAT-006), visit APT-{ids['A']:05d}.")
doctor.post(UNDO, data={"appointment_id": str(ids["D"])})
with app.app_context():
    check("undoing a visit that was Waiting for the Doctor restores that status, and it is on his "
          "list again", status_of(ids["D"]) == WFD
          and ids["D"] in {v["id"] for v in get_waiting_for_doctor()})
check("undoing a visit that is not marked changes nothing and writes nothing",
      staff.post(UNDO, data={"appointment_id": str(ids["E"])}).status_code == 302
      and len(audits("Undid left before being seen")) == undo_before + 2)
staff.post(MARK, data={"appointment_id": str(ids["D"])})
run_sql("UPDATE appointments SET left_at = ? WHERE id = ?",
        (f"{(today - timedelta(days=1)).isoformat()} 16:00:00", ids["D"]))
count_before = len(audits("Undid left before being seen"))
staff.post(UNDO, data={"appointment_id": str(ids["D"])})
check("an undo on a later day is refused: the visit stays closed",
      status_of(ids["D"]) == LEFT and len(audits("Undid left before being seen")) == count_before)
run_sql("UPDATE appointments SET status = ?, left_by = NULL, left_at = NULL, status_before_left = NULL "
        "WHERE id = ?", (WFD, ids["D"]))
other_room = add_visit(22, "Y", 0, "In Consultation", time="15:00")
staff.post(UNDO, data={"appointment_id": str(k)})
check("putting a visit back In Consultation is refused while someone else is in the room",
      status_of(k) == LEFT and status_of(other_room[0]) == "In Consultation")
run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (other_room[0],))
staff.post(UNDO, data={"appointment_id": str(k)})
check("once the room is free the same undo works", status_of(k) == "In Consultation")
run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (k,))

# ---------------------------------------------------------------------------
section("8. THE QUEUE SCREEN: OPEN VISITS FROM EARLIER DAYS, THE DIALOGS, READ-ONLY FOR THE ADMINISTRATOR (Task 10)")
run_sql("UPDATE appointments SET status = 'Waiting', status_before_left = NULL, left_by = NULL, "
        "left_at = NULL WHERE id = ?", (ids["B"],))
OPEN_TITLE = "Open visits from earlier days"


def section_of(page):
    """The earlier-days section: from its title to the end of its table."""
    if OPEN_TITLE not in page:
        return ""
    return page.split(OPEN_TITLE)[1].split("</table>")[0]


def row_in(section_html, marker):
    for chunk in section_html.split("<tr")[1:]:
        r_ = chunk.split("</tr>")[0]
        if marker in r_:
            return r_
    return ""


page = body(staff.get("/staff/appointment_scheduling"))
section_text = section_of(page)
check("Staff see 'Open visits from earlier days' with a waiting visit, its ID and how long",
      NAMES["B"] in section_text and f"APT-{ids['B']:05d}" in section_text and "day" in section_text)
check("it lists every open earlier visit, whatever its status: Waiting, and Waiting for the Doctor",
      all(NAMES[t] in section_text for t in ("D", "S")) and "Waiting for the Doctor" in section_text,
      [t for t in ("D", "S") if NAMES[t] not in section_text])
check("it does not list finished visits, left visits, or today's visits",
      not any(NAMES[t] in section_text for t in ("E", "F", "H", "J", "P", "K", "M", "A")))
check("the old title 'Still waiting from earlier days' is gone", "Still waiting from earlier days" not in page)
d_row = row_in(section_text, f"APT-{ids['D']:05d}")
check("Staff can mark a visit waiting for the Doctor as left before being seen from there, and have no "
      "'See patient'", 'data-bs-target="#leftModal"' in d_row and "See patient" not in d_row, d_row[-300:])
doctor_section = section_of(body(doctor.get("/staff/appointment_scheduling")))
d_row_doctor = row_in(doctor_section, f"APT-{ids['D']:05d}")
check("the Doctor sees 'See patient' on it, as well as Left before being seen",
      f"see-patient/{ids['D']}" in d_row_doctor and "See patient" in d_row_doctor
      and 'data-bs-target="#leftModal"' in d_row_doctor, d_row_doctor[-300:])
check("each earlier row has a Status column", "<th>Status</th>" in section_text
      and "badge-waiting-doctor" in d_row)
check("the dialog asks first, in plain words, and posts to mark-left",
      "Mark this patient as left before being seen?" in page and MARK in page)
admin_queue = body(admin.get("/admin/appointment_scheduling"))
admin_section = section_of(admin_queue)
check("the Administrator sees the section read-only: the visits, no Left button, no See patient, no Start, "
      "no dialog",
      NAMES["D"] in admin_section and "leftModal" not in admin_queue and "mark-left" not in admin_queue
      and "undoLeftForm" not in admin_queue and "See patient" not in admin_queue
      and "Start Consultation" not in admin_section and "sendToDoctorForm" not in admin_queue)
check("the Administrator's queue shows no service for any visit (it names a clinical fact)",
      not any(s in admin_queue for s in ("Prenatal Check-up", "Immunization", "Family Planning", "Dental", HTN, FOLLOW)))
# The counts: the dashboard, the Staff card and the queue (today plus earlier days) agree.
with app.app_context():
    from models.doctor_worklist import count_waiting_for_doctor as _count_wfd
    dashboard_count = _count_wfd()
today_wfd_rows = q1("SELECT COUNT(*) FROM appointments WHERE status = ? AND date(appointment_date) = ?",
                    (WFD, today.isoformat()))[0]
earlier_wfd_rows = q1("SELECT COUNT(*) FROM appointments WHERE status = ? AND date(appointment_date) < ?",
                      (WFD, today.isoformat()))[0]
queue_page = body(doctor.get("/staff/appointment_scheduling"))
today_summary = flat(queue_page.split('id="queueSummary"')[1].split("</div>")[0])
earlier_summary = flat(queue_page.split('id="earlierSummary"')[1].split("</div>")[0])
queue_today = int(re.search(r"Waiting for the Doctor (\d+)", today_summary).group(1))
queue_earlier = int(re.search(r"Waiting for the Doctor (\d+)", earlier_summary).group(1))
dash_text = flat(body(doctor.get("/staff/dashboard")))
check("the dashboard count equals the queue's open visits waiting for the Doctor, today plus earlier days",
      dashboard_count == queue_today + queue_earlier == today_wfd_rows + earlier_wfd_rows
      and earlier_wfd_rows > 0, (dashboard_count, queue_today, queue_earlier))
check("the Staff and Administrator dashboards show the same count",
      f"Waiting for the Doctor: {dashboard_count}" in flat(body(staff.get("/staff/dashboard")))
      and f"Waiting for the Doctor: {dashboard_count}" in flat(body(admin.get("/admin/dashboard"))))
check("the summary under Today's Queue starts with 'Today:' and the earlier-days line has its own counts",
      today_summary.lstrip("> ").startswith("Today: Waiting") and earlier_summary.lstrip("> ").startswith("Earlier days: Waiting")
      and "In Consultation" in earlier_summary)
open_now = {st: q1("SELECT COUNT(*) FROM appointments WHERE status = ? AND date(appointment_date) < ?",
                   (st, today.isoformat()))[0] for st in ("Waiting", "In Consultation", WFD)}
check("the earlier-days counts are the real counts by status",
      all(int(re.search(rf"{st} (\d+)", earlier_summary).group(1)) == n for st, n in open_now.items()),
      (earlier_summary, open_now))
today_wfd = add_visit(23, "W", 0, WFD, consultation=True, time="16:30")[0]
page = body(staff.get("/staff/appointment_scheduling"))
check("today's queue shows the new status as its own badge and counts it in its own footer figure",
      NAMES["W"] in page and "badge-waiting-doctor" in page and "waiting for the Doctor" in page)
run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (ids["B"],))

# ---------------------------------------------------------------------------
section("9. REPORTS: THE NEW STATUS IS ITS OWN FIGURE, 'COMPLETED' IS ONLY FINISHED VISITS (D6)")
rare_day = (today - timedelta(days=30))
while q1("SELECT COUNT(*) FROM appointments WHERE appointment_date = ?", (rare_day.isoformat(),))[0]:
    rare_day -= timedelta(days=1)


def rare(kind, status, n):
    for i in range(n):
        run_sql("INSERT INTO appointments (patient_id, scheduled_by, appointment_type, "
                "appointment_date, status, is_walkin, checked_in_at) VALUES (?, ?, ?, ?, ?, 1, ?)",
                (1 + i % 3, STAFF_ID, kind, rare_day.isoformat(), status,
                 f"{rare_day.isoformat()} 09:{i:02d}:00"))


rare("ZZRARE", LEFT, 3)
rare(CONSULT, "Completed", 1)
rare(CONSULT, WFD, 2)
rare("Prenatal Check-up", "Completed", 5)
from models.appointments import get_appointment_report, get_queue_stats, get_todays_appointments  # noqa: E402

for who, role in (("Administrator", "admin"), ("Barangay Official", "barangay official"),
                  ("Doctor", "doctor")):
    with app.app_context():
        window = get_appointment_report(months=6, start=rare_day, end=rare_day, viewer_role=role)["window"]
    check(f"{who}: 11 appointments: 6 attended (finished), 2 waiting for the Doctor, 3 left; the total "
          "is exact and the rate counts those waiting for the Doctor as having come",
          window["total"] == 11 and window["attended"] == 6 and window["awaiting"] == 2
          and window["left"] == 3 and window["missed"] == 0 and window["resolved"] == 11
          and window["attendance_pct"] == round(8 / 11 * 100, 1), window)
    names = {e["service"]: e["count"] for e in window["services"]}
    if role == "doctor":
        check("Doctor: 'Appointments by service' names every chosen service with its count",
              names.get("ZZRARE") == 3 and names.get(CONSULT) == 3 and names.get("Prenatal Check-up") == 5,
              names)
    else:
        check(f"{who}: services of fewer than 5 fold into Other; Prenatal Check-up (5) is named",
              names.get("Prenatal Check-up") == 5 and "ZZRARE" not in names and CONSULT not in names
              and any(n.startswith("Other") for n in names), names)
RANGE = f"?from={rare_day.isoformat()}&to={rare_day.isoformat()}"
for who, client, url in (("Administrator", admin, "/admin/reports_analytics"),
                         ("Barangay Official", official, "/official/reports_analytics"),
                         ("Doctor", doctor, "/staff/reports_analytics")):
    page = body(client.get(url + RANGE))
    check(f"{who}: Reports & Analytics shows 'Waiting for the Doctor (2)' and 'Left before being "
          "seen (3)' as their own figures",
          "Waiting for the Doctor (2)" in flat(page) and "Left before being seen (3)" in flat(page))
    xlsx = client.get(url + "/export.xlsx" + RANGE)
    book = load_workbook(io.BytesIO(xlsx.data), data_only=True)
    cells = [str(c) for sheet in book.worksheets for row in sheet.iter_rows(values_only=True)
             for c in row if c is not None]
    check(f"{who}: the Excel file has both figures",
          any(cells[i] == "Waiting for the Doctor" and cells[i + 1] == "2" for i in range(len(cells) - 1))
          and any(cells[i] == "Left before being seen" and cells[i + 1] == "3"
                  for i in range(len(cells) - 1)))
    check(f"{who}: the PDF names the new status",
          "Waiting for the Doctor" in pdf_text(client.get(url + "/export.pdf" + RANGE).data))
community = body(official.get("/official/community-reports" + RANGE))
check("Community Reports shows it as its own line",
      re.search(r"Waiting for the Doctor.{0,200}?<span class=\"fig-value\">2</span>", community, re.S)
      is not None)
staff.post(MARK, data={"appointment_id": str(spare[0])})        # one left today, for the figures
with app.app_context():
    entries_today = get_todays_appointments()
    stats = get_queue_stats(entries_today)
check("today's queue figures: 'completed' counts only Completed visits; waiting for the Doctor and "
      "left are their own figures and are not in the waiting count",
      stats["completed"] == sum(1 for e in entries_today if e["status"] == "Completed")
      and stats["waiting_doctor"] == sum(1 for e in entries_today if e["status"] == WFD)
      and stats["waiting_doctor"] >= 1 and stats["left"] >= 1
      and stats["waiting"] == sum(1 for e in entries_today if e["status"] == "Waiting"))
check("the queue's Completed card says it counts finished visits",
      "Finished today" in body(staff.get("/staff/appointment_scheduling")))
# "Appointments by service" reflects the service chosen at check-in.
run_sql("UPDATE appointments SET status = 'Completed' WHERE patient_id = ? AND status = 'Waiting'", (PATIENT,))
before_family = 0
with app.app_context():
    before_family = {e["service"]: e["count"] for e in get_appointment_report(
        months=6, viewer_role="doctor")["window"]["services"]}.get("Family Planning", 0)
staff.post(CHECK_IN, data={"patient_id": str(PATIENT), "reason": "Pills", "service": "Family Planning"})
with app.app_context():
    after_family = {e["service"]: e["count"] for e in get_appointment_report(
        months=6, viewer_role="doctor")["window"]["services"]}.get("Family Planning", 0)
check("a check-in with the service 'Family Planning' adds one to 'Appointments by service' for it",
      after_family == before_family + 1, (before_family, after_family))
run_sql("UPDATE appointments SET status = 'Completed' WHERE patient_id = ? AND status = 'Waiting'", (PATIENT,))
# ---------------------------------------------------------------------------
section("10. THE SCHEMA AND THE MIGRATION")
import tempfile  # noqa: E402
import shutil  # noqa: E402
from pathlib import Path  # noqa: E402
from migrations import migrate_add_left_before_seen as migration  # noqa: E402

SCHEMA = open("schema.sql", encoding="utf-8").read()
fresh = sqlite3.connect(":memory:")
fresh.executescript(SCHEMA)
fresh.execute("INSERT INTO users (username, password_hash, full_name, role) VALUES ('u','h','U','Admin')")
fresh.execute("INSERT INTO patients (patient_id, first_name, last_name, birth_date, sex) "
              "VALUES ('P','F','L','2000-01-01','Male')")
try:
    fresh.execute("INSERT INTO appointments (patient_id, scheduled_by, appointment_date, status) "
                  "VALUES (1,1,'2026-01-01',?)", (LEFT,))
    accepts = True
except sqlite3.IntegrityError:
    accepts = False
try:
    fresh.execute("INSERT INTO appointments (patient_id, scheduled_by, appointment_date, status) "
                  "VALUES (1,1,'2026-01-01','Bogus')")
    refuses = False
except sqlite3.IntegrityError:
    refuses = True
fresh.close()
check("schema.sql builds a table that accepts the new status and refuses a made-up one",
      accepts and refuses)
work = Path(tempfile.mkdtemp(prefix="idsleft_"))
old = work / "old.db"
con = sqlite3.connect(old)
# The table as it was before Task 6: the schema without the new status and columns.
def without_referrals(sql):
    """The schema as it was before Task 19: no 'Referred' status and no referral columns."""
    return (sql.replace(", 'Referred')", ")")
            .replace("    referral_destination TEXT,\n    referral_facility TEXT,\n"
                     "    referral_reason TEXT,\n    referred_by INTEGER,\n    referred_at TEXT,\n"
                     "    status_before_referred TEXT,\n", "")
            .replace("    FOREIGN KEY (left_by) REFERENCES users(id) ON DELETE SET NULL,\n"
                     "    FOREIGN KEY (referred_by) REFERENCES users(id) ON DELETE SET NULL\n",
                     "    FOREIGN KEY (left_by) REFERENCES users(id) ON DELETE SET NULL\n"))


older = (without_referrals(SCHEMA).replace(", 'Left before being seen')", ")")
         .replace("'Waiting for the Doctor', ", "")
         .replace("    left_by INTEGER,\n    left_at TEXT,\n    status_before_left TEXT,\n", "")
         .replace("    FOREIGN KEY (assigned_staff_id) REFERENCES users(id) ON DELETE SET NULL,\n"
                  "    FOREIGN KEY (left_by) REFERENCES users(id) ON DELETE SET NULL\n",
                  "    FOREIGN KEY (assigned_staff_id) REFERENCES users(id) ON DELETE SET NULL\n"))
block = older.split("CREATE TABLE appointments")[1].split("\n);")[0]
assert "left_by" not in block and "Left before being seen" not in block
con.executescript(older)
con.executescript("""
INSERT INTO users (username, password_hash, full_name, role) VALUES ('u','h','U','Admin');
INSERT INTO patients (patient_id, first_name, last_name, birth_date, sex) VALUES ('P','F','L','2000-01-01','Male');
INSERT INTO appointments (patient_id, scheduled_by, appointment_date, status) VALUES (1,1,'2026-01-01','Completed');
INSERT INTO appointments (patient_id, scheduled_by, appointment_date, status, notes) VALUES (1,1,'2026-01-02','Waiting','kept');
""")
con.commit()
before = [tuple(r) for r in con.execute("SELECT * FROM appointments ORDER BY id")]
con.close()
result = migration.run(old, backup_dir=work / "backups")
con = sqlite3.connect(old)
after = [tuple(r[:13]) for r in con.execute("SELECT * FROM appointments ORDER BY id")]
cols_before = len(before[0])
check("the migration backs up first, rebuilds the table and keeps every row and value",
      result["backup"] and Path(result["backup"]).exists() and result["rows"] == 2
      and not result["problems"] and after == before and cols_before == 13, result)
check("the rebuilt table has the three columns and accepts the new status",
      {"left_by", "left_at", "status_before_left"} <= {r[1] for r in con.execute(
          "PRAGMA table_info(appointments)")}
      and "Left before being seen" in con.execute(
          "SELECT sql FROM sqlite_master WHERE name = 'appointments'").fetchone()[0])
con.close()
again = migration.run(old, backup_dir=work / "backups")
check("a second run does nothing: no backup, no change",
      again["backup"] is None and len(list((work / "backups").iterdir())) == 1, again)
shutil.rmtree(work, ignore_errors=True)


# ---------------------------------------------------------------------------
section("11. THE MIGRATION FOR THE NEW STATUS (migrate_visit_service_and_doctor_status.py)")
from migrations import migrate_visit_service_and_doctor_status as m2  # noqa: E402

work2 = Path(tempfile.mkdtemp(prefix="idsdoc_"))
old2 = work2 / "old.db"
con = sqlite3.connect(old2)
before_status = (without_referrals(SCHEMA).replace("'Waiting for the Doctor', ", ""))
con.executescript(before_status)
con.executescript("""
INSERT INTO users (username, password_hash, full_name, role, is_active) VALUES ('u','h','U','Admin',1);
INSERT INTO patients (patient_id, first_name, last_name, birth_date, sex) VALUES ('P','F','L','2000-01-01','Male');
""")


def visit(day, status, kind, diagnosis="<none>", before=None):
    """One old visit. diagnosis '<none>' writes no consultation; None or '' a blank one."""
    cur = con.execute(
        "INSERT INTO appointments (patient_id, scheduled_by, appointment_type, appointment_date, status, "
        "status_before_left) VALUES (1, 1, ?, ?, ?, ?)", (kind, day, status, before))
    if diagnosis != "<none>":
        con.execute("INSERT INTO consultations (patient_id, recorded_by, appointment_id, "
                    "consultation_date, chief_complaint, diagnosis) VALUES (1, 1, ?, ?, 'x', ?)",
                    (cur.lastrowid, day, diagnosis))
    return cur.lastrowid


owed1 = visit("2026-01-01", "Completed", "Walk-in Consultation", None)        # converts
owed2 = visit("2026-01-02", "Completed", "Walk-in Consultation", "   ")       # blank counts as none
owed3 = visit("2026-01-03", "Completed", None, "")                             # no service needs him
diagnosed = visit("2026-01-04", "Completed", "Walk-in Consultation", "Seen")   # stays
immun = visit("2026-01-05", "Completed", "Immunization", None)                 # stays: not his service
closed = visit("2026-01-06", "Completed", "Walk-in Consultation")              # stays: nothing written
waiting = visit("2026-01-07", "Waiting", "Walk-in Consultation")               # stays
left_c = visit("2026-01-08", "Left before being seen", "Walk-in Consultation", None, "Completed")
left_w = visit("2026-01-09", "Left before being seen", "Walk-in Consultation", None, "Waiting")
con.commit()
before_rows = {r[0]: tuple(r) for r in con.execute("SELECT * FROM appointments")}
audit_before = con.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0]
con.close()
result = m2.run(old2, backup_dir=work2 / "backups")
con = sqlite3.connect(old2)
now_rows = {r[0]: tuple(r) for r in con.execute("SELECT * FROM appointments")}
state = {k: v[7] for k, v in now_rows.items()}
check("it backs up first and rebuilds the table with the wider CHECK",
      result["backup"] and Path(result["backup"]).exists() and not result["problems"]
      and "Waiting for the Doctor" in con.execute(
          "SELECT sql FROM sqlite_master WHERE name = 'appointments'").fetchone()[0], result)
check("it converts the Completed visits that have intake, no diagnosis and a service that needs the "
      "Doctor (a blank diagnosis and a blank service count)",
      all(state[i] == "Waiting for the Doctor" for i in (owed1, owed2, owed3)) and result["converted"] == 3,
      state)
check("it leaves alone a diagnosed visit, an immunization, a visit closed with nothing written, a "
      "Waiting visit and a left one",
      state[diagnosed] == "Completed" and state[immun] == "Completed" and state[closed] == "Completed"
      and state[waiting] == "Waiting" and state[left_c] == "Left before being seen"
      and state[left_w] == "Left before being seen" and result["closed_without_record"] == 1)
check("a visit marked left from Completed is set to return to Waiting for the Doctor on undo; one that "
      "was Waiting keeps that",
      now_rows[left_c][-1] == "Waiting for the Doctor" and now_rows[left_w][-1] == "Waiting"
      and result["converted_left"] == 1)
changed_cols = {(i, c) for i in before_rows for c in range(len(before_rows[i]))
                if before_rows[i][c] != now_rows[i][c]}
check("no other value of any row changed, and no row was lost",
      changed_cols == {(owed1, 7), (owed2, 7), (owed3, 7), (left_c, 15)}
      and len(now_rows) == len(before_rows), sorted(changed_cols))
rows_audit = con.execute("SELECT user_id, action, module, details FROM audit_logs").fetchall()
check("it writes ONE record of itself, with the counts, saying the script did it",
      len(rows_audit) == audit_before + 1 and rows_audit[-1][1] == m2.ACTION
      and "3 completed visits" in rows_audit[-1][3] and "migration script" in rows_audit[-1][3], rows_audit)
con.close()
again = m2.run(old2, backup_dir=work2 / "backups")
con = sqlite3.connect(old2)
check("a second run does nothing: no backup, no change, no second record",
      again["backup"] is None and again["converted"] == 0
      and len(list((work2 / "backups").iterdir())) == 1
      and con.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0] == audit_before + 1)
con.close()
fresh = sqlite3.connect(":memory:")
fresh.executescript(SCHEMA)
fresh.execute("INSERT INTO users (username, password_hash, full_name, role) VALUES ('u','h','U','Admin')")
fresh.execute("INSERT INTO patients (patient_id, first_name, last_name, birth_date, sex) "
              "VALUES ('P','F','L','2000-01-01','Male')")
try:
    fresh.execute("INSERT INTO appointments (patient_id, scheduled_by, appointment_date, status) "
                  "VALUES (1,1,'2026-01-01','Waiting for the Doctor')")
    accepts = True
except sqlite3.IntegrityError:
    accepts = False
check("schema.sql accepts the new status, and the migration finds nothing to do on a fresh database",
      accepts and not m2.needs_work(fresh))
fresh.close()
shutil.rmtree(work2, ignore_errors=True)


# ---------------------------------------------------------------------------
section("12. NO AUDIT ENTRY NAMES A VISIT'S SERVICE (L32)")
entries_all = [dict(r) for r in db().execute(
    "SELECT action, details FROM audit_logs WHERE module = 'Appointment Scheduling'")]
STATUS_PHRASE = "The visit is now In Consultation."


def names_service(text):
    """Whether an audit text names a service. Only the full status sentence is ignored."""
    return any(sv in (text or "").replace(STATUS_PHRASE, "") for sv in VISIT_SERVICES)


check("no entry of the Patient Queue names any service, whatever the action",
      entries_all and not [e for e in entries_all if names_service(e["details"])],
      [e for e in entries_all if names_service(e["details"])][:2])
check("the service check ignores only the full status sentence: a service name next to it is still flagged",
      not names_service(f"Corrected the service of visit APT-00001. {STATUS_PHRASE}")
      and all(names_service(f"Corrected the service of visit APT-00001 to {sv}. {STATUS_PHRASE}")
              for sv in VISIT_SERVICES)
      and names_service("In Consultation for Consultation") and names_service("Follow-up Consultation"))
check("the correction entries all read 'Corrected the service of visit APT-nnnnn.', alone or with 'The "
      "visit is now In Consultation.'",
      all(re.fullmatch(r"Corrected the service of visit APT-\d{5}\.( The visit is now In Consultation\.)?",
                       e["details"])
          for e in entries_all if e["action"] == "Changed the service of a visit"))
check("the entry for a visit moved by a service correction says so without naming the service",
      any(e["details"].endswith("The visit is now In Consultation.") for e in entries_all
          if e["action"] == "Changed the service of a visit"))

# ---------------------------------------------------------------------------
section("13. THE MIGRATION THAT REWORDS OLD ENTRIES (migrate_remove_service_names_from_audit.py)")
from migrations import migrate_remove_service_names_from_audit as m3  # noqa: E402

work3 = Path(tempfile.mkdtemp(prefix="idsaud_"))
old3 = work3 / "old.db"
con = sqlite3.connect(old3)
con.executescript(SCHEMA)
con.executescript("INSERT INTO users (username, password_hash, full_name, role, is_active) "
                  "VALUES ('u','h','U','Admin',1);")
rows3 = [
    ("Changed the service of a visit", "Changed the service of visit APT-00081 from Consultation to Immunization."),
    ("Changed the service of a visit", "Changed the service of visit APT-00082 from Prenatal Check-up to Consultation. "
                                      "The visit is now Completed."),
    ("Changed the service of a visit", "Corrected the service of visit APT-00083."),
    ("Marked as left before being seen", "Name (PAT-005), visit APT-00081."),
]
for action, details in rows3:
    con.execute("INSERT INTO audit_logs (user_id, action, module, details) VALUES (1, ?, 'Appointment Scheduling', ?)",
                (action, details))
con.commit()
before3 = [tuple(r) for r in con.execute("SELECT id, user_id, action, module, ip_address, created_at FROM audit_logs")]
con.close()
res3 = m3.run(old3, backup_dir=work3 / "backups")
con = sqlite3.connect(old3)
texts = [r[0] for r in con.execute("SELECT details FROM audit_logs ORDER BY id")]
check("it backs up first and rewords the two old entries, and only them",
      res3["backup"] and Path(res3["backup"]).exists() and res3["reworded"] == 2 and not res3["problems"]
      and texts[:4] == ["Corrected the service of visit APT-00081.",
                        "Corrected the service of visit APT-00082. The visit is now Completed.",
                        "Corrected the service of visit APT-00083.", "Name (PAT-005), visit APT-00081."], texts)
check("who, when, the action and the module of every entry are unchanged",
      [tuple(r) for r in con.execute("SELECT id, user_id, action, module, ip_address, created_at FROM audit_logs "
                                     "WHERE id <= 4 ORDER BY id")] == before3)
rec = con.execute("SELECT action, details FROM audit_logs WHERE id > 4").fetchall()
check("it writes ONE record of itself, with the count, saying the script did it",
      len(rec) == 1 and rec[0][0] == m3.ACTION and "2 earlier entries" in rec[0][1]
      and "migration script" in rec[0][1] and not any(s in rec[0][1] for s in VISIT_SERVICES), rec)
con.close()
res3b = m3.run(old3, backup_dir=work3 / "backups")
con = sqlite3.connect(old3)
check("a second run does nothing: no backup, no change, no second record",
      res3b["backup"] is None and res3b["reworded"] == 0 and len(list((work3 / "backups").iterdir())) == 1
      and con.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0] == 5)
con.close()
shutil.rmtree(work3, ignore_errors=True)

# ---------------------------------------------------------------------------
section("14. THE DOCTOR CAN SEE EVERY WAITING PATIENT, WITH OR WITHOUT AN INTAKE (Task 6d)")
QUEUE_URL, DASH_URL = "/staff/appointment_scheduling", "/staff/dashboard"


def row_of(page, marker, start='<tr'):
    """The table row on `page` that holds `marker` (an APT reference or a patient name)."""
    for chunk in page.split(start)[1:]:
        row = chunk.split("</tr>")[0]
        if marker in row:
            return row
    return ""


def see_patient_href(row):
    """The href of the 'See patient' link in a row, or None."""
    m = re.search(r'<a [^>]*href="([^"]+)"[^>]*>(?:(?!</a>).)*?See patient', row, re.S)
    return m.group(1) if m else None


# The bug: Staff send a Consultation visit to the Doctor with no intake.
name_patient(25, "T6D")
bug = new_visit(CONSULT, "no intake", pid=25)
staff.post(SEND, data={"appointment_id": str(bug)})
check("Task 6d set-up: Staff send a Consultation visit to the Doctor without an intake: it is "
      "Waiting for the Doctor and has no consultation record",
      status_of(bug) == WFD and consultation_of(bug) is None)
queue_row = row_of(body(doctor.get(QUEUE_URL)), f"APT-{bug:05d}")
check("the Doctor's Patient Queue offers 'See patient' on that row", "See patient" in queue_row, queue_row[-600:])
dash_row = row_of(body(doctor.get(DASH_URL)), NAMES["T6D"])
check("the Doctor's dashboard list offers 'See patient' on that row", "See patient" in dash_row, dash_row[-600:])
href = see_patient_href(queue_row)
r = doctor.get(href, follow_redirects=True) if href else None
check("it opens a new consultation form for this visit (the Doctor sees the diagnosis field)",
      r is not None and r.status_code == 200 and 'name="diagnosis"' in body(r)
      and f"/consultation/record/{bug}" in (r.request.path or ""), href)
r = doctor.post(f"/staff/consultation/record/{bug}", data=FULL)
check("saving it with a diagnosis completes the visit, and the record is linked to it",
      status_of(bug) == "Completed" and consultation_of(bug) is not None
      and consultation_of(bug)[1] == "Flow diagnosis", status_of(bug))
with app.app_context():
    check("... and it leaves the Doctor's list", bug not in listed_now())

# The same visit, with an intake record: "See patient" opens that record.
name_patient(24, "T6E")
withrec = new_visit(CONSULT, "with intake", pid=24)
staff.post(f"/staff/consultation/record/{withrec}", data=INTAKE)
rec_id = consultation_of(withrec)[0]
check("Task 6d set-up: Staff do the intake, then the visit is Waiting for the Doctor and has a record",
      status_of(withrec) == WFD and rec_id)
queue_row = row_of(body(doctor.get(QUEUE_URL)), f"APT-{withrec:05d}")
dash_row = row_of(body(doctor.get(DASH_URL)), f"see-patient/{withrec}")
check("the queue and the dashboard both offer the same 'See patient' link for it",
      see_patient_href(queue_row) and see_patient_href(queue_row) == see_patient_href(dash_row)
      and f"see-patient/{withrec}" in see_patient_href(queue_row), (queue_row[-300:], dash_row[-300:]))
r = doctor.get(see_patient_href(queue_row))
check("it opens that record, not a new consultation",
      r.status_code == 302 and r.headers["Location"].endswith(f"/staff/consultation/edit/{rec_id}"),
      r.headers.get("Location"))
r = doctor.get(see_patient_href(queue_row), follow_redirects=True)
check("the Doctor sees the saved intake and the diagnosis field",
      r.status_code == 200 and 'name="diagnosis"' in body(r) and "flow intake" in body(r))
doctor.post(f"/staff/consultation/edit/{rec_id}", data=FULL)
check("saving the diagnosis completes the visit; there is still one record for it",
      status_of(withrec) == "Completed" and consultation_of(withrec)[1] == "Flow diagnosis"
      and q1("SELECT COUNT(*) FROM consultations WHERE appointment_id = ?", (withrec,))[0] == 1)
check("the first visit's new record is linked to it and holds what the Doctor typed",
      tuple(q1("SELECT chief_complaint, diagnosis FROM consultations WHERE appointment_id = ?", (bug,)))
      == ("flow intake", "Flow diagnosis"))
# Staff do not get the button, and the link is refused for them.
name_patient(22, "T6F")
staff_visit = new_visit(CONSULT, "staff view", pid=22)
staff.post(SEND, data={"appointment_id": str(staff_visit)})
check("Staff do not get 'See patient' on the queue", "See patient" not in
      row_of(body(staff.get(QUEUE_URL)), f"APT-{staff_visit:05d}"))
r = staff.get(f"/staff/appointment_scheduling/see-patient/{staff_visit}")
check("... and the link is refused for them on the server",
      r.status_code == 302 and r.headers["Location"].endswith("/staff/appointment_scheduling")
      and status_of(staff_visit) == WFD)
run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (staff_visit,))

# ---------------------------------------------------------------------------
section("15. EVERY STATUS x ROLE: THE QUEUE SHOWS EXACTLY THE MAPPING'S ACTIONS AND EACH ONE WORKS (Task 6d, D3)")
from models.queue_actions import ACTIONS, QUEUE_ACTIONS, allowed_actions  # noqa: E402

START_URL = "/staff/appointment_scheduling/start-consultation"
CHANGE = "/staff/appointment_scheduling/change-service"
PRENATAL = "Prenatal Check-up"
WAIT, INCONS = "Waiting", "In Consultation"
CARE = ("staff", "doctor")
# The decided mapping, written out here on purpose: changing queue_actions.py changes the
# product, so it must change this table too.
DECIDED = {
    WAIT: {"staff": {"start", "mark_left", "refer", "change_service"},
           "doctor": {"start", "mark_left", "refer", "change_service"}},
    INCONS: {"staff": {"open_intake", "send_to_doctor", "mark_completed", "mark_left", "refer", "change_service"},
             "doctor": {"open_intake", "send_to_doctor", "mark_completed", "mark_left", "refer", "change_service"}},
    WFD: {"staff": {"mark_left", "refer", "change_service"},
          "doctor": {"see_patient", "mark_left", "refer", "change_service"}},
    "Completed": {},
    LEFT: {"staff": {"undo_left"}, "doctor": {"undo_left"}},
    "Referred": {"staff": {"referral_slip", "undo_referral"}, "doctor": {"referral_slip", "undo_referral"}},
}
ALL_STATUSES = (WAIT, INCONS, WFD, "Completed", LEFT, "Referred", "Scheduled", "Confirmed", "Cancelled", "No-show")
mapping_diffs = [(st, role, allowed_actions(role, st), DECIDED.get(st, {}).get(role, set()))
                 for st in ALL_STATUSES for role in ("staff", "doctor", "admin", "official")
                 if set(allowed_actions(role, st)) != DECIDED.get(st, {}).get(role, set())]
check("the mapping is the decided one for every status and role; the Administrator and the Official "
      "have no action anywhere", not mapping_diffs and set(QUEUE_ACTIONS) == set(DECIDED), mapping_diffs[:4])

# What the page should show, for the visit as built: a Consultation visit with no diagnosis,
# a Left visit marked today, the room free.
CASES = [
    ("Waiting", dict(status=WAIT), DECIDED[WAIT]),
    ("In Consultation (Consultation)", dict(status=INCONS),
     {r: DECIDED[INCONS][r] - {"mark_completed"} for r in CARE}),
    ("In Consultation (Dental)", dict(status=INCONS, kind="Dental"),
     {r: DECIDED[INCONS][r] - {"send_to_doctor"} for r in CARE}),
    ("Waiting for the Doctor, no intake", dict(status=WFD), DECIDED[WFD]),
    ("Waiting for the Doctor, with intake", dict(status=WFD, intake=True), DECIDED[WFD]),
    ("Completed", dict(status="Completed"), {}),
    ("Left today", dict(status=LEFT, before=WAIT, left_day=0), DECIDED[LEFT]),
    ("Left on an earlier day (undo is for the same day)", dict(status=LEFT, before=WAIT, left_day=1), {}),
    ("Referred today", dict(status="Referred", before=WAIT, ref_day=0), DECIDED["Referred"]),
    ("Referred on an earlier day (undo is for the same day)",
     dict(status="Referred", before=WAIT, ref_day=1), {r: {"referral_slip"} for r in CARE}),
]
MATRIX_PID = 21
ROLE_CLIENTS = (("staff", staff, QUEUE_URL), ("doctor", doctor, QUEUE_URL),
                ("admin", admin, "/admin/appointment_scheduling"), ("official", official, None))


def build(spec, tag):
    """One visit in the described state, with the room free. Returns (id, record id)."""
    run_sql("UPDATE appointments SET status = 'Completed' WHERE status = 'In Consultation'")
    visit_id, record = add_visit(MATRIX_PID, tag, 0, spec["status"], kind=spec.get("kind", CONSULT),
                                 consultation=spec.get("intake"))
    if spec["status"] == LEFT:
        left_at = f"{(today - timedelta(days=spec['left_day'])).isoformat()} 07:00:00"
        run_sql("UPDATE appointments SET status_before_left = ?, left_by = ?, left_at = ? WHERE id = ?",
                (spec["before"], STAFF_ID, left_at, visit_id))
    if spec["status"] == "Referred":
        referred_at = f"{(today - timedelta(days=spec['ref_day'])).isoformat()} 07:00:00"
        run_sql("UPDATE appointments SET status_before_referred = ?, referral_destination = ?, "
                "referred_by = ?, referred_at = ? WHERE id = ?",
                (spec["before"], "Rural Health Unit (RHU)", STAFF_ID, referred_at, visit_id))
    return visit_id, record


def state_of(visit_id):
    row = q1("SELECT status, appointment_type, "
             "(SELECT COUNT(*) FROM consultations WHERE appointment_id = a.id) AS records "
             "FROM appointments a WHERE id = ?", (visit_id,))
    return tuple(row)


def actions_in(row, visit_id):
    """The actions a queue row offers, read off its markup."""
    marks = {
        "start": "Start Consultation",
        "open_intake": f"/staff/consultation/record/{visit_id}",
        "send_to_doctor": 'form="sendToDoctorForm"',
        "mark_completed": 'form="completeConsultationForm"',
        "see_patient": f"/see-patient/{visit_id}",
        "mark_left": 'data-bs-target="#leftModal"',
        "undo_left": 'form="undoLeftForm"',
        "change_service": 'data-bs-target="#serviceModal"',
        "refer": 'data-bs-target="#referModal"',
        "undo_referral": 'form="undoReferralForm"',
        "referral_slip": f"/referral-slip/{visit_id}",
    }
    return {action for action, mark in marks.items() if mark in row}


page_problems, case_n = [], 0
for label, spec, expected in CASES:
    visit_id, _ = build(spec, f"M{case_n}")
    case_n += 1
    for role, client, url in ROLE_CLIENTS:
        if url is None:
            r = client.get("/staff/appointment_scheduling")
            if r.status_code != 302:
                page_problems.append((label, role, "the Official reached the queue", r.status_code))
            continue
        row = row_of(body(client.get(url)), f"APT-{visit_id:05d}")
        shown = actions_in(row, visit_id)
        want = expected.get(role, set())
        if not row or shown != want:
            page_problems.append((label, role, sorted(shown), sorted(want)))
check(f"the queue shows exactly the mapping's actions for all {len(CASES)} visit states x Staff, the Doctor "
      "and the Administrator, and the Official has no queue", not page_problems, page_problems[:6])
primary_visit = build(CASES[3][1], "MP")[0]
check("... with the Doctor's 'See patient' as the filled (primary) button",
      'class="btn-primary-teal"' in row_of(body(doctor.get(QUEUE_URL)), f"/see-patient/{primary_visit}"))


def perform(action, client, visit_id):
    data = {"appointment_id": str(visit_id)}
    if action == "start":
        return client.post(START_URL, data=data)
    if action == "send_to_doctor":
        return client.post(SEND, data=data)
    if action == "mark_completed":
        return client.post(COMPLETE, data=data)
    if action == "mark_left":
        return client.post(MARK, data=data)
    if action == "undo_left":
        return client.post(UNDO, data=data)
    if action == "change_service":
        return client.post(CHANGE, data=dict(data, service=PRENATAL))
    if action == "see_patient":
        return client.get(f"/staff/appointment_scheduling/see-patient/{visit_id}")
    if action == "refer":
        return client.post("/staff/appointment_scheduling/refer",
                           data=dict(data, destination="Rural Health Unit (RHU)"))
    if action == "undo_referral":
        return client.post("/staff/appointment_scheduling/undo-referral", data=data)
    if action == "referral_slip":
        return client.get(f"/staff/appointment_scheduling/referral-slip/{visit_id}")
    return client.get(f"/staff/consultation/record/{visit_id}")      # open_intake


def worked(action, response, visit_id, before, spec, client):
    """Whether the action did what it is for, on the visit built from `spec`."""
    now_state = state_of(visit_id)
    if action == "referral_slip":
        return (now_state == before and response.status_code == 200
                and 'id="referralSlip"' in body(response))
    if action in ("open_intake", "see_patient"):
        if now_state != before:
            return False
        if action == "see_patient":
            where = response.headers.get("Location", "/")
            follow = client.get(where, follow_redirects=True)
            return (response.status_code == 302 and follow.status_code == 200
                    and 'name="diagnosis"' in body(follow)
                    and ("/consultation/edit/" in where) == bool(spec.get("intake")))
        return response.status_code == 200 and 'name="chief_complaint"' in body(response)
    goal = {"start": lambda s: s[0] == INCONS, "send_to_doctor": lambda s: s[0] == WFD,
            "mark_completed": lambda s: s[0] == "Completed", "mark_left": lambda s: s[0] == LEFT,
            "undo_left": lambda s: s[0] == WAIT, "change_service": lambda s: s[1] == PRENATAL,
            "refer": lambda s: s[0] == "Referred", "undo_referral": lambda s: s[0] == WAIT}[action]
    return response.status_code == 302 and goal(now_state)


work_problems, runs = [], 0
for label, spec, expected in CASES:
    for role, client, _ in ROLE_CLIENTS:
        for action in ACTIONS:
            visit_id, record = build(spec, f"W{runs}")
            runs += 1
            before = state_of(visit_id)
            response = perform(action, client, visit_id)
            offered = action in expected.get(role, set())
            if action == "open_intake":
                # The record page is also where "See patient" lands when there is no record.
                offered = offered or ("see_patient" in expected.get(role, set()) and not spec.get("intake"))
            if offered:
                if not worked(action, response, visit_id, before, spec, client):
                    work_problems.append((label, role, action, "offered, but it did not work",
                                          response.status_code, state_of(visit_id)))
            elif state_of(visit_id) != before or (
                    action in ("open_intake", "see_patient", "referral_slip") and response.status_code == 200):
                work_problems.append((label, role, action, "not offered, but it was accepted",
                                      response.status_code, state_of(visit_id)))
check(f"each action the queue shows works for that role, and every other one is refused with no change "
      f"({runs} tries: {len(CASES)} states x 4 roles x {len(ACTIONS)} actions)", not work_problems,
      work_problems[:6])
run_sql("UPDATE appointments SET status = 'Completed' WHERE status = 'In Consultation'")
# The model's own rules (the SQL behind the routes) agree with the mapping, not only the routes do.
from models.appointments import change_visit_service, mark_left  # noqa: E402

model_problems = []
for label, spec, expected in CASES:
    for name, call, action in (
            ("mark_left", lambda v: mark_left(v, STAFF_ID), "mark_left"),
            ("change_visit_service", lambda v: change_visit_service(v, PRENATAL), "change_service")):
        visit_id, _ = build(spec, f"D{runs}")
        runs += 1
        with app.app_context():
            outcome = call(visit_id)
        accepted = outcome in ("marked", "changed", "changed-in-consultation")
        if accepted != (action in expected.get("staff", set())):
            model_problems.append((label, name, outcome))
check("the model's own rules for 'left before being seen' and 'change service' accept exactly the statuses "
      "the mapping allows, called directly", not model_problems, model_problems)
run_sql("UPDATE appointments SET status = 'Completed' WHERE status = 'In Consultation'")

# ---------------------------------------------------------------------------
section("16. THE LEGEND, THE SUMMARY LINE, THE OVERDUE BADGE AND THE WALK-IN TAG (Task 6d, D4)")
from models.immunizations import get_immunization_status  # noqa: E402
from models.patients import get_patient_by_id  # noqa: E402

STATUS_NAMES = ["Waiting", "In Consultation", "Waiting for the Doctor", "Completed", LEFT, "Referred"]
queue_pages = {"staff": body(staff.get(QUEUE_URL)), "doctor": body(doctor.get(QUEUE_URL)),
               "admin": body(admin.get("/admin/appointment_scheduling"))}


def legend_block(page):
    return page.split('id="statusLegend"')[1].split("</div>\n        </div>")[0]


def summary_text(page):
    return re.sub(r"\s+", " ", flat(page.split('id="queueSummary"')[1].split("</div>")[0]))


legends = {role: [re.sub(r"\s+", " ", flat(item)).strip()
                  for item in re.findall(r'<div class="legend-item">(.*?)</div>', legend_block(page), re.S)]
           for role, page in queue_pages.items()}
check("the colour legend lists the six real statuses, in order, for every role",
      all(items == STATUS_NAMES for items in legends.values()), legends)
dots = re.findall(r'<span class="cal-dot (dot-[a-z-]+)"></span>[A-Za-z ]+</div>', legend_block(queue_pages["staff"]))
check("... each with its own colour, and each colour is defined on the page",
      len(dots) == 6 and len(set(dots)) == 6
      and all(f".{d}{{background:" in queue_pages["staff"] for d in dots), dots)
for role, page in queue_pages.items():
    text = " ".join(legends[role]) + " " + summary_text(page)
    check(f"{role}: no 'Scheduled', 'not yet arrived', 'Cancelled or missed' or 'ready to consult' in the "
          "legend or the summary line",
          not any(word.lower() in text.lower() for word in
                  ("Scheduled", "not yet arrived", "Cancelled or missed", "Arrived /", "ready to consult")), text)
    pairs = re.findall(r"(Waiting for the Doctor|Left before being seen|In Consultation|Completed|Referred|Waiting) (\d+)", summary_text(page))
    counts_now = {st: q1("SELECT COUNT(*) FROM appointments WHERE date(appointment_date) = "
                         "date('now', 'localtime') AND status = ?", (st,))[0] for st in STATUS_NAMES}
    check(f"{role}: the summary line uses the six status names, in the legend's order, with today's counts",
          [name for name, _ in pairs] == STATUS_NAMES and all(int(n) == counts_now[name] for name, n in pairs),
          (summary_text(page), counts_now))
check("the queue and the Staff dashboard have no Walk-in tag on any row (every visit is a walk-in)",
      all('class="walkin-tag"' not in page and ">Walk-in</span>" not in page for page in queue_pages.values())
      and ">Walk-in</span>" not in body(staff.get("/staff/dashboard")))
# The no-record badge counts what its words say: a child with doses with no record AND due now.
CHD = q1("SELECT MAX(id) + 1 FROM patients")[0]
run_sql("INSERT INTO patients (id, patient_id, first_name, last_name, birth_date, sex, civil_status, "
        "address, municipality, province, contact_number, status) VALUES (?, ?, 'Chd', 'Patient', "
        "'2026-01-01', 'Female', 'Single', 'Purok 1', 'San Francisco', 'Surigao del Norte', "
        "?, 'Active')", (CHD, f"PAT-{CHD:03d}", f"0900000{CHD:04d}"))
name_patient(CHD, "Chd")
found = {}
with app.app_context():
    for days in range(20, 700):
        run_sql("UPDATE patients SET birth_date = ? WHERE id = ?", ((today - timedelta(days=days)).isoformat(), CHD))
        st = get_immunization_status(get_patient_by_id(CHD))
        key = ("both" if st["no_record_count"] and st["due_count"]
               else "no_record" if st["no_record_count"] else "due" if st["due_count"] else None)
        if key and key not in found:
            found[key] = (days, st["no_record_count"], st["due_count"])
check("set-up: ages exist with doses with no record and due now, with no record only, and with due only",
      {"both", "no_record", "due"} <= set(found), found)


def badge_for(days):
    run_sql("UPDATE patients SET birth_date = ? WHERE id = ?", ((today - timedelta(days=days)).isoformat(), CHD))
    run_sql("UPDATE appointments SET status = 'Completed' WHERE patient_id = ?", (CHD,))
    visit_id = add_visit(CHD, "Chd", 0, "Waiting")[0]
    row = row_of(body(staff.get(QUEUE_URL)), f"APT-{visit_id:05d}")
    run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (visit_id,))
    m = re.search(r'<span class="immun-tag[^"]*"[^>]*>(?:<span[^>]*>[^<]*</span>)?([^<]*)</span>', row)
    return (m.group(1).strip() if m else None), row


def plural(n, word):
    return f"{n} {word}{'' if n == 1 else 's'}"


days, over, due = found["both"]
text_both, row_both = badge_for(days)
check("a child with doses with no record and due now: the badge says only how many vaccines have no "
      f"record ('{plural(over, 'vaccine')}: no record'), not that number plus the due ones",
      text_both == f"{plural(over, 'vaccine')}: no record" and due > 0 and "immun-norecord" in row_both,
      (text_both, found["both"]))
days, over, due = found["no_record"]
text_over = badge_for(days)[0]
check("only doses with no record: '<n> vaccines: no record'", text_over == f"{plural(over, 'vaccine')}: no record", text_over)
days, over, due = found["due"]
text_due, row_due = badge_for(days)
check("only doses due now: it says 'due', not 'no record'",
      text_due == f"{plural(due, 'vaccine')}: due" and "immun-norecord" not in row_due, text_due)

# ---------------------------------------------------------------------------
section("17. A SERVICE CORRECTION NEVER COMPLETES A VISIT; THE BADGE SAYS 'NEXT IN LINE' (Task 6e)")
NON_DOCTOR = [sv for sv in VISIT_SERVICES if sv not in ("Consultation", "Follow-up Consultation",
                                                          "Hypertension / Diabetes Monitoring")]
for label, spec in (("no intake", dict(status=WFD)), ("with intake", dict(status=WFD, intake=True))):
    for role, client in (("Staff", staff), ("the Doctor", doctor)):
        v, _ = build(spec, f"C{label[:2]}{role[:2]}")
        n_audit = len(audits("Changed the service of a visit"))
        client.post(CHANGE, data={"appointment_id": str(v), "service": "Immunization"})
        entry = (audits("Changed the service of a visit") or [{}])[-1]
        check(f"{role}: Waiting for the Doctor ({label}) corrected to Immunization is In Consultation, "
              "not Completed; one entry",
              status_of(v) == "In Consultation" and len(audits("Changed the service of a visit")) == n_audit + 1
              and entry.get("details", "").endswith("The visit is now In Consultation."), status_of(v))
v, _ = build(dict(status=WFD), "Cdoc")
staff.post(CHANGE, data={"appointment_id": str(v), "service": FOLLOW})
check("corrected to a service that does need the Doctor, it stays Waiting for the Doctor (as before)",
      status_of(v) == WFD)
# The room is held by someone else: the move is refused and nothing changes.
v, _ = build(dict(status=WFD), "Cbusy")
other, _ = add_visit(MATRIX_PID, "Cbusy2", 0, "In Consultation")
n_audit = len(audits("Changed the service of a visit"))
staff.post(CHANGE, data={"appointment_id": str(v), "service": "Family Planning"})
check("with the room busy, the correction is refused: status, service and audit unchanged",
      state_of(v)[:2] == (WFD, CONSULT) and len(audits("Changed the service of a visit")) == n_audit,
      state_of(v))
run_sql("UPDATE appointments SET status = 'Completed' WHERE id IN (?, ?)", (v, other))
# No route completes a visit through a service change: every status x service x role.
bad = []
for status in ("Waiting", "In Consultation", WFD):
    for service in VISIT_SERVICES:
        for role, client in (("staff", staff), ("doctor", doctor), ("admin", admin), ("official", official)):
            v, _ = build(dict(status=status, kind="Immunization" if service != "Immunization" else CONSULT),
                         "Call")
            client.post(CHANGE, data={"appointment_id": str(v), "service": service})
            if status_of(v) == "Completed":
                bad.append((status, service, role))
check("no status x service x role of a service change leaves a visit Completed "
      f"(3 x {len(VISIT_SERVICES)} x 4)", not bad, bad[:4])
run_sql("UPDATE appointments SET status = 'Completed' WHERE status = 'In Consultation'")
from models.appointments import change_visit_service as _change  # noqa: E402

v, _ = build(dict(status=WFD), "Cmodel")
with app.app_context():
    outcome = _change(v, "Family Planning")
check("the model itself does it too, called directly",
      outcome == "changed-in-consultation" and status_of(v) == "In Consultation", outcome)
run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (v,))
# The badge.
run_sql("UPDATE appointments SET status = 'Completed' WHERE status IN ('Waiting', 'In Consultation')")
first, _ = build(dict(status="Waiting"), "Bnext")
second, _ = add_visit(MATRIX_PID, "Bnext2", 0, "Waiting", time="23:00")
queue_html = body(staff.get(QUEUE_URL))
badges = re.findall(r'<span class="badge-status [^"]*">([^<]+)</span>', queue_html)
check("the first Waiting patient's badge reads 'Next in line'; the second reads 'Waiting'",
      "Next in line" in row_of(queue_html, f"APT-{first:05d}")
      and ">Waiting<" in row_of(queue_html, f"APT-{second:05d}")
      and "Ready to Consult" not in queue_html, badges)
check("every badge on the queue is one of the six statuses or 'Next in line'",
      set(badges) <= set(STATUS_NAMES) | {"Next in line"}, set(badges))
run_sql("UPDATE appointments SET status = 'Completed' WHERE id IN (?, ?)", (first, second))

# ---------------------------------------------------------------------------
section("18. OPEN VISITS, THE SUMMARY LINES AND AMBER REMINDERS (Task 10)")
# An earlier-day visit Waiting for the Doctor, set up on its own, with and without an intake.
name_patient(21, "T10A")
name_patient(22, "T10B")
earlier_no = add_visit(21, "T10A", 4, WFD, time="08:00")[0]
earlier_rec = add_visit(22, "T10B", 3, WFD, consultation=True, time="08:30")[0]
run_sql("UPDATE appointments SET status = 'Completed' WHERE status = 'In Consultation'")
queue = {name: body(client.get(url)) for name, client, url in (
    ("staff", staff, "/staff/appointment_scheduling"), ("doctor", doctor, "/staff/appointment_scheduling"),
    ("admin", admin, "/admin/appointment_scheduling"))}
rows = {name: {i: row_in(section_of(page), f"APT-{i:05d}") for i in (earlier_no, earlier_rec)}
        for name, page in queue.items()}
check("an earlier-day visit Waiting for the Doctor is in 'Open visits from earlier days', with and without an intake",
      all(rows["doctor"][i] and rows["staff"][i] and rows["admin"][i] for i in (earlier_no, earlier_rec)))
check("the Doctor sees 'See patient' on both, and the link is the same one the dashboard uses",
      all(f"see-patient/{i}" in rows["doctor"][i] and "See patient" in rows["doctor"][i]
          for i in (earlier_no, earlier_rec))
      and f"see-patient/{earlier_no}" in body(doctor.get("/staff/dashboard")))
check("Staff see 'Left before being seen' on both, and no 'See patient'",
      all('data-bs-target="#leftModal"' in rows["staff"][i] and "See patient" not in rows["staff"][i]
          for i in (earlier_no, earlier_rec)))
check("the Administrator sees neither a Left button nor See patient on them, and no check-in reason column",
      all('data-bs-target="#leftModal"' not in rows["admin"][i] and "See patient" not in rows["admin"][i]
          and "see-patient" not in rows["admin"][i] for i in (earlier_no, earlier_rec))
      and "Service" not in section_of(queue["admin"]).split("</thead>")[0])
r = staff.post(MARK, data={"appointment_id": str(earlier_no)})
check("Staff can mark it left from there, and the server accepts it",
      status_of(earlier_no) == LEFT)
staff.post(UNDO, data={"appointment_id": str(earlier_no)})
check("(a left visit from an earlier day is closed: the same-day undo does not apply)",
      status_of(earlier_no) in (LEFT, WFD))
run_sql("UPDATE appointments SET status = ?, status_before_left = NULL, left_by = NULL, left_at = NULL "
        "WHERE id = ?", (WFD, earlier_no))
# The dashboard count equals the queue: today plus earlier days, for every open waiting-for-the-Doctor visit.
page = body(doctor.get("/staff/appointment_scheduling"))
t_n = int(re.search(r"Waiting for the Doctor (\d+)", flat(page.split('id="queueSummary"')[1].split("</div>")[0])).group(1))
e_n = int(re.search(r"Waiting for the Doctor (\d+)", flat(page.split('id="earlierSummary"')[1].split("</div>")[0])).group(1))
with app.app_context():
    from models.doctor_worklist import count_waiting_for_doctor as _cwd, get_waiting_for_doctor as _gwd
    listed = len(_gwd())
    counted = _cwd()
check("the Doctor's list, the dashboard count and the queue (today + earlier days) all agree",
      listed == counted == t_n + e_n, (listed, counted, t_n, e_n))
check("the summary starts with 'Today:' and the earlier-days line counts by status",
      "Today: Waiting" in flat(page) and "Earlier days: Waiting" in flat(page))
# No leftover boxes or tags.
leftovers = ("does not do", "From an earlier day", "Still waiting from earlier days", "advance_booking",
             "Visits cannot be booked ahead")
check("no 'does not do' box, no 'From an earlier day' tag and no booking note on the queue or the dashboards",
      not [w for page in list(queue.values()) + [body(doctor.get("/staff/dashboard")), body(staff.get("/staff/dashboard"))]
           for w in leftovers if w in page])
# Amber, not red, for every immunization reminder.
css = re.search(r"\.immun-tag\.immun-norecord\{[^}]*\}", queue["staff"]).group(0)
check("the queue badge for 'no record' is amber (#FEF3C7 / #92400E), not red",
      "#FEF3C7" in css and "#92400E" in css and "#FEE2E2" not in css and "#B91C1C" not in css, css)
profile = body(staff.get(f"/staff/patient_profile/17"))
alert = profile.split("to follow up")[0].rsplit("<div style=", 1)[-1]
check("the profile alert is amber, not red",
      "#FFFBEB" in alert and "#FEF2F2" not in alert and "#FECACA" not in alert and "#B91C1C" not in alert, alert[-250:])
dash = body(staff.get("/staff/dashboard"))
label_at = dash.index("Children Due for Immunization")
card = dash[dash.rfind('<div class="stat-card">', 0, label_at):label_at + 600].split("</div>\n      </div>")[0]
follow = dash.split("Immunization Follow-up")[1].split("Immunization Follow-up")[0]
follow = follow[:follow.find("panel-card")] if "panel-card" in follow else follow[:2500]
check("the Staff dashboard's immunization card and list use amber, not red",
      "#FEE2E2" not in card and "var(--danger)" not in card and "#FEE2E2" not in follow
      and "var(--danger)" not in follow, card[-300:])
check("the dashboard card says it in plain words: '<N> children have doses with no record', and has no "
      "'doses to follow up' figure",
      re.search(r"\d+ child(?:ren)? ha(?:s|ve) doses with no record", flat(dash)) is not None
      and "to follow up" not in flat(dash).split("Children Due for Immunization")[1][:200])
check("the Administrator's dashboard says 'with no record' in amber, not red",
      "with no record" in body(admin.get("/admin/dashboard"))
      and "var(--danger); font-weight:600;\">" not in body(admin.get("/admin/dashboard")).split("with no record")[0][-120:])
run_sql("UPDATE appointments SET status = 'Completed' WHERE id IN (?, ?)", (earlier_no, earlier_rec))

# ---------------------------------------------------------------------------
section("19. NO BOOKING-ERA FIGURE ANYWHERE; 'OLDER RECORDS' ONLY WHEN SUCH ROWS EXIST (Task 11, L48)")
BOOKING_WORDS = ("Still expected", "Cancelled or no-show", "Cancelled / no-show", "not yet arrived",
                 "Never closed", "Missed rate", "earlier booking system")
OLD_STATUSES = ("Scheduled", "Confirmed", "Cancelled", "No-show")
old_now = q1("SELECT COUNT(*) FROM appointments WHERE status IN ('Scheduled','Confirmed','Cancelled','No-show')")[0]
SCREENS = [(doctor, "/staff/reports_analytics"), (admin, "/admin/reports_analytics"),
           (official, "/official/reports_analytics"), (official, "/official/community-reports"),
           (staff, "/staff/dashboard"), (doctor, "/staff/dashboard"), (admin, "/admin/dashboard"),
           (official, "/official/dashboard"), (staff, "/staff/appointment_scheduling"),
           (admin, "/admin/appointment_scheduling"), (doctor, "/staff/appointment_scheduling")]
FILES = [(doctor, "/staff/reports_analytics/export"), (admin, "/admin/reports_analytics/export"),
         (official, "/official/reports_analytics/export"), (official, "/official/community-reports/export"),
         (staff, "/staff/appointment_scheduling/export"), (admin, "/admin/appointment_scheduling/export")]


def booking_words_found():
    found = []
    for client, url in SCREENS:
        text = body(client.get(url))
        found += [(url, w) for w in BOOKING_WORDS if w in text]
    for client, stem in FILES:
        for ext, convert in ((".xlsx", lambda d: "\n".join(
                str(c.value) for sh in load_workbook(io.BytesIO(d)) for row in sh.iter_rows() for c in row
                if c.value is not None)), (".pdf", pdf_text)):
            text = convert(client.get(stem + ext).get_data())
            found += [(stem + ext, w) for w in BOOKING_WORDS if w in text]
    return found


def older_everywhere():
    """Whether 'Older records' shows on the reports, the community report and the files."""
    return {
        "reports": "Older records: " in body(doctor.get("/staff/reports_analytics")),
        "community": "Older records" in body(official.get("/official/community-reports")),
        "excel": "Older records" in "\n".join(
            str(c.value) for sh in load_workbook(io.BytesIO(official.get("/official/community-reports/export.xlsx").get_data()))
            for row in sh.iter_rows() for c in row if c.value is not None),
    }


found = booking_words_found()
check(f"none of the booking-era figures is on any of {len(SCREENS)} screens or in any of {len(FILES) * 2} Excel and "
      "PDF files (Still expected, Cancelled / no-show, not yet arrived, Never closed, Missed rate, booking system)",
      not found, found[:5])
check("with no old record in the database, 'Older records' is not shown anywhere",
      old_now == 0 and not any(older_everywhere().values()), (old_now, older_everywhere()))
for i, status in enumerate(("No-show", "Cancelled", "Scheduled")):
    run_sql("INSERT INTO appointments (patient_id, scheduled_by, appointment_type, appointment_date, status, is_walkin) "
            "VALUES (?, ?, 'Consultation', ?, ?, 0)",
            (1 + i, STAFF_ID, (today - timedelta(days=20)).isoformat(), status))
shown = older_everywhere()
check("with old records in the window, 'Older records: N' shows on the reports, the community report and the files",
      all(shown.values()), shown)
check("... as one line, with the count of those rows (3)",
      "Older records: 3" in body(doctor.get("/staff/reports_analytics")))
found = booking_words_found()
check("... and the booking-era figures are still not shown", not found, found[:5])
run_sql("DELETE FROM appointments WHERE status IN ('No-show', 'Cancelled', 'Scheduled', 'Confirmed')")
check("once they are gone the line is gone", not any(older_everywhere().values()))

# ---------------------------------------------------------------------------
section("20. START ONLY ON A VISIT CHECKED IN TODAY (Task 11, L49)")
from models.queue_actions import allowed_actions as _allowed  # noqa: E402

run_sql("UPDATE appointments SET status = 'Completed' WHERE status = 'In Consultation'")
check("the mapping holds the day rule: an earlier day's Waiting visit allows only Left (Task 13: "
      "not Start, not Change service)",
      _allowed("staff", "Waiting", on_today=False) == ("mark_left",)
      and _allowed("doctor", "Waiting", on_today=False) == ("mark_left",)
      and "start" in _allowed("staff", "Waiting") and "start" in _allowed("doctor", "Waiting", on_today=True)
      and "see_patient" in _allowed("doctor", "Waiting for the Doctor", on_today=False))
name_patient(21, "T11A")
name_patient(22, "T11B")
old_wait = add_visit(21, "T11A", 5, "Waiting", time="08:00")[0]
old_wfd = add_visit(22, "T11B", 4, WFD, time="08:30")[0]
today_wait = add_visit(23, "T11C", 0, "Waiting", time="09:00")[0]
pages = {n: body(c.get(u)) for n, c, u in (("staff", staff, "/staff/appointment_scheduling"),
                                          ("doctor", doctor, "/staff/appointment_scheduling"),
                                          ("admin", admin, "/admin/appointment_scheduling"))}
for who in ("staff", "doctor"):
    row = row_in(section_of(pages[who]), f"APT-{old_wait:05d}")
    check(f"{who}: an earlier day's Waiting visit offers only Left before being seen (no Start, not even a "
          "disabled one)", row and "Start Consultation" not in row and 'data-bs-target="#leftModal"' in row, row[-300:])
today_row = row_in(pages["staff"], f"APT-{today_wait:05d}")
check("today's Waiting visit still offers Start", "Start Consultation" in today_row)
check("the Administrator sees no Start and no button on an earlier visit",
      "Start Consultation" not in section_of(pages["admin"]) and f"APT-{old_wait:05d}" in section_of(pages["admin"]))
START_URL = "/staff/appointment_scheduling/start-consultation"
for who, client in (("Staff", staff), ("the Doctor", doctor), ("the Administrator", admin), ("the Official", official)):
    n_before = len(audits("Started consultation"))
    r = client.post(START_URL, data={"appointment_id": str(old_wait)})
    check(f"{who}: a Start POST (valid token) on an earlier day's visit is refused: still Waiting, nothing written",
          status_of(old_wait) == "Waiting" and len(audits("Started consultation")) == n_before
          and q1("SELECT COUNT(*) FROM appointments WHERE status = 'In Consultation'")[0] == 0, status_of(old_wait))
staff.post(START_URL, data={"appointment_id": str(old_wait)})
check("the refusal says why, in plain words",
      "from an earlier day" in body(staff.get("/staff/appointment_scheduling")))
with app.app_context():
    from models.appointments import start_consultation as _start
    check("the model itself refuses it too (the rule is in the SQL, not only in the buttons)",
          _start(old_wait) == "earlier-day" and status_of(old_wait) == "Waiting")
r = staff.post(START_URL, data={"appointment_id": str(today_wait)})
check("a visit checked in today still starts", status_of(today_wait) == "In Consultation")
run_sql("UPDATE appointments SET status = 'Completed' WHERE id = ?", (today_wait,))
# See patient and Left still work on an earlier day's visits.
r = doctor.get(f"/staff/appointment_scheduling/see-patient/{old_wfd}")
check("See patient still works on an earlier day's visit waiting for the Doctor",
      r.status_code == 302 and f"/consultation/record/{old_wfd}" in r.headers.get("Location", ""))
staff.post(MARK, data={"appointment_id": str(old_wait)})
staff.post(MARK, data={"appointment_id": str(old_wfd)})
check("Left before being seen still works on an earlier day's Waiting visit and on one waiting for the Doctor",
      status_of(old_wait) == LEFT and status_of(old_wfd) == LEFT)
run_sql("UPDATE appointments SET status = 'Completed' WHERE id IN (?, ?, ?)", (old_wait, old_wfd, today_wait))

# ---------------------------------------------------------------------------
section("21. NO CHANGE SERVICE ON AN EARLIER DAY'S WAITING VISIT (Task 13, D6)")
name_patient(24, "T13A")
name_patient(25, "T13B")
name_patient(26, "T13C")
d6_old_wait = add_visit(24, "T13A", 6, "Waiting", time="08:00")[0]
d6_old_wfd = add_visit(25, "T13B", 5, WFD, time="08:30")[0]
d6_today_wait = add_visit(26, "T13C", 0, "Waiting", time="09:00")[0]
d6_pages = {n: body(c.get(u)) for n, c, u in (("staff", staff, "/staff/appointment_scheduling"),
                                             ("doctor", doctor, "/staff/appointment_scheduling"))}
for who in ("staff", "doctor"):
    old_row = row_in(section_of(d6_pages[who]), f"APT-{d6_old_wait:05d}")
    check(f"{who}: an earlier day's Waiting visit offers Left before being seen and no Change service",
          old_row and 'data-bs-target="#leftModal"' in old_row and 'data-bs-target="#serviceModal"' not in old_row,
          old_row[-300:])
    today_row = row_in(d6_pages[who], f"APT-{d6_today_wait:05d}")
    check(f"{who}: today's Waiting visit still offers Change service",
          'data-bs-target="#serviceModal"' in today_row)
    wfd_row = row_in(section_of(d6_pages[who]), f"APT-{d6_old_wfd:05d}")
    check(f"{who}: an earlier day's visit Waiting for the Doctor keeps Change service",
          'data-bs-target="#serviceModal"' in wfd_row, wfd_row[-200:])
for who, client in (("Staff", staff), ("the Doctor", doctor), ("the Administrator", admin), ("the Official", official)):
    n_before = len(audits("Changed the service of a visit"))
    client.post(CHANGE, data={"appointment_id": str(d6_old_wait), "service": PRENATAL})
    check(f"{who}: a Change service POST (valid token) on an earlier day's Waiting visit is refused: "
          "service unchanged, nothing written",
          q1("SELECT appointment_type FROM appointments WHERE id = ?", (d6_old_wait,))[0] != PRENATAL
          and len(audits("Changed the service of a visit")) == n_before)
staff.post(CHANGE, data={"appointment_id": str(d6_old_wait), "service": PRENATAL})
check("the refusal says why, in plain words",
      "from an earlier day" in body(staff.get("/staff/appointment_scheduling")))
with app.app_context():
    from models.appointments import change_visit_service as _change
    check("the model itself refuses it too (the rule is in the SQL, not only in the buttons)",
          _change(d6_old_wait, PRENATAL) == "not-allowed"
          and q1("SELECT appointment_type FROM appointments WHERE id = ?", (d6_old_wait,))[0] != PRENATAL)
staff.post(CHANGE, data={"appointment_id": str(d6_today_wait), "service": PRENATAL})
check("today's Waiting visit can still have its service changed",
      q1("SELECT appointment_type FROM appointments WHERE id = ?", (d6_today_wait,))[0] == PRENATAL)
staff.post(CHANGE, data={"appointment_id": str(d6_old_wfd), "service": PRENATAL})
check("an earlier day's visit Waiting for the Doctor can still have its service changed",
      q1("SELECT appointment_type FROM appointments WHERE id = ?", (d6_old_wfd,))[0] == PRENATAL)
staff.post(MARK, data={"appointment_id": str(d6_old_wait)})
check("and the earlier day's Waiting visit can still be marked Left before being seen",
      status_of(d6_old_wait) == LEFT)
run_sql("UPDATE appointments SET status = 'Completed' WHERE id IN (?, ?, ?)",
        (d6_old_wait, d6_old_wfd, d6_today_wait))

# ---------------------------------------------------------------------------
passed = sum(1 for _, ok in results if ok)
print("\n" + "=" * 78)
print(f"RESULT: {passed}/{len(results)} checks passed")
failed = [name for name, ok in results if not ok]
for name in failed:
    print(f"  FAILED: {name}")
sys.exit(1 if failed else 0)
