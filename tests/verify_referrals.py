"""Referrals to the RHU or a hospital, the referral slip and the new service list (Task 19).

D1. "Referred" is a visit outcome. Staff and the Doctor set it on a Waiting, In Consultation or
    Waiting for the Doctor visit, from one list of destinations, and undo it the same day. The
    reason is clinical text: care team only, encrypted while the visit is sealed, never in an
    audit entry, an export or a count seen by the Administrator.
D2. The referral slip is a print page for Staff and the Doctor only.
D3. Reports count referrals by destination, with small counts folded for the Administrator and
    the Official.
D4. Dental is no longer offered; Family Planning is, and it does not need the Doctor.
"""
import io
import re
import shutil
import sqlite3
import sys
import tempfile
import zlib
from datetime import date, timedelta
from pathlib import Path

import _support  # noqa: E402  (first: builds the test database, points the app at it)

from app import create_app  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

sys.stdout.reconfigure(errors="replace")

app = create_app()
_support.attach(app)
TEST_DB = _support.DB_PATH
ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID = 1, 2, 3, 5

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok)))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}")
    if detail and not ok:
        for line in str(detail).splitlines()[:12]:
            print(f"       {line}")
    return ok


def section(title):
    print(f"\n--- {title} ---")


def db():
    conn = sqlite3.connect(TEST_DB)
    conn.row_factory = sqlite3.Row
    return conn


def q1(sql, args=()):
    conn = db()
    try:
        return conn.execute(sql, args).fetchone()
    finally:
        conn.close()


def run_sql(sql, args=()):
    conn = db()
    try:
        cur = conn.execute(sql, args)
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def client_as(user_id):
    c = app.test_client()
    with c.session_transaction() as sess:
        sess["user_id"] = user_id
        _support.open_session(sess)
    return c


def body(response):
    return response.get_data(as_text=True)


def flat(text):
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", text))


def xlsx_text(data):
    book = load_workbook(io.BytesIO(data), data_only=True)
    out = []
    for sheet in book.worksheets:
        out.append(sheet.title)
        for row in sheet.iter_rows(values_only=True):
            out.extend(str(cell) for cell in row if cell is not None)
    return "\n".join(out)


def pdf_text(data):
    out = []
    for match in re.finditer(rb"stream\r?\n(.*?)endstream", data, re.S):
        try:
            out.append(zlib.decompress(match.group(1)).decode("latin-1"))
        except Exception:
            out.append(match.group(1).decode("latin-1", "replace"))
    return "\n".join(out).replace("\\(", "(").replace("\\)", ")")


admin, staff, official, doctor = (client_as(i) for i in
                                  (ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID))
today = date.today()
QUEUE = "/staff/appointment_scheduling"
ADMIN_QUEUE = "/admin/appointment_scheduling"
REFER, UNDO = f"{QUEUE}/refer", f"{QUEUE}/undo-referral"
WAIT, INCONS, WFD, DONE, LEFT = ("Waiting", "In Consultation", "Waiting for the Doctor",
                                 "Completed", "Left before being seen")

from models.referrals import (  # noqa: E402
    CARAGA, OTHER, REFERRAL_DESTINATIONS, RHU, SLIP_CONFIDENTIAL,
)

REASON = "ZZREASONMARK needs the RHU doctor"
FACILITY = "ZZFACILITYMARK Medical Center"
COMPLAINT = "ZZCOMPLAINTMARK cough"

print("=" * 78)
print("IDS-BPHCSN — referrals, the referral slip and the service list (Task 19)")
print("=" * 78)

# ---------------------------------------------------------------------------
section("1. SET UP VISITS IN KNOWN STATES")
NAMES = {}
for extra in range(21, 31):
    run_sql("INSERT INTO patients (id, patient_id, first_name, last_name, birth_date, sex, civil_status, "
            "address, municipality, province, contact_number, status) VALUES (?, ?, 'Extra', 'Patient', "
            "'1985-01-01', 'Female', 'Single', 'Purok 1', 'San Francisco', 'Surigao del Norte', ?, 'Active')",
            (extra, f"PAT-{extra:03d}", f"0900000{extra:04d}"))
run_sql("UPDATE appointments SET status = 'Completed' WHERE status IN ('In Consultation', 'Waiting')")
_next_patient = [100]


def add_visit(status, tag, kind="Consultation", days_ago=0, consultation=False, restricted=False,
              diagnosis=None, complaint=COMPLAINT, vitals=True):
    """One visit for a fresh patient. `consultation` writes an intake with vitals."""
    pid = _next_patient[0]
    _next_patient[0] += 1
    run_sql("INSERT INTO patients (id, patient_id, first_name, last_name, birth_date, sex, civil_status, "
            "address, municipality, province, contact_number, status) VALUES (?, ?, ?, 'Refmark', "
            "'1985-01-01', 'Female', 'Single', 'Purok 1', 'San Francisco', 'Surigao del Norte', ?, 'Active')",
            (pid, f"PAT-{pid:03d}", f"Zzrf{tag}", f"0900000{pid:04d}"))
    NAMES[tag] = f"Zzrf{tag} Refmark"
    day = (today - timedelta(days=days_ago)).isoformat()
    appointment = run_sql(
        "INSERT INTO appointments (patient_id, scheduled_by, appointment_type, appointment_date, "
        "status, is_walkin, checked_in_at) VALUES (?, ?, ?, ?, ?, 1, ?)",
        (pid, STAFF_ID, kind, day, status, f"{day} 09:30:00"))
    cid = None
    if consultation:
        cid = run_sql(
            "INSERT INTO consultations (patient_id, recorded_by, appointment_id, consultation_date, "
            "chief_complaint, diagnosis, is_restricted, temperature, blood_pressure, pulse_rate) "
            "VALUES (?,?,?,?,?,?,?,?,?,?)",
            (pid, STAFF_ID, appointment, day, complaint, diagnosis, 1 if restricted else 0,
             38.2 if vitals else None, "120/80" if vitals else None, 88 if vitals else None))
    return appointment, cid, pid


def row_of(appointment):
    return q1("SELECT * FROM appointments WHERE id = ?", (appointment,))


def status_of(appointment):
    return row_of(appointment)["status"]


def audits(action):
    conn = db()
    try:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM audit_logs WHERE action = ? ORDER BY id", (action,))]
    finally:
        conn.close()


def refer(client, appointment, destination=RHU, facility="", reason=""):
    return client.post(REFER, data={"appointment_id": str(appointment), "destination": destination,
                                    "facility": facility, "reason": reason})


def free_room():
    run_sql("UPDATE appointments SET status = 'Completed' WHERE status = 'In Consultation'")


check("the destination list is the three decided places, in order",
      REFERRAL_DESTINATIONS == ("Rural Health Unit (RHU)", "Caraga Regional Hospital",
                                "Other hospital or facility"))

# ---------------------------------------------------------------------------
section("2. WHO CAN REFER, FROM WHICH STATUS (D1)")
REFER_AUDIT = "Referred a patient"
for who, client, uid in (("Staff", staff, STAFF_ID), ("the Doctor", doctor, DOCTOR_ID)):
    for status in (WAIT, INCONS, WFD):
        free_room()
        v, _, _ = add_visit(status, f"A{who[:3]}{status[:3]}", consultation=(status == WFD))
        n = len(audits(REFER_AUDIT))
        refer(client, v, RHU, reason=REASON)
        row = row_of(v)
        today_stamp = row["referred_at"] or ""
        check(f"{who} can refer a {status} visit: it is Referred, with who, when and where from",
              row["status"] == "Referred" and row["referral_destination"] == RHU
              and row["referred_by"] == uid and today_stamp[:10] == today.isoformat()
              and re.fullmatch(r"\d\d:\d\d:\d\d", today_stamp[11:] or "")
              and row["status_before_referred"] == status
              and len(audits(REFER_AUDIT)) == n + 1, dict(row))
free_room()

for destination in (RHU, CARAGA, OTHER):
    v, _, _ = add_visit(WAIT, "Dest")
    refer(staff, v, destination, facility=FACILITY if destination == OTHER else "")
    row = row_of(v)
    check(f"destination '{destination}' is stored as chosen; the facility name only for Other",
          row["referral_destination"] == destination
          and row["referral_facility"] == (FACILITY if destination == OTHER else None), dict(row))
v, _, _ = add_visit(WAIT, "Fac")
refer(staff, v, RHU, facility="Typed anyway")
check("a facility name typed with a listed destination is not stored",
      row_of(v)["referral_facility"] is None and status_of(v) == "Referred")

for label, client in (("Administrator", admin), ("Barangay Official", official)):
    v, _, _ = add_visit(WAIT, "Deny")
    n = len(audits(REFER_AUDIT))
    r = refer(client, v, RHU, reason=REASON)
    check(f"{label}: POST refer (valid token) is refused on the server; nothing changes, no audit entry",
          r.status_code in (302, 403) and status_of(v) == WAIT and len(audits(REFER_AUDIT)) == n,
          (r.status_code, status_of(v)))
check("the Administrator has no /admin/ route for it, nor for the undo or the slip",
      admin.post(f"{ADMIN_QUEUE}/refer", data={"appointment_id": "1"}).status_code == 404
      and admin.post(f"{ADMIN_QUEUE}/undo-referral", data={"appointment_id": "1"}).status_code == 404
      and admin.get(f"{ADMIN_QUEUE}/referral-slip/1").status_code == 404)

n = len(audits(REFER_AUDIT))
done_v, _, _ = add_visit(DONE, "Done", consultation=True, diagnosis="Seen by the Doctor")
refer(staff, done_v)
closed_v, _, _ = add_visit(DONE, "Closed")          # finished with nothing written: no diagnosis to stop it
refer(staff, closed_v)
left_v, _, _ = add_visit(WAIT, "Lft")
run_sql("UPDATE appointments SET status = ?, status_before_left = 'Waiting', left_by = 2, left_at = ? "
        "WHERE id = ?", (LEFT, f"{today.isoformat()} 08:00:00", left_v))
refer(staff, left_v)
free_room()
diag_v, _, _ = add_visit(INCONS, "Diag", consultation=True, diagnosis="A diagnosis is written")
refer(doctor, diag_v)
check("a Completed visit (with or without a diagnosis), a visit marked left and a visit that has a "
      "diagnosis cannot be referred",
      status_of(done_v) == DONE and status_of(closed_v) == DONE and status_of(left_v) == LEFT
      and status_of(diag_v) == INCONS
      and len(audits(REFER_AUDIT)) == n)
free_room()
earlier, _, _ = add_visit(WAIT, "Early", days_ago=2)
refer(staff, earlier)
check("an earlier day's Waiting visit cannot be referred (it offers only Left before being seen)",
      status_of(earlier) == WAIT and len(audits(REFER_AUDIT)) == n)
earlier_wfd, _, _ = add_visit(WFD, "EarlyW", days_ago=2, consultation=True)
refer(doctor, earlier_wfd)
check("... but an earlier day's visit waiting for the Doctor can be", status_of(earlier_wfd) == "Referred")
again = len(audits(REFER_AUDIT))
refer(staff, earlier_wfd)
check("referring a referred visit again changes nothing and writes nothing",
      len(audits(REFER_AUDIT)) == again)

# ---------------------------------------------------------------------------
section("3. THE FORM IS CHECKED ON THE SERVER (D1)")
v, _, _ = add_visit(WAIT, "Val")
n = len(audits(REFER_AUDIT))
bad = [("no destination", dict(destination="")), ("a made-up destination", dict(destination="Clinic X")),
       ("the right name in the wrong case", dict(destination="rural health unit (rhu)")),
       ("'Other' with no facility name", dict(destination=OTHER, facility="")),
       ("'Other' with a blank facility name", dict(destination=OTHER, facility="   ")),
       ("a facility name over 80 characters", dict(destination=OTHER, facility="F" * 81)),
       ("a reason over 500 characters", dict(destination=RHU, reason="R" * 501))]
for label, change in bad:
    data = {"appointment_id": str(v), "destination": RHU, "facility": "", "reason": ""}
    data.update(change)
    staff.post(REFER, data=data)
    check(f"{label} is refused: the visit stays Waiting and nothing is audited",
          status_of(v) == WAIT and len(audits(REFER_AUDIT)) == n)
r = refer(staff, v, OTHER, facility="Brgy Hospital")
check("'Other' with a facility name is accepted", status_of(v) == "Referred"
      and row_of(v)["referral_facility"] == "Brgy Hospital")
r = staff.post(REFER, data={"destination": RHU})
check("a POST with no visit is refused calmly", r.status_code == 302)
r = staff.post(REFER, data={"appointment_id": "999999", "destination": RHU})
check("a POST for a visit that does not exist is refused calmly", r.status_code == 302)
page = body(staff.get(QUEUE))
modal = page.split('id="referModal"')[1].split("</form>")[0] if 'id="referModal"' in page else ""
check("the form offers exactly the three destinations, from the one list, a name field and a reason",
      all(f'<option value="{d}">' in modal for d in REFERRAL_DESTINATIONS)
      and modal.count("<option value=") == 4 and 'name="facility"' in modal and 'name="reason"' in modal)
check("the Administrator's queue has no referral form and no Refer button",
      'id="referModal"' not in body(admin.get(ADMIN_QUEUE)) and "referModal" not in body(admin.get(ADMIN_QUEUE)))

# ---------------------------------------------------------------------------
section("4. UNDO WORKS THE SAME DAY ONLY (D1)")
free_room()
UNDO_AUDIT = "Undid a referral"
v, _, _ = add_visit(WFD, "Und", consultation=True)
refer(doctor, v, OTHER, facility=FACILITY, reason=REASON)
n = len(audits(UNDO_AUDIT))
admin.post(UNDO, data={"appointment_id": str(v)})
official.post(UNDO, data={"appointment_id": str(v)})
check("the Administrator and the Official cannot undo a referral", status_of(v) == "Referred")
staff.post(UNDO, data={"appointment_id": str(v)})
row = row_of(v)
check("Staff undo it the same day: the visit is back, and the destination, facility, reason, who and "
      "when are cleared",
      row["status"] == WFD and all(row[c] is None for c in (
          "referral_destination", "referral_facility", "referral_reason", "referred_by", "referred_at",
          "status_before_referred")) and len(audits(UNDO_AUDIT)) == n + 1, dict(row))
refer(staff, v, RHU)
doctor.post(UNDO, data={"appointment_id": str(v)})
check("the Doctor can undo it too", status_of(v) == WFD and len(audits(UNDO_AUDIT)) == n + 2)
refer(staff, v, RHU)
run_sql("UPDATE appointments SET referred_at = ? WHERE id = ?",
        (f"{(today - timedelta(days=1)).isoformat()} 09:00:00", v))
n = len(audits(UNDO_AUDIT))
staff.post(UNDO, data={"appointment_id": str(v)})
check("referred on an earlier day, it can no longer be undone", status_of(v) == "Referred"
      and len(audits(UNDO_AUDIT)) == n)
free_room()
v2, _, _ = add_visit(INCONS, "UndB")
refer(staff, v2)
holder, _, _ = add_visit(INCONS, "Holder")
staff.post(UNDO, data={"appointment_id": str(v2)})
check("a visit referred from In Consultation cannot go back while another patient holds the room",
      status_of(v2) == "Referred")
free_room()
staff.post(UNDO, data={"appointment_id": str(v2)})
check("... and goes back once the room is free", status_of(v2) == INCONS)
free_room()
n = len(audits(UNDO_AUDIT))
plain, _, _ = add_visit(WAIT, "Plain")
staff.post(UNDO, data={"appointment_id": str(plain)})
check("undoing a visit that is not referred changes nothing", status_of(plain) == WAIT
      and len(audits(UNDO_AUDIT)) == n)

# ---------------------------------------------------------------------------
section("5. THE DOCTOR'S LIST AND THE COUNTS AGREE (D1, D3)")
from models.appointments import get_appointment_report, get_queue_stats  # noqa: E402
from models.doctor_worklist import count_waiting_for_doctor, get_waiting_for_doctor  # noqa: E402

free_room()
w1, _, _ = add_visit(WFD, "ListA", consultation=True)
w2, _, _ = add_visit(WFD, "ListB", consultation=True)
with app.app_context():
    before = count_waiting_for_doctor()
    listed_before = {v["id"] for v in get_waiting_for_doctor()}
doctor.post(REFER, data={"appointment_id": str(w1), "destination": CARAGA})
with app.app_context():
    after = count_waiting_for_doctor()
    listed_after = {v["id"] for v in get_waiting_for_doctor()}
check("a referred visit leaves the Doctor's list, and the count follows",
      w1 in listed_before and w1 not in listed_after and w2 in listed_after and after == before - 1)
doc_dash = body(doctor.get("/staff/dashboard"))
check("the Doctor's dashboard no longer lists the visit, and still lists the other one",
      f"/see-patient/{w1}" not in doc_dash and f"/see-patient/{w2}" in doc_dash)
sql_today = q1("SELECT COUNT(*) FROM appointments WHERE date(appointment_date) = date('now', 'localtime') "
               "AND status = 'Referred'")[0]
with app.app_context():
    stats = get_queue_stats()
    report = get_appointment_report(months=6, viewer_role="doctor")["window"]
sql_range = q1("SELECT COUNT(*) FROM appointments WHERE status = 'Referred' AND substr(appointment_date,1,10) "
               "BETWEEN ? AND ?", (report["start"].isoformat(), report["end"].isoformat()))[0]
check("the queue figure, the SQL count and the report all agree on the referred visits",
      stats["referred"] == sql_today and report["referred"] == sql_range and sql_today > 0
      and report["referred"] >= stats["referred"], (stats["referred"], sql_today, report["referred"], sql_range))
pages = {"staff": body(staff.get(QUEUE)), "doctor": body(doctor.get(QUEUE)), "admin": body(admin.get(ADMIN_QUEUE))}
for role, page in pages.items():
    summary = flat(page.split('id="queueSummary"')[1].split("</div>")[0])
    check(f"{role}: the queue summary line says 'Referred {sql_today}' and the legend lists Referred",
          f"Referred {sql_today}" in summary and "Referred" in page.split('id="statusLegend"')[1][:2500])
check("the Administrator sees the status only: 'Referred', and no destination",
      "Referred" in pages["admin"] and "To: " not in pages["admin"] and CARAGA not in pages["admin"]
      and RHU not in pages["admin"])
check("Staff and the Doctor see where the patient was sent", f"To: {CARAGA}" in pages["doctor"]
      and f"To: {CARAGA}" in pages["staff"])
resolved = (report["attended"] + report["awaiting"] + report["referred"] + report["missed"] + report["left"])
check("a referred visit is resolved and counts as having come (attendance)", report["resolved"] == resolved
      and report["attendance_pct"] is not None, report)

# ---------------------------------------------------------------------------
section("6. THE AUDIT NAMES THE DESTINATION AND NEVER THE REASON (D1)")
free_room()
a1, _, p1 = add_visit(WAIT, "AudA")
refer(staff, a1, RHU, reason=REASON)
a2, _, p2 = add_visit(WAIT, "AudB")
refer(doctor, a2, OTHER, facility=FACILITY, reason=REASON)
a3, _, p3 = add_visit(WAIT, "AudC")
refer(staff, a3, CARAGA)
mine = {e["details"] for e in audits(REFER_AUDIT)}
check("the entry reads 'Referred <Name> (PAT-005), visit APT-00081, to the RHU.'",
      f"Referred {NAMES['AudA']} (PAT-{p1:03d}), visit APT-{a1:05d}, to the RHU." in mine, sorted(mine)[-3:])
check("... and for Caraga Regional Hospital",
      f"Referred {NAMES['AudC']} (PAT-{p3:03d}), visit APT-{a3:05d}, to Caraga Regional Hospital." in mine)
check("... and for Other, the list entry only: not the typed facility name",
      f"Referred {NAMES['AudB']} (PAT-{p2:03d}), visit APT-{a2:05d}, to another hospital or facility." in mine)
everything = " ".join(str(v) for r in db().execute("SELECT * FROM audit_logs") for v in tuple(r))
check("no audit entry, anywhere, holds a reason, a typed facility name or the complaint",
      "ZZREASONMARK" not in everything and "ZZFACILITYMARK" not in everything
      and "ZZCOMPLAINTMARK" not in everything)
check("the entry is filed under Appointment Scheduling, by the account that did it",
      all(e["module"] == "Appointment Scheduling" for e in audits(REFER_AUDIT))
      and {e["user_id"] for e in audits(REFER_AUDIT)} <= {STAFF_ID, DOCTOR_ID})
undone = audits(UNDO_AUDIT)
check("the undo entry names the patient and the visit, and never a reason or destination",
      undone and all(re.fullmatch(r".+ \(PAT-\d+\), visit APT-\d{5}\.", e["details"]) for e in undone))

# ---------------------------------------------------------------------------
section("7. THE ADMINISTRATOR NEVER SEES THE REASON (D1)")
free_room()
r_v, _, _ = add_visit(WFD, "Priv", consultation=True)
refer(doctor, r_v, OTHER, facility=FACILITY, reason=REASON)
from models.settings import KEY_NAME, get_settings  # noqa: E402

with app.app_context():
    CENTER = get_settings()[KEY_NAME]
stored = row_of(r_v)["referral_reason"]
check("positive control: an unsealed visit keeps the reason as written", stored == REASON)
admin_pages = [ADMIN_QUEUE, "/admin/dashboard", "/admin/reports_analytics", "/admin/patient_records"]
leaks = []
for url in admin_pages + [f"/admin/patient_profile/{q1('SELECT patient_id FROM appointments WHERE id = ?', (r_v,))[0]}"]:
    page = body(admin.get(url, follow_redirects=True))
    if "ZZREASONMARK" in page or "ZZFACILITYMARK" in page:
        leaks.append(url)
check("no Administrator screen carries the reason or the facility name", not leaks, leaks)
for label, url, kind in (("queue Excel", f"{ADMIN_QUEUE}/export.xlsx", "xlsx"),
                         ("queue PDF", f"{ADMIN_QUEUE}/export.pdf", "pdf"),
                         ("reports Excel", "/admin/reports_analytics/export.xlsx", "xlsx"),
                         ("reports PDF", "/admin/reports_analytics/export.pdf", "pdf")):
    resp = admin.get(url)
    text = xlsx_text(resp.data) if kind == "xlsx" else pdf_text(resp.data)
    check(f"Administrator {label}: generated, readable, and holds neither the reason nor the facility name",
          resp.status_code == 200 and len(text) > 200 and "ZZREASONMARK" not in text
          and "ZZFACILITYMARK" not in text, (resp.status_code, len(text)))
for label, url in (("Official Reports & Analytics", "/official/reports_analytics"),
                   ("Official Community Reports", "/official/community-reports"),
                   ("Official dashboard", "/official/dashboard")):
    page = body(official.get(url))
    check(f"{label}: no reason, no facility name", "ZZREASONMARK" not in page and "ZZFACILITYMARK" not in page)
check("the care team's queue does not print the reason either (only the slip does)",
      "ZZREASONMARK" not in body(staff.get(QUEUE)) and "ZZREASONMARK" not in body(doctor.get(QUEUE)))
check("no route sends data as JSON apart from the session checks (nothing to leak the reason through)",
      not [r.rule for r in app.url_map.iter_rules() if "json" in r.rule.lower()])

# ---------------------------------------------------------------------------
section("8. A SEALED VISIT: THE REASON IS ENCRYPTED, AND FOLLOWS THE SEAL (D1)")
from models.consultations import set_consultation_restriction  # noqa: E402

free_room()
s_v, s_cid, s_pid = add_visit(WFD, "Seal", consultation=True, restricted=True)
from utils import sealed_crypto  # noqa: E402

run_sql("UPDATE consultations SET chief_complaint = ? WHERE id = ?",
        (sealed_crypto.encrypt(COMPLAINT, s_cid, "chief_complaint"), s_cid))
doctor.post(REFER, data={"appointment_id": str(s_v), "destination": RHU, "reason": REASON})
stored = row_of(s_v)["referral_reason"]
check("on a sealed visit the reason is stored encrypted, and the plain text is not in the database",
      sealed_crypto.is_encrypted(stored) and "ZZREASONMARK" not in stored)
with app.app_context():
    out = set_consultation_restriction(s_cid, False, DOCTOR_ID, role="Doctor")
check("lifting the seal opens the reason again", out == "lifted" and row_of(s_v)["referral_reason"] == REASON, out)
with app.app_context():
    out = set_consultation_restriction(s_cid, True, DOCTOR_ID, role="Doctor")
check("sealing the record afterwards encrypts a referral that was already written",
      out == "restricted" and sealed_crypto.is_encrypted(row_of(s_v)["referral_reason"]), out)
free_room()
late_v, late_cid, _ = add_visit(WFD, "Late", consultation=True)
refer(staff, late_v, RHU, reason=REASON)
check("an unsealed referral keeps the reason as written (Staff may refer too)",
      row_of(late_v)["referral_reason"] == REASON)

# ---------------------------------------------------------------------------
section("9. THE REFERRAL SLIP (D2)")
free_room()
sl_v, sl_cid, sl_pid = add_visit(WFD, "Slip", consultation=True)
refer(doctor, sl_v, OTHER, facility=FACILITY, reason=REASON)
SLIP = f"{QUEUE}/referral-slip/{sl_v}"
for label, client, uid in (("Staff", staff, STAFF_ID), ("the Doctor", doctor, DOCTOR_ID)):
    r = client.get(SLIP)
    page = body(r)
    text = flat(page)
    from models.patients import format_age  # noqa: E402
    from utils.dates import display_date  # noqa: E402

    patient = q1("SELECT * FROM patients WHERE id = ?", (sl_pid,))
    row = row_of(sl_v)
    stamp = row["referred_at"]
    check(f"{label} can open the slip, an A4 page with its own print rules and a Print button",
          r.status_code == 200 and "size:A4 portrait" in page and "window.print()" in page)
    check(f"{label}: the slip prints the health center, the date and time, name, ID, age and sex",
          CENTER in text and display_date(stamp[:10]) in text and stamp[11:16] in text
          and NAMES["Slip"] in text and patient["patient_id"] in text
          and (format_age(patient["birth_date"]) or "?") in text and patient["sex"] in text, text[:400])
    check(f"{label}: ... the destination with the facility name, the complaint, the vitals and the reason",
          f"{OTHER}: {FACILITY}" in text and COMPLAINT in text and "38.2" in text and "120/80" in text
          and "88" in text and REASON in text)
    check(f"{label}: ... the name of the referring user and a line for a signature",
          q1("SELECT full_name FROM users WHERE id = ?", (DOCTOR_ID,))[0] in text and "Signature" in text)
check("the slip reads from the visit: it is dated the day of the referral",
      row_of(sl_v)["referred_at"][:10] == today.isoformat())

# A sealed visit: the sealed fields print as the confidential line, for everyone.
free_room()
ss_v, ss_cid, _ = add_visit(WFD, "SlipS", consultation=True, restricted=True)
run_sql("UPDATE consultations SET chief_complaint = ? WHERE id = ?",
        (sealed_crypto.encrypt(COMPLAINT, ss_cid, "chief_complaint"), ss_cid))
doctor.post(REFER, data={"appointment_id": str(ss_v), "destination": RHU, "reason": REASON})
for label, client in (("Staff", staff), ("the Doctor", doctor)):
    r = client.get(f"{QUEUE}/referral-slip/{ss_v}")
    page = body(r)
    text = flat(page)
    check(f"{label}: on a sealed visit the complaint and the reason print as "
          "'Confidential — discuss directly with the receiving doctor.'",
          r.status_code == 200 and text.count(SLIP_CONFIDENTIAL) == 2, text.count(SLIP_CONFIDENTIAL))
    check(f"{label}: ... the sealed text, the encrypted value and the plain reason are nowhere in the page",
          "ZZREASONMARK" not in page and "ZZCOMPLAINTMARK" not in page and "enc:v1:" not in page)
    check(f"{label}: ... the vitals, the destination and the patient still print",
          "38.2" in text and RHU in text and NAMES["SlipS"] in text)
check("the confidential line is the one the task decided",
      SLIP_CONFIDENTIAL == "Confidential — discuss directly with the receiving doctor.")
for label, client in (("Administrator", admin), ("Barangay Official", official)):
    r = client.get(SLIP)
    check(f"{label} cannot open the slip", r.status_code in (302, 403, 404) and "Referral Slip" not in body(r),
          r.status_code)
r = admin.get(f"{ADMIN_QUEUE}/referral-slip/{sl_v}")
check("the Administrator has no slip route at all", r.status_code == 404)
plain_v, _, _ = add_visit(WAIT, "NotRef")
r = staff.get(f"{QUEUE}/referral-slip/{plain_v}")
check("the slip of a visit that is not referred is refused", r.status_code == 302
      and "Referral Slip" not in body(r))
r = staff.get(f"{QUEUE}/referral-slip/999999")
check("the slip of a visit that does not exist is refused calmly", r.status_code == 302)
page = body(staff.get(QUEUE))
check("the Referred row offers the slip and the undo", f"/referral-slip/{sl_v}" in page
      and 'form="undoReferralForm"' in page)

# ---------------------------------------------------------------------------
section("10. REPORTS COUNT REFERRALS BY DESTINATION, SMALL COUNTS FOLDED (D3)")
RANGE_DAY = (today - timedelta(days=3)).isoformat()
RANGE = f"?from={RANGE_DAY}&to={RANGE_DAY}"


def add_referred(destination, n, day=RANGE_DAY):
    for _ in range(n):
        run_sql("INSERT INTO appointments (patient_id, scheduled_by, appointment_type, appointment_date, "
                "status, is_walkin, checked_in_at, referral_destination, referred_by, referred_at, "
                "status_before_referred) VALUES (6, 2, 'Consultation', ?, 'Referred', 1, ?, ?, 2, ?, 'Waiting')",
                (day, f"{day} 09:00:00", destination, f"{day} 10:00:00"))


add_referred(RHU, 6)
add_referred(CARAGA, 1)
add_referred(OTHER, 2)
run_sql("INSERT INTO appointments (patient_id, scheduled_by, appointment_type, appointment_date, status, "
        "is_walkin, checked_in_at) VALUES (6, 2, 'Consultation', ?, 'Completed', 1, ?)",
        (RANGE_DAY, f"{RANGE_DAY} 09:00:00"))
start = end = today - timedelta(days=3)
with app.app_context():
    views = {role: get_appointment_report(months=1, start=start, end=end, viewer_role=role)["window"]
             for role in ("doctor", "health staff", "admin", "official")}
for role in ("doctor", "health staff"):
    w = views[role]
    check(f"{role}: the report counts 9 referred and breaks them down by destination, all three shown",
          w["referred"] == 9 and [(r["destination"], r["count"]) for r in w["referral_destinations"]]
          == [(RHU, 6), (CARAGA, 1), (OTHER, 2)], w["referral_destinations"])
for role in ("admin", "official"):
    w = views[role]
    rows = w["referral_destinations"]
    check(f"{role}: the total is exact (9); the destinations with fewer than 5 are folded into one Other row",
          w["referred"] == 9 and [(r["destination"], r["count"]) for r in rows][0] == (RHU, 6)
          and len(rows) == 2 and rows[1].get("other") and rows[1]["count"] == 3
          and CARAGA not in str(rows) and OTHER not in str(rows), rows)
check("the report counts the Completed visit as attended and the referred ones as their own outcome",
      views["doctor"]["attended"] == 1 and views["doctor"]["total"] == 10, views["doctor"])
admin_page = body(admin.get(f"/admin/reports_analytics{RANGE}"))
official_page = body(official.get(f"/official/reports_analytics{RANGE}"))
community = body(official.get(f"/official/community-reports{RANGE}"))
for label, page in (("Administrator", admin_page), ("Official", official_page)):
    text = flat(page)
    check(f"{label}'s Reports & Analytics: Referred (9), the RHU row, and the folded row; the small "
          "destinations are not named",
          "Referred (9)" in text and f"{RHU} 6" in text and "Other (each fewer than 5) 3" in text
          and f"{CARAGA} 1" not in text and f"{OTHER} 2" not in text, text[text.find("Referred ("):][:200])
text = flat(community)
check("the Official's Community Reports: the Referred figure, the RHU row and the folded row",
      re.search(r"Referred\s*9", text) and f"{RHU} 6" in text and "Other (each fewer than 5) 3" in text
      and f"{CARAGA} 1" not in text, text[text.find("Referred"):][:200])
staff_report = body(doctor.get(f"/staff/reports_analytics{RANGE}"))
check("the Doctor's Reports & Analytics names every destination with its count",
      f"{CARAGA} 1" in flat(staff_report) and f"{OTHER} 2" in flat(staff_report), flat(staff_report)[:100])
for label, client, url, kind in (
        ("Administrator Excel", admin, f"/admin/reports_analytics/export.xlsx{RANGE}", "xlsx"),
        ("Administrator PDF", admin, f"/admin/reports_analytics/export.pdf{RANGE}", "pdf"),
        ("Official Excel", official, f"/official/reports_analytics/export.xlsx{RANGE}", "xlsx"),
        ("Official Community Reports Excel", official, f"/official/community-reports/export.xlsx{RANGE}", "xlsx")):
    resp = client.get(url)
    text = xlsx_text(resp.data) if kind == "xlsx" else pdf_text(resp.data)
    check(f"{label}: 'Referrals by destination' with the RHU row and the folded row; small destinations "
          "not named",
          resp.status_code == 200 and "Referrals by destination" in text and RHU in text
          and "Other (each fewer than 5)" in text and CARAGA not in text, (resp.status_code, len(text)))
resp = doctor.get(f"/staff/reports_analytics/export.xlsx{RANGE}")
check("the Doctor's Excel names every destination", resp.status_code == 200
      and CARAGA in xlsx_text(resp.data) and OTHER in xlsx_text(resp.data))
with app.app_context():
    from models.appointments import get_appointment_day_summary

    day = get_appointment_day_summary(start, end)[RANGE_DAY]
check("the calendar's day summary counts Referred visits as referred, not as cancelled",
      day["referred"] == 9 and day["cancelled"] == 0, day)

# ---------------------------------------------------------------------------
section("11. DENTAL IS NOT OFFERED; FAMILY PLANNING IS (D4)")
from models.doctor_worklist import VISIT_SERVICES, clean_service, needs_doctor  # noqa: E402

check("the service list is the six decided services, in order",
      VISIT_SERVICES == ("Consultation", "Follow-up Consultation", "Hypertension / Diabetes Monitoring",
                         "Prenatal Check-up", "Family Planning", "Immunization"))
check("Dental is not a choice; Family Planning is", clean_service("Dental") is None
      and clean_service("Family Planning") == "Family Planning")
page = body(staff.get(QUEUE))
check("the queue's change-service dialog offers Family Planning and not Dental",
      '<option value="Family Planning">' in page and '<option value="Dental">' not in page)
CHECK_IN = f"{QUEUE}/check-in"
free_room()
run_sql("UPDATE appointments SET status = 'Completed' WHERE patient_id IN (27, 28, 29) AND status != 'Completed'")
n_visits = q1("SELECT COUNT(*) FROM appointments WHERE patient_id = 27")[0]
r = staff.post(CHECK_IN, data={"patient_id": "27", "reason": "x", "service": "Dental"})
check("a new check-in for Dental is refused on the server",
      q1("SELECT COUNT(*) FROM appointments WHERE patient_id = 27")[0] == n_visits)
r = staff.post(CHECK_IN, data={"patient_id": "27", "reason": "x", "service": "Family Planning"})
fp = q1("SELECT * FROM appointments WHERE patient_id = 27 ORDER BY id DESC LIMIT 1")
check("a check-in for Family Planning is accepted", fp["appointment_type"] == "Family Planning"
      and fp["status"] == WAIT)
staff.post(f"{QUEUE}/start-consultation", data={"appointment_id": str(fp["id"])})
check("... and it needs no Doctor: not in the services that do", not needs_doctor("Family Planning")
      and status_of(fp["id"]) == INCONS)
staff.post(f"/staff/consultation/record/{fp['id']}", data={"chief_complaint": "pills", "remarks": "",
                                                           "follow_up_date": ""})
with app.app_context():
    listed = {v["id"] for v in get_waiting_for_doctor()}
check("a Family Planning visit is Completed by Staff's intake, never Waiting for the Doctor, and is not on the "
      "Doctor's list", status_of(fp["id"]) == DONE and fp["id"] not in listed, status_of(fp["id"]))
free_room()
old_v, _, old_pid = add_visit(WAIT, "OldD", kind="Dental")
run_sql("UPDATE patients SET first_name = 'Zzolddental' WHERE id = ?", (old_pid,))
page = body(staff.get(QUEUE))
profile = body(staff.get(f"/staff/patient_profile/{old_pid}"))
check("an old Dental visit still displays on the queue",
      "Dental" in page.split(f"APT-{old_v:05d}")[1][:2500])
check("... and in the patient's appointment history", "Dental" in profile)
with app.app_context():
    services = {e["service"]: e["count"] for e in get_appointment_report(
        months=6, viewer_role="doctor")["window"]["services"]}
check("... and it is still counted in 'Appointments by service'", services.get("Dental", 0) >= 1, services)
check("... and it still behaves as a visit that does not need the Doctor", not needs_doctor("Dental"))
admin_queue = body(admin.get(ADMIN_QUEUE))
check("the Administrator's queue shows no service names (Family Planning, Dental or any other)",
      "Family Planning" not in admin_queue.split("<tbody")[1].split("</tbody>")[0]
      and "Dental" not in admin_queue.split("<tbody")[1].split("</tbody>")[0])

# ---------------------------------------------------------------------------
section("12. THE MIGRATION (backs up, keeps every row, safe to run twice, one audit entry)")
from migrations import migrate_add_referrals as migration  # noqa: E402

SCHEMA = (Path(_support.REPO) / "schema.sql").read_text(encoding="utf-8")


def without_referrals(sql):
    return (sql.replace(", 'Referred')", ")")
            .replace("    referral_destination TEXT,\n    referral_facility TEXT,\n"
                     "    referral_reason TEXT,\n    referred_by INTEGER,\n    referred_at TEXT,\n"
                     "    status_before_referred TEXT,\n", "")
            .replace("    FOREIGN KEY (left_by) REFERENCES users(id) ON DELETE SET NULL,\n"
                     "    FOREIGN KEY (referred_by) REFERENCES users(id) ON DELETE SET NULL\n",
                     "    FOREIGN KEY (left_by) REFERENCES users(id) ON DELETE SET NULL\n"))


old_schema = without_referrals(SCHEMA)
fresh = sqlite3.connect(":memory:")
fresh.executescript(SCHEMA)
fresh.executescript("INSERT INTO users (username, password_hash, full_name, role) VALUES ('u','h','U','Admin');"
                    "INSERT INTO patients (patient_id, first_name, last_name, birth_date, sex) "
                    "VALUES ('P','F','L','2000-01-01','Male');")
accepts = refuses = False
try:
    fresh.execute("INSERT INTO appointments (patient_id, scheduled_by, appointment_date, status) "
                  "VALUES (1,1,'2026-01-01','Referred')")
    accepts = True
except sqlite3.IntegrityError:
    pass
try:
    fresh.execute("INSERT INTO appointments (patient_id, scheduled_by, appointment_date, status) "
                  "VALUES (1,1,'2026-01-01','Bogus')")
except sqlite3.IntegrityError:
    refuses = True
fresh.close()
check("schema.sql builds a table that accepts 'Referred' and refuses a made-up status", accepts and refuses)
old_block = old_schema.split("CREATE TABLE appointments")[1].split("\n);")[0]
check("the old schema used for the migration test really lacks the status and the columns",
      "Referred" not in old_block and "referral_reason" not in old_block and "referred_by" not in old_block)
work = Path(tempfile.mkdtemp(prefix="idsref_"))
old = work / "old.db"
con = sqlite3.connect(old)
con.executescript(old_schema)
con.executescript("""
INSERT INTO users (username, password_hash, full_name, role) VALUES ('u','h','U','Admin');
INSERT INTO patients (patient_id, first_name, last_name, birth_date, sex) VALUES ('P','F','L','2000-01-01','Male');
INSERT INTO appointments (patient_id, scheduled_by, appointment_date, status) VALUES (1,1,'2026-01-01','Completed');
INSERT INTO appointments (patient_id, scheduled_by, appointment_date, status, notes, left_at, status_before_left)
  VALUES (1,1,'2026-01-02','Left before being seen','kept','2026-01-02 08:00:00','Waiting');
""")
con.commit()
before = [tuple(r) for r in con.execute("SELECT * FROM appointments ORDER BY id")]
audit_before = con.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0]
con.close()
result = migration.run(old, backup_dir=work / "backups")
con = sqlite3.connect(old)
after = [tuple(r) for r in con.execute("SELECT * FROM appointments ORDER BY id")]
columns = [r[1] for r in con.execute("PRAGMA table_info(appointments)")]
check("the migration backs up first, rebuilds the table and keeps every row and value",
      result["backup"] and Path(result["backup"]).exists() and result["rows"] == 2
      and not result["problems"] and [row[:len(before[0])] for row in after] == before, result)
check("the six new columns are there and empty, and the table accepts 'Referred'",
      {"referral_destination", "referral_facility", "referral_reason", "referred_by", "referred_at",
       "status_before_referred"} <= set(columns)
      and all(v is None for row in after for v in row[len(before[0]):])
      and "'Referred'" in con.execute("SELECT sql FROM sqlite_master WHERE name = 'appointments'").fetchone()[0])
check("it writes ONE audit entry of itself, saying the script did it",
      con.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0] == audit_before + 1
      and "migration script" in con.execute("SELECT details FROM audit_logs ORDER BY id DESC").fetchone()[0])
con.close()
again = migration.run(old, backup_dir=work / "backups")
con = sqlite3.connect(old)
check("a second run does nothing: no backup, no change, no second audit entry",
      again["backup"] is None and len(list((work / "backups").iterdir())) == 1
      and [tuple(r)[:len(before[0])] for r in con.execute("SELECT * FROM appointments ORDER BY id")] == before
      and con.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0] == audit_before + 1, again)
con.close()
shutil.rmtree(work, ignore_errors=True)

# ---------------------------------------------------------------------------
failed = [name for name, ok in results if not ok]
print("\n" + "=" * 78)
print(f"RESULT: {len(results) - len(failed)}/{len(results)} checks passed")
for name in failed:
    print(f"  FAILED: {name}")
print("=" * 78)
sys.exit(1 if failed else 0)
