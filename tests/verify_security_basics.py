"""Security basics: secret key, debug mode, session limits, forced password
change, the seed, the must-change migration, and no URL text on screens.

Runs on a fresh test database built by tests/_support.py. Time is simulated by
replacing utils.security.now_ts, so nothing here sleeps.
"""
import contextlib
import io
import os
import re
import shutil
import sqlite3
import stat
import subprocess
import sys

import _support  # noqa: E402  (first: builds the test database, points the app at it)

import config  # noqa: E402
import utils.security as sec  # noqa: E402
from app import create_app  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

app = create_app()
_support.attach(app)
TEST_DB = _support.DB_PATH
REPO = str(_support.REPO)

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok)))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}")
    if detail and not ok:
        for line in str(detail).splitlines()[:12]:
            print(f"       {line}")
    return ok


def section(title):
    print(f"\n--- {title} ---")


def q1(sql, args=()):
    conn = sqlite3.connect(TEST_DB)
    conn.row_factory = sqlite3.Row
    try:
        return conn.execute(sql, args).fetchone()
    finally:
        conn.close()


def sign_in(username, password):
    client = app.test_client()
    response = client.post("/login", data={"username": username, "password": password})
    return client, response


def location(response):
    return response.headers.get("Location", "")


def fake_phone(n):
    """A test-only mobile number in the 09XXXXXXXXX shape. Not anyone's number."""
    return "0917" + f"{9000000 + n:07d}"


# A clock the whole application reads (utils.security.now_ts).
clock = [1_800_000_000.0]
sec.now_ts = lambda: clock[0]
FETCH = {"X-Requested-With": "fetch"}

print("=" * 78)
print("IDS-BPHCSN — security basics")
print("=" * 78)

# ===========================================================================
section("1. SECRET KEY AND DEBUG MODE")


def start_app(env_changes):
    """Import app in a new Python with these environment changes. No request is made."""
    env = dict(os.environ)
    env.update(env_changes)
    env["PYTHONIOENCODING"] = "utf-8"
    code = ("import app, config; "
            "print('STARTED debug=%s app.debug=%s' % (config.DEBUG, app.app.debug))")
    proc = subprocess.run([sys.executable, "-c", code], cwd=REPO, env=env,
                          capture_output=True, text=True, encoding="utf-8")
    return proc.returncode, proc.stdout + proc.stderr


GOOD_KEY = "a" * 64
for label, value in (("missing", ""), ("too short", "short-key-123"),
                     ("the old public default", "ids-bphcsn-dev-secret")):
    code, out = start_app({"SECRET_KEY": value})
    check(f"refuses to start when SECRET_KEY is {label}",
          code != 0 and "STARTED" not in out and "cannot start" in out
          and "setup_env.py" in out, out[-300:])
check("secret_key_problem() accepts a 64-character key",
      config.secret_key_problem(GOOD_KEY) is None)
code, out = start_app({"SECRET_KEY": GOOD_KEY, "IDS_DEBUG": ""})
check("starts with a proper key, and debug is off without IDS_DEBUG",
      code == 0 and "STARTED debug=False app.debug=False" in out, out[-300:])
for value in ("0", "true", "yes"):
    code, out = start_app({"SECRET_KEY": GOOD_KEY, "IDS_DEBUG": value})
    check(f"debug stays off for IDS_DEBUG={value}", "debug=False" in out, out[-300:])
code, out = start_app({"SECRET_KEY": GOOD_KEY, "IDS_DEBUG": "1"})
check("debug is on only for IDS_DEBUG=1", "debug=True app.debug=True" in out, out[-300:])

import setup_env  # noqa: E402

env_path = _support.work_path("setup_env_test.env")
setup_env.ENV_FILE = __import__("pathlib").Path(env_path)
out = io.StringIO()
with contextlib.redirect_stdout(out):
    setup_env.main()
written = open(env_path, encoding="utf-8").read()
key_line = re.search(r"^SECRET_KEY=([0-9a-f]+)$", written, re.M)
check("setup_env.py writes a 64-character hex SECRET_KEY and IDS_DEBUG=0",
      key_line is not None and len(key_line.group(1)) == 64
      and re.search(r"^IDS_DEBUG=0$", written, re.M) is not None)
check("setup_env.py never prints the key",
      key_line is not None and key_line.group(1) not in out.getvalue())
out2 = io.StringIO()
with contextlib.redirect_stdout(out2):
    setup_env.main()
check("setup_env.py changes nothing when .env exists, and says so",
      open(env_path, encoding="utf-8").read() == written
      and "Nothing was changed" in out2.getvalue())
example = open(os.path.join(REPO, ".env.example"), encoding="utf-8").read()
check(".env.example names SECRET_KEY with no value",
      re.search(r"^SECRET_KEY=\s*$", example, re.M) is not None)
ignored = subprocess.run(["git", "check-ignore", "-q", ".env"], cwd=REPO).returncode == 0
tracked = subprocess.run(["git", "check-ignore", "-q", ".env.example"], cwd=REPO).returncode != 0
check(".env is ignored by git and .env.example is not", ignored and tracked,
      f".env ignored={ignored} .env.example visible={tracked}")

# ===========================================================================
section("2. SESSION LIMITS")
STAFF_PW = _support.PASSWORDS["staff"]
c, r = sign_in("staff", STAFF_PW)
with c.session_transaction() as s:
    started = (s.get("signed_in_at"), s.get("last_active"))
check("sign-in records the sign-in time and the last activity",
      started == (int(clock[0]), int(clock[0])), started)

clock[0] += 60
r = c.post("/session/keep-alive", headers=FETCH)
with c.session_transaction() as s:
    after = (s.get("signed_in_at"), s.get("last_active"))
check("keep-alive answers 204 when signed in", r.status_code == 204, r.status_code)
check("keep-alive moves the last activity but not the sign-in time",
      after == (started[0], int(clock[0])), after)

clock[0] += 60
c.post("/session/keep-alive", headers=dict(FETCH, **{"X-Session-Probe": "1"}))
with c.session_transaction() as s:
    probed = s.get("last_active")
check("a probe is not activity", probed == after[1], probed)

mark = q1("SELECT COALESCE(MAX(id), 0) FROM audit_logs")[0]
clock[0] += config.IDLE_TIMEOUT.total_seconds() + 60
r = c.get("/staff/dashboard")
check("after 30 idle minutes, a page sends the person to sign in with the 'idle' key",
      r.status_code == 302 and location(r).endswith("/login?error=idle"), location(r))
entry = q1("SELECT action, details FROM audit_logs WHERE id > ? ORDER BY id DESC", (mark,))
check("the idle sign-out is audited",
      entry is not None and entry["action"] == "Signed out after 30 minutes without activity"
      and '"staff"' in entry["details"], entry and dict(entry))
page = app.test_client().get(location(r)).get_data(as_text=True)
check("the login form shows the known idle message", sec.SESSION_IDLE_MESSAGE in page)
with c.session_transaction() as s:
    check("the session was cleared", "user_id" not in s)

c, r = sign_in("staff", STAFF_PW)
mark = q1("SELECT COALESCE(MAX(id), 0) FROM audit_logs")[0]
steps, r = 0, None
while steps < 40:
    clock[0] += 29 * 60
    steps += 1
    r = c.post("/session/keep-alive", headers=FETCH)
    if r.status_code != 204:
        break
hours = steps * 29 / 60
check("keep-alives every 29 minutes do not stretch the 12-hour limit",
      r.status_code == 401 and 12 <= hours < 12.5, f"ended after {hours:.1f} h, {r.status_code}")
check("a fetch request after the limit gets 401 JSON with the 'expired' key",
      r.is_json and r.get_json().get("login_url", "").endswith("/login?error=expired"),
      r.get_data(as_text=True)[:200])
entry = q1("SELECT action FROM audit_logs WHERE id > ? ORDER BY id DESC", (mark,))
check("the 12-hour sign-out is audited",
      entry is not None and entry["action"] == "Signed out after 12 hours", entry and dict(entry))
page = app.test_client().get("/login?error=expired").get_data(as_text=True)
check("the login form shows the known 12-hour message", sec.SESSION_EXPIRED_MESSAGE in page)

r = app.test_client().post("/session/keep-alive", headers=FETCH)
check("keep-alive without a session answers 401 JSON, not the login page",
      r.status_code == 401 and r.is_json and "<html" not in r.get_data(as_text=True).lower())
forged = app.test_client()
with forged.session_transaction() as s:
    s["user_id"] = 2
r = forged.get("/staff/dashboard")
check("a session without its times is refused", r.status_code == 302
      and "/login" in location(r), location(r))

c, r = sign_in("staff", STAFF_PW)
page = c.get("/staff/dashboard").get_data(as_text=True)
check("a signed-in page carries the timings and loads session.js",
      'data-session-idle-ms="1800000"' in page and 'data-session-warn-ms="120000"' in page
      and 'data-session-ping-ms="300000"' in page and "js/session.js" in page)
check("a signed-in page has the warning dialog in plain words",
      "Are you still there?" in page
      and "You will be signed out in 2 minutes because there was no activity." in page
      and "Stay signed in" in page)
check("session.js is served", app.test_client().get("/static/js/session.js").status_code == 200)
check("the login page loads no session script",
      "session.js" not in app.test_client().get("/login").get_data(as_text=True))
a, r = sign_in("admin", _support.PASSWORDS["admin"])
page = a.get("/admin/settings?tab=security").get_data(as_text=True)
pane = page.split('id="pane-security"')[1].split('id="pane-pwd"')[0] if 'id="pane-security"' in page else ""
check("Settings, Security states the limits as a read-only line",
      sec.SESSION_RULE_TEXT in pane
      and sec.SESSION_RULE_TEXT == ("Users are signed out after 30 minutes without "
                                    "activity, and 12 hours after they sign in."))
check("Settings, Security has no disabled controls and no two-factor switch",
      pane and "disabled" not in pane and "Two-Factor" not in pane
      and 'type="checkbox"' not in pane)

# ===========================================================================
section("3. FORCED PASSWORD CHANGE")
FIRST = "Kalapati-Umaga-47"
OWN = "Bundok-Tahimik-82"
a, r = sign_in("admin", _support.PASSWORDS["admin"])
page = a.get("/admin/user_management").get_data(as_text=True)
check("Add User and Reset Password say the user must change it at first sign-in",
      page.count("The user must change this password the first time they sign in.") >= 2)
a.post("/admin/user_management/add", data={
    "full_name": "Forced Change Probe", "username": "forcedprobe",
    "role": "Health Center Staff", "password": FIRST, "confirm_password": FIRST})
row = q1("SELECT id, must_change_password FROM users WHERE username = 'forcedprobe'")
check("a new account is created with the flag set", row is not None and row[1] == 1, row and tuple(row))
u, r = sign_in("forcedprobe", FIRST)
check("its first sign-in goes to My Account", location(r).endswith("/account"), location(r))
for path in ("/staff/dashboard", "/staff/patient_records", "/staff/medicine_inventory",
             "/staff/appointment_scheduling", "/dashboard"):
    r = u.get(path)
    check(f"while flagged, {path} goes to My Account",
          r.status_code == 302 and location(r).endswith("/account"), f"{r.status_code} {location(r)}")
r = u.get("/account")
check("My Account opens and says what to do", r.status_code == 200
      and "Please set your own password before you continue." in r.get_data(as_text=True))
check("keep-alive still works while flagged",
      u.post("/session/keep-alive", headers=FETCH).status_code == 204)
r = u.post("/account/password", data={"current_password": FIRST, "new_password": FIRST,
                                      "confirm_new_password": FIRST})
check("the new password must differ from the current one",
      r.status_code == 200 and "It must be different from the one you have now." in r.get_data(as_text=True))
check("that refusal leaves the flag set",
      q1("SELECT must_change_password FROM users WHERE username='forcedprobe'")[0] == 1)
# One character under the rule's own minimum, so the check follows the rule.
from models.users import PASSWORD_MIN_LENGTH  # noqa: E402

too_short = OWN[:PASSWORD_MIN_LENGTH - 1]
r = u.post("/account/password", data={"current_password": FIRST, "new_password": too_short,
                                      "confirm_new_password": too_short})
check("the new password must follow the existing rule (one character too short is refused)",
      r.status_code == 200 and "Too short" in r.get_data(as_text=True))
mark = q1("SELECT COALESCE(MAX(id), 0) FROM audit_logs")[0]
r = u.post("/account/password", data={"current_password": FIRST, "new_password": OWN,
                                      "confirm_new_password": OWN})
check("a new password of their own is accepted", r.status_code == 302, r.status_code)
check("the flag is cleared after the change",
      q1("SELECT must_change_password FROM users WHERE username='forcedprobe'")[0] == 0)
entry = q1("SELECT action, details FROM audit_logs WHERE id > ? AND action = "
           "'Changed own account password'", (mark,))
check("the change is audited, and says the requirement is met",
      entry is not None and "requirement is now met and cleared" in entry["details"]
      and OWN not in entry["details"] and FIRST not in entry["details"])
check("after the change, the dashboard opens", u.get("/staff/dashboard").status_code == 200)

u2, r = sign_in("forcedprobe", OWN)
a.post("/admin/user_management/reset_password", data={
    "user_id": str(row["id"]), "password": FIRST, "confirm_password": FIRST})
check("an Administrator's reset sets the flag again",
      q1("SELECT must_change_password FROM users WHERE username='forcedprobe'")[0] == 1)
r = u2.get("/staff/dashboard")
check("an open session is sent to My Account at its next click",
      location(r).endswith("/account"), location(r))
r = u2.post("/logout")
check("logout still works while flagged", r.status_code == 302 and "/login" in location(r))
with u2.session_transaction() as s:
    check("and it ends the session", "user_id" not in s)
check("the routes allowed while flagged are exactly these",
      sec.PASSWORD_CHANGE_ENDPOINTS == {"auth.my_account", "auth.change_own_password",
                                        "auth.keep_alive"})

# ===========================================================================
section("4. THE SEED HAS NO FIXED PASSWORD")
from database.schema import init_db  # noqa: E402
from migrate_add_must_change_password import old_demo_passwords  # noqa: E402
from models.users import DEMO_USERS, check_password_strength  # noqa: E402
from werkzeug.security import check_password_hash  # noqa: E402

installs = []
for n in range(2):
    path = _support.work_path(f"fresh_install_{n}.db")
    saved = config.DATABASE
    config.DATABASE = path
    buf = io.StringIO()
    try:
        with contextlib.redirect_stdout(buf):
            init_db()
    finally:
        config.DATABASE = saved
    printed = dict(re.findall(r"^  (\w+)\s+(\S+)$", buf.getvalue(), re.M))
    conn = sqlite3.connect(path)
    rows = conn.execute("SELECT username, password_hash, must_change_password FROM users").fetchall()
    audit = conn.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0]
    conn.close()
    installs.append((buf.getvalue(), printed, rows, audit))

out0, printed0, rows0, audit0 = installs[0]
check("DEMO_USERS holds no passwords (username, full name, role only)",
      all(len(entry) == 3 for entry in DEMO_USERS))
check("a new install creates the four accounts, each flagged",
      sorted(r[0] for r in rows0) == ["admin", "doctor", "official", "staff"]
      and all(r[2] == 1 for r in rows0))
check("it prints each username and password once, with the note",
      sorted(printed0) == ["admin", "doctor", "official", "staff"]
      and "Write these down now. They will not be shown again." in out0
      and "Each user must change it at first sign-in." in out0)
check("each printed password is the one stored (as a hash only)",
      all(check_password_hash(h, printed0[u]) and printed0[u] not in h for u, h, _ in rows0))
check("each printed password passes the password rule",
      all(check_password_strength(printed0[u], u) is None for u in printed0))
check("two new installs get different passwords",
      all(installs[0][1][u] != installs[1][1][u] for u in printed0))
check("the seed writes nothing to the audit trail", audit0 == 0)
# The H6 cleanup took the old passwords out of this repository's history, so
# the reader is tested on a small history made here: one version of
# models/users.py with a password, then one with the text the rewrite left.
hist = _support.work_path("old_history")
os.makedirs(os.path.join(hist, "models"))
GIT = ["git", "-c", "user.name=Test", "-c", "user.email=test@example.invalid"]
subprocess.run(GIT + ["init", "-q", hist], check=True)
for held in ("Lumang-Demo-Pass-31", "(retired demo password)"):
    with open(os.path.join(hist, "models", "users.py"), "w", encoding="utf-8") as fh:
        fh.write(f'DEMO_USERS = [("staff", "{held}", "Test Staff", "Health Center Staff")]\n')
    subprocess.run(GIT + ["-C", hist, "add", "-A"], check=True)
    subprocess.run(GIT + ["-C", hist, "commit", "-q", "-m", "x"], check=True)
check("the reader finds a password DEMO_USERS held in history, and skips the retired text",
      old_demo_passwords(hist) == {"Lumang-Demo-Pass-31"})
# git makes its object files read-only, which would stop the suite's own cleanup.
shutil.rmtree(hist, onexc=lambda func, path, exc: (os.chmod(path, stat.S_IWRITE), func(path)))
try:
    old_demo_passwords()
    history_clean = False
except RuntimeError as exc:
    history_clean = "no old demo password" in str(exc)
check("this repository's history holds no old demo password any more", history_clean)

# ===========================================================================
section("5. THE MUST-CHANGE MIGRATION")
import migrate_add_must_change_password as migration  # noqa: E402
from werkzeug.security import generate_password_hash  # noqa: E402

GIVEN_OLD = "Lumang-Demo-Pass-31"
mig_db = _support.work_path("before_migration.db")
conn = sqlite3.connect(mig_db)
conn.executescript(open(os.path.join(REPO, "schema.sql"), encoding="utf-8").read())
conn.execute("ALTER TABLE users DROP COLUMN must_change_password")
conn.execute("INSERT INTO users (username, password_hash, full_name, role) VALUES "
             "('onold', ?, 'Old Password Holder', 'Health Center Staff')",
             (generate_password_hash(GIVEN_OLD),))
conn.execute("INSERT INTO users (username, password_hash, full_name, role) VALUES "
             "('onown', ?, 'Own Password Holder', 'Health Center Staff')",
             (generate_password_hash("Sariling-Pili-Nila-66"),))
conn.commit()
conn.close()
backups = _support.work_path("migration_backups")
first = migration.run(mig_db, {GIVEN_OLD}, backup_dir=backups)
check("first run adds the column and flags the account on the given old password",
      first["column_added"] and first["flagged"] == ["onold"] and not first["problems"], first)
check("first run made a backup before changing anything",
      first["backup"] is not None and os.path.exists(first["backup"])
      and "must_change_password" not in [r[1] for r in sqlite3.connect(first["backup"])
                                         .execute("PRAGMA table_info(users)")])
second = migration.run(mig_db, {GIVEN_OLD}, backup_dir=backups)
check("second run changes nothing and makes no backup",
      not second["column_added"] and second["flagged"] == [] and second["backup"] is None)
conn = sqlite3.connect(mig_db)
flags = dict(conn.execute("SELECT username, must_change_password FROM users"))
conn.close()
check("after both runs, only the old-password account is flagged",
      flags == {"onold": 1, "onown": 0}, flags)
stop_db = _support.work_path("stop_migration.db")
conn = sqlite3.connect(stop_db)
conn.executescript(open(os.path.join(REPO, "schema.sql"), encoding="utf-8").read())
conn.execute("ALTER TABLE users DROP COLUMN must_change_password")
conn.commit()
conn.close()
before = open(stop_db, "rb").read()
buf = io.StringIO()
with contextlib.redirect_stdout(buf):
    code = migration.main([stop_db])
check("with no old password in history, the migration stops, changes nothing and names --flag-all",
      code == 1 and open(stop_db, "rb").read() == before and "--flag-all" in buf.getvalue(),
      buf.getvalue())

# ===========================================================================
section("6. NO TEXT FROM THE URL ON SCREEN (L8)")
EVIL = "ZZURLTEXT<b>x</b>"
a, r = sign_in("admin", _support.PASSWORDS["admin"])
page = a.get(f"/admin/reports_analytics?type={EVIL}&from={EVIL}&to={EVIL}").get_data(as_text=True)
check("Reports: a bad type, start and end date are not printed back", "ZZURLTEXT" not in page)
check("Reports: fixed messages are shown instead",
      "The start date in the link is not a date this screen can read." in page
      and "The report type in the link is not one this screen offers" in page)
data = a.get(f"/admin/reports_analytics/export.xlsx?from={EVIL}&to={EVIL}").data
cells = " ".join(str(v) for ws in load_workbook(io.BytesIO(data)).worksheets
                 for row in ws.iter_rows(values_only=True) for v in row if v is not None)
check("Reports: the Excel export does not carry the URL text either", "ZZURLTEXT" not in cells)
o, r = sign_in("official", _support.PASSWORDS["official"])
page = o.get(f"/official/community-reports?from={EVIL}&to=2026-01-01").get_data(as_text=True)
check("Community Reports: the URL text is not printed back", "ZZURLTEXT" not in page)

s, r = sign_in("staff", STAFF_PW)
page = s.get(f"/staff/register-patient?check_in=1&check_in_reason={EVIL}").get_data(as_text=True)
check("Register Patient: a check-in reason in the URL is ignored", "ZZURLTEXT" not in page)
page = s.get(f"/staff/register-patient?address={EVIL}&zip_code={EVIL}").get_data(as_text=True)
check("Register Patient: address text in the URL is ignored", "ZZURLTEXT" not in page)
r = s.post("/staff/register-patient", data={"check_in_start": "1",
                                            "check_in_reason": "Ubo at sipon ng tatlong araw"})
check("the Check In dialog hands the reason over by POST, and the URL has no reason",
      r.status_code == 302 and location(r).endswith("/staff/register-patient?check_in=1"),
      location(r))
check("the form shows the reason the dialog handed over",
      "Ubo at sipon ng tatlong araw" in s.get(location(r)).get_data(as_text=True))
form = {
    "first_name": "Lorena", "last_name": "Pagsanjan", "birth_date": "1991-04-04",
    "sex": "Female", "civil_status": "Single", "address": "Purok 5, tabi ng kapilya",
    "municipality": "San Francisco", "province": "Surigao del Norte", "zip_code": "8405",
    "emergency_contact_name": "Test Relative", "emergency_contact_relationship": "Sibling",
    "emergency_contact_number": fake_phone(1), "confirm_duplicate": "yes",
    "save_action": "register_another",
}
r = s.post("/staff/register-patient", data=form)
check("Save & Register Another carries a patient id, not the address",
      r.status_code == 302 and "prefill_from=" in location(r) and "Purok" not in location(r),
      location(r))
page = s.get(location(r)).get_data(as_text=True)
check("the fresh form reads the address from the saved record",
      "Purok 5, tabi ng kapilya" in page and 'value="8405"' in page)
page = s.get("/staff/register-patient?prefill_from=999999").get_data(as_text=True)
check("an id that names no patient fills nothing", "Purok 5, tabi ng kapilya" not in page)
page = app.test_client().get(f"/login?error={EVIL}").get_data(as_text=True)
check("the login form shows nothing for an unknown message key", "ZZURLTEXT" not in page)

# ===========================================================================
passed = sum(1 for _, ok in results if ok)
print("\n" + "=" * 78)
print(f"RESULT: {passed}/{len(results)} checks passed")
failed = [name for name, ok in results if not ok]
for name in failed:
    print("  FAILED:", name)
print("=" * 78)
sys.exit(0 if not failed else 1)
