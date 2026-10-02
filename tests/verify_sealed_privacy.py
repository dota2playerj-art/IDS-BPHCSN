"""Sealed record privacy (H4, L2, Task 5b; RA 11166 Section 44, RA 10173).

Checks that every free-text field of a sealed visit is encrypted in the database
(diagnosis, treatment plan, chief complaint, additional notes, and the check-in
reason of its queue visit), that only the Doctor's screens ever show them, that
no export, print, report, audit entry or backup carries them, that the
Administrator sees neither the check-in reason on the queue nor the medicine
names of a sealed visit, and that a wrong key fails with a plain message instead
of a crash.

Runs on a fresh test database built by tests/_support.py, with a test-only key.
Every "absent" check has a positive control: an unsealed sentinel that must be
found by the same extraction. Sentinels are single tokens, so a PDF line break
cannot split one and hide a leak.
"""
import ast
import base64
import contextlib
import io
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import zlib
from pathlib import Path

import _support  # noqa: E402  (first: builds the test database, points the app at it)

from app import create_app  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

sys.stdout.reconfigure(errors="replace")

app = create_app()
_support.attach(app)
TEST_DB = _support.DB_PATH
REPO = Path(_support.REPO)
ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID = 1, 2, 3, 5
UNOPENABLE = "This sealed record cannot be opened on this computer."
PLACEHOLDER = "Restricted — Doctor access only"

DX_REC, TX_REC = "ZZSEALREC1DX", "ZZSEALREC1TX"
DX_LATE, TX_LATE = "ZZSEALLATE2DX", "ZZSEALLATE2TX"
DX_LATE_EDIT = "ZZSEALLATE2EDITDX"
DX_CTRL = "ZZOPENCTRL3DX"
# Task 5b: the rest of a sealed visit's free text.
CC_REC, NOTE_REC, REASON_REC = "ZZSEALREC1CC", "ZZSEALREC1NOTE", "ZZSEALREC1REASON"
CC_LATE, NOTE_LATE, REASON_LATE = "ZZSEALLATE2CC", "ZZSEALLATE2NOTE", "ZZSEALLATE2REASON"
CC_LATE_EDIT = "ZZSEALLATE2EDITCC"
# Unsealed controls: the same fields on visits nobody sealed.
CC_CTRL, NOTE_CTRL, REASON_CTRL = "ZZOPENCTRL3CC", "ZZOPENCTRL3NOTE", "ZZOPENCTRL3REASON"
REASON_QUEUE = "ZZOPENQUEUE4REASON"  # a visit still waiting in today's queue
CC_STAFF, NOTE_STAFF = "ZZOPENSTAFF5CC", "ZZOPENSTAFF5NOTE"  # saved by Staff, unsealed
MED_SEALED, MED_OPEN = "ZZSEALEDMED5", "ZZOPENMED6"  # medicine names, snapshot text
SEALED_SENTINELS = (DX_REC, TX_REC, DX_LATE, TX_LATE, DX_LATE_EDIT,
                    CC_REC, NOTE_REC, REASON_REC, CC_LATE, NOTE_LATE, REASON_LATE,
                    CC_LATE_EDIT)
# Free clinical text that must not be in any audit entry, sealed or not (D5).
FREE_TEXT = SEALED_SENTINELS + (DX_CTRL, CC_CTRL, NOTE_CTRL, REASON_CTRL, REASON_QUEUE,
                                CC_STAFF, NOTE_STAFF)

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok)))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}")
    if detail and not ok:
        for line in str(detail).splitlines()[:12]:
            print(f"       {line}")
    return ok


def section(title):
    print(f"\n--- {title} ---")


def db():
    conn = sqlite3.connect(TEST_DB)
    conn.row_factory = sqlite3.Row
    return conn


def q1(sql, args=()):
    conn = db()
    try:
        return conn.execute(sql, args).fetchone()
    finally:
        conn.close()


def client_as(user_id):
    c = app.test_client()
    with c.session_transaction() as sess:
        sess["user_id"] = user_id
        _support.open_session(sess)
    return c


def body(response):
    return response.get_data(as_text=True)


def pdf_text(data):
    out = []
    for match in re.finditer(rb"stream\r?\n(.*?)endstream", data, re.S):
        chunk = match.group(1)
        try:
            out.append(zlib.decompress(chunk).decode("latin-1"))
        except Exception:
            out.append(chunk.decode("latin-1", "replace"))
    return "\n".join(out)


def xlsx_text(data):
    book = load_workbook(io.BytesIO(data), data_only=True)
    out = []
    for sheet in book.worksheets:
        for row in sheet.iter_rows(values_only=True):
            out.extend(str(cell) for cell in row if cell is not None)
    return "\n".join(out)


def readable(response):
    """Everything a person could read in a response: page text, workbook, PDF."""
    data = response.data
    kind = response.headers.get("Content-Type", "")
    if "spreadsheetml" in kind:
        return xlsx_text(data)
    if "pdf" in kind:
        return pdf_text(data) + data.decode("latin-1")
    return data.decode("utf-8", "replace")


def leaks(text):
    return [s for s in SEALED_SENTINELS if s in text]


def new_queue_entry(patient_id, reason=None, status="In Consultation"):
    conn = db()
    cur = conn.execute(
        "INSERT INTO appointments (patient_id, scheduled_by, appointment_date, "
        "status, is_walkin, checked_in_at, notes) VALUES (?, ?, date('now','localtime'), "
        "?, 1, datetime('now','localtime'), ?)", (patient_id, STAFF_ID, status, reason))
    conn.commit()
    new_id = cur.lastrowid
    conn.close()
    return new_id


def run_sql(sql, args=()):
    conn = db()
    try:
        conn.execute(sql, args)
        conn.commit()
    finally:
        conn.close()


def visible(text, who):
    """What a person sees on screen. The Doctor's screens hold sealed text inside
    no-print spans, which never reach paper, so those are taken out for the paper
    test; every other role has none to take out."""
    return re.sub(r'<span class="no-print">.*?</span>', "", text, flags=re.S)


def edit_form(cid, **overrides):
    row = dict(q1("SELECT * FROM consultations WHERE id=?", (cid,)))
    form = {}
    for field in ("chief_complaint", "diagnosis", "treatment_plan", "remarks",
                  "temperature", "blood_pressure", "pulse_rate", "respiratory_rate",
                  "weight", "height", "head_circumference", "follow_up_date"):
        value = row.get(field)
        # The Doctor's form shows a sealed record decrypted, so this does too.
        if row.get("is_restricted") and sc.is_encrypted(value):
            try:
                value = sc.decrypt(value, cid, field)
            except sc.SealedDataError:
                pass  # a wrong or missing key: the form would not carry the field
        form[field] = "" if value is None else (f"{value:g}" if isinstance(value, float)
                                                 else str(value))
    if row.get("is_restricted"):
        form["is_restricted"] = "1"
    form.update(overrides)
    return form


def audit_max():
    return q1("SELECT COALESCE(MAX(id), 0) FROM audit_logs")[0]


admin, staff, official, doctor = (client_as(i) for i in
                                  (ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID))
from utils import sealed_crypto as sc  # noqa: E402

print("=" * 78)
print("IDS-BPHCSN — sealed record privacy (H4, L2)")
print("=" * 78)

# ===========================================================================
section("1. THE ENCRYPTION MODULE")
token = sc.encrypt("Sample sealed text", 7, "diagnosis")
check("a value encrypts to the versioned format enc:v1:", token.startswith("enc:v1:"))
check("and decrypts back for the same row and field",
      sc.decrypt(token, 7, "diagnosis") == "Sample sealed text")
raw = base64.b64decode(token[len("enc:v1:"):])
check("the stored value holds a 96-bit nonce, the text and a 16-byte tag",
      len(raw) == 12 + len("Sample sealed text".encode()) + 16, len(raw))
check("the same text encrypts differently each time (fresh nonce)",
      sc.encrypt("Sample sealed text", 7, "diagnosis") != token)


def refuses(call):
    try:
        call()
    except sc.SealedDataError:
        return True
    return False


check("a value moved to another row does not decrypt",
      refuses(lambda: sc.decrypt(token, 8, "diagnosis")))
check("a value moved to another field does not decrypt",
      refuses(lambda: sc.decrypt(token, 7, "treatment_plan")))
check("a damaged value does not decrypt",
      refuses(lambda: sc.decrypt(token[:-6] + "AAAAAA", 7, "diagnosis")))
visit_token = sc.encrypt("Sample reason", 7, "notes", "appointments")
check("a queue visit's reason is bound to its own row and table",
      sc.decrypt(visit_token, 7, "notes", "appointments") == "Sample reason"
      and refuses(lambda: sc.decrypt(visit_token, 8, "notes", "appointments"))
      and refuses(lambda: sc.decrypt(visit_token, 7, "notes"))
      and refuses(lambda: sc.decrypt(visit_token, 7, "notes", "consultations")))
check("None and empty text are left as they are",
      sc.encrypt(None, 7, "diagnosis") is None and sc.encrypt("", 7, "diagnosis") == "")
check("key_problem() refuses a missing, non-base64 or 16-byte key",
      sc.key_problem("") and sc.key_problem("not base64!!")
      and sc.key_problem(base64.b64encode(os.urandom(16)).decode()))
check("key_problem() accepts a 32-byte key", sc.key_problem(sc.new_key()) is None)

# ===========================================================================
section("2. STARTUP, setup_env.py AND export_sealed_key.py")


def start_app(env_changes):
    env = dict(os.environ)
    env.update(env_changes)
    env["PYTHONIOENCODING"] = "utf-8"
    proc = subprocess.run([sys.executable, "-c", "import app; print('STARTED')"],
                          cwd=REPO, env=env, capture_output=True, text=True,
                          encoding="utf-8")
    return proc.returncode, proc.stdout + proc.stderr


for label, value in (("missing", ""),
                     ("16 bytes", base64.b64encode(os.urandom(16)).decode()),
                     ("not base64", "%%%%")):
    code, out = start_app({"SEALED_DATA_KEY": value})
    check(f"the app refuses to start when SEALED_DATA_KEY is {label}",
          code != 0 and "STARTED" not in out and "cannot start" in out
          and "SEALED_DATA_KEY" in out, out[-300:])
code, out = start_app({"SEALED_DATA_KEY": sc.new_key()})
check("the app starts with a proper 32-byte key", code == 0 and "STARTED" in out, out[-300:])

import setup_env  # noqa: E402

env_path = Path(_support.work_path("existing.env"))
old_secret = "f" * 64
env_path.write_text(f"# old file\nSECRET_KEY={old_secret}\nIDS_DEBUG=0\n", encoding="utf-8")
setup_env.ENV_FILE = env_path
printed = io.StringIO()
with contextlib.redirect_stdout(printed):
    setup_env.main()
written = env_path.read_text(encoding="utf-8")
key_line = re.search(r"^SEALED_DATA_KEY=(\S+)$", written, re.M)
check("setup_env.py adds SEALED_DATA_KEY to an existing .env",
      key_line is not None and sc.key_problem(key_line.group(1)) is None, written)
check("and leaves SECRET_KEY and the rest exactly as they were",
      written.startswith(f"# old file\nSECRET_KEY={old_secret}\nIDS_DEBUG=0\n"))
check("and never prints the key",
      key_line is not None and key_line.group(1) not in printed.getvalue())
again = io.StringIO()
with contextlib.redirect_stdout(again):
    setup_env.main()
check("a second run changes nothing", env_path.read_text(encoding="utf-8") == written
      and "Nothing was changed" in again.getvalue())
new_path = Path(_support.work_path("new.env"))
setup_env.ENV_FILE = new_path
with contextlib.redirect_stdout(io.StringIO()):
    setup_env.main()
fresh = new_path.read_text(encoding="utf-8")
check("a new .env gets both keys",
      re.search(r"^SECRET_KEY=[0-9a-f]{64}$", fresh, re.M) is not None
      and re.search(r"^SEALED_DATA_KEY=\S+$", fresh, re.M) is not None)
example = (REPO / ".env.example").read_text(encoding="utf-8")
check(".env.example names SEALED_DATA_KEY with no value",
      re.search(r"^SEALED_DATA_KEY=\s*$", example, re.M) is not None)

import export_sealed_key  # noqa: E402

export_sealed_key.ENV_FILE = env_path
the_key = key_line.group(1) if key_line else "missing"
folder = sorted(os.listdir(env_path.parent))
shown = io.StringIO()
with contextlib.redirect_stdout(shown):
    refused = export_sealed_key.main(read=lambda _prompt: "no")
check("export_sealed_key.py shows nothing unless YES is typed",
      refused == 1 and the_key not in shown.getvalue())
shown = io.StringIO()
with contextlib.redirect_stdout(shown):
    accepted = export_sealed_key.main(read=lambda _prompt: "YES")
check("after YES it prints the key with the envelope instructions",
      accepted == 0 and the_key in shown.getvalue() and "envelope" in shown.getvalue()
      and "Doctor" in shown.getvalue())
check("and writes nothing to disk", sorted(os.listdir(env_path.parent)) == folder)

# ===========================================================================
section("3. SEALING, UNSEALING AND EDITING THROUGH THE REAL ROUTES")
P_REC = 9
appt = new_queue_entry(P_REC, REASON_REC)
r = doctor.post(f"/staff/consultation/record/{appt}", data={
    "chief_complaint": CC_REC, "remarks": NOTE_REC, "temperature": "36.9", "weight": "60",
    "diagnosis": DX_REC, "treatment_plan": TX_REC, "is_restricted": "1"})
CID_REC = q1("SELECT id FROM consultations WHERE appointment_id = ?", (appt,))[0]
row = q1("SELECT * FROM consultations WHERE id = ?", (CID_REC,))
check("sealed at record time: both fields are stored encrypted",
      row["is_restricted"] == 1 and row["restricted_by"] == DOCTOR_ID
      and sc.is_encrypted(row["diagnosis"]) and sc.is_encrypted(row["treatment_plan"]),
      dict(row))
check("and they decrypt to what the Doctor typed",
      sc.decrypt(row["diagnosis"], CID_REC, "diagnosis") == DX_REC
      and sc.decrypt(row["treatment_plan"], CID_REC, "treatment_plan") == TX_REC)
check("Task 5b: the chief complaint and additional notes are encrypted too, and decrypt",
      sc.is_encrypted(row["chief_complaint"]) and sc.is_encrypted(row["remarks"])
      and sc.decrypt(row["chief_complaint"], CID_REC, "chief_complaint") == CC_REC
      and sc.decrypt(row["remarks"], CID_REC, "remarks") == NOTE_REC, dict(row))
visit = q1("SELECT notes FROM appointments WHERE id = ?", (appt,))
check("Task 5b: the check-in reason of the linked queue visit is encrypted, bound to that row",
      sc.is_encrypted(visit["notes"])
      and sc.decrypt(visit["notes"], appt, "notes", "appointments") == REASON_REC
      and refuses(lambda: sc.decrypt(visit["notes"], appt, "notes")), dict(visit))
check("what is not free text stays readable: vitals, date and the seal itself",
      row["temperature"] == 36.9 and row["weight"] == 60.0 and row["consultation_date"]
      and row["is_restricted"] == 1)
profile = body(doctor.get(f"/staff/patient_profile/{P_REC}"))
check("L2: the seal message says the text is hidden from everyone except the Doctor",
      "hidden from everyone except the Doctor" in profile, profile[:200])
check("the Doctor's profile shows the sealed text", DX_REC in profile and TX_REC in profile)

CID_LATE = 10
P_LATE = q1("SELECT patient_id FROM consultations WHERE id = ?", (CID_LATE,))[0]
doctor.post(f"/staff/consultation/edit/{CID_LATE}",
            data=edit_form(CID_LATE, diagnosis=DX_LATE, treatment_plan=TX_LATE,
                           chief_complaint=CC_LATE, remarks=NOTE_LATE))
APPT_LATE = q1("SELECT appointment_id FROM consultations WHERE id=?", (CID_LATE,))[0]
run_sql("UPDATE appointments SET notes = ? WHERE id = ?", (REASON_LATE, APPT_LATE))
row = q1("SELECT * FROM consultations WHERE id=?", (CID_LATE,))
check("before sealing, an ordinary record holds plain text",
      row["diagnosis"] == DX_LATE and row["chief_complaint"] == CC_LATE
      and row["remarks"] == NOTE_LATE
      and q1("SELECT notes FROM appointments WHERE id=?", (APPT_LATE,))[0] == REASON_LATE)
restrict_url = f"/staff/patient_profile/{P_LATE}/restrict-consultation"
doctor.post(restrict_url, data={"consultation_id": str(CID_LATE), "restrict_action": "restrict"})
row = q1("SELECT * FROM consultations WHERE id = ?", (CID_LATE,))
first_cipher = row["diagnosis"]
check("sealing encrypts both fields in the database",
      row["is_restricted"] == 1 and sc.is_encrypted(row["diagnosis"])
      and sc.is_encrypted(row["treatment_plan"]) and DX_LATE not in row["diagnosis"])
check("Task 5b: sealing a record already on file also encrypts its chief complaint, "
      "notes and check-in reason, in the same save",
      sc.is_encrypted(row["chief_complaint"]) and sc.is_encrypted(row["remarks"])
      and sc.decrypt(row["chief_complaint"], CID_LATE, "chief_complaint") == CC_LATE
      and sc.decrypt(row["remarks"], CID_LATE, "remarks") == NOTE_LATE
      and sc.decrypt(q1("SELECT notes FROM appointments WHERE id=?", (APPT_LATE,))[0],
                     APPT_LATE, "notes", "appointments") == REASON_LATE)
doctor.post(restrict_url, data={"consultation_id": str(CID_LATE), "restrict_action": "lift"})
row = q1("SELECT * FROM consultations WHERE id = ?", (CID_LATE,))
check("lifting the seal restores the plain text",
      row["is_restricted"] == 0 and row["diagnosis"] == DX_LATE
      and row["treatment_plan"] == TX_LATE, dict(row))
check("Task 5b: and restores the chief complaint, the notes and the check-in reason",
      row["chief_complaint"] == CC_LATE and row["remarks"] == NOTE_LATE
      and q1("SELECT notes FROM appointments WHERE id=?", (APPT_LATE,))[0] == REASON_LATE)
doctor.post(restrict_url, data={"consultation_id": str(CID_LATE), "restrict_action": "restrict"})
row = q1("SELECT * FROM consultations WHERE id = ?", (CID_LATE,))
check("sealing again encrypts again, with a new nonce",
      sc.is_encrypted(row["diagnosis"]) and row["diagnosis"] != first_cipher)

plan_before = row["treatment_plan"]
notes_before = row["remarks"]
doctor.post(f"/staff/consultation/edit/{CID_LATE}",
            data=edit_form(CID_LATE, diagnosis=DX_LATE_EDIT, treatment_plan=TX_LATE,
                           chief_complaint=CC_LATE_EDIT))
row = q1("SELECT * FROM consultations WHERE id = ?", (CID_LATE,))
check("Task 5b: a Doctor edit of the chief complaint on a sealed record stays encrypted",
      sc.is_encrypted(row["chief_complaint"]) and CC_LATE_EDIT not in row["chief_complaint"]
      and sc.decrypt(row["chief_complaint"], CID_LATE, "chief_complaint") == CC_LATE_EDIT)
check("Task 5b: the notes the Doctor left alone keep their stored ciphertext",
      row["remarks"] == notes_before)
check("a Doctor edit of a sealed record stays encrypted",
      row["is_restricted"] == 1 and sc.is_encrypted(row["diagnosis"])
      and sc.decrypt(row["diagnosis"], CID_LATE, "diagnosis") == DX_LATE_EDIT)
check("a field the Doctor left alone keeps its stored ciphertext",
      row["treatment_plan"] == plan_before)
sealed_now = (row["diagnosis"], row["treatment_plan"])
sealed_all = tuple(row[f] for f in ("diagnosis", "treatment_plan", "chief_complaint", "remarks"))
reason_now = q1("SELECT notes FROM appointments WHERE id=?", (APPT_LATE,))[0]
staff.post(f"/staff/consultation/edit/{CID_LATE}",
           data={k: v for k, v in edit_form(CID_LATE, temperature="37.5").items()
                 if k not in ("diagnosis", "treatment_plan", "is_restricted",
                              "chief_complaint", "remarks")})
row = q1("SELECT * FROM consultations WHERE id = ?", (CID_LATE,))
check("a staff edit of a sealed record leaves both encrypted fields byte for byte",
      (row["diagnosis"], row["treatment_plan"]) == sealed_now and row["is_restricted"] == 1
      and row["temperature"] == 37.5)
check("Task 5b: and the encrypted chief complaint, notes and check-in reason too",
      tuple(row[f] for f in ("diagnosis", "treatment_plan", "chief_complaint", "remarks"))
      == sealed_all
      and q1("SELECT notes FROM appointments WHERE id=?", (APPT_LATE,))[0] == reason_now)
mark = audit_max()
staff.post(f"/staff/consultation/edit/{CID_LATE}",
           data={k: v for k, v in edit_form(CID_LATE, chief_complaint="ZZFORGEDCC").items()
                 if k not in ("diagnosis", "treatment_plan", "is_restricted", "remarks")})
row = q1("SELECT * FROM consultations WHERE id = ?", (CID_LATE,))
check("Task 5b: a forged chief complaint from Staff on a sealed record is refused and recorded",
      tuple(row[f] for f in ("diagnosis", "treatment_plan", "chief_complaint", "remarks"))
      == sealed_all
      and db().execute("SELECT COUNT(*) FROM audit_logs WHERE id > ? AND action = "
                       "'Refused clinical judgement from an unlicensed role'", (mark,)
                       ).fetchone()[0] == 1
      and not any("ZZFORGEDCC" in (d[0] or "") for d in
                  db().execute("SELECT details FROM audit_logs WHERE id > ?", (mark,))))

CID_CTRL = 11
P_CTRL = q1("SELECT patient_id FROM consultations WHERE id = ?", (CID_CTRL,))[0]
doctor.post(f"/staff/consultation/edit/{CID_CTRL}", data=edit_form(CID_CTRL, diagnosis=DX_CTRL))
# The unsealed controls for the new fields: plain text, so the same extraction that
# must not find a sealed sentinel is shown to find these.
APPT_CTRL = q1("SELECT appointment_id FROM consultations WHERE id=?", (CID_CTRL,))[0]
run_sql("UPDATE consultations SET chief_complaint = ?, remarks = ? WHERE id = ?",
        (CC_CTRL, NOTE_CTRL, CID_CTRL))
run_sql("UPDATE appointments SET notes = ? WHERE id = ?", (REASON_CTRL, APPT_CTRL))
check("control: an unsealed record keeps plain text",
      q1("SELECT diagnosis FROM consultations WHERE id=?", (CID_CTRL,))[0] == DX_CTRL)
row = q1("SELECT * FROM consultations WHERE id=?", (CID_CTRL,))
check("control: an unsealed record's chief complaint, notes and check-in reason stay plain",
      row["chief_complaint"] == CC_CTRL and row["remarks"] == NOTE_CTRL
      and q1("SELECT notes FROM appointments WHERE id=?", (APPT_CTRL,))[0] == REASON_CTRL)
# A waiting visit nobody has written up, with a reason: the queue's own control.
P_QUEUE = 13
QUEUE_APPT = new_queue_entry(P_QUEUE, REASON_QUEUE, status="Waiting")
# Medicine names handed over at a sealed and at an unsealed visit (D3). The rows are
# written straight to the table: the name is a snapshot of text, which is what the
# readers under test show.
run_sql("INSERT INTO medicine_dispensing (patient_id, consultation_id, inventory_item_id, "
        "medicine_name, unit, quantity_dispensed, date_dispensed, dispensed_by) "
        "VALUES (?, ?, 2, ?, 'capsule', 7, date('now','localtime'), ?)",
        (P_REC, CID_REC, MED_SEALED, STAFF_ID))
run_sql("INSERT INTO medicine_dispensing (patient_id, consultation_id, inventory_item_id, "
        "medicine_name, unit, quantity_dispensed, date_dispensed, dispensed_by) "
        "VALUES (?, ?, 3, ?, 'tablet', 5, date('now','localtime'), ?)",
        (P_CTRL, CID_CTRL, MED_OPEN, STAFF_ID))

# ===========================================================================
section("4. ONLY THE DOCTOR'S SCREENS SHOW SEALED TEXT")
plain_gets = sorted({rule.rule for rule in app.url_map.iter_rules()
                     if "GET" in rule.methods and not rule.arguments
                     and not rule.rule.startswith("/static")
                     and rule.rule not in ("/logout",)})
check("the crawl covers the app's GET screens and exports", len(plain_gets) > 40,
      len(plain_gets))
for who, client, prefix in (("Administrator", admin, "/admin"), ("Health Center Staff", staff, "/staff"),
                            ("Barangay Official", official, "/official"), ("Doctor", doctor, "/staff")):
    found = []
    for url in plain_gets:
        response = client.get(url)
        # The Doctor reads sealed text on screen only. A no-print span is on the
        # screen and not on paper, so it is set aside here; an export has none.
        hits = leaks(visible(readable(response), who))
        if hits:
            found.append(f"{url}: {hits}")
    check(f"{who}: no sealed text on any screen, export or download without an id",
          not found, "\n".join(found))
    if who in ("Barangay Official",):
        continue
    for pid in (P_REC, P_LATE):
        for url in (f"{prefix}/patient_profile/{pid}", f"{prefix}/patient_profile/{pid}?print=1"):
            page = body(client.get(url))
            if who == "Doctor":
                screen_only = re.sub(r'<span class="no-print">.*?</span>', "", page, flags=re.S)
                check(f"Doctor: {url} shows sealed text on screen, never on paper",
                      leaks(page) and not leaks(screen_only)
                      and f'<span class="print-only">{PLACEHOLDER}</span>' in page,
                      leaks(screen_only))
            elif who == "Administrator":
                # Task 5c: his profile is demographics only. There is no clinical
                # section to hold a placeholder, only one plain line.
                check(f"{who}: {url} shows demographics only, not the text",
                      not leaks(page) and "Clinical details are shown only to the Doctor and Health Center Staff." in page, leaks(page))
            else:
                check(f"{who}: {url} shows the placeholder, not the text",
                      not leaks(page) and PLACEHOLDER in page, leaks(page))
    for cid in (CID_REC, CID_LATE):
        response = client.get(f"{prefix}/consultation/edit/{cid}")
        page = body(response)
        if who == "Doctor":
            check(f"Doctor: the edit form for #{cid} shows the sealed text to edit",
                  any(s in page for s in SEALED_SENTINELS))
        elif who == "Administrator":
            # Task 5c: he has no edit form at all now.
            check(f"{who}: there is no edit form for #{cid} (404), so no sealed text",
                  response.status_code == 404 and not leaks(page), response.status_code)
        else:
            check(f"{who}: the edit form for #{cid} carries no sealed text",
                  not leaks(page), leaks(page))
    for kind in ("xlsx", "pdf"):
        response = client.post(f"{prefix}/patient_records/export-selection.{kind}",
                               data={"patient_ids": [str(P_REC), str(P_LATE)]})
        check(f"{who}: the selection export ({kind}) carries no sealed text",
              not leaks(readable(response)), response.status_code)
r = doctor.post("/session/keep-alive", headers={"X-Requested-With": "fetch"})
check("the keep-alive (JSON) response carries no sealed text", not leaks(body(r)))
# Task 5c: the Administrator sees no clinical text at all, so the control that an
# unsealed diagnosis IS shown now uses Health Center Staff.
page = body(staff.get(f"/staff/patient_profile/{P_CTRL}"))
check("control: an unsealed diagnosis is shown to Health Center Staff", DX_CTRL in page)
page = body(admin.get(f"/admin/patient_profile/{P_CTRL}"))
check("and not to the Administrator, sealed or not (Task 5c)", DX_CTRL not in page)

# ===========================================================================
section("5. THE KEY TEST: THE BACKUP DOWNLOAD HOLDS NO SEALED TEXT")
# Task 5c (M25): the download is an encrypted .idsbackup file. Its bytes hold no
# text at all, sealed or not. The checks below that need the database inside it
# open it the way the health center does, with restore_backup.py's decryptor.
import restore_backup  # noqa: E402

response = admin.get("/admin/settings/backup.idsbackup")
encrypted = response.data
check("the Administrator's backup route answers with an encrypted backup file",
      response.status_code == 200 and encrypted[:7] == b"IDSBAK1"
      and encrypted[:16] != b"SQLite format 3\x00", response.status_code)
check("*** the downloaded file holds no sentinel at all, sealed or unsealed ***",
      not [t for t in SEALED_SENTINELS + (DX_CTRL, CC_CTRL, NOTE_CTRL, REASON_CTRL)
           if t.encode() in encrypted])
data = restore_backup.decrypt_backup(encrypted)
check("opened with the key, it is a database file", data[:16] == b"SQLite format 3\x00")
check("control: the unsealed sentinel IS found in the opened backup", DX_CTRL.encode() in data)
check("control: an unsealed chief complaint, note and check-in reason ARE found in the "
      "opened backup",
      all(t.encode() in data for t in (CC_CTRL, NOTE_CTRL, REASON_CTRL)))
check("*** no sealed chief complaint, note or check-in reason is in the backup bytes ***",
      not [t for t in (CC_REC, NOTE_REC, REASON_REC, CC_LATE, NOTE_LATE, REASON_LATE,
                       CC_LATE_EDIT) if t.encode() in data])
check("*** no sealed sentinel appears anywhere in the backup bytes ***",
      not [s for s in SEALED_SENTINELS if s.encode() in data],
      [s for s in SEALED_SENTINELS if s.encode() in data])
restored = Path(_support.work_path("restored.db"))
restored.write_bytes(data)
conn = sqlite3.connect(restored)
row = conn.execute("SELECT diagnosis, is_restricted FROM consultations WHERE id = ?",
                   (CID_REC,)).fetchone()
conn.close()
check("the backup restores fully: the sealed row is there, encrypted, and opens with the key",
      row is not None and row[1] == 1 and sc.decrypt(row[0], CID_REC, "diagnosis") == DX_REC)
live = Path(TEST_DB).read_bytes()
check("the database file itself holds no sealed text either (secure_delete)",
      not [s for s in SEALED_SENTINELS if s.encode() in live])
settings_page = body(admin.get("/admin/settings"))
check("Settings → Backup says the backup file is encrypted (Task 5c wording)",
      "The backup file is encrypted. It can only be opened with this system&#39;s key, "
      "using restore_backup.py on the health-center computer." in settings_page
      or "The backup file is encrypted. It can only be opened with this system's key, "
      "using restore_backup.py on the health-center computer." in settings_page)

# ===========================================================================
section("6. NO AUDIT ENTRY HOLDS DIAGNOSIS OR TREATMENT TEXT")
details = [r[0] or "" for r in db().execute("SELECT details FROM audit_logs")]
check("no entry quotes a sealed diagnosis or treatment plan",
      not [d for d in details if leaks(d)])
check("no entry quotes an unsealed diagnosis either (D2)",
      not [d for d in details if DX_CTRL in d])
mark = audit_max()
staff.post(f"/staff/consultation/edit/13",
           data=edit_form(13, chief_complaint=CC_STAFF, remarks=NOTE_STAFF))
staff_entry = q1("SELECT details FROM audit_logs WHERE id > ? AND action = "
                 "'Updated consultation record'", (mark,))
check("D5: Staff changing an unsealed chief complaint and notes is saved",
      q1("SELECT chief_complaint FROM consultations WHERE id=13")[0] == CC_STAFF)
check("D5: and its audit entry names the fields only: Changed: Chief complaint, Additional notes",
      staff_entry is not None and "Changed: Chief complaint, Additional notes" in staff_entry[0]
      and CC_STAFF not in staff_entry[0] and NOTE_STAFF not in staff_entry[0],
      staff_entry[0] if staff_entry else None)
details = [r[0] or "" for r in db().execute("SELECT details FROM audit_logs")]
check("D5: no audit entry holds any chief complaint, note, check-in reason or other "
      "free clinical text, sealed or not",
      not [d for d in details if any(t in d for t in FREE_TEXT)],
      [t for d in details for t in FREE_TEXT if t in d])
edit_entry = q1("SELECT details FROM audit_logs WHERE action = 'Updated consultation record' "
                "AND details LIKE ? ORDER BY id DESC", (f"Consultation #{CID_CTRL},%",))
check("an edit entry names the field instead: Changed: Diagnosis",
      edit_entry is not None and "Changed: Diagnosis" in edit_entry[0],
      edit_entry[0] if edit_entry else None)

from migrations import migrate_redact_clinical_audit_text as migration  # noqa: E402

copy = Path(_support.work_path("migrate.db"))
shutil.copyfile(TEST_DB, copy)
OLD_DX, OLD_TX, OLD_SEALED = "ZZOLDAUDIT9DX", "ZZOLDAUDIT9TX", "ZZOLDSEALED8DX"
conn = sqlite3.connect(copy)
conn.execute("INSERT INTO audit_logs (user_id, action, module, details, created_at) VALUES "
             "(1, 'Updated consultation record', 'Consultation Records', ?, '2026-08-01 01:00:00')",
             (f'Consultation #3 for X (PAT-003). 2 fields changed — Temperature: 36 °C → 37 °C; '
              f'Diagnosis: "{OLD_DX} a; b" → "{OLD_DX} c"; Treatment plan / medicines given: '
              f'not recorded → "{OLD_TX}"; Additional notes: "x" → "y"',))
conn.execute("UPDATE consultations SET diagnosis = ?, is_restricted = 1, restricted_by = 5, "
             "restricted_at = '2026-08-01 02:00:00' WHERE id = 12", (OLD_SEALED,))
conn.commit()
before_entries = conn.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0]
conn.close()
backups = Path(_support.work_path("migration-backups"))
with contextlib.redirect_stdout(io.StringIO()):
    first = migration.run(copy, backup_dir=backups)
conn = sqlite3.connect(copy)
texts = [r[0] or "" for r in conn.execute("SELECT details FROM audit_logs")]
records = conn.execute("SELECT user_id, details FROM audit_logs WHERE action = ?",
                       (migration.ACTION,)).fetchall()
sealed_12 = conn.execute("SELECT diagnosis FROM consultations WHERE id = 12").fetchone()[0]
kept = conn.execute("SELECT details FROM audit_logs WHERE details LIKE "
                    "'Consultation #3 for X (PAT-003)%'").fetchone()[0]
after_entries = conn.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0]
conn.close()
check("the migration made a backup first, in the backups folder",
      first["backup"] and Path(first["backup"]).exists()
      and Path(first["backup"]).parent == backups, first)
check("it removed the old clinical text from the entries",
      first["entries_redacted"] >= 1 and not [t for t in texts if OLD_DX in t or OLD_TX in t],
      first)
check("and left who, what changed and the other fields in place",
      "Temperature: 36 °C → 37 °C" in kept and "Diagnosis: (clinical text removed)" in kept
      and "Treatment plan / medicines given: (clinical text removed)" in kept
      and 'Additional notes: "x" → "y"' in kept, kept)
check("it encrypted the already-sealed record",
      first["records_encrypted"] == 1 and sc.is_encrypted(sealed_12)
      and sc.decrypt(sealed_12, 12, "diagnosis") == OLD_SEALED)
check("it wrote exactly one record of itself, with the reason and no clinical text",
      len(records) == 1 and after_entries == before_entries + 1
      and "RA 11166" in records[0][1] and "RA 10173" in records[0][1]
      and OLD_DX not in records[0][1], records)
check("after VACUUM the file holds none of the removed text",
      not [s for s in (OLD_DX, OLD_TX, OLD_SEALED) if s.encode() in copy.read_bytes()])
with contextlib.redirect_stdout(io.StringIO()):
    second = migration.run(copy, backup_dir=backups)
conn = sqlite3.connect(copy)
records_after = conn.execute("SELECT COUNT(*) FROM audit_logs WHERE action = ?",
                             (migration.ACTION,)).fetchone()[0]
conn.close()
check("running it twice changes nothing and makes no second backup or record",
      second["backup"] is None and records_after == 1, second)

# ===========================================================================
section("7. REPORTS COUNT SEALED VISITS UNDER CONFIDENTIAL")
with app.app_context():
    from models.consultations import get_consultation_stats
    stats = get_consultation_stats(viewer_role="Doctor")
sealed_total = q1("SELECT COUNT(*) FROM consultations WHERE is_restricted = 1")[0]
confidential = [d for d in stats["top_diagnoses"] if d["label"] == "Confidential"]
check("the all-time ranking has one Confidential row counting every sealed visit",
      len(confidential) == 1 and confidential[0]["count"] == sealed_total,
      (confidential, sealed_total))
check("and the total still includes them",
      stats["total"] == q1("SELECT COUNT(*) FROM consultations")[0])
confidential_cc = [d for d in stats["top_complaints"] if d["label"] == "Confidential"]
check("Task 5b: the complaints ranking has one Confidential row counting every sealed "
      "visit, and ranks no sealed complaint text",
      len(confidential_cc) == 1 and confidential_cc[0]["count"] == sealed_total
      and not [d for d in stats["top_complaints"] if leaks(d["label"]) or d["label"].startswith("enc:v1:")],
      stats["top_complaints"])
day = q1("SELECT substr(consultation_date, 1, 10) FROM consultations WHERE id=?", (CID_REC,))[0]
RANGE = f"?from={day}&to={day}"
# Task 5d: the Administrator and the Official read a ranking row of fewer than 5 as
# "Other (each fewer than 5)", the Confidential row included. The Doctor reads the
# exact row. So each viewer is checked against the label THEIR page must carry.
from utils.small_counts import OTHER_LABEL
sealed_that_day = q1("SELECT COUNT(*) FROM consultations WHERE is_restricted = 1 "
                     "AND substr(consultation_date, 1, 10) = ?", (day,))[0]
check("the sealed visits on that day are counted (the label checks below need them)",
      sealed_that_day >= 1, sealed_that_day)
def expected_label(who):
    if who == "Doctor" or sealed_that_day >= 5:
        return "Confidential"
    return OTHER_LABEL
for who, client, url in (("Administrator", admin, "/admin/reports_analytics"),
                         ("Doctor", doctor, "/staff/reports_analytics"),
                         ("Barangay Official", official, "/official/reports_analytics"),
                         ("Barangay Official", official, "/official/community-reports")):
    page = body(client.get(url + RANGE))
    check(f"{who}: {url} shows {expected_label(who)} and no sealed text",
          expected_label(who) in page and not leaks(page))
for who, client, url in (("Administrator", admin, "/admin/reports_analytics/export.xlsx"),
                         ("Doctor", doctor, "/staff/reports_analytics/export.pdf"),
                         ("Barangay Official", official, "/official/community-reports/export.xlsx")):
    text = readable(client.get(url + RANGE))
    check(f"{who}: {url} lists {expected_label(who)} and no sealed text",
          expected_label(who) in text and not leaks(text))

# ===========================================================================
section("8. A WRONG OR MISSING KEY GIVES A PLAIN MESSAGE, NOT A CRASH")
right_key = os.environ["SEALED_DATA_KEY"]
stored = q1("SELECT diagnosis, treatment_plan FROM consultations WHERE id=?", (CID_REC,))
for label, value in (("a wrong key", sc.new_key()), ("no key", None)):
    if value is None:
        os.environ.pop("SEALED_DATA_KEY")
    else:
        os.environ["SEALED_DATA_KEY"] = value
    mark = audit_max()
    response = doctor.get(f"/staff/patient_profile/{P_REC}")
    page = body(response)
    entries = db().execute("SELECT action, details FROM audit_logs WHERE id > ?",
                           (mark,)).fetchall()
    check(f"with {label}: the Doctor's profile still opens and says so plainly",
          response.status_code == 200 and UNOPENABLE in page and not leaks(page),
          response.status_code)
    check(f"with {label}: one audit entry, without clinical text",
          sum(1 for e in entries if e[0] == "Could not open sealed consultation") == 1
          and not any(leaks(e[1] or "") for e in entries), [tuple(e) for e in entries])
    page = body(doctor.get(f"/staff/consultation/edit/{CID_REC}"))
    check(f"with {label}: the edit form shows the line and no diagnosis box",
          UNOPENABLE in page and 'name="diagnosis"' not in page)
    form = {k: v for k, v in edit_form(CID_REC, temperature="37.1").items()
            if k not in ("diagnosis", "treatment_plan", "chief_complaint", "remarks")}
    form["is_restricted"] = "1"
    doctor.post(f"/staff/consultation/edit/{CID_REC}", data=form)
    doctor.post(f"/staff/patient_profile/{P_REC}/restrict-consultation",
                data={"consultation_id": str(CID_REC), "restrict_action": "lift"})
    after = q1("SELECT diagnosis, treatment_plan, is_restricted FROM consultations "
               "WHERE id=?", (CID_REC,))
    check(f"with {label}: saving or lifting cannot overwrite or unseal the record",
          (after[0], after[1]) == (stored[0], stored[1]) and after[2] == 1)
os.environ["SEALED_DATA_KEY"] = right_key
check("with the right key back, the Doctor reads the record again",
      DX_REC in body(doctor.get(f"/staff/patient_profile/{P_REC}")))

# ===========================================================================
section("9. ONE PLACE DECRYPTS, AND IT REFUSES EVERY ROLE BUT THE DOCTOR")
calls = []
for path in sorted(list((REPO / "models").glob("*.py")) + list((REPO / "routes").glob("*.py"))
                   + list((REPO / "utils").glob("*.py")) + [REPO / "app.py"]):
    tree = ast.parse(path.read_text(encoding="utf-8"))
    for func in ast.walk(tree):
        if not isinstance(func, ast.FunctionDef):
            continue
        for node in ast.walk(func):
            if isinstance(node, ast.Call):
                name = getattr(node.func, "attr", None) or getattr(node.func, "id", None)
                # The AES call inside sealed_crypto.decrypt() is the primitive
                # itself, not a caller of it.
                if name == "decrypt" and not (path.name == "sealed_crypto.py"
                                              and func.name == "decrypt"):
                    calls.append(f"{path.name}:{func.name}")
check("decrypt() is called from exactly one function: open_sealed_value()",
      sorted(set(calls)) == ["consultations.py:open_sealed_value"], calls)
from models.consultations import open_sealed_value, set_consultation_restriction  # noqa: E402

cipher = q1("SELECT diagnosis FROM consultations WHERE id=?", (CID_REC,))[0]
for role in ("admin", "health staff", "barangay official", None):
    try:
        open_sealed_value(cipher, CID_REC, "diagnosis", role)
        refused = False
    except PermissionError:
        refused = True
    check(f"open_sealed_value() refuses the role {role!r}", refused)
check("open_sealed_value() opens it for the Doctor",
      open_sealed_value(cipher, CID_REC, "diagnosis", "doctor") == (DX_REC, True))
visit_cipher = q1("SELECT notes FROM appointments WHERE id=?", (appt,))[0]
for role in ("admin", "health staff", "barangay official", None):
    try:
        open_sealed_value(visit_cipher, appt, "notes", role, "appointments")
        refused = False
    except PermissionError:
        refused = True
    check(f"Task 5b: the check-in reason is refused to the role {role!r} by the same function",
          refused)
check("Task 5b: the same function opens the check-in reason for the Doctor",
      open_sealed_value(visit_cipher, appt, "notes", "doctor", "appointments")
      == (REASON_REC, True))
with app.app_context():
    try:
        set_consultation_restriction(CID_REC, False, ADMIN_ID, "admin")
        raised = False
    except PermissionError:
        raised = True
check("the model refuses to unseal for a role that may not decrypt, and nothing changes",
      raised and q1("SELECT is_restricted FROM consultations WHERE id=?", (CID_REC,))[0] == 1)


# ===========================================================================
section("10. THE CHECK-IN REASON ON THE PATIENT QUEUE (Task 5b, D1, D4)")


def queue_header_and_rows(response):
    """The queue sheet's header row and body rows, from an Excel download."""
    book = load_workbook(io.BytesIO(response.data), data_only=True)
    for sheet in book.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        for index, row in enumerate(rows):
            if row and row[0] == "#":
                header = [str(c) for c in row if c is not None]
                body_rows = [[str(c) for c in r if c is not None] for r in rows[index + 1:]
                             if r and any(c is not None for c in r)]
                return header, body_rows
    return [], []


QUEUE_PAGES = {"Administrator": (admin, "/admin"), "Health Center Staff": (staff, "/staff"),
               "Doctor": (doctor, "/staff")}
queue_pages = {who: body(client.get(f"{prefix}/appointment_scheduling"))
               for who, (client, prefix) in QUEUE_PAGES.items()}
check("control: Health Center Staff see the reason of an ordinary waiting visit",
      REASON_QUEUE in queue_pages["Health Center Staff"])
check("control: the Doctor sees it too",
      REASON_QUEUE in queue_pages["Doctor"])
check("D4: the Administrator's Patient Queue carries no check-in reason, for any visit",
      not any(t in queue_pages["Administrator"]
              for t in (REASON_QUEUE, REASON_REC, REASON_LATE, REASON_CTRL))
      and 'class="reason-text"' not in queue_pages["Administrator"]
      and "patient ID, or reason" not in queue_pages["Administrator"],
      [t for t in (REASON_QUEUE, REASON_REC, REASON_LATE, REASON_CTRL)
       if t in queue_pages["Administrator"]]
      + ['class="reason-text"' in queue_pages["Administrator"],
         "patient ID, or reason" in queue_pages["Administrator"]])
check("D4: the Administrator still sees each patient and their status in the queue",
      "Ramos" in queue_pages["Administrator"] and "badge-status" in queue_pages["Administrator"])
check("D1/D2: Health Center Staff see the placeholder, not the reason, of a sealed visit",
      REASON_REC not in queue_pages["Health Center Staff"]
      and PLACEHOLDER in queue_pages["Health Center Staff"])
check("D1/D2: the Doctor reads a sealed visit's reason on screen, never on paper",
      REASON_REC in queue_pages["Doctor"]
      and REASON_REC not in visible(queue_pages["Doctor"], "Doctor")
      and f'<span class="print-only">{PLACEHOLDER}</span>' in queue_pages["Doctor"])
check("a sealed reason is not offered to the queue's search box",
      f'data-search="' in queue_pages["Doctor"]
      and not re.search(r'data-search="[^"]*' + REASON_REC, queue_pages["Doctor"]))
for who in ("Administrator", "Health Center Staff", "Doctor"):
    client, prefix = QUEUE_PAGES[who]
    for kind in ("xlsx", "pdf"):
        response = client.get(f"{prefix}/appointment_scheduling/export.{kind}")
        text = readable(response)
        if who == "Administrator":
            check(f"D4: the Administrator's queue export ({kind}) has no check-in reason",
                  response.status_code == 200
                  and not [t for t in (REASON_QUEUE, REASON_REC, REASON_LATE, REASON_CTRL)
                           if t in text])
        else:
            check(f"{who}: the queue export ({kind}) has the ordinary reason, never a sealed one",
                  response.status_code == 200 and REASON_QUEUE in text and not leaks(text),
                  leaks(text))
header_admin, _ = queue_header_and_rows(admin.get("/admin/appointment_scheduling/export.xlsx"))
header_staff, rows_staff = queue_header_and_rows(staff.get("/staff/appointment_scheduling/export.xlsx"))
check("D4: the Administrator's workbook has no Notes column at all",
      header_admin and "Notes" not in header_admin and "Patient" in header_admin, header_admin)
check("control: Staff's workbook keeps the Notes column, with the placeholder for a sealed visit",
      "Notes" in header_staff and any(PLACEHOLDER in r for r in rows_staff), header_staff)
response = admin.get("/staff/appointment_scheduling")
check("D4 is enforced on the server: the Administrator cannot reach the care team's "
      "queue URL to get the reason",
      response.status_code in (302, 401, 403) or REASON_QUEUE not in body(response))
with app.app_context():
    from models.appointments import attach_visit_reasons, get_todays_appointments
    rows = get_todays_appointments()
    check("a queue row carries no reason unless it went through attach_visit_reasons()",
          rows and all("notes" not in r for r in rows))
    shaped = {role: attach_visit_reasons(get_todays_appointments(), role, queue_view=True)
              for role in ("admin", "barangay official", "health staff", "doctor")}
    check("the model hides the reason from the Administrator and the Official, for every visit",
          all(r["notes"] is None and r["reason_hidden"]
              for role in ("admin", "barangay official") for r in shaped[role]))
    check("and gives it to the care team, the sealed one as placeholder or decrypted",
          any(r["notes"] == REASON_QUEUE for r in shaped["health staff"])
          and any(r["notes"] == PLACEHOLDER for r in shaped["health staff"])
          and any(r["notes"] == REASON_REC for r in shaped["doctor"])
          and not any(r["notes"] == REASON_REC for r in shaped["health staff"]))
# The database backup is a copy of the whole file, so it holds every unsealed value as
# it is stored. That is a known limit recorded in PROJECT_STATUS.md (M25, M26), not a
# screen or an export of this page, so it is left out of this crawl.
FILES_OF_THE_WHOLE_DATABASE = ("/admin/settings/backup.idsbackup",)
for who, client in (("Administrator", admin), ("Barangay Official", official)):
    found = []
    for url in plain_gets:
        if url in FILES_OF_THE_WHOLE_DATABASE:
            continue
        if REASON_QUEUE in readable(client.get(url)):
            found.append(url)
    check(f"D4: {who}: no screen, export or JSON anywhere carries the waiting visit's reason",
          not found, found)

# ===========================================================================
section("11. MEDICINE NAMES AT A SEALED VISIT (Task 5b, D3)")
admin_profile = body(admin.get(f"/admin/patient_profile/{P_REC}"))
admin_print = body(admin.get(f"/admin/patient_profile/{P_REC}?print=1"))
staff_profile = body(staff.get(f"/staff/patient_profile/{P_REC}"))
doctor_profile = body(doctor.get(f"/staff/patient_profile/{P_REC}"))
# Task 5c: his profile has no Dispensing History at all, so there is no
# placeholder to show either -- one plain line instead.
check("D3: the Administrator's profile of a sealed visit names no medicine (demographics only)",
      MED_SEALED not in admin_profile and MED_SEALED not in admin_print
      and "Clinical details are shown only to the Doctor and Health Center Staff." in admin_profile)
check("D3: Health Center Staff and the Doctor still see what was dispensed",
      MED_SEALED in staff_profile and MED_SEALED in doctor_profile)
check("control: Health Center Staff see the medicine dispensed at an unsealed visit",
      MED_OPEN in body(staff.get(f"/staff/patient_profile/{P_CTRL}")))
check("and since Task 5c the Administrator does not, sealed or not",
      MED_OPEN not in body(admin.get(f"/admin/patient_profile/{P_CTRL}")))
for who, client, prefix in (("Administrator", admin, "/admin"), ("Health Center Staff", staff, "/staff"),
                            ("Doctor", doctor, "/staff")):
    response = client.get(f"{prefix}/consultation/edit/{CID_REC}")
    page = body(response)
    if who == "Administrator":
        check("D3: the Administrator has no edit form to list a medicine on (404, Task 5c)",
              response.status_code == 404 and MED_SEALED not in page)
    else:
        check(f"{who}: the edit form of a sealed visit still lists what was dispensed",
              MED_SEALED in page)
found = []
for url in plain_gets:
    if url in FILES_OF_THE_WHOLE_DATABASE:
        continue
    if MED_SEALED in readable(admin.get(url)):
        found.append(url)
# Task 5c (M25): Reports used to rank a sealed visit's medicine by name. It is
# now counted under "Confidential", so nothing he can reach names it.
check("D3: no screen, export or dashboard of the Administrator names it, Reports included",
      not found, found)
check("Reports count it under Confidential instead",
      "Confidential" in readable(admin.get("/admin/reports_analytics/export.xlsx")))
dash = [u for u in plain_gets if "dashboard" in u]
for url in dash:
    if url.startswith("/admin"):
        page = body(admin.get(url))
        check(f"D3: the Administrator's dashboard recent-consultations panel hides the "
              f"medicine of a sealed visit ({url})",
              MED_SEALED not in page and not leaks(page))
with app.app_context():
    from models.dispensing import get_dispensing_stats
    totals = get_dispensing_stats(viewer_role="Doctor")
units = q1("SELECT COALESCE(SUM(quantity_dispensed),0) FROM medicine_dispensing")[0]
check("stock and dispensing totals are unchanged by hiding the names",
      totals["total_units"] == units
      and totals["total_events"] == q1("SELECT COUNT(*) FROM medicine_dispensing")[0])
inventory = body(admin.get("/admin/medicine_inventory"))
check("the Administrator's Medicine Inventory still lists the shelf and its stock",
      "Amoxicillin" in inventory and "Paracetamol" in inventory)

# Dispensing at a sealed visit through the real form: audit entries and flash text.
losartan = q1("SELECT id FROM inventory_items WHERE name LIKE 'Losartan%'")[0]
metformin = q1("SELECT id FROM inventory_items WHERE name LIKE 'Metformin%'")[0]
NO_SEALED_FORM = ("diagnosis", "treatment_plan", "is_restricted", "chief_complaint", "remarks")
mark = audit_max()
form = {k: v for k, v in edit_form(CID_LATE).items() if k not in NO_SEALED_FORM}
form.update({"dispense_item_0": str(losartan), "dispense_qty_0": "2"})
response = staff.post(f"/staff/consultation/edit/{CID_LATE}", data=form, follow_redirects=True)
entries = [e[0] or "" for e in db().execute("SELECT details FROM audit_logs WHERE id > ?", (mark,))]
check("Health Center Staff dispensing at a sealed visit works and still sees the name",
      q1("SELECT COUNT(*) FROM medicine_dispensing WHERE consultation_id=? AND "
         "inventory_item_id=?", (CID_LATE, losartan))[0] == 1 and "Losartan" in body(response))
check("D3: but no audit entry for it names the medicine",
      entries and not [e for e in entries if "Losartan" in e], entries)
mark = audit_max()
form = {k: v for k, v in edit_form(CID_LATE).items() if k not in NO_SEALED_FORM}
form.update({"dispense_item_0": str(metformin), "dispense_qty_0": "3"})
# Task 5c: the Administrator no longer dispenses at all. Both URLs refuse him,
# nothing leaves the shelf, and no entry is written.
response = admin.post(f"/admin/consultation/edit/{CID_LATE}", data=form)
response2 = admin.post(f"/staff/consultation/edit/{CID_LATE}", data=form)
entries = [e[0] or "" for e in db().execute("SELECT details FROM audit_logs WHERE id > ?", (mark,))]
check("the Administrator cannot dispense at a sealed visit: refused on both URLs, "
      "and no medicine name comes back",
      response.status_code == 404 and response2.status_code == 302
      and q1("SELECT COUNT(*) FROM medicine_dispensing WHERE consultation_id=? AND "
             "inventory_item_id=?", (CID_LATE, metformin))[0] == 0
      and "Metformin" not in body(response) + body(response2),
      (response.status_code, response2.status_code))
check("and nothing was written to the audit trail for it", entries == [], entries)

# ===========================================================================
section("12. THE EDIT CONSULTATION PAGE PRINTS SEALED FIELDS AS 'Sealed — not printed' (D6)")
page = body(doctor.get(f"/staff/consultation/edit/{CID_REC}"))
areas = re.findall(r'<textarea[^>]*name="(chief_complaint|diagnosis|treatment_plan|remarks)"[^>]*>',
                   page)
tags = {m.group(1): m.group(0) for m in re.finditer(
    r'<textarea[^>]*name="(chief_complaint|diagnosis|treatment_plan|remarks)"[^>]*>', page)}
check("the Doctor's edit page of a sealed visit still shows all four fields to edit",
      sorted(tags) == ["chief_complaint", "diagnosis", "remarks", "treatment_plan"]
      and CC_REC in page and NOTE_REC in page and DX_REC in page and TX_REC in page)
check("each of the four fields carries the sealed-field class",
      len(tags) == 4 and all("sealed-field" in t for t in tags.values()), tags)
check("the page's print styles hide sealed fields and show the plain line",
      "@media print" in page and ".sealed-field{display:none !important;}" in page
      and ".sealed-print-note{display:block;" in page)
check("each hidden field has the line 'Sealed — not printed' beside it",
      page.count('<div class="sealed-print-note">Sealed — not printed</div>') == 4,
      page.count("sealed-print-note"))
paper = re.sub(r'<textarea[^>]*sealed-field[^>]*>.*?</textarea>', "", page, flags=re.S)
check("with those fields removed, as the print styles do, no sealed text is left on the page",
      not leaks(paper), leaks(paper))
control = body(doctor.get(f"/staff/consultation/edit/{CID_CTRL}"))
check("control: an unsealed record's edit page hides nothing when printed",
      "sealed-field" not in re.sub(r"<style.*?</style>", "", control, flags=re.S)
      and "Sealed — not printed" not in control and CC_CTRL in control)
page = body(staff.get(f"/staff/consultation/edit/{CID_REC}"))
check("Health Center Staff: the sealed visit's edit page has no chief complaint or notes box, "
      "only the placeholder",
      'name="chief_complaint"' not in page and 'name="remarks"' not in page
      and PLACEHOLDER in page and not leaks(page), leaks(page))
response = admin.get(f"/admin/consultation/edit/{CID_REC}")
check("Administrator: there is no edit page to print at all (404, Task 5c)",
      response.status_code == 404 and not leaks(body(response))
      and 'name="chief_complaint"' not in body(response))

# ===========================================================================
section("13. THE SECOND MIGRATION: AUDIT ENTRIES AND ALREADY-SEALED VISITS (Task 5b)")
from migrations import migrate_seal_free_text_fields as migration2  # noqa: E402

copy2 = Path(_support.work_path("migrate2.db"))
shutil.copyfile(TEST_DB, copy2)
OLD_CC, OLD_NOTE, OLD_REASON = "ZZOLDAUDIT7CC", "ZZOLDAUDIT7NOTE", "ZZOLDSEALED6REASON"
OLD_SEALED_CC, OLD_SEALED_NOTE = "ZZOLDSEALED6CC", "ZZOLDSEALED6NOTE"
conn = sqlite3.connect(copy2)
conn.execute("INSERT INTO audit_logs (user_id, action, module, details, created_at) VALUES "
             "(1, 'Updated consultation record', 'Consultation Records', ?, '2026-08-01 01:00:00')",
             ('Consultation #3 for X (PAT-003). 3 fields changed — Temperature: 36 °C → 37 °C; '
              f'Chief complaint: "{OLD_CC} a; b" → "{OLD_CC} c"; Additional notes: not '
              f'recorded → "{OLD_NOTE}"',))
oc = conn.execute("SELECT appointment_id FROM consultations WHERE id = 14").fetchone()[0]
conn.execute("UPDATE consultations SET chief_complaint = ?, remarks = ?, is_restricted = 1, "
             "restricted_by = 5, restricted_at = '2026-08-01 02:00:00' WHERE id = 14",
             (OLD_SEALED_CC, OLD_SEALED_NOTE))
conn.execute("UPDATE appointments SET notes = ? WHERE id = ?", (OLD_REASON, oc))
conn.commit()
before_entries = conn.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0]
before_other = conn.execute("SELECT chief_complaint, remarks FROM consultations WHERE id = 20"
                            ).fetchone()
conn.close()
backups2 = Path(_support.work_path("migration2-backups"))
with contextlib.redirect_stdout(io.StringIO()):
    first2 = migration2.run(copy2, backup_dir=backups2)
conn = sqlite3.connect(copy2)
texts = [r[0] or "" for r in conn.execute("SELECT details FROM audit_logs")]
records2 = conn.execute("SELECT user_id, details FROM audit_logs WHERE action = ?",
                        (migration2.ACTION,)).fetchall()
row14 = conn.execute("SELECT chief_complaint, remarks, is_restricted FROM consultations "
                     "WHERE id = 14").fetchone()
reason14 = conn.execute("SELECT notes FROM appointments WHERE id = ?", (oc,)).fetchone()[0]
kept = conn.execute("SELECT details FROM audit_logs WHERE details LIKE "
                    "'Consultation #3 for X (PAT-003)%'").fetchone()[0]
after_entries = conn.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0]
after_other = conn.execute("SELECT chief_complaint, remarks FROM consultations WHERE id = 20"
                           ).fetchone()
conn.close()
check("the second migration made a backup first, in the backups folder",
      first2["backup"] and Path(first2["backup"]).exists()
      and Path(first2["backup"]).parent == backups2, first2)
check("it removed the free text from the old audit entries",
      first2["entries_redacted"] >= 1
      and not [t for t in texts if OLD_CC in t or OLD_NOTE in t], first2)
check("and kept who, when, the action and the other fields",
      "Temperature: 36 °C → 37 °C" in kept and "Chief complaint: (clinical text removed)" in kept
      and "Additional notes: (clinical text removed)" in kept, kept)
check("it encrypted the already-sealed visit's chief complaint, notes and check-in reason",
      first2["consultations_encrypted"] == 1 and first2["visits_encrypted"] == 1
      and sc.decrypt(row14[0], 14, "chief_complaint") == OLD_SEALED_CC
      and sc.decrypt(row14[1], 14, "remarks") == OLD_SEALED_NOTE
      and sc.decrypt(reason14, oc, "notes", "appointments") == OLD_REASON
      and row14[2] == 1, first2)
check("and left an unsealed visit's text alone", before_other == after_other)
check("it wrote exactly one record of itself, with the reason and no free text",
      len(records2) == 1 and after_entries == before_entries + 1
      and "RA 11166 Section 44" in records2[0][1]
      and not [t for t in (OLD_CC, OLD_NOTE, OLD_REASON, OLD_SEALED_CC) if t in records2[0][1]],
      records2)
check("after VACUUM the file holds none of the removed or encrypted text",
      not [t for t in (OLD_CC, OLD_NOTE, OLD_REASON, OLD_SEALED_CC, OLD_SEALED_NOTE)
           if t.encode() in copy2.read_bytes()])
with contextlib.redirect_stdout(io.StringIO()):
    second2 = migration2.run(copy2, backup_dir=backups2)
conn = sqlite3.connect(copy2)
records_after2 = conn.execute("SELECT COUNT(*) FROM audit_logs WHERE action = ?",
                              (migration2.ACTION,)).fetchone()[0]
conn.close()
check("running it twice changes nothing and makes no second backup or record",
      second2["backup"] is None and records_after2 == 1, second2)
check("no audit entry in the migrated copy holds any free clinical text",
      not [t for t in texts if any(x in t for x in (OLD_CC, OLD_NOTE, OLD_REASON))])

# ===========================================================================
passed = sum(1 for _, ok in results if ok)
print("\n" + "=" * 78)
print(f"RESULT: {passed}/{len(results)} checks passed")
failed = [name for name, ok in results if not ok]
for name in failed:
    print("  FAILED:", name)
print("=" * 78)
sys.exit(0 if not failed else 1)
