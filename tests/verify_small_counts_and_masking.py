"""Small counts fold into "Other", and phone numbers are masked (Task 5d).

D1. For the Administrator and the Barangay Official, any row of a ranking or
breakdown by diagnosis, complaint, condition or dispensed medicine with a count
of 1 to 4 folds into one row, "Other (each fewer than 5)". The "Confidential" row
folds the same way when its count is under 5. Totals stay exact. The Doctor sees
every count. The threshold is SMALL_COUNT_THRESHOLD in config.py.
"""
import io
import json
import re
import sqlite3
import sys
import zlib
from datetime import date, timedelta

import _support  # noqa: E402  (first: builds the test database, points the app at it)

import config  # noqa: E402
from app import create_app  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

sys.stdout.reconfigure(errors="replace")

app = create_app()
_support.attach(app)
TEST_DB = _support.DB_PATH
ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID = 1, 2, 3, 5

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok)))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}")
    if detail and not ok:
        for line in str(detail).splitlines()[:12]:
            print(f"       {line}")
    return ok


def section(title):
    print(f"\n--- {title} ---")


def db():
    conn = sqlite3.connect(TEST_DB)
    conn.row_factory = sqlite3.Row
    return conn


def q1(sql, args=()):
    conn = db()
    try:
        return conn.execute(sql, args).fetchone()
    finally:
        conn.close()


def run_sql(sql, args=()):
    conn = db()
    try:
        cur = conn.execute(sql, args)
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def client_as(user_id):
    c = app.test_client()
    with c.session_transaction() as sess:
        sess["user_id"] = user_id
        _support.open_session(sess)
    return c


def body(response):
    return response.get_data(as_text=True)


def pdf_text(data):
    """Text of every PDF stream. A PDF escapes parentheses, so they are put back."""
    out = []
    for match in re.finditer(rb"stream\r?\n(.*?)endstream", data, re.S):
        try:
            out.append(zlib.decompress(match.group(1)).decode("latin-1"))
        except Exception:
            out.append(match.group(1).decode("latin-1", "replace"))
    return "\n".join(out).replace("\\(", "(").replace("\\)", ")")


def workbook(data):
    return load_workbook(io.BytesIO(data), data_only=True)


def sheet_rows(data, title):
    book = workbook(data)
    for sheet in book.worksheets:
        if sheet.title == title:
            return [list(row) for row in sheet.iter_rows(values_only=True)]
    return []


def xlsx_text(data):
    return "\n".join(str(cell) for sheet in workbook(data).worksheets
                     for row in sheet.iter_rows(values_only=True) for cell in row
                     if cell is not None)


def readable(response):
    kind = response.headers.get("Content-Type", "")
    if "spreadsheetml" in kind:
        return xlsx_text(response.data)
    if "pdf" in kind:
        return pdf_text(response.data) + response.data.decode("latin-1")
    return response.data.decode("utf-8", "replace")


admin, staff, official, doctor = (client_as(i) for i in
                                  (ADMIN_ID, STAFF_ID, OFFICIAL_ID, DOCTOR_ID))

from utils.small_counts import (  # noqa: E402
    OTHER_LABEL, SMALL_COUNT_TEXT, fold_small_counts, shown_count,
)

print("=" * 78)
print("IDS-BPHCSN — small counts fold into Other; phone numbers are masked (Task 5d)")
print("=" * 78)

# ===========================================================================
section("1. THE ONE SHARED HELPER")
check("the threshold is one named constant in config.py, and it is 5",
      config.SMALL_COUNT_THRESHOLD == 5)
check("the Other label carries the threshold: 'Other (each fewer than 5)'",
      OTHER_LABEL == "Other (each fewer than 5)", OTHER_LABEL)
rows = [{"label": "twelve", "count": 12}, {"label": "five", "count": 5},
        {"label": "four", "count": 4}, {"label": "one", "count": 1},
        {"label": "none", "count": 0}]
for role in ("Administrator", "admin", "Barangay Official", "official", "", None, "unknown"):
    folded = fold_small_counts(rows, role, label_key="label", sum_keys=("count",))
    check(f"viewer {role!r}: 12 and 5 stay named, 4 and 1 fold to one Other row of 5, "
          "0 stays named",
          [(r["label"], r["count"]) for r in folded]
          == [("twelve", 12), ("five", 5), ("none", 0), (OTHER_LABEL, 5)], folded)
for role in ("Doctor", "doctor", "Health Center Staff", "health staff"):
    folded = fold_small_counts(rows, role, label_key="label", sum_keys=("count",))
    check(f"viewer {role!r} (care team): every row stays, untouched", folded == rows)
check("4 folds and 5 does not: the boundary is exactly the threshold",
      [r["label"] for r in fold_small_counts(
          [{"label": "a", "count": 4}, {"label": "b", "count": 5}], "admin",
          label_key="label", sum_keys=("count",))] == ["b", OTHER_LABEL])
only_small = fold_small_counts([{"label": "a", "count": 1}, {"label": "b", "count": 2}],
                               "admin", label_key="label", sum_keys=("count",))
check("when folding leaves no named row, only the Other row remains, with the total",
      [(r["label"], r["count"]) for r in only_small] == [(OTHER_LABEL, 3)], only_small)
check("a restricted-visit count under 5 reads 'Fewer than 5' for the Administrator, "
      "its number for the Doctor",
      shown_count(3, "admin") == SMALL_COUNT_TEXT == "Fewer than 5"
      and shown_count(3, "doctor") == 3 and shown_count(5, "admin") == 5
      and shown_count(0, "admin") == 0)

# ===========================================================================
section("2. SEED COUNTS OF 1, 4, 5 AND 12")


def free_days(count):
    """Days inside the default six-month window with no consultation or dispensing."""
    days, day = [], date.today() - timedelta(days=20)
    while len(days) < count:
        text = day.isoformat()
        used = (q1("SELECT COUNT(*) FROM consultations WHERE substr(consultation_date,1,10)=?",
                   (text,))[0]
                + q1("SELECT COUNT(*) FROM medicine_dispensing "
                     "WHERE substr(date_dispensed,1,10)=?", (text,))[0])
        if not used:
            days.append(text)
        day -= timedelta(days=1)
    return days


DAY1, DAY2 = free_days(2)
GROUPS = (("ONE", 1), ("FOUR", 4), ("FIVE", 5), ("TWELVE", 12))
patient_ids = [r[0] for r in db().execute("SELECT id FROM patients ORDER BY id")]
_n = [0]


def add_visit(day, sealed, diagnosis=None, complaint=None, medicine=None, item=1):
    _n[0] += 1
    cid = run_sql(
        "INSERT INTO consultations (patient_id, recorded_by, consultation_date, "
        "chief_complaint, diagnosis, is_restricted, restricted_by) VALUES (?,?,?,?,?,?,?)",
        (patient_ids[_n[0] % len(patient_ids)], DOCTOR_ID, day, complaint, diagnosis,
         1 if sealed else 0, DOCTOR_ID if sealed else None))
    if medicine:
        run_sql("INSERT INTO medicine_dispensing (patient_id, consultation_id, "
                "inventory_item_id, medicine_name, unit, quantity_dispensed, date_dispensed, "
                "dispensed_by) VALUES ((SELECT patient_id FROM consultations WHERE id=?),"
                "?,?,?,'tablet',2,?,?)", (cid, cid, item, medicine, day, STAFF_ID))
    return cid


# Day 1: groups of 1, 4, 5 and 12, and THREE sealed visits (a Confidential row under 5).
for name, count in GROUPS:
    for _ in range(count):
        add_visit(DAY1, False, f"ZZDX{name}", f"ZZCC{name}", f"ZZMED{name}")
for _ in range(3):
    add_visit(DAY1, True, medicine="ZZMEDSEALED", item=2)
# Day 2: FIVE sealed visits (Confidential stays named) beside a group of 5.
for _ in range(5):
    add_visit(DAY2, False, "ZZDXPAIR", "ZZCCPAIR", "ZZMEDPAIR")
for _ in range(5):
    add_visit(DAY2, True, medicine="ZZMEDSEALEDFIVE", item=2)
RANGE1, RANGE2 = f"?from={DAY1}&to={DAY1}", f"?from={DAY2}&to={DAY2}"
check("seeded: 22 visible visits and 3 sealed on day 1; 5 and 5 on day 2",
      q1("SELECT COUNT(*) FROM consultations WHERE substr(consultation_date,1,10)=?",
         (DAY1,))[0] == 25
      and q1("SELECT COUNT(*) FROM consultations WHERE substr(consultation_date,1,10)=?",
             (DAY2,))[0] == 10)

# The register: 12 with Hypertension, 5 with Diabetes, 4 with Asthma, 1 with Tuberculosis.
run_sql("UPDATE patients SET medical_conditions = '[]'")
for index, pid in enumerate(patient_ids):
    held = [name for name, limit in (("Hypertension", 12), ("Diabetes", 5),
                                     ("Asthma", 4), ("Tuberculosis", 1)) if index < limit]
    run_sql("UPDATE patients SET medical_conditions = ? WHERE id = ?",
            (json.dumps(held), pid))
check("seeded: 12 / 5 / 4 / 1 patients hold the four conditions",
      [q1("SELECT COUNT(*) FROM patients WHERE medical_conditions LIKE ?", (f'%"{c}"%',))[0]
       for c in ("Hypertension", "Diabetes", "Asthma", "Tuberculosis")] == [12, 5, 4, 1])

from models.consultations import get_consultation_stats  # noqa: E402
from models.dispensing import get_dispensing_stats  # noqa: E402
from models.patients import get_patient_report  # noqa: E402
from utils.dates import resolve_range  # noqa: E402

ROLES = {"Administrator": "admin", "Barangay Official": "barangay official",
         "Doctor": "doctor"}


def stats(day, role):
    d = date.fromisoformat(day)
    with app.app_context():
        return (get_consultation_stats(months=6, start=d, end=d, viewer_role=role),
                get_dispensing_stats(months=6, start=d, end=d, viewer_role=role))


def pairs(entries, label_key, count_key):
    return [(e[label_key], e[count_key]) for e in entries]


# ===========================================================================
section("3. THE MODEL ANSWERS: DIAGNOSES, COMPLAINTS, MEDICINES")
for who in ("Administrator", "Barangay Official"):
    consult, dispense = stats(DAY1, ROLES[who])
    window = consult["window"]
    check(f"{who}: diagnoses show 12 and 5 named; 1, 4 and the 3 sealed fold to Other = 8",
          pairs(window["top_diagnoses"], "label", "count")
          == [("ZZDXTWELVE", 12), ("ZZDXFIVE", 5), (OTHER_LABEL, 8)],
          window["top_diagnoses"])
    check(f"{who}: complaints fold the same way",
          pairs(window["top_complaints"], "label", "count")
          == [("ZZCCTWELVE", 12), ("ZZCCFIVE", 5), (OTHER_LABEL, 8)],
          window["top_complaints"])
    items = pairs(window["top_diagnoses"], "label", "count") + pairs(
        window["top_complaints"], "label", "count")
    check(f"{who}: no label of a group of 1 or 4 is in the model's answer",
          not any(n in label for label, _ in items for n in ("ONE", "FOUR")))
    check(f"{who}: the Confidential row (3 sealed) is folded, not shown",
          not any(e.get("confidential") for e in window["top_diagnoses"]))
    check(f"{who}: the restricted-visit count reads 'Fewer than 5'",
          window["diagnoses_withheld_shown"] == "Fewer than 5"
          and window["diagnoses_withheld"] == 3)
    named = pairs(dispense["window"]["all_medicines"], "medicine_name", "events")
    check(f"{who}: medicines fold on issues: 12 and 5 named, 1 + 4 + 3 sealed = 8 in Other",
          named == [("ZZMEDTWELVE", 12), ("ZZMEDFIVE", 5), (OTHER_LABEL, 8)], named)
    check(f"{who}: the Other medicine row adds up units too (8 issues x 2 units)",
          dispense["window"]["all_medicines"][-1]["units"] == 16)
    check(f"{who}: the chart list and the table list are the same rows",
          [e["medicine_name"] for e in dispense["window"]["top_medicines"]]
          == [e["medicine_name"] for e in dispense["window"]["all_medicines"]])

consult_d, dispense_d = stats(DAY1, "doctor")
check("Doctor: every diagnosis is named with its own count, Confidential included",
      pairs(consult_d["window"]["top_diagnoses"], "label", "count")
      == [("ZZDXTWELVE", 12), ("ZZDXFIVE", 5), ("ZZDXFOUR", 4), ("Confidential", 3),
          ("ZZDXONE", 1)], consult_d["window"]["top_diagnoses"])
check("Doctor: every medicine is named with its issues, Confidential included",
      sorted(pairs(dispense_d["window"]["all_medicines"], "medicine_name", "events"))
      == sorted([("ZZMEDTWELVE", 12), ("ZZMEDFIVE", 5), ("ZZMEDFOUR", 4),
                 ("ZZMEDONE", 1), ("Confidential", 3)]))
check("Doctor: the restricted-visit count is the exact number",
      consult_d["window"]["diagnoses_withheld_shown"] == 3)

# Totals are exact for every role.
totals = {}
for who, role in ROLES.items():
    consult, dispense = stats(DAY1, role)
    totals[who] = (consult["window"]["total"], consult["window"]["patients_seen"],
                   consult["window"]["diagnoses_withheld"], consult["total"],
                   dispense["window"]["units"], dispense["window"]["events"],
                   dispense["window"]["medicines"], dispense["total_units"],
                   sum(e["count"] for e in consult["window"]["top_diagnoses"]),
                   sum(e["events"] for e in dispense["window"]["all_medicines"]))
check("totals are identical for the Administrator, the Official and the Doctor",
      len(set(totals.values())) == 1, totals)
check("and the ranking rows still add up to every consultation and every issue",
      totals["Administrator"][8] == 25 and totals["Administrator"][9] == 25
      and totals["Administrator"][0] == 25, totals["Administrator"])

# Day 2: the Confidential count is 5, so it stays named.
for who in ("Administrator", "Barangay Official"):
    consult, dispense = stats(DAY2, ROLES[who])
    check(f"{who}: a Confidential count of 5 is shown by name",
          ("Confidential", 5) in pairs(consult["window"]["top_diagnoses"], "label", "count")
          and ("Confidential", 5) in pairs(dispense["window"]["all_medicines"],
                                           "medicine_name", "events"),
          (consult["window"]["top_diagnoses"], dispense["window"]["all_medicines"]))
    check(f"{who}: with nothing small left, no Other row is added",
          not any(e.get("other") for e in consult["window"]["top_diagnoses"])
          and not any(e.get("other") for e in dispense["window"]["all_medicines"]))
    check(f"{who}: the restricted-visit count of 5 is shown as 5",
          consult["window"]["diagnoses_withheld_shown"] == 5)

# Conditions, whole register.
with app.app_context():
    report_a = get_patient_report(viewer_role="admin")
    report_o = get_patient_report(viewer_role="barangay official")
    report_d = get_patient_report(viewer_role="doctor")
for who, report in (("Administrator", report_a), ("Barangay Official", report_o)):
    check(f"{who}: conditions show 12 and 5 named, and 4 + 1 folded into Other = 5",
          pairs(report["conditions"]["counts"], "condition", "count")
          == [("Hypertension", 12), ("Diabetes", 5), (OTHER_LABEL, 5)],
          report["conditions"]["counts"])
    check(f"{who}: condition totals are not changed",
          report["conditions"]["cohort_size"] == report_d["conditions"]["cohort_size"]
          and report["conditions"]["with_any"] == report_d["conditions"]["with_any"]
          == 12)
check("Doctor: all four conditions with their own counts",
      pairs(report_d["conditions"]["counts"], "condition", "count")
      == [("Hypertension", 12), ("Diabetes", 5), ("Asthma", 4), ("Tuberculosis", 1)])

# ===========================================================================
section("4. SCREEN, CHART DATA, PRINT TABLE, EXCEL AND PDF AGREE, PER VIEWER")
SURFACES = {
    "Administrator": {
        "page": "/admin/reports_analytics", "xlsx": "/admin/reports_analytics/export.xlsx",
        "pdf": "/admin/reports_analytics/export.pdf", "client": admin},
    "Barangay Official": {
        "page": "/official/reports_analytics",
        "xlsx": "/official/reports_analytics/export.xlsx",
        "pdf": "/official/reports_analytics/export.pdf", "client": official},
    "Doctor": {
        "page": "/staff/reports_analytics", "xlsx": "/staff/reports_analytics/export.xlsx",
        "pdf": "/staff/reports_analytics/export.pdf", "client": doctor},
}
RANK_ROW = re.compile(r'rank-label">(?:<span[^>]*>lock</span> )?([^<]+)</span>\s*'
                      r'<span class="rank-count">(\d+)</span>')


def print_table(page, caption):
    """{label: number} from the print-only table that sits under one chart."""
    match = re.search(r"<caption>" + re.escape(caption) + r"</caption>(.*?)</table>",
                      page, re.S)
    if not match:
        return None
    return [(m.group(1), int(m.group(2))) for m in re.finditer(
        r'<th scope="row">([^<]+)</th>\s*<td class="num">(\d+)</td>', match.group(1))]


def js_series(page, name, value_key):
    match = re.search(name + r":\s*\{\s*labels:\s*(\[.*?\]),\s*" + value_key + r":\s*(\[.*?\])",
                      page, re.S)
    return list(zip(json.loads(match.group(1)), json.loads(match.group(2)))) if match else None


for who, urls in SURFACES.items():
    client, role = urls["client"], ROLES[who]
    consult, dispense = stats(DAY1, role)
    expect_dx = pairs(consult["window"]["top_diagnoses"], "label", "count")
    expect_cc = pairs(consult["window"]["top_complaints"], "label", "count")
    expect_med = pairs(dispense["window"]["all_medicines"], "medicine_name", "units")
    page = body(client.get(urls["page"] + RANGE1))
    named = [(label.strip(), int(n)) for label, n in RANK_ROW.findall(page)
             if label.strip().startswith(("ZZDX", "Other (each", "Confidential"))]
    check(f"{who}: the screen's diagnosis ranking is the model's, row for row", named == expect_dx,
          (named, expect_dx))
    check(f"{who}: the chart data for the most-dispensed items is the model's units",
          js_series(page, "topMedicines", "units") == expect_med,
          (js_series(page, "topMedicines", "units"), expect_med))
    check(f"{who}: the print table under that chart matches the chart",
          print_table(page, "Most dispensed items in this range") == expect_med,
          print_table(page, "Most dispensed items in this range"))
    check(f"{who}: the screen says 'Across the 25 consultations' (the exact total)",
          "Across the 25 consultations" in re.sub(r"\s+", " ", page))
    if role == "doctor":
        check("Doctor: the screen names the groups of 1 and 4",
              "ZZDXONE" in page and "ZZDXFOUR" in page and "ZZMEDONE" in page)
        check("Doctor: no 'Other (each fewer than 5)' row in a ranking and no folded-rows note",
              OTHER_LABEL not in "".join(l for l, _ in named)
              and "are grouped under Other" not in page.split("Most Common Diagnoses")[1][:3000])
    else:
        check(f"{who}: the screen carries no label of a group of 1 or 4, in any panel",
              not any(w in page for w in ("ZZDXONE", "ZZDXFOUR", "ZZCCONE", "ZZCCFOUR",
                                          "ZZMEDONE", "ZZMEDFOUR", "ZZMEDSEALED")))
        check(f"{who}: the screen says why rows are grouped under Other",
              "Rows with fewer than 5 are grouped under Other" in re.sub(r"\s+", " ", page))

    book = client.get(urls["xlsx"] + RANGE1)
    check(f"{who}: Excel export is generated", book.status_code == 200, book.status_code)
    # The Diagnoses sheet holds the diagnosis table and then the complaint table.
    # Rows that are a ranking label with a whole-number count are read in order.
    ranked = [(r[0], r[1]) for r in sheet_rows(book.data, "Diagnoses")
              if r and isinstance(r[1], int)
              and str(r[0]).startswith(("ZZ", "Other (each", "Confidential"))]
    dx_only = ranked[:len(expect_dx)]
    check(f"{who}: the Excel diagnosis ranking is the model's, row for row",
          dx_only == expect_dx, (ranked, expect_dx))
    check(f"{who}: the Excel complaint ranking that follows is the model's, row for row",
          ranked[len(expect_dx):] == [("ZZCC" + l[4:] if l.startswith("ZZDX") else l, n)
                                      for l, n in expect_dx], (ranked, expect_cc))
    med_rows = [(r[0], r[1]) for r in sheet_rows(book.data, "Items Dispensed")
                if r and r[0] and isinstance(r[1], int)
                and str(r[0]).startswith(("ZZMED", "Other", "Confid"))]
    check(f"{who}: the Excel most-dispensed table is the model's, row for row",
          med_rows == expect_med, (med_rows, expect_med))
    text = xlsx_text(book.data)
    complaint_names = {label for label, _ in expect_cc}
    check(f"{who}: the Excel complaints ranking has exactly the model's labels",
          {m for m in re.findall(r"ZZCC[A-Z]+", text)} == {l for l in complaint_names
                                                           if l.startswith("ZZCC")})
    pdf = client.get(urls["pdf"] + RANGE1)
    ptext = pdf_text(pdf.data)
    check(f"{who}: PDF export is generated", pdf.status_code == 200, pdf.status_code)
    named_labels = {l for l, _ in expect_dx if l.startswith("ZZDX")}
    check(f"{who}: the PDF names exactly the model's diagnosis labels",
          set(re.findall(r"ZZDX[A-Z]+", ptext)) == named_labels,
          (set(re.findall(r"ZZDX[A-Z]+", ptext)), named_labels))
    check(f"{who}: the PDF names exactly the model's medicine labels",
          set(re.findall(r"ZZMED[A-Z]+", ptext))
          == {l for l, _ in expect_med if l.startswith("ZZMED")})
    if role != "doctor":
        check(f"{who}: Excel and PDF carry the Other label and the grouping note",
              OTHER_LABEL in text and OTHER_LABEL in ptext
              and "grouped under Other" in text)
        check(f"{who}: neither file has a group of 1 or 4, or sealed medicine, by name",
              not any(w in text + ptext for w in ("ZZDXONE", "ZZDXFOUR", "ZZCCONE",
                                                  "ZZCCFOUR", "ZZMEDONE", "ZZMEDFOUR",
                                                  "ZZMEDSEALED")))
        check(f"{who}: the Excel basis line for restricted visits reads 'Fewer than 5'",
              "Fewer than 5" in text)
    else:
        check("Doctor: Excel and PDF name every group, with no Other row",
              "ZZDXONE" in text and "ZZDXFOUR" in text and OTHER_LABEL not in text
              and OTHER_LABEL not in ptext)

# Day 2 on screen: the Confidential row of 5 is named for the Administrator.
page = body(admin.get(SURFACES["Administrator"]["page"] + RANGE2))
check("Administrator: on day 2 the screen names the Confidential row with 5",
      ("Confidential", 5) in [(l.strip(), int(n)) for l, n in RANK_ROW.findall(page)])
check("Administrator: on day 2 the Confidential medicine row is in the chart data",
      ("Confidential", 10) in (js_series(page, "topMedicines", "units") or []))

# Conditions: screen chart, print table, Excel, PDF.
COND_SURFACES = (
    ("Barangay Official", official, "/official/community-reports",
     "/official/community-reports/export.xlsx", "/official/community-reports/export.pdf"),
)
for who, client, page_url, xlsx_url, pdf_url in COND_SURFACES:
    page = body(client.get(page_url + RANGE1))
    expected = pairs(report_o["conditions"]["counts"], "condition", "count")
    expected = sorted([p for p in expected if p[0] != OTHER_LABEL],
                      key=lambda p: (-p[1], p[0].lower())) + [(OTHER_LABEL, 5)]
    chart = js_series(page, "concerns", "counts")
    check(f"{who}: the Community Reports condition chart is the folded list, Other last",
          chart == expected, (chart, expected))
    check(f"{who}: its print table matches the chart",
          print_table(page, "Chronic conditions recorded at registration (whole register)")
          == expected)
    book = client.get(xlsx_url + RANGE1)
    rows_x = []
    for r in sheet_rows(book.data, "Health Concerns"):
        if r and r[0] in ("Hypertension", "Diabetes", "Asthma", "Tuberculosis", OTHER_LABEL)                 and isinstance(r[1], int):
            rows_x.append((r[0], r[1]))
            if r[0] == OTHER_LABEL:
                break  # the diagnoses table that follows has its own Other row
    check(f"{who}: the Excel condition table matches",
          rows_x == expected, (rows_x, expected))
    ptext = pdf_text(client.get(pdf_url + RANGE1).data)
    check(f"{who}: the PDF lists the Other row",
          OTHER_LABEL in ptext)
    check(f"{who}: the Community Reports diagnosis list is folded too",
          "ZZDXONE" not in page and "ZZDXFOUR" not in page
          and ("Other (each fewer than 5)" in page))
for who, client, urls in (("Administrator", admin, SURFACES["Administrator"]),
                          ("Barangay Official", official, SURFACES["Barangay Official"])):
    page = body(client.get(urls["page"] + RANGE1))
    senior = js_series(page, "seniorConditions", "counts") or []
    check(f"{who}: the senior-condition chart has no condition under 5 by name, "
          "and its print table agrees",
          all(n >= 5 or n == 0 or l == OTHER_LABEL for l, n in senior)
          and print_table(page, "Senior patients by recorded chronic condition, whole register")
          == senior, senior)
    book_rows = [(r[0], r[1]) for r in sheet_rows(
        client.get(urls["xlsx"] + RANGE1).data, "Senior Citizens")
                 if r and r[0] in ("Hypertension", "Diabetes", "Asthma", "Tuberculosis",
                                   OTHER_LABEL) and isinstance(r[1], int)]
    check(f"{who}: the Excel senior-condition table agrees with the chart", book_rows == senior,
          (book_rows, senior))

# ===========================================================================
section("5. THE ADMINISTRATOR'S AND THE OFFICIAL'S WHOLE CRAWL SHOWS NO SMALL GROUP")
SMALL = ("ZZDXONE", "ZZDXFOUR", "ZZCCONE", "ZZCCFOUR", "ZZMEDONE", "ZZMEDFOUR",
         "ZZMEDSEALED")
plain_gets = sorted({rule.rule for rule in app.url_map.iter_rules()
                     if "GET" in rule.methods and not rule.arguments
                     and not rule.rule.startswith("/static")
                     and rule.rule != "/logout" and not rule.rule.endswith(".idsbackup")})
for who, client, prefix in (("Administrator", admin, "/admin"),
                            ("Barangay Official", official, "/official")):
    found, pages = [], 0
    for url in plain_gets:
        if not url.startswith(prefix):
            continue
        for suffix in ("", RANGE1, RANGE2):
            text = readable(client.get(url + suffix))
            pages += 1
            hit = [w for w in SMALL if w in text]
            if hit:
                found.append(f"{url}{suffix}: {hit}")
    check(f"{who}: no group of 1 or 4 is named on any of the {pages} pages, exports or print "
          "views reachable", not found, "\n".join(found))
    check(f"{who}: the crawl is wide", pages >= 30, pages)

# ===========================================================================
section("6. PHONE NUMBERS ARE MASKED ON EVERY SMS SCREEN, EXPORT AND JSON RESPONSE")
from utils.phone import mask_numbers_in_text, mask_phone  # noqa: E402

NUMBERS = {"broadcast_one": "09123450001", "broadcast_two": "09123450002",
           "appointment": "09123450003", "audit": "09123450004"}
MASKED = {k: f"•••• ••• {v[-4:]}" for k, v in NUMBERS.items()}
check("mask_phone keeps the last four digits and the fixed prefix",
      mask_phone("09123450001") == "•••• ••• 0001"
      and mask_phone("+63 912 345 0001") == "•••• ••• 0001"
      and mask_phone("639123450001") == "•••• ••• 0001")
check("mask_phone shows nothing of a short value, keeps an empty one empty, "
      "and is safe to apply twice",
      mask_phone("1234") == "•••• •••" and mask_phone("") == "" and mask_phone(None) == ""
      and mask_phone(mask_phone("09123450001")) == "•••• ••• 0001")
check("mask_numbers_in_text masks a number inside a sentence and leaves other digits",
      mask_numbers_in_text("Texted 09123450004 on 20260905, ref 12345678901.")
      == "Texted •••• ••• 0004 on 20260905, ref 12345678901.")

broadcast_id = run_sql(
    "INSERT INTO sms_broadcasts (reference, message, category, audience, audience_label, "
    "recipient_count, sent_count, failed_count, segments, encoding, status, delivery_mode, "
    "sent_by) VALUES ('BC-ZZ5D01', 'Clinic notice.', 'General Reminder', 'all_patients', "
    "'All patients', 2, 2, 0, 1, 'GSM-7', 'Sent', 'simulated', ?)", (ADMIN_ID,))
for name, number in (("ZZ Recipient One", NUMBERS["broadcast_one"]),
                     ("ZZ Recipient Two", NUMBERS["broadcast_two"])):
    run_sql("INSERT INTO sms_broadcast_recipients (broadcast_id, recipient_name, "
            "mobile_number, status, gateway_response) VALUES (?,?,?,'Sent',?)",
            (broadcast_id, name, number, f"accepted for {number}"))
run_sql("INSERT INTO appointment_sms_messages (reference, patient_id, event, recipient_name, "
        "mobile_number, message, segments, encoding, status, delivery_mode, sent_by) "
        "VALUES ('SMS-ZZ5D01', ?, 'booked', 'ZZ Appointment Person', ?, 'Booked.', 1, "
        "'GSM-7', 'Sent', 'simulated', ?)", (patient_ids[0], NUMBERS["appointment"], STAFF_ID))
from models.audit import ACTION_APPOINTMENT_SMS_SENT, MODULE_SMS  # noqa: E402
run_sql("INSERT INTO audit_logs (user_id, action, module, details) VALUES (?,?,?,?)",
        (STAFF_ID, ACTION_APPOINTMENT_SMS_SENT, MODULE_SMS,
         f"Texted {NUMBERS['audit']} about a booking."))

ROLE_CLIENTS = (("Administrator", admin, "/admin/sms_notifications"),
                ("Doctor", doctor, "/staff/sms_notifications"),
                ("Health Center Staff", staff, "/staff/sms_notifications"),
                ("Barangay Official", official, "/official/sms_notifications"))
FULL = re.compile(r"(?<!\d)(?:\+?63|0)?912345000[1-4](?!\d)")
for who, client, url in ROLE_CLIENTS:
    page = body(client.get(url))
    check(f"{who}: SMS Notifications renders, with the recipient names and counts",
          "ZZ Recipient One" in page or who == "Administrator" and "ZZ Recipient One" in page,
          page[:200])
    check(f"{who}: no full number on SMS Notifications", not FULL.search(page))
    if who == "Administrator":
        check("Administrator: the recipient list has names and results and no number at all",
              "ZZ Recipient One" in page and "•••• •••" not in page)
    else:
        check(f"{who}: the recipient list shows only the last four digits",
              MASKED["broadcast_one"] in page and MASKED["broadcast_two"] in page)
    if who in ("Doctor", "Health Center Staff"):
        check(f"{who}: the appointment-confirmation panel shows only the last four digits",
              MASKED["appointment"] in page)
    else:
        check(f"{who}: the appointment-confirmation panel is not on the page",
              "ZZ Appointment Person" not in page)

audit_urls = sorted(r.rule for r in app.url_map.iter_rules()
                    if r.rule.startswith("/admin") and "audit" in r.rule
                    and "GET" in r.methods and not r.arguments)
audit_page = next((u for u in audit_urls if not u.endswith((".xlsx", ".pdf"))), None)
audit_text = body(admin.get(audit_page)) if audit_page else ""
check("Administrator: Audit Logs shows the SMS entry with the number masked",
      MASKED["audit"] in audit_text and not FULL.search(audit_text), audit_urls)
for url in audit_urls:
    if url.endswith((".xlsx", ".pdf")):
        text = readable(admin.get(url))
        check(f"Administrator: {url} carries no full number and shows the masked one",
              not FULL.search(text) and (MASKED["audit"] in text or url.endswith(".pdf")))

crawled = {"pages": 0, "exports": 0, "json": 0}
leaks = []
for who, client, _ in ROLE_CLIENTS:
    prefixes = {"Administrator": "/admin", "Doctor": "/staff", "Health Center Staff": "/staff",
                "Barangay Official": "/official"}
    for url in plain_gets + ["/session/keep-alive"]:
        if url != "/session/keep-alive" and not url.startswith(prefixes[who]):
            continue
        response = (client.post(url, headers={"X-Requested-With": "fetch"})
                    if url == "/session/keep-alive" else client.get(url))
        kind = response.headers.get("Content-Type", "")
        crawled["exports" if ("spreadsheetml" in kind or "pdf" in kind)
                else "json" if url == "/session/keep-alive" else "pages"] += 1
        if FULL.search(readable(response)):
            leaks.append(f"{who} {url}")
print(f"       crawled {crawled['pages']} pages, {crawled['exports']} exports, "
      f"{crawled['json']} JSON responses, over four roles")
check("*** no full marker number on any page, export, print view or JSON response, "
      "for any role ***", not leaks, "\n".join(leaks))
check("the crawl is wide: every screen and export of four roles, and the keep-alive JSON",
      crawled["pages"] >= 35 and crawled["exports"] >= 20 and crawled["json"] >= 4, crawled)

# The send path still gets full numbers.
import models.broadcasts as broadcasts_model  # noqa: E402

sent_to = []
real_deliver = broadcasts_model.deliver_announcement_sms


def spy(message, number, api_key=None):
    sent_to.append(number)
    return real_deliver(message, number, api_key)


broadcasts_model.deliver_announcement_sms = spy
try:
    response = admin.post("/admin/sms_notifications/send", data={
        "message": "Health center open on Saturday.", "audience": "all_patients",
        "category": "General Reminder"})
finally:
    broadcasts_model.deliver_announcement_sms = real_deliver
check("a broadcast still sends: the gateway is handed whole numbers",
      response.status_code == 302 and len(sent_to) >= 1
      and all(re.fullmatch(r"(?:63|0)9\d{9}|\+?639\d{9}", n) for n in sent_to),
      (response.status_code, len(sent_to)))
newest = q1("SELECT id FROM sms_broadcasts ORDER BY id DESC LIMIT 1")[0]
stored = [r[0] for r in db().execute(
    "SELECT mobile_number FROM sms_broadcast_recipients WHERE broadcast_id = ?", (newest,))]
check("the stored recipient rows keep the full numbers, and match what was sent",
      sorted(stored) == sorted(sent_to) and all("•" not in n for n in stored))
after = body(official.get("/official/sms_notifications"))
check("and the next page view of that broadcast shows only masked numbers",
      "•••• •••" in after and not FULL.search(after)
      and not any(n in after for n in stored))
audit_row = q1("SELECT details FROM audit_logs WHERE action = 'Sent community SMS broadcast' "
               "ORDER BY id DESC LIMIT 1")
check("the broadcast's audit entry names no number", audit_row
      and not re.search(r"(?<!\d)(?:\+?63|0)9\d{9}(?!\d)", audit_row[0]))

# ===========================================================================
section("7. TASK 5e: SERVICE BREAKDOWN FOLDS, SINGLE COUNTS READ 'FEWER THAN 5'")
from models.appointments import get_appointment_report  # noqa: E402
from models.patients import get_condition_counts  # noqa: E402


def free_appointment_day():
    day = date.today() - timedelta(days=40)
    while q1("SELECT COUNT(*) FROM appointments WHERE appointment_date = ?",
             (day.isoformat(),))[0]:
        day -= timedelta(days=1)
    return day.isoformat()


SVC_DAY = free_appointment_day()
SVC_RANGE = f"?from={SVC_DAY}&to={SVC_DAY}"
for name, count in (("ONE", 1), ("FOUR", 4), ("FIVE", 5), ("TWELVE", 12)):
    for _ in range(count):
        _n[0] += 1
        run_sql("INSERT INTO appointments (patient_id, scheduled_by, appointment_type, "
                "appointment_date, status, is_walkin) VALUES (?,?,?,?,'Completed',1)",
                (patient_ids[_n[0] % len(patient_ids)], STAFF_ID, f"ZZSVC{name}", SVC_DAY))
day_obj = date.fromisoformat(SVC_DAY)


def services(role):
    with app.app_context():
        return get_appointment_report(months=6, start=day_obj, end=day_obj,
                                      viewer_role=role)["window"]


for who in ("Administrator", "Barangay Official"):
    window = services(ROLES[who])
    check(f"{who}: services show 12 and 5 named, and 1 + 4 fold into Other = 5",
          pairs(window["services"], "service", "count")
          == [("ZZSVCTWELVE", 12), ("ZZSVCFIVE", 5), (OTHER_LABEL, 5)], window["services"])
window_d = services("doctor")
check("Doctor: every service is named with its own count, no Other row",
      pairs(window_d["services"], "service", "count")
      == [("ZZSVCTWELVE", 12), ("ZZSVCFIVE", 5), ("ZZSVCFOUR", 4), ("ZZSVCONE", 1)])
check("service totals are equal for all roles and the rows add up to the total",
      len({services(r)["total"] for r in ROLES.values()}) == 1
      and window_d["total"] == 22
      and sum(e["count"] for e in services("admin")["services"]) == 22)
check("with no prenatal visits the count reads 0 for every role",
      services("admin")["prenatal_visits"] == 0 and window_d["prenatal_visits"] == 0)
for _ in range(2):
    _n[0] += 1
    run_sql("INSERT INTO appointments (patient_id, scheduled_by, appointment_type, "
            "appointment_date, status, is_walkin) VALUES (?,?,'Prenatal Check-up',?,"
            "'Completed',1)", (patient_ids[_n[0] % len(patient_ids)], STAFF_ID, SVC_DAY))
check("two prenatal visits: 'Fewer than 5' for the Administrator, 2 for the Doctor",
      services("admin")["prenatal_visits"] == "Fewer than 5"
      and services("doctor")["prenatal_visits"] == 2)
for who, client, urls in (("Administrator", admin, SURFACES["Administrator"]),
                          ("Barangay Official", official, SURFACES["Barangay Official"]),
                          ("Doctor", doctor, SURFACES["Doctor"])):
    page = body(client.get(urls["page"] + SVC_RANGE))
    table_rows = print_table(page, "Appointments by service requested")
    expected = pairs(services(ROLES[who])["services"], "service", "count")
    check(f"{who}: the screen's service chart data and print table are the model's",
          table_rows == expected and js_series(page, "services", "counts") == expected,
          (table_rows, expected))
    xl = [(r[0], r[1]) for r in sheet_rows(client.get(urls["xlsx"] + SVC_RANGE).data,
                                          "Health Services")
          if r and isinstance(r[1], int)
          and str(r[0]).startswith(("ZZSVC", "Other (each", "Prenatal"))]
    check(f"{who}: the Excel service table is the model's", xl == expected, (xl, expected))
    text = pdf_text(client.get(urls["pdf"] + SVC_RANGE).data)
    check(f"{who}: the PDF names exactly the model's services",
          set(re.findall(r"ZZSVC[A-Z]+", text))
          == {l for l, _ in expected if l.startswith("ZZSVC")})
    if ROLES[who] != "doctor":
        check(f"{who}: no service of 1 to 4 appointments is named on the screen",
              "ZZSVCONE" not in page and "ZZSVCFOUR" not in page)


def other_only(role, free_text_patients):
    """Condition counts with N patients holding only free-text conditions."""
    run_sql("UPDATE patients SET medical_conditions = '[]'")
    for pid in patient_ids[:free_text_patients]:
        run_sql("UPDATE patients SET medical_conditions = ? WHERE id = ?",
                ('["ZZFREETEXT"]', pid))
    with app.app_context():
        return get_condition_counts(viewer_role=role)


for n, expected_admin in ((0, 0), (1, "Fewer than 5"), (4, "Fewer than 5"), (5, 5), (12, 12)):
    for who in ("admin", "barangay official"):
        counts = other_only(who, n)
        check(f"{who}: {n} patient(s) with only free-text conditions read {expected_admin!r}",
              counts["other_only_shown"] == expected_admin and counts["other_only"] == n
              and counts["with_any_shown"] == expected_admin)
    check(f"doctor: {n} patient(s) with only free-text conditions read exactly {n}",
          other_only("doctor", n)["other_only_shown"] == n)
other_only("official", 3)
page = body(official.get("/official/community-reports"))
check("Community Reports prints 'Fewer than 5 patients', not the number, for the Official",
      "Fewer than 5 patients" in re.sub(r"\s+", " ", page))
check("and its Excel file carries 'Fewer than 5' for the free-text-only count",
      "Fewer than 5" in xlsx_text(official.get("/official/community-reports/export.xlsx").data))

# ===========================================================================
section("8. TASK 5e: NO AUDIT ENTRY STORES A WHOLE PHONE NUMBER")
from models.audit import log_action  # noqa: E402

with app.app_context():
    entry_id = log_action(STAFF_ID, "Sent appointment SMS to patient", MODULE_SMS,
                          "Texted 09123450009 and +639123450008 about a booking.")
    plain_id = log_action(STAFF_ID, "Test entry", MODULE_SMS, "Ref 20260905 id 12345678901")
check("a new audit entry that quotes numbers stores them as their last four digits",
      q1("SELECT details FROM audit_logs WHERE id = ?", (entry_id,))[0]
      == "Texted •••• ••• 0009 and •••• ••• 0008 about a booking.")
check("an entry with no number is stored unchanged",
      q1("SELECT details FROM audit_logs WHERE id = ?", (plain_id,))[0]
      == "Ref 20260905 id 12345678901")

import shutil as _shutil  # noqa: E402
import tempfile  # noqa: E402
from pathlib import Path  # noqa: E402
from migrations import migrate_mask_audit_phone_numbers as migration  # noqa: E402

work = Path(tempfile.mkdtemp(prefix="idsmask_"))
copy = work / "copy.db"
_shutil.copy(TEST_DB, copy)
c = sqlite3.connect(copy)
for text in ("Sent to 09123450011 on a booking.", "No number here, ref 20260905.",
             "Two: 639123450012 and +639123450013."):
    c.execute("INSERT INTO audit_logs (user_id, action, module, details) VALUES (1,'Old entry',"
              "'SMS',?)", (text,))
c.commit()
before_rows = c.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0]
c.close()
result = migration.run(copy, backup_dir=work / "backups")
c = sqlite3.connect(copy)
rows = [r[0] for r in c.execute("SELECT details FROM audit_logs WHERE action = 'Old entry' "
                                "ORDER BY id")]
check("the migration masks the numbers in old entries and leaves the rest of the text",
      rows == ["Sent to •••• ••• 0011 on a booking.", "No number here, ref 20260905.",
               "Two: •••• ••• 0012 and •••• ••• 0013."], rows)
# Three: the two numbered entries added here and the one section 6 inserted raw.
check("it reports counts only: 3 entries masked, one record written, no problems",
      result["entries_masked"] == 3 and result["record_written"] and not result["problems"],
      result)
check("it backed the database up first",
      result["backup"] and Path(result["backup"]).exists()
      and len(list((work / "backups").iterdir())) == 1)
check("it wrote exactly ONE audit entry recording the run",
      c.execute("SELECT COUNT(*) FROM audit_logs WHERE action = ?",
                (migration.ACTION,)).fetchone()[0] == 1
      and c.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0] == before_rows + 1)
check("that entry holds no number", not re.search(
    r"\d{9,}", c.execute("SELECT details FROM audit_logs WHERE action = ?",
                         (migration.ACTION,)).fetchone()[0]))
c.close()
again = migration.run(copy, backup_dir=work / "backups")
c = sqlite3.connect(copy)
check("a second run changes nothing: no backup, no entry, no masking",
      again["backup"] is None and not again["entries_masked"]
      and c.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0] == before_rows + 1
      and len(list((work / "backups").iterdir())) == 1)
c.close()
_shutil.rmtree(work, ignore_errors=True)

# ===========================================================================
passed = sum(1 for _, ok in results if ok)
print("\n" + "=" * 78)
print(f"RESULT: {passed}/{len(results)} checks passed")
failed = [name for name, ok in results if not ok]
for name in failed:
    print(f"  FAILED: {name}")
sys.exit(1 if failed else 0)
