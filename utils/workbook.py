"""Writing a .xlsx file, with no knowledge of what is in it.

The mechanics only: sheets, headed tables, wrapped explanatory notes, column
widths, and the one house style they are all drawn in. What the numbers *are*
belongs to routes/report_export.py, which is the module that knows about
consultations and reorder levels; this one would write a workbook of anything.

Kept apart for the same reason utils/dates.py is: two report screens produce a
workbook each, and a third would produce a third. A shared writer means every
export the project ever grows has the same headings, the same freeze row and the
same way of stating a caveat, rather than each screen inventing its own.

WHY openpyxl
------------
It is the standard library for the .xlsx format in Python, it is pure Python
with one small pure-Python dependency (et-xmlfile), and it writes a real
OOXML workbook offline -- nothing in it reaches the network, which is the
constraint this project works under. The alternatives were weighed:

  * xlsxwriter writes beautifully but cannot read, and a future task that wants
    to check a produced file would need a second library to open it. openpyxl
    reads and writes, so the verification of an export is done with the same
    dependency that produced it.
  * pandas.to_excel would pull in pandas and numpy -- tens of megabytes of
    compiled extensions -- to format a few dozen rows this project already holds
    as plain dicts. It also needs openpyxl underneath regardless.
  * Writing CSV needs no dependency at all, but a CSV is one flat table with no
    sheets, no headings and no room for the notes that say which figures a date
    range narrowed and which it did not. Those notes are the difference between
    a report and a pile of numbers, and this project does not ship the second.

A NOTE ON NOTES
---------------
Long explanatory text is wrapped in Python and written one line per row, rather
than put in a merged cell with wrap_text. Excel does not auto-fit the height of
a merged cell, so that route produces a paragraph clipped to one visible line --
the caveat would be in the file and invisible in the application that opens it,
which is the worst of both. One line per row prints and displays correctly
everywhere, and the wrapped rows are excluded from the column-width calculation
so a paragraph cannot stretch column A to the width of the screen.
"""
from datetime import date, datetime
from io import BytesIO
from textwrap import wrap

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# The MIME type a browser needs to hand the file to Excel, LibreOffice or Google
# Sheets rather than showing it as a download of unknown type.
XLSX_MIMETYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

# 02_UI_DESIGN_SYSTEM's teal, so an exported sheet is recognisably from the same
# system as the screen it was taken off.
_TEAL = "FF169B92"
_TEAL_DARK = "FF0F6E56"
_MUTED = "FF6B7280"

_TITLE_FONT = Font(bold=True, size=14, color=_TEAL_DARK)
_BLOCK_FONT = Font(bold=True, size=11.5)
_HEAD_FONT = Font(bold=True, size=10.5, color="FFFFFFFF")
_NOTE_FONT = Font(italic=True, size=9.5, color=_MUTED)
_HEAD_FILL = PatternFill("solid", fgColor=_TEAL)
_HEAD_BORDER = Border(bottom=Side(style="thin", color="FFD1D5DB"))

_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)
_RIGHT = Alignment(horizontal="right", vertical="center")
_HEAD_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_HEAD_RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)

# How wide a wrapped note line may get. Comfortably inside a printed page and
# narrow enough that a caveat reads as prose rather than as a ruler.
_NOTE_WIDTH = 104

# Column width bounds, in Excel's character units.
_MIN_WIDTH, _MAX_WIDTH = 11, 46

# Excel forbids these in a sheet name and caps it at 31 characters.
_FORBIDDEN = set("[]:*?/\\")
_TITLE_LIMIT = 31

# What an absent figure is written as. The same em dash utils/dates.py uses for
# a date it cannot render, so a blank in an export reads the same as a blank on
# a screen -- and is never written as 0, which would be a claim.
MISSING = "—"


def table(title, columns, rows, note=None, empty="Nothing recorded in this range."):
    """A headed table. `rows` is a list of lists, parallel to `columns`.

    `empty` is written in place of the rows when there are none, so a section
    with nothing in it says so instead of arriving as a heading over blank space
    that a reader could take for a rendering fault.
    """
    return {
        "kind": "table",
        "title": title,
        "columns": list(columns),
        "rows": [list(row) for row in rows],
        "note": note,
        "empty": empty,
    }


def figures(title, rows, note=None):
    """A Figure / Value / Basis block -- the spreadsheet form of a stat card.

    `rows` are (label, value) or (label, value, basis). The Basis column is
    dropped when no row fills it, so a block of plain counts is two columns wide
    rather than two columns and a stripe of blanks.
    """
    normalized = [list(row) + [""] * (3 - len(row)) for row in rows]
    if any(str(row[2]).strip() for row in normalized):
        return table(title, ("Figure", "Value", "Basis"), normalized, note=note,
                     empty="No figures in this section.")
    return table(title, ("Figure", "Value"), [row[:2] for row in normalized],
                 note=note, empty="No figures in this section.")


def sheet(title, blocks, heading=None, intro=None):
    """One worksheet: a heading, an optional paragraph, then blocks in order."""
    return {
        "title": title,
        "heading": heading or title,
        "intro": intro,
        "blocks": [block for block in blocks if block],
    }


def build(sheets):
    """The sheets, as the bytes of a real .xlsx file."""
    book = Workbook()
    # Workbook() starts with one empty sheet. Removed rather than reused so the
    # first real sheet is created by the same code path as the rest.
    book.remove(book.active)
    for spec in sheets:
        _write_sheet(book, spec)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def safe_filename(text):
    """A download name a Windows, macOS or Linux filesystem will all accept."""
    cleaned = "".join(
        " " if character in '<>:"/\\|?*' else character
        for character in str(text)
    )
    return " ".join(cleaned.split()).strip(" .") or "export"


# ---------------------------------------------------------------------------
# Writing
# ---------------------------------------------------------------------------

def _safe_title(text, used=None):
    """A sheet name Excel will accept, and that no other sheet already holds."""
    cleaned = "".join(" " if ch in _FORBIDDEN else ch for ch in str(text)).strip()
    cleaned = " ".join(cleaned.split())[:_TITLE_LIMIT] or "Sheet"
    if used is None:
        return cleaned
    candidate, suffix = cleaned, 2
    while candidate.lower() in used:
        tail = f" {suffix}"
        candidate = cleaned[: _TITLE_LIMIT - len(tail)] + tail
        suffix += 1
    used.add(candidate.lower())
    return candidate


def _wrapped(text):
    """A paragraph as the lines it should occupy, one per row."""
    if not text:
        return []
    lines = []
    for paragraph in str(text).split("\n"):
        lines.extend(wrap(paragraph.strip(), _NOTE_WIDTH) or [""])
    return lines


def _is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _write_sheet(book, spec):
    used = getattr(book, "_ids_used_titles", None)
    if used is None:
        used = set()
        book._ids_used_titles = used
    worksheet = book.create_sheet(_safe_title(spec["title"], used))

    # Rows holding wrapped prose. Excluded from the width calculation below so a
    # 104-character caveat does not decide how wide column A is.
    prose_rows = set()
    row = 1

    cell = worksheet.cell(row=row, column=1, value=spec["heading"])
    cell.font = _TITLE_FONT
    prose_rows.add(row)
    row += 2

    for line in _wrapped(spec.get("intro")):
        worksheet.cell(row=row, column=1, value=line).font = _NOTE_FONT
        prose_rows.add(row)
        row += 1
    if spec.get("intro"):
        row += 1

    for block in spec["blocks"]:
        row = _write_block(worksheet, block, row, prose_rows)

    _autosize(worksheet, prose_rows)
    # The heading stays put while the reader scrolls a long table.
    worksheet.freeze_panes = "A2"
    # So a sheet sent to a printer from Excel comes out on one page across
    # rather than with its last column on a sheet of its own.
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    return worksheet


def _write_block(worksheet, block, row, prose_rows):
    if block["kind"] == "note":
        for line in _wrapped(block["text"]):
            worksheet.cell(row=row, column=1, value=line).font = _NOTE_FONT
            prose_rows.add(row)
            row += 1
        return row + 1

    if block["title"]:
        worksheet.cell(row=row, column=1, value=block["title"]).font = _BLOCK_FONT
        prose_rows.add(row)
        row += 1

    for line in _wrapped(block["note"]):
        worksheet.cell(row=row, column=1, value=line).font = _NOTE_FONT
        prose_rows.add(row)
        row += 1

    # A column is a number column when its first filled cell is a number, so the
    # header sits over its figures rather than away from them.
    numeric = []
    for index in range(len(block["columns"])):
        column_values = [
            data[index] for data in block["rows"] if index < len(data)
        ]
        first = next((value for value in column_values if value is not None), None)
        numeric.append(_is_number(first))

    for index, name in enumerate(block["columns"], start=1):
        cell = worksheet.cell(row=row, column=index, value=name)
        cell.font = _HEAD_FONT
        cell.fill = _HEAD_FILL
        cell.border = _HEAD_BORDER
        cell.alignment = _HEAD_RIGHT if numeric[index - 1] else _HEAD_LEFT
    row += 1

    if not block["rows"]:
        worksheet.cell(row=row, column=1, value=block["empty"]).font = _NOTE_FONT
        prose_rows.add(row)
        return row + 2

    for data in block["rows"]:
        for index, value in enumerate(data, start=1):
            cell = worksheet.cell(row=row, column=index, value=_coerce(value))
            if _is_number(value):
                cell.alignment = _RIGHT
                cell.number_format = "#,##0" if isinstance(value, int) else "#,##0.0"
            elif isinstance(value, (date, datetime)):
                cell.number_format = "yyyy-mm-dd"
            else:
                cell.alignment = _LEFT
        row += 1

    return row + 1


def _coerce(value):
    """What actually goes in the cell.

    None becomes the em dash rather than an empty cell or a zero: an absent
    figure and a figure of nought are different claims, and this project states
    the difference everywhere else it reports one.
    """
    if value is None:
        return MISSING
    if isinstance(value, (int, float, str, date, datetime)):
        return value
    return str(value)


def _autosize(worksheet, prose_rows):
    widths = {}
    for row_cells in worksheet.iter_rows():
        for cell in row_cells:
            if cell.value is None or cell.row in prose_rows:
                continue
            length = len(f"{cell.value:,}") if _is_number(cell.value) else len(str(cell.value))
            widths[cell.column] = max(widths.get(cell.column, 0), length)
    for column, width in widths.items():
        worksheet.column_dimensions[get_column_letter(column)].width = min(
            _MAX_WIDTH, max(_MIN_WIDTH, width + 3)
        )
